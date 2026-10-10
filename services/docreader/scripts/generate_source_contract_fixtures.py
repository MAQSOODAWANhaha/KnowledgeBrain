"""Generate recorded *real Python parser responses* consumed by Rust tests.

Run from the repository root with PYTHONPATH=services and the docreader test
requirements installed. The fixture deliberately mixes native text, a blank
page, a scanned page, and a merged DOCX grid with multiline/tabbed Unicode.
"""
import argparse
import base64
import hashlib
import json
from io import BytesIO
from pathlib import Path

from docx import Document as DocxDocument
from docx.oxml import OxmlElement
from PIL import Image, ImageDraw
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject, NumberObject

from docreader.parser.output_inventory import parse_output_inventory


def mixed_pdf():
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    font = writer._add_object(DictionaryObject({
        NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'),
        NameObject('/BaseFont'): NameObject('/Helvetica'),
    }))
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 10 270 Td (Native evidence 10 and 20) Tj 0 -30 Td (Repeated header) Tj ET")
    page[NameObject('/Contents')] = writer._add_object(stream)
    writer.add_blank_page(width=300, height=300)
    page = writer.add_blank_page(width=300, height=300)
    image = Image.new("RGB", (300, 300), "white")
    ImageDraw.Draw(image).text((30, 100), "Scanned evidence: 10 / 20", fill="black")
    image_object = DecodedStreamObject()
    image_object.set_data(image.tobytes())
    image_object.update({NameObject('/Type'): NameObject('/XObject'), NameObject('/Subtype'): NameObject('/Image'),
                         NameObject('/Width'): NumberObject(300), NameObject('/Height'): NumberObject(300),
                         NameObject('/ColorSpace'): NameObject('/DeviceRGB'), NameObject('/BitsPerComponent'): NumberObject(8)})
    compressed = image_object.flate_encode()
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/XObject'): DictionaryObject({NameObject('/Im0'): writer._add_object(compressed)})})
    stream = DecodedStreamObject()
    stream.set_data(b"q 300 0 0 300 0 0 cm /Im0 Do Q")
    page[NameObject('/Contents')] = writer._add_object(stream)
    raw = BytesIO()
    writer.write(raw)
    return raw.getvalue()


def multiline_docx():
    document = DocxDocument()
    document.add_heading("资格要求", 1)
    document.add_paragraph("重复的正文。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "项目名称"
    header = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(header)
    cell = table.cell(1, 0)
    cell.text = "10"
    cell.add_paragraph("20")
    paragraph = table.cell(1, 1).paragraphs[0]
    paragraph.add_run("中").add_text("文")
    paragraph.add_run().add_break()
    paragraph.add_run("承诺\t否")
    document.add_paragraph("重复的正文。")
    raw = BytesIO()
    document.save(raw)
    return raw.getvalue()


def records():
    from docreader.config import CONFIG
    from dataclasses import replace
    from unittest.mock import patch
    config = replace(CONFIG, pdf_render_parallelism=1, pdf_render_dpi=72, pdf_render_max_edge=300)
    output = []
    for name, source in [("mixed.pdf", mixed_pdf()), ("multiline.docx", multiline_docx())]:
        media = name.rsplit('.', 1)[-1]
        with patch("docreader.config.CONFIG", config), patch("docreader.parser.pdf_parser.CONFIG", config):
            parsed = parse_output_inventory(name, media, source)
        output.append({
            "name": name, "file_type": media, "file_sha256": hashlib.sha256(source).hexdigest(),
            "markdown": parsed.content,
            "metadata": {key: value if isinstance(value, str) else json.dumps(value, ensure_ascii=False)
                         for key, value in parsed.metadata.items()},
            "structured_source_units": [unit.model_dump(mode="json", exclude={
                "physical_locator", "physical_path", "source_section_id", "source_issues", "header_cells",
            }) for unit in parsed.structured_source_units],
            "images": [{"original_ref": ref, "hex": base64.b64decode(data).hex()}
                       for ref, data in parsed.images.items()],
        })
    return output


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('output', type=Path)
    args = parser.parse_args()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(records(), ensure_ascii=False, indent=2) + '\n')
