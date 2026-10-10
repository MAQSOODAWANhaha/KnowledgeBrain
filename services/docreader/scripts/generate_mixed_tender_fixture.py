"""Four synthetic native carriers, production parsers/wire; no OCR model calls."""
import base64
from datetime import datetime
from io import BytesIO
import json
from pathlib import Path
import runpy
import sys
import zipfile

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from docx import Document
from docx.shared import Inches
from PIL import Image, ImageDraw
from docreader.parser.docx_parser import DocxParser
from docreader.parser.image_parser import ImageParser
from docreader.main import _structured_units_to_proto
from docreader.proto.docreader_pb2 import ImageRef, ReadResponse


def records():
    result = runpy.run_path(str(Path(__file__).with_name('generate_related_tender_fixture.py')))['records'](mixed=True)
    image = Image.new('RGB', (480, 120), 'white')
    ImageDraw.Draw(image).text((20, 40), 'Signed schedule required', fill='black')
    png = BytesIO()
    image.save(png, format='PNG')
    document = Document()
    document.core_properties.created = datetime(2000, 1, 1)
    document.core_properties.modified = datetime(2000, 1, 1)
    document.add_heading('Prescribed response', 1)
    document.add_paragraph('Keep this fixed wording. Submit the signed schedule.')
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = 'Response form'
    table.cell(1, 0).text = 'Bidder name: ____'
    table.cell(1, 1).text = 'Authorized signature: ____'
    document.add_paragraph('Signature instructions:').add_run().add_picture(BytesIO(png.getvalue()), width=Inches(3))
    raw = BytesIO()
    document.save(raw)
    deterministic = BytesIO()
    with zipfile.ZipFile(BytesIO(raw.getvalue())) as source, zipfile.ZipFile(deterministic, 'w') as target:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, (2000, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(info, source.read(name))
    for key, name, parser, source in [
        ('response_docx', 'response.docx', DocxParser, deterministic.getvalue()),
        ('signature_png', 'signature.png', ImageParser, png.getvalue()),
    ]:
        parsed = parser(file_name=name, file_type=name.rsplit('.', 1)[-1]).parse(source)
        response = ReadResponse(markdown_content=parsed.content,
            metadata={k: str(v) for k, v in parsed.metadata.items()},
            structured_source_units=_structured_units_to_proto(parsed.structured_source_units),
            image_refs=[ImageRef(filename=ref, original_ref=ref, mime_type='image/png',
                image_data=base64.b64decode(data)) for ref, data in parsed.images.items()])
        result[key] = response.SerializeToString(deterministic=True).hex()
        result[key + '_source_hex'] = source.hex()
    result['ocr_contract_stub'] = 'Signed schedule required'
    result['ocr_scope'] = 'Golden synthetic text only; not OCR or model-semantic validation.'
    return result


if __name__ == '__main__':
    Path(sys.argv[1]).write_text(json.dumps(records(), sort_keys=True) + '\n')
