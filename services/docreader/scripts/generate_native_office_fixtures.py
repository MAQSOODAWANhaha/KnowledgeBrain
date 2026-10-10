"""Record real native DOCX/XLSX producer results without optional OCR imports.

Run from the repository root:
  python services/docreader/scripts/generate_native_office_fixtures.py OUTPUT.json

Only the native producers are loaded. The package's convenience initializer
eagerly imports unrelated engines; this standalone harness intentionally does
not run that initializer. No producer implementation or dependency is mocked.
"""
import argparse
import base64
from datetime import datetime, timezone
import hashlib
import importlib.abc
import importlib.machinery
import json
from io import BytesIO
from pathlib import Path
import sys
import types
from xml.etree import ElementTree
import zipfile


class NativeOnlyImports(importlib.abc.MetaPathFinder):
    def find_spec(self, fullname, path=None, target=None):
        if fullname.split(".")[0] in {"onnxruntime", "magika", "markitdown"}:
            raise ImportError(f"optional network-capable engine outside native fixture scope: {fullname}")
        return None


sys.meta_path.insert(0, NativeOnlyImports())
SERVICE_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SERVICE_ROOT.parent))
package = types.ModuleType("docreader.parser")
package.__path__ = [str(SERVICE_ROOT / "parser")]
package.__package__ = "docreader.parser"
package.__spec__ = importlib.machinery.ModuleSpec("docreader.parser", loader=None, is_package=True)
sys.modules["docreader.parser"] = package

from docx import Document as DocxDocument
from docx.oxml import OxmlElement
from openpyxl import Workbook
from docreader.parser.docx_parser import DocxParser
from docreader.parser.excel_parser import ExcelParser


def deterministic_archive(raw):
    """Remove archive and core-property clocks, without changing source content."""
    output = BytesIO()
    with zipfile.ZipFile(BytesIO(raw)) as source, zipfile.ZipFile(output, "w") as target:
        for name in sorted(source.namelist()):
            data = source.read(name)
            if name == "docProps/core.xml":
                root = ElementTree.fromstring(data)
                for child in root:
                    if child.tag.rsplit("}", 1)[-1] in {"created", "modified"}:
                        child.text = "2000-01-01T00:00:00Z"
                data = ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)
            info = zipfile.ZipInfo(name, (2000, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            target.writestr(info, data)
    return output.getvalue()


def native_docx():
    document = DocxDocument()
    document.core_properties.created = datetime(2000, 1, 1, tzinfo=timezone.utc)
    document.core_properties.modified = datetime(2000, 1, 1, tzinfo=timezone.utc)
    document.add_heading("第一章 资格要求", 1)
    document.add_paragraph("表前正文😀")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "资格材料"
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    table.cell(1, 0).text = "必须提交营业执照"
    table.cell(1, 0).add_paragraph("10\t20")
    table.cell(1, 1).text = "有效期内"
    document.add_paragraph("表后正文😀")
    raw = BytesIO()
    document.save(raw)
    return deterministic_archive(raw.getvalue())


def native_xlsx():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "资格材料"
    sheet.append(["资格条件", "证明文件"])
    sheet.append(["必须提交营业执照", "有效期内"])
    sheet.append(["多行\n中文😀", "10\t20"])
    sheet.merge_cells("A4:B4")
    sheet["A4"] = "合并单元格"
    workbook.create_sheet("空白附表")
    raw = BytesIO()
    workbook.save(raw)
    return deterministic_archive(raw.getvalue())


def records():
    output = []
    for name, producer, source in [
        ("native.docx", DocxParser, native_docx()),
        ("native.xlsx", ExcelParser, native_xlsx()),
    ]:
        media = name.rsplit(".", 1)[-1]
        parsed = producer(file_name=name, file_type=media).parse(source)
        contract = json.loads(parsed.metadata["source_contract"])
        if any(unit["completeness"] != "complete" for unit in contract["units"]):
            raise ValueError(f"{name} has incomplete native carriers")
        output.append({
            "name": name, "file_type": media,
            "file_sha256": hashlib.sha256(source).hexdigest(), "source_hex": source.hex(),
            "markdown": parsed.content,
            "metadata": {key: value if isinstance(value, str) else json.dumps(value, ensure_ascii=False)
                         for key, value in parsed.metadata.items()},
            "structured_source_units": [unit.model_dump(mode="json", exclude={
                "physical_locator", "physical_path", "source_section_id", "source_issues", "header_cells",
            }) for unit in parsed.structured_source_units],
            "images": [{"original_ref": ref, "hex": base64.b64decode(data).hex()}
                       for ref, data in parsed.images.items()],
        })
    assert not any(name.split(".")[0] in {"onnxruntime", "magika", "markitdown"} for name in sys.modules)
    return output


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(records(), ensure_ascii=False, indent=2) + "\n")
