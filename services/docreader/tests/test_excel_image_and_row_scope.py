import base64
import io
import json

import pytest
from openpyxl import Workbook
from openpyxl.drawing.image import Image
from PIL import Image as PillowImage

from docreader.parser.excel_parser import ExcelParser, _preflight_xlsx_structure


def workbook_bytes(workbook):
    raw = io.BytesIO()
    workbook.save(raw)
    return raw.getvalue()


def test_embedded_png_is_registered_with_payload_and_sheet_source():
    workbook = Workbook()
    workbook.active['A1'] = 'Synthetic image source'
    pixels = io.BytesIO()
    PillowImage.new('RGB', (12, 8), 'blue').save(pixels, format='PNG')
    expected = pixels.getvalue()
    workbook.active.add_image(Image(io.BytesIO(expected)), 'C2')
    document = ExcelParser(file_name='synthetic.xlsx', file_type='xlsx').parse(workbook_bytes(workbook))
    images = [u for u in document.structured_source_units if u.kind.value == 'image_region']
    assert len(images) == 1
    unit = images[0]
    assert base64.b64decode(document.images[unit.locator.original_ref]) == expected
    assert (unit.locator.width, unit.locator.height) == (12, 8)
    assert unit.locator.original_ref in document.content
    contract = json.loads(document.metadata['source_contract'])
    source = next(u for u in contract['units'] if u['unit_id'] == unit.key)
    assert source['section_id'] == 'sheet:0'
    assert source['rendered_spans']
    assert 'image_payload_missing' not in source['reasons']


def test_same_row_on_two_sheets_has_independent_cell_budget():
    workbook = Workbook()
    workbook.active.append([f'cell-{i}' for i in range(1200)])
    workbook.create_sheet().append([f'cell-{i}' for i in range(1200)])
    document = ExcelParser(file_name='synthetic.xlsx', file_type='xlsx').parse(workbook_bytes(workbook))
    rows = [u for u in document.structured_source_units if u.kind.value == 'table_row']
    assert len(rows) == 2
    assert all(len(row.locator.cells) == 1200 for row in rows)


def test_one_sheet_overfull_row_still_rejected(monkeypatch):
    from docreader.parser import excel_parser
    monkeypatch.setattr(excel_parser, 'XLSX_MAX_LOGICAL_COLUMN', 4096)
    workbook = Workbook()
    workbook.active.append([str(i) for i in range(2049)])
    with pytest.raises(ValueError, match='cells_per_row:2049>2048'):
        _preflight_xlsx_structure(workbook_bytes(workbook))


def test_discarded_embedded_image_is_explicitly_unsupported():
    import zipfile
    workbook = Workbook()
    pixels = io.BytesIO()
    PillowImage.new('RGB', (12, 8), 'blue').save(pixels, format='PNG')
    workbook.active.add_image(Image(io.BytesIO(pixels.getvalue())), 'C2')
    original = workbook_bytes(workbook)
    damaged = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(original)) as source, zipfile.ZipFile(damaged, 'w') as target:
        for name in source.namelist():
            target.writestr(name, b'unsupported image bytes' if name.startswith('xl/media/') else source.read(name))
    with pytest.raises(ValueError, match='XLSX_IMAGE_UNSUPPORTED'):
        ExcelParser(file_name='synthetic.xlsx', file_type='xlsx').parse(damaged.getvalue())


@pytest.mark.parametrize('placements', [1, 2])
def test_image_only_sheet_preserves_each_occurrence_with_shared_payload(placements):
    workbook = Workbook()
    pixels = io.BytesIO()
    PillowImage.new('RGB', (12, 8), 'blue').save(pixels, format='PNG')
    for anchor in ['A1', 'D4'][:placements]:
        workbook.active.add_image(Image(io.BytesIO(pixels.getvalue())), anchor)
    document = ExcelParser(file_name='synthetic.xlsx', file_type='xlsx').parse(workbook_bytes(workbook))
    units = [u for u in document.structured_source_units if u.kind.value == 'image_region']
    assert len(units) == placements
    assert len({u.key for u in units}) == placements
    assert len({u.physical_path for u in units}) == placements
    assert len(document.images) == 1
    assert len({u.locator.original_ref for u in units}) == 1
    contract = json.loads(document.metadata['source_contract'])
    identities = [u for u in contract['units'] if u['kind'] == 'image_region']
    assert len(identities) == placements
    assert all(u['section_id'] == 'sheet:0' and u['rendered_spans'] for u in identities)
    assert len({u['rendered_spans'][0]['start_byte'] for u in identities}) == placements


@pytest.mark.parametrize('merge', ['A1:A2', 'A4:C4', 'A3:C4'])
def test_row_preserves_complete_intersecting_merges_outside_defined_table(merge):
    from openpyxl.worksheet.table import Table

    workbook = Workbook()
    sheet = workbook.active
    for address, value in {'D1': '项目', 'E1': '响应', 'D2': '表内合并',
                           'D3': '服务', 'E3': '提供'}.items():
        sheet[address] = value
    sheet.merge_cells('D2:E2')
    sheet.add_table(Table(displayName='DefinedForm', ref='D1:E3'))
    anchor = merge.split(':')[0]
    sheet[anchor] = '必须保留合并格式的证明说明。'
    sheet.merge_cells(merge)
    row_number = sheet[anchor].row
    # This blank merge shares the row but lies beyond its physical columns.
    sheet.merge_cells(f'G{row_number}:H{row_number}')
    document = ExcelParser(file_name='merge.xlsx', file_type='xlsx').parse(workbook_bytes(workbook))
    row = next(unit for unit in document.structured_source_units if unit.key == f'sheet:0:row:{row_number}')
    assert [region.a1_range for region in row.locator.merged_ranges] == [merge]
    assert row.locator.region.start_row == row.locator.region.end_row == row_number
    assert row.locator.region.end_column == (1 if row_number == 4 else 5)
    identity = next(unit for unit in json.loads(document.metadata['source_contract'])['units']
                    if unit['unit_id'] == row.key)
    assert identity['physical_locator']['merged_ranges'] == [region.model_dump() for region in row.locator.merged_ranges]
    table = next(unit for unit in document.structured_source_units if unit.grid)
    assert [region.a1_range for region in table.locator.merged_ranges] == ['D2:E2']
    assert (table.grid.row_count, table.grid.column_count) == (3, 2)
    merged = next(cell for cell in table.grid.cells if (cell.row, cell.column) == (1, 0))
    assert merged.col_span == 2 and merged.text == '表内合并'
    assert not any((cell.row, cell.column) == (1, 1) for cell in table.grid.cells)
