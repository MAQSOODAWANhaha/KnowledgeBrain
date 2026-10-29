"""Project authorization and durable bridge behavior with no model requests."""
from io import BytesIO
import asyncio
import hashlib
import os
from types import SimpleNamespace
from uuid import uuid4

from docx import Document
import psycopg
from psycopg import sql
from psycopg.conninfo import make_conninfo
from psycopg.types.json import Jsonb
import pytest

from docreader.tests.database import initialize_schema

from docreader.integrations.authoring import read_authoring_input
from docreader.integrations.store import Conflict, InvalidConfirmation
from docreader.integrations.web import WebBridge, WebError, _request_key, dispatch_authoring, execute_authoring_job, handle_command
from docreader.jobs.settings import Settings
from docreader.llm.config import ModelConfig, ModelEndpoint


def document_bytes():
    document = Document()
    document.add_heading("项目背景", level=1)
    document.add_paragraph("本项目建设公共服务平台。实施计划和最终报价由人工确认。")
    data = BytesIO()
    document.save(data)
    return data.getvalue()


@pytest.fixture
async def runtime(tmp_path):
    dsn = os.environ.get("TENDER_TEST_DATABASE_URL")
    if not dsn:
        pytest.skip("TENDER_TEST_DATABASE_URL is required for isolated PostgreSQL bridge tests")
    schema = "web_bridge_" + uuid4().hex
    async with await psycopg.AsyncConnection.connect(dsn, autocommit=True) as conn:
        await conn.execute(sql.SQL("CREATE SCHEMA {}").format(sql.Identifier(schema)))
    isolated = make_conninfo(dsn, options=f"-c search_path={schema}")
    settings = Settings(isolated, os.environ.get("TENDER_TEST_REDIS_URL", "redis://unused.invalid"), tmp_path / "artifacts")
    bridge = WebBridge(settings)
    project, other, owner = (str(uuid4()) for _ in range(3))
    async with bridge.store.transaction() as conn:
        await conn.execute("CREATE TABLE tender_projects(id uuid PRIMARY KEY, owner_user_id uuid NOT NULL, status text NOT NULL)")
        await conn.execute("CREATE TABLE tender_documents(id uuid PRIMARY KEY,project_id uuid NOT NULL REFERENCES tender_projects(id),"
                           "file_name text,media_type text,content_sha256 text,byte_length bigint,UNIQUE(id,project_id))")
        await conn.execute("INSERT INTO tender_projects VALUES(%s,%s,'open'),(%s,%s,'open')", (project, owner, other, owner))
    await initialize_schema(isolated, sections=("analysis", "authoring", "web"))
    await bridge.setup()
    try:
        yield SimpleNamespace(bridge=bridge, project=project, other=other, owner=owner, settings=settings)
    finally:
        async with await psycopg.AsyncConnection.connect(dsn, autocommit=True) as conn:
            await conn.execute(sql.SQL("DROP SCHEMA {} CASCADE").format(sql.Identifier(schema)))


def payload(runtime, action="workbench", **extra):
    return {"action": action, "project_id": runtime.project, "owner_user_id": runtime.owner, **extra}


async def upload(runtime, *, kind="tender"):
    data, document_id = document_bytes(), str(uuid4())
    async with runtime.bridge.store.transaction() as conn:
        await conn.execute("INSERT INTO tender_documents VALUES(%s,%s,%s,%s,%s,%s)",
            (document_id, runtime.project, "示例来源.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
             hashlib.sha256(data).hexdigest(), len(data)))
    request = payload(runtime, "document.put", document_id=document_id, kind=kind, role="main")
    await runtime.bridge.put_document(request, data)
    return request, data


def configured(monkeypatch):
    monkeypatch.setattr(ModelConfig, "from_env", classmethod(lambda cls: SimpleNamespace()))


async def test_owner_check_precedes_any_artifact_read(runtime, monkeypatch):
    monkeypatch.setattr(Settings, "from_env", classmethod(lambda cls: runtime.settings))
    def forbidden(*args):
        raise AssertionError("unauthorized command must not read artifacts")
    monkeypatch.setattr(runtime.bridge.artifacts.__class__, "read_bytes", forbidden)
    with pytest.raises(WebError, match="无权访问"):
        await handle_command(payload(runtime, "artifact.get", owner_user_id=str(uuid4()), execution_id="other", kind="docx"))


async def test_upload_rejects_bytes_not_matching_registered_identity(runtime):
    request, data = await upload(runtime)
    await runtime.bridge.put_document(request, data)  # Exact replay is harmless.
    with pytest.raises(WebError, match="身份不一致"):
        await runtime.bridge.put_document(request, b"different bytes")
    with pytest.raises(WebError, match="尚未登记"):
        await runtime.bridge.put_document({**request, "project_id": runtime.other}, data)


async def test_source_and_job_relations_cannot_cross_projects(runtime):
    request, _ = await upload(runtime)
    with pytest.raises(psycopg.errors.ForeignKeyViolation):
        async with runtime.bridge.store.transaction() as conn:
            await conn.execute("UPDATE tender_document_sources SET project_id=%s WHERE document_id=%s",
                               (runtime.other, request["document_id"]))
    run = await runtime.bridge.source_run(runtime.project)
    with pytest.raises(psycopg.errors.ForeignKeyViolation):
        async with runtime.bridge.store.transaction() as conn:
            await conn.execute("INSERT INTO tender_authoring_jobs(execution_id,project_id,run_id,payload) "
                               "VALUES(%s,%s,%s,'{}')", (uuid4().hex, runtime.other, run["run_id"]))


async def test_source_only_snapshot_has_no_hidden_analysis_command_or_model_calls(runtime):
    await upload(runtime)
    run = await runtime.bridge.source_run(runtime.project)
    assert run["status"] == "pending" and run["commands"] == [] and run["usage"] == []
    assert await runtime.bridge.latest_run(runtime.project) is None
    assert (await runtime.bridge.source_run(runtime.project))["run_id"] == run["run_id"]
    snapshot = await read_authoring_input(runtime.bridge.store, runtime.bridge.artifacts,
        run_id=run["run_id"], source_version=1, domain_revision=1)
    assert snapshot["frozen_ref"] is not None
    assert snapshot["verification_ref"] is None and snapshot["confirmation_ref"] is None
    source = runtime.bridge.artifacts.read_json(snapshot["frozen_ref"])
    assert source["native_range_count"] > 0
    assert source["stats"]["ocr_request_count"] == 0
    view = await runtime.bridge.analysis_view(run)
    assert not view["confirmed"] and view["status"] == "pending" and view["issues"]


def outline_tree(outline, *, title=None):
    return [{"chapter_id": chapter["chapter_id"], "title": title or chapter["title"],
             "children": outline_tree({"chapters": chapter["children"]})}
            for chapter in outline["chapters"]]


async def finish_outline(runtime, monkeypatch, *, key="outline-candidate", target_id="chapter-1:target-1"):
    configured(monkeypatch)
    started = await runtime.bridge.start_outline(payload(runtime, request_key=key))
    execution_id = started["generation"]["execution_id"]
    job = await runtime.bridge.job(runtime.project, execution_id)
    snapshot = await read_authoring_input(runtime.bridge.store, runtime.bridge.artifacts, **job["payload"]["identity"])
    outline = {"schema_version": 2, "authoring_input": snapshot, "issues": [], "chapters": [{
        "chapter_id": "chapter-1", "title": "实施方案", "children": [], "targets": [{
            "target_id": target_id, "kind": "narrative", "purpose": "说明实施方法",
            "rule_refs": [], "source_refs": []}]}]}
    reference = runtime.bridge.artifacts.put_json(outline)
    await runtime.bridge.authoring.submit(execution_id, snapshot, {"protocol": "outline-test"})
    lease = await runtime.bridge.authoring.acquire(execution_id)
    await runtime.bridge.authoring.release(lease, "succeeded", result_ref=reference)
    return execution_id, reference, outline


async def test_outline_candidate_is_independent_from_word_results(runtime, monkeypatch):
    await upload(runtime)
    execution_id, reference, outline = await finish_outline(runtime, monkeypatch)
    again = await runtime.bridge.start_outline(payload(runtime, request_key="outline-candidate"))
    assert again["generation"]["execution_id"] == execution_id
    view = await runtime.bridge.outline_view(runtime.project)
    assert view["draft"] is None and view["candidate"] == {
        "execution_id": execution_id, "sha256": reference["sha256"], "outline": outline}
    assert view["generation"]["operation"] == "outline" and view["generation"]["status"] == "succeeded"
    workbench = await runtime.bridge.workbench(runtime.project)
    assert workbench["context_execution_id"] is None
    assert workbench["artifacts"] == {"docx_available": False, "package_available": False}
    assert workbench["authoring"] is None and workbench["executions"] == []
    with pytest.raises(WebError, match="Word"):
        await runtime.bridge.result(runtime.project, execution_id)


async def test_outline_save_cas_replay_and_provenance_are_preserved(runtime, monkeypatch):
    await upload(runtime)
    execution_id, reference, outline = await finish_outline(runtime, monkeypatch)
    request = payload(runtime, request_key="save-outline", expected_revision_id=None,
                      source_execution_id=execution_id, chapters=outline_tree(outline, title="人工实施方案"))
    first = await runtime.bridge.save_outline(request)
    draft = first["draft"]
    assert draft["revision"] == 1 and draft["parent_id"] is None and draft["sources_current"] is True
    assert draft["outline"]["chapters"][0]["title"] == "人工实施方案"
    assert draft["outline"]["chapters"][0]["targets"] == outline["chapters"][0]["targets"]
    assert draft["outline"]["authoring_input"] == outline["authoring_input"]
    second_request = payload(runtime, request_key="save-next", expected_revision_id=draft["revision_id"],
        source_revision_id=draft["revision_id"], chapters=outline_tree(draft["outline"], title="下一修订"))
    second = await runtime.bridge.save_outline(second_request)
    assert second["draft"]["revision"] == 2 and second["draft"]["parent_id"] == draft["revision_id"]
    assert await runtime.bridge.save_outline(request) == first  # Replay works after head advances.
    with pytest.raises(Conflict):
        await runtime.bridge.save_outline({**request, "chapters": outline_tree(outline, title="另一个内容")})
    with pytest.raises(WebError) as stale:
        await runtime.bridge.save_outline({**second_request, "request_key": "stale-save"})
    assert stale.value.status == 409 and stale.value.code == "OUTLINE_REVISION_CONFLICT"
    assert (await runtime.bridge.outline_view(runtime.project))["draft"] == second["draft"]
    with pytest.raises(WebError) as other:
        await runtime.bridge.save_outline({**request, "project_id": runtime.other, "request_key": "other-project"})
    assert other.value.status == 404


async def test_concurrent_outline_saves_advance_only_one_head(runtime, monkeypatch):
    await upload(runtime)
    execution_id, _, outline = await finish_outline(runtime, monkeypatch)
    requests = [payload(runtime, request_key=f"save-{i}", expected_revision_id=None,
        source_execution_id=execution_id, chapters=outline_tree(outline, title=f"方案{i}")) for i in range(2)]
    results = await asyncio.gather(*(runtime.bridge.save_outline(request) for request in requests), return_exceptions=True)
    assert sum(isinstance(result, dict) for result in results) == 1
    assert sum(isinstance(result, WebError) and result.status == 409 for result in results) == 1
    assert (await runtime.bridge.outline_view(runtime.project))["draft"]["revision"] == 1


async def test_outline_history_restores_responses_and_appends_a_new_head(runtime, monkeypatch):
    from docreader.integrations.web_outline import outline_detail, outline_history

    await upload(runtime)
    execution_id, _, outline = await finish_outline(runtime, monkeypatch)
    first = (await runtime.bridge.save_outline(payload(runtime, request_key='history-first', expected_revision_id=None,
        source_execution_id=execution_id, chapters=outline_tree(outline))))['draft']
    deleted = (await runtime.bridge.save_outline(payload(runtime, request_key='history-delete',
        expected_revision_id=first['revision_id'], source_revision_id=first['revision_id'], chapters=[])))['draft']
    detail = await outline_detail(runtime.bridge, runtime.project, deleted['revision_id'])
    assert detail['available_targets'] == outline['chapters'][0]['targets']
    manual = 'manual-328f852d-3d53-4b7a-851d-fc0572420173'
    restored = (await runtime.bridge.save_outline(payload(runtime, request_key='history-reassign',
        expected_revision_id=deleted['revision_id'], source_revision_id=deleted['revision_id'],
        chapters=[{'chapter_id': manual, 'title': '恢复响应', 'target_ids': ['chapter-1:target-1'], 'children': []}])))['draft']
    assert restored['outline']['chapters'][0]['targets'] == outline['chapters'][0]['targets']
    assert restored['parent_id'] == deleted['revision_id']
    assert not restored['outline']['issues']
    history = (await outline_history(runtime.bridge, runtime.project))['revisions']
    assert [entry['revision'] for entry in history] == [3, 2, 1]
    assert (await outline_detail(runtime.bridge, runtime.project, first['revision_id']))['draft'] == first
    with pytest.raises(WebError) as other:
        await outline_detail(runtime.bridge, runtime.other, first['revision_id'])
    assert other.value.status == 404


async def test_web_job_freezes_analysis_before_the_queue_is_consumed(runtime, monkeypatch):
    configured(monkeypatch)
    await upload(runtime)
    submitted = await runtime.bridge.start_outline(payload(runtime, request_key='queue-snapshot'))
    job = await runtime.bridge.job(runtime.project, submitted['generation']['execution_id'])
    fixed = runtime.bridge.artifacts.read_json(job['payload']['authoring_input_ref'])
    assert all(fixed[key] == value for key, value in job['payload']['identity'].items())
    async with runtime.bridge.store.transaction() as conn:
        await conn.execute('UPDATE tender_analysis_runs SET domain_revision=domain_revision+1 WHERE run_id=%s', (fixed['run_id'],))
    with pytest.raises(InvalidConfirmation):
        await read_authoring_input(runtime.bridge.store, runtime.bridge.artifacts, **job['payload']['identity'])
    assert runtime.bridge.artifacts.read_json(job['payload']['authoring_input_ref']) == fixed


async def test_response_catalog_does_not_merge_another_generation_from_the_same_analysis(runtime, monkeypatch):
    from docreader.integrations.web_outline import outline_detail

    await upload(runtime)
    first_execution, _, first_outline = await finish_outline(runtime, monkeypatch, key='catalog-first')
    first = (await runtime.bridge.save_outline(payload(runtime, request_key='catalog-save-first', expected_revision_id=None,
        source_execution_id=first_execution, chapters=outline_tree(first_outline))))['draft']
    second_execution, _, second_outline = await finish_outline(runtime, monkeypatch, key='catalog-second', target_id='second-response')
    assert first_outline['authoring_input'] == second_outline['authoring_input']
    second = (await runtime.bridge.save_outline(payload(runtime, request_key='catalog-save-second',
        expected_revision_id=first['revision_id'], source_execution_id=second_execution, chapters=[])))['draft']
    detail = await outline_detail(runtime.bridge, runtime.project, second['revision_id'])
    assert [target['target_id'] for target in detail['available_targets']] == ['second-response']
    with pytest.raises(ValueError, match='不属于'):
        await runtime.bridge.save_outline(payload(runtime, request_key='catalog-cross-generation',
            expected_revision_id=second['revision_id'], source_revision_id=second['revision_id'],
            chapters=[{'chapter_id': 'manual-328f852d-3d53-4b7a-851d-fc0572420173', 'title': '错误引用',
                       'target_ids': ['chapter-1:target-1'], 'children': []}]))


async def test_web_knowledge_choice_is_fixed_in_the_job_and_worker_without_credentials(runtime, monkeypatch):
    from docreader.tender import document_execution

    await upload(runtime)
    configured(monkeypatch)
    selection = {'product_id': str(uuid4()), 'version_id': str(uuid4())}
    request = payload(runtime, request_key='knowledge-job', knowledge_selections=[selection])
    started = await runtime.bridge.start_authoring(request)
    job = await runtime.bridge.job(runtime.project, started['authoring']['execution_id'])
    assert job['payload']['knowledge_selections'] == [selection]
    assert job['payload']['request']['knowledge_selections'] == [selection]
    assert 'bearer' not in str(job['payload']) and 'api_key' not in str(job['payload'])
    assert await runtime.bridge.start_authoring(request) == started
    with pytest.raises(Conflict):
        await runtime.bridge.start_authoring({**request, 'knowledge_selections': []})
    config = ModelConfig(ModelEndpoint('http://fixture.invalid', 'test-only', 'fixture', 131072), None)
    monkeypatch.setattr(ModelConfig, 'from_env', classmethod(lambda cls: config))
    observed = []
    async def capture(*args, **kwargs):
        observed.append(kwargs)
        raise RuntimeError('offline boundary')
    monkeypatch.setattr(document_execution, 'execute_document', capture)
    await execute_authoring_job(job['execution_id'], runtime.settings)
    assert observed[0]['knowledge'].spec()['selections'] == [selection]
    assert observed[0]['authoring_input_ref'] == job['payload']['authoring_input_ref']


async def test_document_submission_pins_saved_outline_and_rejects_changed_sources(runtime, monkeypatch):
    await upload(runtime)
    execution_id, _, outline = await finish_outline(runtime, monkeypatch)
    saved = await runtime.bridge.save_outline(payload(runtime, request_key="save-for-document", expected_revision_id=None,
        source_execution_id=execution_id, chapters=outline_tree(outline)))
    draft = saved["draft"]
    request = payload(runtime, request_key="body-from-outline", outline_revision_id=draft["revision_id"],
                      outline_sha256=draft["sha256"], outline_ref={"sha256": "f" * 64, "path": "/etc/passwd"})
    submitted = await runtime.bridge.start_authoring(request)
    job = await runtime.bridge.job(runtime.project, submitted["authoring"]["execution_id"])
    assert job["payload"]["outline_ref"]["sha256"] == draft["sha256"]
    assert job["payload"]["identity"] == {key: outline["authoring_input"][key]
                                          for key in ("run_id", "source_version", "domain_revision")}
    with pytest.raises(WebError) as bad_hash:
        await runtime.bridge.start_authoring({**request, "request_key": "wrong-hash", "outline_sha256": "0" * 64})
    assert bad_hash.value.status == 409
    await upload(runtime)
    assert (await runtime.bridge.outline_view(runtime.project))["draft"]["sources_current"] is False
    assert await runtime.bridge.start_authoring(request) == submitted
    with pytest.raises(WebError) as changed:
        await runtime.bridge.start_authoring({**request, "request_key": "new-click-on-stale-outline"})
    assert changed.value.code == "SOURCE_CHANGED_REQUIRES_NEW_OUTLINE"


async def test_outline_worker_and_document_worker_receive_the_explicit_saved_input(runtime, monkeypatch):
    from docreader.tender import document_execution

    await upload(runtime)
    execution_id, _, outline = await finish_outline(runtime, monkeypatch)
    draft = (await runtime.bridge.save_outline(payload(runtime, request_key="worker-outline-save", expected_revision_id=None,
        source_execution_id=execution_id, chapters=outline_tree(outline))))["draft"]
    next_outline = (await runtime.bridge.start_outline(payload(runtime, request_key="outline-worker")))["generation"]
    document = (await runtime.bridge.start_authoring(payload(runtime, request_key="document-worker",
        outline_revision_id=draft["revision_id"], outline_sha256=draft["sha256"])))["authoring"]
    config = ModelConfig(ModelEndpoint("http://fixture.invalid", "test-only", "fixture", 131072), None)
    monkeypatch.setattr(ModelConfig, "from_env", classmethod(lambda cls: config))
    observed = []
    async def capture(*args, **kwargs):
        observed.append(kwargs)
        raise RuntimeError("offline worker boundary")
    monkeypatch.setattr(document_execution, "execute_document", capture)
    await execute_authoring_job(next_outline["execution_id"], runtime.settings)
    await execute_authoring_job(document["execution_id"], runtime.settings)
    assert observed[0]["outline_only"] is True and "outline_ref" not in observed[0]
    assert observed[1]["outline_ref"]["sha256"] == draft["sha256"] and "outline_only" not in observed[1]
    assert "outline_revision_id" not in observed[1] and "outline_sha256" not in observed[1]
    resumed = await runtime.bridge.update_job(payload(runtime, execution_id=next_outline["execution_id"], request_key="retry-outline"))
    assert resumed["authoring"]["operation"] == "outline"
    await execute_authoring_job(next_outline["execution_id"], runtime.settings)
    assert observed[2]["execution_id"] == observed[0]["execution_id"]
    workbench = await runtime.bridge.workbench(runtime.project)
    assert workbench["authoring"]["execution_id"] == document["execution_id"]
    assert all(view["operation"] == "document" for view in workbench["executions"])
    assert workbench["artifacts"] == {"docx_available": False, "package_available": False}


async def test_outline_predecessor_is_persisted_and_forwarded_to_the_worker(runtime, monkeypatch):
    from docreader.tender import document_execution

    await upload(runtime)
    previous, _, _ = await finish_outline(runtime, monkeypatch)
    request = payload(runtime, request_key="outline-inherit", previous_execution_id=previous)
    started = await runtime.bridge.start_outline(request)
    assert await runtime.bridge.start_outline(request) == started
    job = await runtime.bridge.job(runtime.project, started["generation"]["execution_id"])
    assert job["payload"]["operation"] == "outline"
    assert job["payload"]["request"]["previous_execution_id"] == previous
    config = ModelConfig(ModelEndpoint("http://fixture.invalid", "test-only", "fixture", 131072), None)
    monkeypatch.setattr(ModelConfig, "from_env", classmethod(lambda cls: config))
    observed = []

    async def capture(*_args, **kwargs):
        observed.append(kwargs)
        raise RuntimeError("offline worker boundary")

    monkeypatch.setattr(document_execution, "execute_document", capture)
    await execute_authoring_job(job["execution_id"], runtime.settings)
    assert observed[0]["previous_execution_id"] == previous and observed[0]["outline_only"] is True
    assert observed[0]["authoring_input_ref"] == job["payload"]["authoring_input_ref"]
    assert "outline_ref" not in observed[0] and observed[0]["review_request"] is None


@pytest.mark.parametrize("options", [
    {"review_request": {"target_ids": ["target"], "reason": "核对"}},
    {"outline_revision_id": str(uuid4()), "outline_sha256": "a" * 64},
    {"knowledge_selections": [{"product_id": str(uuid4()), "version_id": str(uuid4())}]},
])
async def test_outline_predecessor_does_not_enable_document_options(runtime, monkeypatch, options):
    await upload(runtime)
    previous, _, _ = await finish_outline(runtime, monkeypatch)
    with pytest.raises(WebError):
        await runtime.bridge.start_outline(payload(runtime, request_key="invalid-outline-options",
            previous_execution_id=previous, **options))


async def test_web_outline_predecessor_cannot_cross_operation_project_or_current_run(runtime, monkeypatch):
    from unittest.mock import AsyncMock

    await upload(runtime)
    previous, _, _ = await finish_outline(runtime, monkeypatch)
    document = await runtime.bridge.start_authoring(payload(runtime, request_key="document-before-outline"))
    with pytest.raises(Conflict, match="predecessor"):
        await runtime.bridge.start_outline(payload(runtime, request_key="wrong-operation",
            previous_execution_id=document["authoring"]["execution_id"]))
    other_project = SimpleNamespace(**{**vars(runtime), "project": runtime.other})
    await upload(other_project)
    with pytest.raises(WebError) as other:
        await runtime.bridge.start_outline(payload(other_project, request_key="wrong-project",
            previous_execution_id=previous))
    assert other.value.code == "EXECUTION_NOT_FOUND"
    prior_job = await runtime.bridge.job(runtime.project, previous)
    prior_run = await runtime.bridge.store.get_run(prior_job["run_id"])
    new = await runtime.bridge.store.submit(prior_run["manifest_ref"], "another-analysis-run")
    frozen = next(entry["artifact_ref"] for entry in prior_run["artifacts"] if entry["stage"] == "freeze")
    async with runtime.bridge.store.transaction() as conn:
        await conn.execute("INSERT INTO tender_analysis_artifacts "
            "(run_id,source_version,domain_revision,stage,artifact_ref,epoch,metadata) "
            "VALUES(%s,1,0,'freeze',%s,0,'{}')", (new["run_id"], Jsonb(frozen)))
    current_run = await runtime.bridge.store.get_run(new["run_id"])
    monkeypatch.setattr(runtime.bridge, "latest_run", AsyncMock(return_value=current_run))
    with pytest.raises(Conflict, match="predecessor"):
        await runtime.bridge.start_outline(payload(runtime, request_key="wrong-run", previous_execution_id=previous))


@pytest.mark.parametrize("action", ["outline.start", "outline.save", "outline.status"])
async def test_outline_commands_authorize_before_reading_or_generating(runtime, monkeypatch, action):
    monkeypatch.setattr(Settings, "from_env", classmethod(lambda cls: runtime.settings))
    def forbidden(*args):
        raise AssertionError("unauthorized outline command must not read artifacts or model configuration")
    monkeypatch.setattr(ModelConfig, "from_env", forbidden)
    monkeypatch.setattr(runtime.bridge.artifacts.__class__, "read_json", forbidden)
    with pytest.raises(WebError) as denied:
        await handle_command(payload(runtime, action, owner_user_id=str(uuid4()), request_key="denied"))
    assert denied.value.status == 404


async def test_ended_project_rejects_outline_mutations_without_creating_jobs(runtime, monkeypatch):
    monkeypatch.setattr(Settings, "from_env", classmethod(lambda cls: runtime.settings))
    async with runtime.bridge.store.transaction() as conn:
        await conn.execute("UPDATE tender_projects SET status='ended' WHERE id=%s", (runtime.project,))
    for action in ("outline.start", "outline.save"):
        with pytest.raises(WebError) as ended:
            await handle_command(payload(runtime, action, request_key="ended"))
        assert ended.value.code == "PROJECT_ENDED"
    assert (await handle_command(payload(runtime, "outline.status")))[0] == {"draft": None, "candidate": None, "generation": None}


async def test_authoring_submission_is_idempotent_and_does_not_wait_for_analysis(runtime, monkeypatch):
    configured(monkeypatch)
    await upload(runtime)
    request = payload(runtime, "authoring.start", request_key="same-click")
    first = await runtime.bridge.start_authoring(request)
    second = await runtime.bridge.start_authoring(request)
    assert first == second and first["authoring"]["status"] == "pending"
    execution_id = first["authoring"]["execution_id"]
    job = await runtime.bridge.job(runtime.project, execution_id)
    assert (await runtime.bridge.store.get_run(job["run_id"]))["commands"] == []
    async with runtime.bridge.store.transaction() as conn:
        assert (await (await conn.execute("SELECT count(*) AS n FROM tender_authoring_jobs")).fetchone())["n"] == 1
    with pytest.raises(Conflict):
        await runtime.bridge.start_authoring({**request, "review_request": {"target_ids": ["target-1"], "reason": "请核对"}})


async def test_concurrent_authoring_request_key_cannot_accept_different_options(runtime, monkeypatch):
    configured(monkeypatch)
    await upload(runtime)
    await runtime.bridge.source_run(runtime.project)
    arrived, count = asyncio.Event(), 0
    async def materials(_project_id):
        nonlocal count
        count += 1
        if count == 2:
            arrived.set()
        await asyncio.wait_for(arrived.wait(), 5)
        return []
    monkeypatch.setattr(runtime.bridge, "materials", materials)
    request = payload(runtime, request_key="concurrent")
    results = await asyncio.gather(runtime.bridge.start_authoring(request),
        runtime.bridge.start_authoring({**request, "review_request": {"target_ids": ["target-1"], "reason": "请核对"}}),
        return_exceptions=True)
    assert sum(isinstance(result, dict) for result in results) == 1
    assert sum(isinstance(result, Conflict) for result in results) == 1


async def test_new_tender_upload_uses_current_free_sources_without_overwriting_old_analysis(runtime, monkeypatch):
    configured(monkeypatch)
    await upload(runtime)
    analyzed = await runtime.bridge.source_run(runtime.project)
    async with runtime.bridge.store.transaction() as conn:
        await conn.execute("UPDATE tender_project_analyses SET source_only=false WHERE run_id=%s", (analyzed["run_id"],))
    first = await runtime.bridge.start_authoring(payload(runtime, request_key="before-new-upload"))
    first_job = await runtime.bridge.job(runtime.project, first["authoring"]["execution_id"])
    assert first_job["run_id"] == analyzed["run_id"]
    await upload(runtime)
    second = await runtime.bridge.start_authoring(payload(runtime, request_key="after-new-upload"))
    second_job = await runtime.bridge.job(runtime.project, second["authoring"]["execution_id"])
    assert second_job["run_id"] != analyzed["run_id"]
    current = await runtime.bridge.store.get_run(second_job["run_id"])
    snapshot = await read_authoring_input(runtime.bridge.store, runtime.bridge.artifacts, **second_job["payload"]["identity"])
    frozen = runtime.bridge.artifacts.read_json(snapshot["frozen_ref"])
    assert len(frozen["documents"]) == 2 and frozen["stats"]["ocr_request_count"] == 0
    assert current["commands"] == [] and current["usage"] == []
    assert (await runtime.bridge.job(runtime.project, first_job["execution_id"]))["payload"] == first_job["payload"]
    analysis = (await runtime.bridge.workbench(runtime.project))["analysis"]
    assert analysis["run_id"] == analyzed["run_id"] and analysis["sources_current"] is False
    with pytest.raises(WebError, match="招标资料已更新"):
        await runtime.bridge.start_authoring(payload(runtime, request_key="old-predecessor",
            previous_execution_id=first_job["execution_id"]))


async def test_pending_analysis_sources_are_compared_before_freeze_is_ready(runtime):
    await upload(runtime)
    manifest = runtime.bridge.artifacts.put_json(await runtime.bridge.manifest(runtime.project))
    created = await runtime.bridge.store.submit(manifest, "pending-fixture")
    run = await runtime.bridge.store.get_run(created["run_id"])
    assert (await runtime.bridge.analysis_view(run, runtime.project))["sources_current"] is True
    await upload(runtime)
    assert (await runtime.bridge.analysis_view(run, runtime.project))["sources_current"] is False


async def test_enterprise_materials_are_fixed_at_submission_and_new_uploads_reach_new_jobs(runtime, monkeypatch):
    configured(monkeypatch)
    await upload(runtime)
    await upload(runtime, kind="enterprise")
    first = await runtime.bridge.start_authoring(payload(runtime, request_key="old-materials"))
    first_job = await runtime.bridge.job(runtime.project, first["authoring"]["execution_id"])
    await upload(runtime, kind="enterprise")
    second = await runtime.bridge.start_authoring(payload(runtime, request_key="new-materials"))
    second_job = await runtime.bridge.job(runtime.project, second["authoring"]["execution_id"])
    assert len(second_job["payload"]["materials"]) == 2 * len(first_job["payload"]["materials"]) > 0
    assert (await runtime.bridge.job(runtime.project, first_job["execution_id"]))["payload"] == first_job["payload"]


async def test_dispatch_failure_is_recoverable_and_preserves_execution_id(runtime, monkeypatch):
    configured(monkeypatch)
    await upload(runtime)
    result = await runtime.bridge.start_authoring(payload(runtime, request_key="dispatch"))
    execution_id = result["authoring"]["execution_id"]
    def unavailable(_execution_id):
        raise ConnectionError("broker unavailable")
    assert await dispatch_authoring(runtime.bridge.store, unavailable) == 0
    assert (await runtime.bridge.job(runtime.project, execution_id))["status"] == "pending"
    sent = []
    assert await dispatch_authoring(runtime.bridge.store, sent.append) == 1
    assert sent == [execution_id]
    assert await dispatch_authoring(runtime.bridge.store, sent.append) == 0


async def test_resume_uses_same_job_and_replaying_resume_does_not_requeue_again(runtime, monkeypatch):
    configured(monkeypatch)
    await upload(runtime)
    result = await runtime.bridge.start_authoring(payload(runtime, request_key="initial"))
    execution_id = result["authoring"]["execution_id"]
    async with runtime.bridge.store.transaction() as conn:
        await conn.execute("UPDATE tender_authoring_jobs SET status='failed' WHERE execution_id=%s", (execution_id,))
    request = payload(runtime, execution_id=execution_id, request_key="resume-once")
    await runtime.bridge.update_job(request)
    assert (await runtime.bridge.job(runtime.project, execution_id))["status"] == "pending"
    async with runtime.bridge.store.transaction() as conn:
        await conn.execute("UPDATE tender_authoring_jobs SET status='failed' WHERE execution_id=%s", (execution_id,))
    await runtime.bridge.update_job(request)
    assert (await runtime.bridge.job(runtime.project, execution_id))["status"] == "failed"
    with pytest.raises(WebError, match="不属于此项目"):
        await runtime.bridge.update_job({**request, "project_id": runtime.other})


async def test_cancelled_job_cannot_be_resumed_or_dispatched(runtime, monkeypatch):
    configured(monkeypatch)
    await upload(runtime)
    result = await runtime.bridge.start_authoring(payload(runtime, request_key="cancel"))
    execution_id = result["authoring"]["execution_id"]
    request = payload(runtime, execution_id=execution_id, request_key="resume-cancelled")
    assert (await runtime.bridge.update_job(request, cancel=True))["authoring"]["status"] == "cancelled"
    with pytest.raises(Conflict, match="cancelled"):
        await runtime.bridge.update_job(request)
    assert await dispatch_authoring(runtime.bridge.store, lambda _: pytest.fail("cancelled job dispatched")) == 0


async def completed_job(runtime, *, materials=()):
    run = await runtime.bridge.source_run(runtime.project)
    snapshot = await read_authoring_input(runtime.bridge.store, runtime.bridge.artifacts,
        run_id=run["run_id"], source_version=1, domain_revision=1)
    execution_id = "saved-" + uuid4().hex
    await runtime.bridge.authoring.submit(execution_id, snapshot, {"test_fixture": True})
    lease = await runtime.bridge.authoring.acquire(execution_id)
    reference = runtime.bridge.artifacts.put_json({"schema_version": 2, "authoring_input": snapshot, "chapters": [], "issues": []})
    result = runtime.bridge.artifacts.put_json({"outline_ref": reference, "unit_refs": [], "gaps": [],
        "basis_ref": runtime.bridge.artifacts.put_json({"materials": list(materials)}),
        "docx_ref": runtime.bridge.artifacts.put_bytes(document_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        "package_ref": runtime.bridge.artifacts.put_bytes(b"synthetic zip", "application/zip")})
    await runtime.bridge.authoring.release(lease, "succeeded", result_ref=result)
    async with runtime.bridge.store.transaction() as conn:
        await conn.execute("INSERT INTO tender_authoring_jobs(execution_id,project_id,run_id,payload,status) "
                           "VALUES(%s,%s,%s,%s,'succeeded')", (execution_id, runtime.project, run["run_id"], Jsonb({})))
    return execution_id


async def test_artifacts_are_bound_to_successful_project_execution(runtime, monkeypatch):
    await upload(runtime)
    execution_id = await completed_job(runtime)
    monkeypatch.setattr(Settings, "from_env", classmethod(lambda cls: runtime.settings))
    metadata, data = await handle_command(payload(runtime, "artifact.get", execution_id=execution_id, kind="docx",
        path="/etc/passwd", artifact_ref={"sha256": "a" * 64}))
    assert data.startswith(b"PK") and metadata["sha256"] == hashlib.sha256(data).hexdigest()
    with pytest.raises(WebError, match="不属于此项目"):
        await handle_command(payload(runtime, "artifact.get", project_id=runtime.other, execution_id=execution_id, kind="docx"))
    with pytest.raises(WebError, match="只能读取"):
        await handle_command(payload(runtime, "artifact.get", execution_id=execution_id, kind="arbitrary"))


async def test_existing_document_outline_can_be_saved_without_another_model_call(runtime):
    await upload(runtime)
    execution_id = await completed_job(runtime)
    view = await runtime.bridge.outline_view(runtime.project)
    assert view["draft"] is None and view["generation"] is None
    assert view["candidate"]["execution_id"] == execution_id
    saved = await runtime.bridge.save_outline(payload(runtime, request_key="adopt-document-outline", expected_revision_id=None,
        source_execution_id=execution_id, chapters=[]))
    assert saved["draft"]["execution_id"] == execution_id and saved["draft"]["revision"] == 1
    assert (await runtime.bridge.workbench(runtime.project))["context_execution_id"] == execution_id


async def test_completed_outline_never_replaces_existing_document_context(runtime, monkeypatch):
    await upload(runtime)
    document_id = await completed_job(runtime)
    outline_id, _, _ = await finish_outline(runtime, monkeypatch)
    view = await runtime.bridge.workbench(runtime.project, document_id)
    assert view["context_execution_id"] == document_id
    assert view["artifacts"]["execution_id"] == document_id
    assert [execution["execution_id"] for execution in view["executions"]] == [document_id]
    assert (await runtime.bridge.outline_view(runtime.project))["candidate"]["execution_id"] == outline_id


async def test_failed_new_job_preserves_previous_successful_artifacts(runtime, monkeypatch):
    configured(monkeypatch)
    await upload(runtime)
    saved_id = await completed_job(runtime)
    new = await runtime.bridge.start_authoring(payload(runtime, request_key="new-attempt"))
    async with runtime.bridge.store.transaction() as conn:
        await conn.execute("UPDATE tender_authoring_jobs SET status='failed' WHERE execution_id=%s",
                           (new["authoring"]["execution_id"],))
    view = await runtime.bridge.workbench(runtime.project)
    assert view["authoring"]["status"] == "failed"
    assert view["artifacts"] == {"execution_id": saved_id, "docx_available": True, "package_available": True}
    assert view["analysis"] is None


async def test_workbench_context_follows_selected_word_instead_of_new_candidate(runtime):
    await upload(runtime)
    current_id = await completed_job(runtime)
    candidate_id = await completed_job(runtime)
    view = await runtime.bridge.workbench(runtime.project, current_id)
    assert view["context_execution_id"] == current_id
    assert view["artifacts"]["execution_id"] == candidate_id
    assert view["documents"][0]["kind"] == "tender"
    with pytest.raises(WebError, match="不属于此项目"):
        await runtime.bridge.workbench(runtime.other, current_id)


async def test_workbench_materials_restore_source_text_and_preserve_origin(runtime):
    await upload(runtime)
    text = "营业执照原始文字"
    material = {"material_id": "license", "title": "营业执照", "text": "不得用生成正文替代原文",
        "content_ref": runtime.bridge.artifacts.put_bytes(text.encode(), "text/plain"),
        "original_ref": runtime.bridge.artifacts.put_bytes(document_bytes(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        "locator": {"file_name": "营业执照.docx", "page": 1}, "usage": "proof", "attachment": True}
    execution_id = await completed_job(runtime, materials=[material])
    view = await runtime.bridge.workbench(runtime.project, execution_id)
    assert view["materials"] == [{**material, "text": text}]
    with pytest.raises(WebError, match="不属于此项目"):
        await runtime.bridge.workbench(runtime.other, execution_id)


async def test_running_outline_is_visible_only_from_committed_receipt_without_saved_word(runtime, monkeypatch):
    configured(monkeypatch)
    await upload(runtime)
    started = await runtime.bridge.start_authoring(payload(runtime, request_key="running-outline"))
    execution_id = started["authoring"]["execution_id"]
    job = await runtime.bridge.job(runtime.project, execution_id)
    snapshot = await read_authoring_input(runtime.bridge.store, runtime.bridge.artifacts, **job["payload"]["identity"])
    await runtime.bridge.authoring.submit(execution_id, snapshot, {"test_fixture": True})
    lease = await runtime.bridge.authoring.acquire(execution_id)
    outline = {"chapters": [{"chapter_id": "working", "title": "已完成组织的结构", "targets": [], "children": []}]}
    reference = runtime.bridge.artifacts.put_json(outline)
    assert (await runtime.bridge.workbench(runtime.project))["outline"] is None
    await runtime.bridge.authoring.commit_artifact(lease, "document-outline", reference)
    view = await runtime.bridge.workbench(runtime.project)
    assert view["outline"] == outline and view["context_execution_id"] is None
    assert view["units"] == [] and not view["artifacts"]["docx_available"]
    saved_id = await completed_job(runtime)
    view = await runtime.bridge.workbench(runtime.project, saved_id)
    assert view["outline"]["chapters"] == [] and view["context_execution_id"] == saved_id


async def test_material_original_is_reached_only_through_same_project_execution_basis(runtime, monkeypatch):
    await upload(runtime)
    original = document_bytes()
    material = {"material_id": "license", "title": "营业执照", "text": "原始文字",
        "original_ref": runtime.bridge.artifacts.put_bytes(original,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        "locator": {"file_name": "营业执照.docx"}}
    execution_id = await completed_job(runtime, materials=[material, {"material_id": "text-only", "text": "只有文字"}])
    monkeypatch.setattr(Settings, "from_env", classmethod(lambda cls: runtime.settings))
    request = payload(runtime, "material.original", execution_id=execution_id, material_id="license",
        path="/etc/passwd", original_ref={"sha256": "a" * 64})
    metadata, data = await handle_command(request)
    assert data == original and metadata["file_name"] == "营业执照.docx"
    assert metadata["media_type"] == material["original_ref"]["media_type"]
    with pytest.raises(WebError, match="不属于此项目"):
        await handle_command({**request, "project_id": runtime.other})
    with pytest.raises(WebError, match="不存在"):
        await handle_command({**request, "material_id": "unknown"})
    with pytest.raises(WebError, match="没有可读取的原件"):
        await handle_command({**request, "material_id": "text-only"})


async def test_authoring_status_defaults_to_latest_job_or_null(runtime, monkeypatch):
    configured(monkeypatch)
    monkeypatch.setattr(Settings, "from_env", classmethod(lambda cls: runtime.settings))
    assert await handle_command(payload(runtime, "authoring.status")) == ({"authoring": None}, b"")
    await upload(runtime)
    started = await runtime.bridge.start_authoring(payload(runtime, request_key="status"))
    assert (await handle_command(payload(runtime, "authoring.status")))[0] == started


async def test_worker_failure_is_durable_without_persisting_exception_text(runtime, monkeypatch):
    from docreader.tender import document_execution

    config = ModelConfig(ModelEndpoint("http://fixture.invalid", "test-only", "fixture", 131072), None)
    monkeypatch.setattr(ModelConfig, "from_env", classmethod(lambda cls: config))
    await upload(runtime)
    started = await runtime.bridge.start_authoring(payload(runtime, request_key="worker"))
    execution_id = started["authoring"]["execution_id"]
    observed = []
    async def fail(*args, **kwargs):
        observed.append(kwargs["execution_id"])
        raise RuntimeError("sensitive fixture exception body")
    monkeypatch.setattr(document_execution, "execute_document", fail)
    await execute_authoring_job(execution_id, runtime.settings)
    job = await runtime.bridge.job(runtime.project, execution_id)
    assert job["status"] == "failed" and job["error"] == {"type": "RuntimeError"}
    await runtime.bridge.update_job(payload(runtime, execution_id=execution_id, request_key="recover"))
    await execute_authoring_job(execution_id, runtime.settings)
    assert observed == [execution_id, execution_id]


async def test_cancel_during_worker_initialization_does_not_publish_success(runtime, monkeypatch):
    from docreader.tender import document_execution

    config = ModelConfig(ModelEndpoint("http://fixture.invalid", "test-only", "fixture", 131072), None)
    monkeypatch.setattr(ModelConfig, "from_env", classmethod(lambda cls: config))
    await upload(runtime)
    started = await runtime.bridge.start_authoring(payload(runtime, request_key="cancel-worker"))
    execution_id = started["authoring"]["execution_id"]
    entered, stopped = asyncio.Event(), asyncio.Event()
    async def waiting(*args, **kwargs):
        entered.set()
        try:
            await asyncio.Event().wait()
        finally:
            stopped.set()
    monkeypatch.setattr(document_execution, "execute_document", waiting)
    task = asyncio.create_task(execute_authoring_job(execution_id, runtime.settings))
    try:
        await asyncio.wait_for(entered.wait(), 5)
        await runtime.bridge.update_job(payload(runtime, execution_id=execution_id), cancel=True)
        await asyncio.wait_for(task, 5)
        assert stopped.is_set()
        assert (await runtime.bridge.job(runtime.project, execution_id))["status"] == "cancelled"
    finally:
        task.cancel()
        await asyncio.gather(task, return_exceptions=True)


def test_request_keys_are_project_scoped():
    first = {"request_key": "same", "project_id": str(uuid4())}
    assert _request_key(first) == _request_key(first)
    assert _request_key(first) != _request_key({**first, "project_id": str(uuid4())})
    with pytest.raises(WebError):
        _request_key({"project_id": first["project_id"]})
