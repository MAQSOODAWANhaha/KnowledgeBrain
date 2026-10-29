"""Bidder responses require applicable knowledge; gaps must still export a draft."""

from io import BytesIO
import json

from docx import Document
import pytest

from docreader.integrations.artifacts import ArtifactStore
from docreader.llm.budget import RequestBudget
from docreader.llm.request import ModelWorkUnitExhausted
from docreader.tender.authoring_pipeline import bind_writing_response, generate_document, prepare_materials
from docreader.tests.submission_responses import assemble_independent_candidates, provided_template_ids


def task(*, manual=False):
    return {"task_id": "w1", "chapter_id": "chapter-1", "title_path": ["技术响应"], "targets": [
        {"target_id": "target-supported", "kind": "narrative", "purpose": "角色权限管理",
         "manual_fill": manual, "rule_refs": [], "source_refs": []},
        {"target_id": "target-missing", "kind": "narrative", "purpose": "异地容灾",
         "rule_refs": [], "source_refs": []}]}


def material(artifacts, *, targets=("target-supported",)):
    return prepare_materials(artifacts, [{"material_id": "permissions", "title": "产品功能说明",
        "text": "本产品支持角色权限管理。", "locator": {"target_ids": list(targets)}}])[0]


def block(text="本产品支持角色权限管理。", **changes):
    return {"kind": "paragraph", "text": text, "target_ids": ["t1"],
            "support_ids": ["k1"], **changes}


def budget():
    return RequestBudget(model="fixture", context_tokens=131072, output_reserve=16384)


def snapshot():
    return {"schema_version": 2, "run_id": "grounded-writing", "source_version": 1, "domain_revision": 1}


def payload(request):
    return json.loads(request["messages"][1]["content"][0]["text"])


def outline(*, targets=1):
    return {"schema_version": 2, "chapters": [{"title": "技术响应", "children": [], "targets": [
        {"kind": "narrative", "purpose": purpose, "rule_ids": [], "range_ids": []}
        for purpose in ("角色权限管理", "异地容灾")[:targets]]}], "findings": []}


def document_text(artifacts, result):
    document = Document(BytesIO(artifacts.read_bytes(result["docx_ref"])))
    return "\n".join([paragraph.text for paragraph in document.paragraphs] +
                     [cell.text for table in document.tables for row in table.rows for cell in row.cells])


@pytest.mark.parametrize("changes", [
    {"support_ids": []}, {"support_ids": ["not-provided"]},
    {"target_ids": []}, {"target_ids": ["not-provided"]},
    {"target_ids": ["t2"]}, {"target_ids": ["t1", "t2"]},
])
def test_invalid_or_cross_target_block_is_dropped_without_losing_the_location(tmp_path, changes):
    artifacts = ArtifactStore(tmp_path)
    unit = bind_writing_response(task(), {"blocks": [block(**changes)]}, {}, material(artifacts))
    assert unit["blocks"] == [] and unit["status"] == "placeholder"
    assert unit["gaps"] and unit["target_ids"] == ["target-supported", "target-missing"]


def test_partially_supported_task_keeps_only_supported_content_and_a_local_gap(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    unit = bind_writing_response(task(), {"blocks": [block(), block("自动支持异地容灾。", target_ids=["t2"])]},
                                 {}, material(artifacts))
    assert unit["status"] == "generated"
    assert [item["text"] for item in unit["blocks"]] == ["本产品支持角色权限管理。"]
    assert unit["blocks"][0]["target_ids"] == ["target-supported"]
    assert unit["blocks"][0]["material_ids"] == ["permissions"]
    assert unit["gaps"], "One supported target must not silently mark its sibling as answered."


def test_manual_target_cannot_receive_knowledge_supported_model_text(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    unit = bind_writing_response(task(manual=True), {"blocks": [block("报价123456元。")], "cell_fills": []},
                                 {}, material(artifacts))
    assert unit["blocks"] == [] and unit["status"] == "placeholder" and unit["gaps"]


def test_unread_original_cannot_be_attached_or_support_generated_prose(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    materials, _ = prepare_materials(artifacts, [{"material_id": "unread-proof", "title": "扫描证明",
        "text": "", "usage": "proof", "attachment": True,
        "original_ref": artifacts.put_bytes(b"synthetic unread scan")}])
    response = {"blocks": [block("扫描件证明支持全部功能。", support_ids=["k1"])],
                "attachment_ids": ["k1"]}
    unit = bind_writing_response(task(), response, {}, materials)
    assert unit["blocks"] == [] and unit["gaps"]
    assert unit["attachments"] == []


def test_fill_cannot_use_another_targets_template_even_with_valid_material(tmp_path):
    artifacts, fixed = ArtifactStore(tmp_path), task()
    first = {"source_index_ref": artifacts.put_json({"synthetic": "first"}), "unit_id": "grid1"}
    second = {"source_index_ref": artifacts.put_json({"synthetic": "second"}), "unit_id": "grid2"}
    fixed["targets"][0].update(kind="form", source_refs=[first],
                                template_refs=[{"source_ref": first, "selectors": []}])
    fixed["targets"][1].update(kind="form", source_refs=[second],
                                template_refs=[{"source_ref": second, "selectors": []}])
    fill = {"source_id": "s2", "row": 0, "column": 1, "text": "角色权限管理",
            "target_ids": ["t1"], "support_ids": ["k1"]}
    unit = bind_writing_response(fixed, {"blocks": [], "cell_fills": [fill]},
                                 {"s1": first, "s2": second}, material(artifacts))
    assert unit["cell_fills"] == [] and unit["status"] == "placeholder" and unit["gaps"]


@pytest.mark.parametrize("references", [
    {"support_ids": []}, {"support_ids": ["not-provided"]}, {"target_ids": ["t2"]},
])
def test_invalid_or_cross_target_fill_is_left_empty(tmp_path, references):
    artifacts = ArtifactStore(tmp_path)
    fixed = task()
    source = {"source_index_ref": artifacts.put_json({"synthetic": "template"}), "unit_id": "grid"}
    fixed["targets"][0].update(kind="form", source_refs=[source],
                                template_refs=[{"source_ref": source, "selectors": []}])
    fill = {"source_id": "s1", "row": 0, "column": 1, "text": "不应写入的值",
            "target_ids": ["t1"], "support_ids": ["k1"], **references}
    unit = bind_writing_response(fixed, {"blocks": [], "cell_fills": [fill]}, {"s1": source}, material(artifacts))
    assert unit["cell_fills"] == [] and unit["status"] == "placeholder" and unit["gaps"]


@pytest.mark.parametrize("other_scope", [False, True])
async def test_no_applicable_material_skips_proposal_and_writing_but_exports(tmp_path, other_scope):
    artifacts, calls = ArtifactStore(tmp_path), []

    async def execute(stage, request, validate):
        value = payload(request)
        calls.append(value.get("operation", stage))
        if value.get("operation") == "assemble":
            return validate(assemble_independent_candidates(value))
        if value.get("operation") == "structure":
            return validate({"findings": []})
        if stage == "outline":
            return validate(outline())
        pytest.fail("Absent evidence must not spend a proposal, writing or body-review call.")

    materials = material(artifacts, targets=("another-target",)) if other_scope else []
    result = artifacts.read_json(await generate_document(artifacts, snapshot(), budget(), execute, materials=materials))
    assert calls == ["outline", "assemble", "structure"]
    assert result["docx_ref"] and result["package_ref"] and result["gaps"]
    assert all(artifacts.read_json(ref)["status"] == "placeholder" for ref in result["unit_refs"])
    assert "技术响应" in document_text(artifacts, result) and "待补" in document_text(artifacts, result)


async def test_unrelated_retrieval_can_be_semantically_rejected_without_free_writing(tmp_path):
    from test_authoring_pipeline_boundaries import snapshot as frozen_snapshot

    artifacts, calls = ArtifactStore(tmp_path), []
    fixed = frozen_snapshot(artifacts, [{"text": "应支持角色权限管理。"}])

    class Knowledge:
        async def search(self, queries, _artifacts):
            return {"materials": [{"material_id": "unrelated", "title": "绿化养护手册",
                "text": "每月修剪树木。", "locator": {"target_ids": [queries[0]["target_id"]]}}], "issues": []}

    async def execute(stage, request, validate):
        value = payload(request)
        calls.append(value.get("operation", stage))
        if value.get("operation") == "assemble":
            return validate(assemble_independent_candidates(value))
        if value.get("operation") == "structure":
            return validate({"findings": []})
        if stage == "outline":
            organized = outline()
            sources = [entry["range_id"] for entry in value["sources"]]
            organized["chapters"][0]["targets"][0]["range_ids"] = sources
            organized["search_intents"] = [{"query": "角色权限管理", "for_ids": sources}]
            return validate(organized)
        assert value["grounding_policy"] == "knowledge_required_v2"
        if stage == "authoring_write":
            assert any(entry["origin"] == "knowledge" and entry["text"] == "每月修剪树木。"
                       for entry in value["supports"])
            return validate({"blocks": [], "gaps": ["检索材料不支持权限要求，响应留空。"]})
        if stage == "authoring_basis":
            return validate({"shared_arrangements": [], "gaps": ["未发现适用资料。"]})
        return validate({"findings": []})

    result = artifacts.read_json(await generate_document(artifacts, fixed, budget(), execute, knowledge=Knowledge()))
    assert calls.count("authoring_write") == 1
    assert calls.count("authoring_review") == 1, "A complete semantic no-match still needs independent review."
    assert artifacts.read_json(result["unit_refs"][0])["status"] == "placeholder"
    assert "修剪树木" not in document_text(artifacts, result) and result["gaps"]


async def test_partial_knowledge_exports_supported_response_and_keeps_the_other_target(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    supported, invented = "本产品支持角色权限管理。", "自动支持异地容灾。"

    async def execute(stage, request, validate):
        value = payload(request)
        if value.get("operation") == "assemble":
            return validate(assemble_independent_candidates(value))
        if value.get("operation") == "structure":
            return validate({"findings": []})
        if stage == "outline":
            return validate(outline(targets=2))
        if stage == "authoring_basis":
            return validate({"shared_arrangements": [{"topic": "权限", "text": supported,
                "nature": "provided", "support_ids": ["k1"]}]})
        if stage == "authoring_write":
            assert [target["target_id"] for target in value["targets"]] == ["t1", "t2"]
            return validate({"blocks": [block(supported), block(invented, target_ids=["t2"])]})
        return validate({"findings": []})

    materials = material(artifacts, targets=("chapter-1:target-1",))
    result = artifacts.read_json(await generate_document(artifacts, snapshot(), budget(), execute, materials=materials))
    unit = artifacts.read_json(result["unit_refs"][0])
    assert unit["status"] == "generated" and len(unit["blocks"]) == 1 and unit["gaps"]
    persisted = artifacts.read_json(result["outline_ref"])
    assert len(persisted["chapters"][0]["targets"]) == 2
    text = document_text(artifacts, result)
    assert supported in text and invented not in text
    assert result["gaps"] and result["package_ref"]


async def test_no_knowledge_keeps_manual_price_position_without_writing_or_revision(tmp_path):
    artifacts, calls = ArtifactStore(tmp_path), []

    async def execute(stage, request, validate):
        value = payload(request)
        calls.append(value.get("operation", stage))
        if value.get("operation") == "assemble":
            return validate(assemble_independent_candidates(value))
        if value.get("operation") == "structure":
            return validate({"findings": []})
        if stage == "outline":
            organized = outline()
            organized["chapters"][0].update(title="投标报价")
            organized["chapters"][0]["targets"][0].update(purpose="最终投标报价由人工填写。", manual_fill=True)
            return validate(organized)
        pytest.fail("A manual price position must not invoke writing or body review.")

    result = artifacts.read_json(await generate_document(artifacts, snapshot(), budget(), execute))
    unit = artifacts.read_json(result["unit_refs"][0])
    assert calls == ["outline", "assemble", "structure"] and unit["status"] == "placeholder"
    text = document_text(artifacts, result)
    assert "投标报价" in text and "人工填写" in text
    assert "________" in text and not unit["cell_fills"] and result["package_ref"]


async def test_unavailable_semantic_review_preserves_supported_draft_and_still_exports(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    unreviewed = "本产品支持角色权限管理。"

    async def execute(stage, request, validate):
        value = payload(request)
        if value.get("operation") == "assemble":
            return validate(assemble_independent_candidates(value))
        if value.get("operation") == "structure":
            return validate({"findings": []})
        if stage == "outline":
            return validate(outline())
        if stage == "authoring_basis":
            return validate({"shared_arrangements": [{"topic": "权限", "text": unreviewed,
                "nature": "provided", "support_ids": ["k1"]}]})
        if stage == "authoring_write":
            return validate({"blocks": [block(unreviewed)]})
        raise ModelWorkUnitExhausted(stage)

    result = artifacts.read_json(await generate_document(artifacts, snapshot(), budget(), execute,
        materials=material(artifacts, targets=())))
    unit = artifacts.read_json(result["unit_refs"][0])
    assert unit["status"] == "generated" and unit["blocks"][0]["text"] == unreviewed
    assert unreviewed in document_text(artifacts, result)
    assert any(gap["code"] == "review_unavailable" for gap in result["gaps"])
    assert result["docx_ref"] and result["package_ref"]


async def test_review_of_unsupported_fill_preserves_original_template_with_empty_value(tmp_path):
    from test_authoring_pipeline_boundaries import snapshot as frozen_snapshot

    artifacts = ArtifactStore(tmp_path)
    fixed = frozen_snapshot(artifacts, [{"heading": "项目人员表", "grid": {
        "row_count": 1, "column_count": 2, "cells": [
            {"row": 0, "column": 0, "text": "项目经理姓名", "row_span": 1, "col_span": 1},
            {"row": 0, "column": 1, "text": "", "row_span": 1, "col_span": 1}]}}])

    async def execute(stage, request, validate):
        value = payload(request)
        if value.get("operation") == "assemble":
            return validate(assemble_independent_candidates(value))
        if value.get("operation") == "structure":
            return validate({"findings": []})
        if stage == "outline":
            organized = outline()
            organized["chapters"][0]["targets"][0].update(kind="form", purpose="填写项目经理姓名。",
                range_ids=[source["range_id"] for source in value["sources"]],
                template_part_ids=provided_template_ids(value))
            return validate(organized)
        if stage == "authoring_basis":
            return validate({"shared_arrangements": []})
        if stage == "authoring_write":
            source = next(entry for entry in value["supports"] if entry.get("grid"))
            return validate({"blocks": [], "cell_fills": [{"source_id": source["support_id"],
                "row": 0, "column": 1, "text": "杜撰的项目经理",
                "target_ids": [value["targets"][0]["target_id"]], "support_ids": ["k1"]}]})
        unit = value["units"][0]
        assert unit["content"]["cell_fills"][0]["content_id"] == "f1"
        return validate({"findings": [{"code": "unsupported_content", "detail": "给定资料没有项目经理姓名。",
            "task_ids": [unit["task_id"]], "content_ids": ["f1"]}]})

    result = artifacts.read_json(await generate_document(artifacts, fixed, budget(), execute, materials=[
        {"material_id": "personnel", "title": "人员安排", "text": "项目经理尚待指派，姓名未确定。"}]))
    unit = artifacts.read_json(result["unit_refs"][0])
    assert not unit["cell_fills"] and unit["status"] == "placeholder"
    document = Document(BytesIO(artifacts.read_bytes(result["docx_ref"])))
    table = next(table for table in document.tables if table.cell(0, 0).text == "项目经理姓名")
    assert table.cell(0, 1).text == ""
    assert "杜撰的项目经理" not in document_text(artifacts, result) and result["package_ref"]


@pytest.mark.parametrize("identified", [True, False])
async def test_semantic_review_removes_unsupported_content_without_recovering_old_draft(tmp_path, identified):
    artifacts, writes = ArtifactStore(tmp_path), []
    supported = "本产品支持角色权限管理。"
    invented = "本产品已建成三地容灾中心。"

    async def execute(stage, request, validate):
        value = payload(request)
        if value.get("operation") == "assemble":
            return validate(assemble_independent_candidates(value))
        if value.get("operation") == "structure":
            return validate({"findings": []})
        if stage == "outline":
            return validate(outline())
        if stage == "authoring_basis":
            return validate({"shared_arrangements": [{"topic": "功能", "text": supported,
                "nature": "provided", "support_ids": ["k1"]}]})
        if stage == "authoring_write":
            writes.append(value)
            if "previous_content" in value:
                return validate({"blocks": [], "gaps": ["没有支持补充内容的资料。"]})
            return validate({"blocks": [block(supported), block(invented)]})
        unit = value["units"][0]
        if not any(item["text"] == invented for item in unit["content"]["blocks"]):
            return validate({"findings": []})
        finding = {"code": "unsupported_content", "detail": "知识资料没有证明已建成容灾中心。",
                   "task_ids": [unit["task_id"]], "can_revise": True}
        if identified:
            finding["content_ids"] = ["b2"]
        return validate({"findings": [finding]})

    result = artifacts.read_json(await generate_document(artifacts, snapshot(), budget(), execute,
        materials=material(artifacts, targets=())))
    text = document_text(artifacts, result)
    assert invented not in text and result["gaps"]
    if identified:
        assert supported in text
    else:
        assert supported not in text
    assert len(writes) <= 2, "Semantic rejection must not create an unbounded rewrite loop."


async def test_short_knowledge_queries_preserve_full_conditions_in_writing_and_review(tmp_path):
    from test_outline_organization import context

    requirement = "一标段部署在内网时，应支持独立角色权限配置。"
    artifacts, fixed, _ = await context(tmp_path, documents=[("main", "main", [(requirement, "obligation")])])
    queries, writing, reviewed = [], [], []

    class Knowledge:
        async def search(self, submitted, _artifacts):
            queries.extend(submitted)
            return {"materials": [{"material_id": "permissions", "title": "产品功能说明",
                "text": "本产品支持角色权限配置。",
                "locator": {"target_ids": [submitted[0]["target_id"]]}}], "issues": []}

    async def execute(stage, request, validate):
        value = payload(request)
        if value.get("operation") == "assemble":
            return validate(assemble_independent_candidates(value))
        if value.get("operation") == "structure":
            return validate({"findings": []})
        if stage == "outline":
            organized = outline()
            organized["chapters"][0]["targets"][0]["rule_ids"] = [rule["rule_id"] for rule in value["rules"]]
            organized["search_intents"] = [{"query": "独立角色权限配置", "for_ids": [
                rule["rule_id"] for rule in value["rules"]]}]
            return validate(organized)
        if stage == "authoring_basis":
            return validate({"shared_arrangements": []})
        if stage == "authoring_write":
            writing.append(value)
            return validate({"blocks": [block("本产品支持角色权限配置。")]})
        assert stage == "authoring_review"
        reviewed.append(value)
        return validate({"findings": []})

    result = artifacts.read_json(await generate_document(artifacts, fixed, budget(), execute, knowledge=Knowledge()))
    assert queries and queries[0]["query"] == "独立角色权限配置"
    assert writing and requirement in json.dumps(writing[0]["rules"], ensure_ascii=False)
    assert reviewed and requirement in json.dumps(reviewed[0]["units"][0]["rules"], ensure_ascii=False)
    assert result["docx_ref"] and result["package_ref"]


async def test_shared_basis_receives_the_candidates_applicable_project_requirements(tmp_path):
    from test_outline_organization import context

    requirement = "一标段部署在内网时，应支持独立角色权限配置。"
    artifacts, fixed, _ = await context(tmp_path, documents=[("main", "main", [(requirement, "obligation")])])
    basis_calls = []

    async def execute(stage, request, validate):
        value = payload(request)
        if value.get("operation") == "assemble":
            return validate(assemble_independent_candidates(value))
        if value.get("operation") == "structure":
            return validate({"findings": []})
        if stage == "outline":
            organized = outline()
            organized["chapters"][0]["targets"][0]["rule_ids"] = [rule["rule_id"] for rule in value["rules"]]
            return validate(organized)
        if stage == "authoring_basis":
            basis_calls.append(value)
            assert requirement in json.dumps(value["rules"], ensure_ascii=False)
            assert value["targets"][0]["rule_ids"] and value["targets"][0]["support_ids"] == ["k1"]
            return validate({"shared_arrangements": []})
        if stage == "authoring_write":
            return validate({"blocks": [block()]})
        return validate({"findings": []})

    result = artifacts.read_json(await generate_document(artifacts, fixed, budget(), execute,
        materials=material(artifacts, targets=())))
    assert basis_calls and result["docx_ref"]


async def test_split_rereview_replaces_multitask_unsupported_finding_without_deleting_new_content(tmp_path, monkeypatch):
    import docreader.tender.authoring_pipeline as pipeline

    artifacts = ArtifactStore(tmp_path)
    original_batches = pipeline._batches
    review_plans, reviews, writes = [], [], []

    def split_followup(records, admission, stage, make_payload, **options):
        if stage == "authoring_review":
            review_plans.append(len(records))
            if len(review_plans) > 1:
                # Exercise a valid change in batch boundaries without changing
                # either the request budget or the model's semantic findings.
                for record in records:
                    yield [record]
                return
        yield from original_batches(records, admission, stage, make_payload, **options)

    monkeypatch.setattr(pipeline, "_batches", split_followup)

    async def execute(stage, request, validate):
        value = payload(request)
        if value.get("operation") == "assemble":
            return validate(assemble_independent_candidates(value))
        if value.get("operation") == "structure":
            return validate({"findings": []})
        if stage == "outline":
            return validate({"schema_version": 2, "chapters": [
                {"title": title, "children": [], "targets": [{"kind": "narrative",
                    "purpose": "依据知识资料说明本部分实施方法。", "rule_ids": [], "range_ids": []}]}
                for title in ("A部分", "B部分")], "findings": []})
        if stage == "authoring_basis":
            return validate({"shared_arrangements": []})
        if stage == "authoring_write":
            title = value["title_path"][-1]
            revised = "previous_content" in value
            writes.append((title, revised))
            text = title + ("有依据修订正文" if revised else "不受支持原稿")
            return validate({"blocks": [block(text, target_ids=[value["targets"][0]["target_id"]])]})
        reviews.append(value)
        if len(reviews) == 1:
            assert len(value["units"]) == 2
            return validate({"findings": [{"code": "unsupported_content",
                "detail": "两个部分的初稿均未按给定资料描述实施方法。",
                "task_ids": [unit["task_id"] for unit in value["units"]], "content_ids": ["b1"],
                "suggestion": "依据给定知识资料重写该正文。", "can_revise": True}]})
        assert len(value["units"]) == 1
        assert value["units"][0]["content"]["blocks"][0]["text"].endswith("有依据修订正文")
        return validate({"findings": []})

    result = artifacts.read_json(await generate_document(artifacts, snapshot(), budget(), execute, materials=[
        {"material_id": "permissions", "title": "给定实施方法",
         "text": "A部分有依据修订正文；B部分有依据修订正文。"}]))
    assert len(writes) == 4 and sum(revised for _, revised in writes) == 2
    assert review_plans == [2, 2] and len(reviews) == 3
    final_review = artifacts.read_json(result["review_ref"])
    assert final_review["findings"] == [], "A completed per-task rereview must replace that task's initial rejection."
    units = [artifacts.read_json(ref) for ref in result["unit_refs"]]
    assert all(unit["status"] == "generated" for unit in units)
    text = document_text(artifacts, result)
    assert "A部分有依据修订正文" in text and "B部分有依据修订正文" in text
    assert "不受支持原稿" not in text
