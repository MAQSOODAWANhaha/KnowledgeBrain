"""Known tender requirements survive blank or removed bidder responses in physical exports."""

from copy import deepcopy
from io import BytesIO
import json
from zipfile import ZipFile

from docx import Document
import pytest

from docreader.integrations.artifacts import ArtifactStore
from docreader.tender.authoring_export import export_authoring_document

from test_submission_export import chapter, grid, part, source, target, writing


def rule_ref(artifacts, identity="requirement", *, decision=None, index=None):
    reference = {"item_ref": {"artifact_ref": artifacts.put_json({"fixture": identity}), "item_id": identity}}
    if decision is not None:
        reference.update(decision_id=decision, rule_index=index)
    return reference


@pytest.mark.parametrize("disposition", ["constraint", "excluded"])
def test_non_body_dispositions_preserve_rules_without_false_unassigned_gap(tmp_path, disposition):
    artifacts = ArtifactStore(tmp_path)
    reference = rule_ref(artifacts)
    outline = {"chapters": [chapter("提交安排")], "constraint_rule_refs": [], "issues": []}
    if disposition == "constraint":
        outline["constraint_rule_refs"] = [reference]
    else:
        outline["issues"] = [{"code": "outline_target_excluded", "rule_refs": [reference]}]
    data = requirements(requirement(reference, "按原文限定处理提交安排。", targets=()))
    result = export_authoring_document(artifacts, outline, [], requirements=data)
    assert not any(gap["code"] == "export_requirement_unassigned" for gap in result["gaps"])
    with ZipFile(BytesIO(artifacts.read_bytes(result["package_ref"]))) as archive:
        assert json.loads(archive.read("招标要求底稿.json")) == data
    pending = rule_ref(artifacts, "truly-pending")
    data["rules"].append(requirement(pending, "该要求尚未安排。", targets=()))
    result = export_authoring_document(artifacts, outline, [], requirements=data)
    assert any(gap["code"] == "export_requirement_unassigned" for gap in result["gaps"])


def requirement(reference, text="设备应支持原文规定的功能。", *, targets=("technical",), **fields):
    return {"rule_ref": reference, "rule": {"item_id": reference["item_ref"]["item_id"], "kind": "obligation",
        "text": text, "evidence": [{"range_id": "source-1", "quote": "设备应支持原文规定的功能。"}],
        **fields}, "target_ids": list(targets)}


def requirements(*rules, origin="verified_projection"):
    return {"schema_version": 2, "authoring_input": {"fixture": "frozen-authoring-input"},
            "rule_origin": origin, "rules": list(rules), "issues": []}


def outline_target(identity, refs, **fields):
    return {**target(identity, kind="narrative"), "rule_refs": list(refs), **fields}


def word(artifacts, result):
    return Document(BytesIO(artifacts.read_bytes(result["docx_ref"])))


@pytest.mark.parametrize("legacy_unit", [False, True])
def test_resolved_reference_only_displays_actual_location_without_pending_prose(tmp_path, legacy_unit):
    artifacts = ArtifactStore(tmp_path)
    actual = outline_target("actual", [])
    reference = outline_target("reference", [], related_target_ids=["actual"])
    outline = {"chapters": [chapter("内容", [actual]), chapter("引用", [reference])]}
    units = [writing("body", "内容", ["actual"], blocks=[{"kind": "paragraph", "text": "已有正文。"}])]
    if legacy_unit:
        units.append(writing("old-empty", "引用", ["reference"], status="placeholder", gaps=["旧运行的资料缺口。"]))

    result = export_authoring_document(artifacts, outline, units)

    paragraphs = [p.text for p in word(artifacts, result).paragraphs]
    assert "参见 1 内容" in paragraphs
    assert not any("待补" in p for p in paragraphs)
    assert not any(g['code'] == 'export_content_pending' for g in result['gaps'])
    if legacy_unit:
        assert any(g['detail'] == '旧运行的资料缺口。' for g in result['gaps'])


@pytest.mark.parametrize("related", [[], ["unavailable"], ["reference"], ["actual", "unavailable"]])
def test_unresolved_reference_keeps_pending_position(tmp_path, related):
    artifacts = ArtifactStore(tmp_path)
    reference = outline_target("reference", [], related_target_ids=related)
    outline = {"chapters": [chapter("内容", [outline_target("actual", [])]), chapter("引用", [reference])]}
    result = export_authoring_document(artifacts, outline, [])
    assert any(g['code'] == 'export_content_pending' and g.get('target_id') == 'reference' for g in result['gaps'])
    if related:
        assert any(g['code'] == 'export_related_target_unavailable' and g.get('target_id') == 'reference'
                   for g in result['gaps'])


def test_pure_reference_does_not_hide_existing_written_content(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    reference = outline_target("reference", [], related_target_ids=["actual"])
    outline = {"chapters": [chapter("内容", [outline_target("actual", [])]), chapter("引用", [reference])]}
    units = [writing("legacy", "引用", ["reference"], blocks=[{"kind": "paragraph", "text": "先前已保留的说明。"}])]
    result = export_authoring_document(artifacts, outline, units)
    assert [p.text for p in word(artifacts, result).paragraphs].count("先前已保留的说明。") == 1


def test_reference_does_not_hide_an_empty_unit_shared_with_actual_work(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    reference = outline_target("reference", [], related_target_ids=["actual"])
    outline = {"chapters": [chapter("内容", [reference, outline_target("actual", [])])]}
    units = [writing("shared", "内容", ["reference", "actual"], status="placeholder")]
    result = export_authoring_document(artifacts, outline, units)
    assert "（本部分内容待补）" in [p.text for p in word(artifacts, result).paragraphs]
    assert any(g['code'] == 'export_content_pending' for g in result['gaps'])


@pytest.mark.parametrize("mode", ["no_unit", "no_knowledge", "response_removed_by_review"])
def test_requirements_remain_before_blank_or_removed_enterprise_response(tmp_path, mode):
    artifacts = ArtifactStore(tmp_path)
    reference = rule_ref(artifacts)
    outline = {"chapters": [chapter("技术", [outline_target("technical", [reference])])]}
    units = [] if mode == "no_unit" else [writing("w1", "技术", ["technical"],
        status="placeholder", gaps=["没有真实知识依据。" if mode == "no_knowledge" else "无依据应答已移除。"])]

    result = export_authoring_document(artifacts, outline, units,
                                      requirements=requirements(requirement(reference)))

    document = word(artifacts, result)
    assert len(document.tables) == 1
    row = document.tables[0].rows[1]
    assert "设备应支持原文规定的功能。" in row.cells[0].text
    assert row.cells[-1].text == ""
    assert "招标要求依据（企业应答另行填写）" in [p.text for p in document.paragraphs]
    assert not any("完全满足" in cell.text for cell in row.cells)
    body = list(document.element.body)
    pending = next(p for p in document.paragraphs if "待补" in p.text)
    assert body.index(document.tables[0]._tbl) < body.index(pending._p)


def test_requirement_scope_conditions_exceptions_and_ordered_quotes_are_lossless(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    reference = rule_ref(artifacts)
    entry = requirement(reference, "需要定制适配。", scope="仅适用于已明确的联动部分。",
        conditions=["不能提供该部分证明时", "按实际项目范围"],
        exceptions=["SD-WAN兼容性证明另行提交"],
        evidence=[{"range_id": "p1", "quote": "如不能提供证明，需要给出定制开发承"},
                  {"range_id": "p2", "quote": "诺；其余部分另行提供证明。"}])
    data = requirements(entry)
    before = deepcopy(data)
    outline = {"chapters": [chapter("技术", [outline_target("technical", [reference])])]}

    result = export_authoring_document(artifacts, outline, [], requirements=data)

    document = word(artifacts, result)
    cells = document.tables[0].rows[1].cells
    for text in [entry["rule"]["text"], entry["rule"]["scope"], *entry["rule"]["conditions"],
                 *entry["rule"]["exceptions"]]:
        assert text in cells[0].text
    quotes = [e["quote"] for e in entry["rule"]["evidence"]]
    assert all(quote in cells[1].text for quote in quotes)
    assert cells[1].text.index(quotes[0]) < cells[1].text.index(quotes[1])
    assert any("原文依据" in p.text and "有效要求" in p.text for p in document.paragraphs)
    assert data == before


@pytest.mark.parametrize("details", [
    {"numbers": [{"value": "30", "unit": "天", "quote": "30天"}]},
    {"kind": "scoring_rule", "scoring": {"points": 12, "max_points": 20, "method": "逐项评分"}},
    {"kind": "consequence_rule", "consequence": {"effect": "否决投标", "trigger": "缺少规定证明"}},
    {"kind": "submission_requirement", "submission": {"fields": ["具体参数"], "documents": ["原厂证明"]}},
])
def test_typed_source_facts_are_retained_without_calculation(tmp_path, details):
    artifacts = ArtifactStore(tmp_path)
    reference = rule_ref(artifacts)
    entry = requirement(reference, **details)
    outline = {"chapters": [chapter("要求", [outline_target("technical", [reference])])]}

    result = export_authoring_document(artifacts, outline, [], requirements=requirements(entry))

    content = word(artifacts, result).tables[0].cell(1, 0).text
    for value in details.values():
        if isinstance(value, (dict, list)):
            def leaves(item):
                if isinstance(item, dict):
                    return [leaf for child in item.values() for leaf in leaves(child)]
                if isinstance(item, list):
                    return [leaf for child in item for leaf in leaves(child)]
                return [str(item)]
            assert all(leaf in content for leaf in leaves(value))
    assert word(artifacts, result).tables[0].cell(1, 2).text == ""


def test_effective_rule_matches_full_reference_and_does_not_revive_original(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    reference = rule_ref(artifacts, decision="accepted-amendment", index=1)
    different = {**reference, "rule_index": 0}
    effective = {"rule_ref": reference, "rule": {"text": "仅保留范围乙的现行要求。", "scope": "范围乙",
        "evidence": [{"range_id": "old", "quote": "原文件曾要求范围甲。"}]}, "target_ids": ["technical"]}
    old = requirement(different, "已失效的范围甲要求。", targets=())
    outline = {"chapters": [chapter("技术", [outline_target("technical", [reference])])]}

    result = export_authoring_document(artifacts, outline, [], requirements=requirements(old, effective))

    document = word(artifacts, result)
    assert len(document.tables[0].rows) == 2
    assert "仅保留范围乙的现行要求。" in document.tables[0].cell(1, 0).text
    assert "已失效的范围甲要求。" not in document.tables[0].cell(1, 0).text
    assert "原文件曾要求范围甲。" in document.tables[0].cell(1, 1).text


def test_requirement_basis_precedes_generated_response_and_matches_volume(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    reference = rule_ref(artifacts)
    outline = {"chapters": [chapter("技术", [outline_target("technical", [reference])], volume=True)]}
    units = [writing("w", "技术", ["technical"], blocks=[{"kind": "paragraph", "text": "有依据的企业应答。"}])]

    result = export_authoring_document(artifacts, outline, units,
                                      requirements=requirements(requirement(reference)))

    with ZipFile(BytesIO(artifacts.read_bytes(result["package_ref"]))) as archive:
        for content in [artifacts.read_bytes(result["docx_ref"]), archive.read("volumes/001.docx")]:
            document = Document(BytesIO(content))
            body = list(document.element.body)
            response = next(p for p in document.paragraphs if p.text == "有依据的企业应答。")
            assert body.index(document.tables[0]._tbl) < body.index(response._p)
            assert document.tables[0].cell(1, 2).text == ""


def test_all_requirements_precede_a_shared_writing_unit(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    first, second = rule_ref(artifacts, "first"), rule_ref(artifacts, "second")
    outline = {"chapters": [chapter("技术", [outline_target("technical", [first]),
                                               outline_target("service", [second])])]}
    units = [writing("w", "技术", ["technical", "service"],
                     blocks=[{"kind": "paragraph", "text": "两个位置共享的企业应答。"}])]
    data = requirements(requirement(first), requirement(second, targets=("service",)))

    result = export_authoring_document(artifacts, outline, units, requirements=data)

    document = word(artifacts, result)
    body = list(document.element.body)
    response = next(p for p in document.paragraphs if p.text == "两个位置共享的企业应答。")
    assert len(document.tables) == 1
    assert len(document.tables[0].rows) == 3
    assert all(body.index(table._tbl) < body.index(response._p) for table in document.tables)


def test_forms_and_attachments_are_not_recreated_and_manual_prices_stay_blank(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    reference = rule_ref(artifacts)
    template = source(artifacts, "original-form", grid=grid())
    form = {**target("price", templates=[part(template)], manual=True), "rule_refs": [reference]}
    attachment = {**target("proof", kind="attachment"), "rule_refs": [reference]}
    outline = {"chapters": [chapter("报价及证明", [form, attachment])]}
    data = requirements(requirement(reference, targets=("price", "proof")))

    result = export_authoring_document(artifacts, outline, [], requirements=data)

    document = word(artifacts, result)
    assert len(document.tables) == 2
    assert len(document.tables[0].rows) == 2
    assert document.tables[0].cell(1, 2).text == ""
    assert document.tables[1].cell(2, 1).text == ""
    assert document.tables[1].cell(0, 0).text == "规定表头"
    assert "招标要求依据（企业应答另行填写）" in [p.text for p in document.paragraphs]


def test_attachment_and_form_requirements_survive_without_enterprise_materials(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    proof_ref, deviation_ref = rule_ref(artifacts, "proof"), rule_ref(artifacts, "deviation")
    template = source(artifacts, "deviation-form", grid=grid())
    targets = [outline_target("proof", [proof_ref], kind="attachment"),
               outline_target("deviation", [deviation_ref], kind="form", template_refs=[part(template)])]
    entries = [requirement(proof_ref, "各项证明材料须加盖原厂公章。", targets=("proof",),
        evidence=[{"range_id": "proof-source", "quote": "提供证明材料，并加盖原厂商公章。"}]),
        requirement(deviation_ref, "结构性偏差须有投运业绩证明。", targets=("deviation",),
            conditions=["经评审同意后"], evidence=[{"range_id": "deviation-source",
                                                    "quote": "经评审同意，并有投运业绩证明。"}])]
    outline = {"chapters": [chapter("技术", targets, volume=True)]}
    units = [writing("proof-unit", "技术", ["proof"], status="placeholder", gaps=["缺少真实证明。"])]

    result = export_authoring_document(artifacts, outline, units, requirements=requirements(*entries))

    with ZipFile(BytesIO(artifacts.read_bytes(result["package_ref"]))) as archive:
        for content in [artifacts.read_bytes(result["docx_ref"]), archive.read("volumes/001.docx")]:
            document = Document(BytesIO(content))
            assert len(document.tables) == 2
            basis, original = document.tables
            assert len(basis.rows) == 3
            for row, entry in zip(basis.rows[1:], entries):
                assert entry["rule"]["text"] in row.cells[0].text
                assert entry["rule"]["evidence"][0]["quote"] == row.cells[1].text
                assert row.cells[2].text == ""
            assert "经评审同意后" in basis.cell(2, 0).text
            assert original.cell(0, 0).text == "规定表头"
            assert original.cell(2, 1).text == ""
            body = list(document.element.body)
            pending = next(p for p in document.paragraphs if "待补" in p.text)
            assert body.index(basis._tbl) < body.index(pending._p) < body.index(original._tbl)


def test_chapter_deduplicates_full_rule_identity_without_merging_form_instances(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    shared = rule_ref(artifacts, "shared", decision="current", index=0)
    other = {**shared, "rule_index": 1}
    template = source(artifacts, "shared-form", grid=grid())
    targets = [outline_target("description", [shared]),
               outline_target("comparison", [shared, other]),
               outline_target("proof", [shared], kind="attachment"),
               outline_target("bidder", [shared], kind="form", purpose="投标人分别填写。",
                              template_refs=[part(template)]),
               outline_target("manufacturer", [shared], kind="form", purpose="原厂分别填写。",
                              template_refs=[part(template)])]
    data = requirements(requirement(shared, "原厂证明须盖章。", targets=tuple(t["target_id"] for t in targets)),
                        requirement(other, "另一个有效规则必须保留。", targets=("comparison",)))
    before = deepcopy(data)

    result = export_authoring_document(artifacts, {"chapters": [chapter("技术", targets)]}, [], requirements=data)

    document = word(artifacts, result)
    assert len(document.tables) == 3
    basis, bidder, manufacturer = document.tables
    assert len(basis.rows) == 3
    assert basis.cell(1, 0).text == "原厂证明须盖章。"
    assert basis.cell(2, 0).text == "另一个有效规则必须保留。"
    assert all(row.cells[2].text == "" for row in basis.rows[1:])
    assert bidder.cell(0, 0).text == manufacturer.cell(0, 0).text == "规定表头"
    assert bidder.cell(2, 1).text == manufacturer.cell(2, 1).text == ""
    captions = [p.text for p in document.paragraphs]
    assert "投标人分别填写。" in captions and "原厂分别填写。" in captions
    assert captions.count("招标要求依据（企业应答另行填写）") == 1
    assert data == before


def test_cover_requirements_remain_in_submission_notes_without_preceding_the_original_cover(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    cover_ref, body_ref = rule_ref(artifacts, "cover"), rule_ref(artifacts, "body")
    original_text = "投标文件\n投标人名称：________\n法定代表人签字：________"
    template = source(artifacts, "original-cover", text=original_text)
    cover = {**chapter("封面", [outline_target("cover-form", [cover_ref], kind="form",
        template_refs=[part(template)])]), "presentation_role": "cover"}
    body = chapter("技术要求", [outline_target("technical", [body_ref])])
    cover_text = "封面应填写投标人名称并由法定代表人签字。"
    cover_quote = "投标文件封面须填写投标人全称，法定代表人签字。"
    data = requirements(requirement(cover_ref, cover_text, targets=("cover-form",),
        evidence=[{"range_id": "cover-source", "quote": cover_quote}]),
        requirement(body_ref, "接口应支持原文规定的协议。"))
    before = deepcopy(data)

    result = export_authoring_document(artifacts, {"chapters": [cover, body]}, [], requirements=data)

    document = word(artifacts, result)
    assert len(document.tables) == 1
    assert document.tables[0].cell(1, 0).text == "接口应支持原文规定的协议。"
    assert document.tables[0].cell(1, 2).text == ""
    paragraphs = [p.text for p in document.paragraphs]
    assert paragraphs.index(original_text) < paragraphs.index("目录")
    assert paragraphs.index("目录") < paragraphs.index("招标要求依据（企业应答另行填写）")
    assert cover_text not in "\n".join(paragraphs)
    with ZipFile(BytesIO(artifacts.read_bytes(result["package_ref"]))) as archive:
        notes = archive.read("提交要求与待办.md").decode()
        assert "封面填写要求" in notes
        assert cover_text in notes and cover_quote in notes
        assert json.loads(archive.read("招标要求底稿.json")) == before
    assert not any(gap["code"] == "export_requirement_unassigned" for gap in result["gaps"])
    assert data == before


def test_complete_requirements_are_packaged_with_unassigned_and_unknown_rules(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    assigned = rule_ref(artifacts, "assigned")
    unassigned = rule_ref(artifacts, "unassigned")
    unknown = rule_ref(artifacts, "unknown")
    data = requirements(requirement(assigned), requirement(unassigned, "尚未安排的要求。", targets=()),
                        {"rule_ref": unknown, "rule": None, "target_ids": ["pending"]})
    outline = {"chapters": [chapter("技术", [outline_target("technical", [assigned]),
                                               outline_target("pending", [unknown])])]}
    before = deepcopy(data)

    result = export_authoring_document(artifacts, outline, [], requirements=data)

    with ZipFile(BytesIO(artifacts.read_bytes(result["package_ref"]))) as archive:
        manifest = json.loads(archive.read("manifest.json"))
        record = manifest["requirements"]
        assert record["package_path"] == "招标要求底稿.json"
        assert json.loads(archive.read(record["package_path"])) == before
        assert artifacts.read_json(record["artifact_ref"]) == before
    document = word(artifacts, result)
    assert len(document.tables) == 1
    assert "待核" in document.tables[0].cell(2, 0).text
    assert document.tables[0].cell(2, 2).text == ""
    assert "尚未安排的要求。" not in "\n".join(c.text for table in document.tables for row in table.rows for c in row.cells)
    assert any(g["code"] == "export_requirement_unassigned" for g in result["gaps"])
    assert any(g["code"] == "export_requirement_unavailable" for g in result["gaps"])
    assert data == before


def test_source_only_does_not_copy_context_or_change_legacy_export(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    context = source(artifacts, "context", text="来源上下文不得被整段复制到投标正文。", grid=grid())
    outline = {"chapters": [chapter("技术", [outline_target("technical", [], source_refs=[context])])]}

    legacy = export_authoring_document(artifacts, outline, [])
    result = export_authoring_document(artifacts, outline, [], requirements=requirements(origin="source_only"))

    assert not word(artifacts, result).tables
    assert [p.text for p in word(artifacts, legacy).paragraphs] == [p.text for p in word(artifacts, result).paragraphs]
    with ZipFile(BytesIO(artifacts.read_bytes(legacy["package_ref"]))) as archive:
        assert "requirements" not in json.loads(archive.read("manifest.json"))
        assert "招标要求底稿.json" not in archive.namelist()


def test_legacy_requirement_shape_keeps_document_and_complete_raw_snapshot(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    reference = rule_ref(artifacts)
    entry = requirement(reference, "已有可读要求。", kind=["旧事项标签"],
        scope={"旧范围字段": "特定区域"}, conditions=["原条件", {"旧条件字段": "附带条件"}],
        exceptions="旧协议的单字符串例外", evidence=[{"range_id": "legacy-without-quote"}, "旧引文片段", None])
    data = requirements(entry)
    outline = {"chapters": [chapter("技术", [outline_target("technical", [reference])])]}

    result = export_authoring_document(artifacts, outline, [], requirements=data)

    document = word(artifacts, result)
    row = document.tables[0].rows[1]
    for value in ["已有可读要求。", "特定区域", "原条件", "附带条件", "旧协议的单字符串例外"]:
        assert value in row.cells[0].text
    assert "旧引文片段" in row.cells[1].text and "待核" in row.cells[1].text
    assert row.cells[2].text == ""
    with ZipFile(BytesIO(artifacts.read_bytes(result["package_ref"]))) as archive:
        assert json.loads(archive.read("招标要求底稿.json")) == data
    assert any(g["code"] == "export_requirement_evidence_unverified" for g in result["gaps"])


@pytest.mark.parametrize("details,expected", [
    ({"numbers": [{"value": "30", "unit": "天", "quote": "30天"},
                  {"value": 0, "unit": "元", "quote": "0元"}]},
     ["数值：30", "单位：天", "原文：30天", "数值：0", "单位：元", "原文：0元"]),
    ({"kind": "scoring_rule", "scoring": {"points": 12, "max_points": 20, "method": "逐项评分"}},
     ["分值：12", "最高分：20", "评分方法：逐项评分"]),
    ({"kind": "consequence_rule", "consequence": {"effect": "否决投标", "trigger": "缺少规定证明"}},
     ["后果：否决投标", "触发条件：缺少规定证明"]),
    ({"kind": "submission_requirement", "submission": {"fields": ["具体参数", "品牌型号"],
                                                        "documents": ["原厂证明", "检测报告"]}},
     ["填写字段：", "具体参数", "品牌型号", "提交材料：", "原厂证明", "检测报告"]),
])
def test_structured_requirement_details_are_readable_business_text(tmp_path, details, expected):
    artifacts = ArtifactStore(tmp_path)
    reference = rule_ref(artifacts)
    data = requirements(requirement(reference, **details))
    outline = {"chapters": [chapter("要求", [outline_target("technical", [reference])])]}

    result = export_authoring_document(artifacts, outline, [], requirements=data)

    content = word(artifacts, result).tables[0].cell(1, 0).text
    assert all(value in content for value in expected)
    assert [content.index(value) for value in expected] == sorted(content.index(value) for value in expected)
    assert "后果：后果：" not in content
    assert not any(value in content for value in ['"max_points"', '"points"', '"value"', '"unit"',
                                                  '"quote"', '"effect"', '"trigger"', '"fields"',
                                                  '"documents"', '{', '}', '[', ']'])
    with ZipFile(BytesIO(artifacts.read_bytes(result["package_ref"]))) as archive:
        assert json.loads(archive.read("招标要求底稿.json")) == data


def test_original_business_label_text_and_values_are_preserved(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    reference = rule_ref(artifacts)
    original = "原文示例：后果：后果：否决投标"
    data = requirements(requirement(reference, original,
        consequence={"effect": "后果：否决投标", "trigger": "条件：来源值"},
        evidence=[{"range_id": "original", "quote": original}]))
    before = deepcopy(data)
    outline = {"chapters": [chapter("要求", [outline_target("technical", [reference])])]}

    result = export_authoring_document(artifacts, outline, [], requirements=data)

    row = word(artifacts, result).tables[0].rows[1]
    assert row.cells[0].text == original + "\n后果：后果：否决投标\n触发条件：条件：来源值"
    assert row.cells[1].text == original
    assert row.cells[2].text == ""
    assert data == before
    with ZipFile(BytesIO(artifacts.read_bytes(result["package_ref"]))) as archive:
        assert json.loads(archive.read("招标要求底稿.json")) == before


def test_unknown_legacy_fields_are_pending_without_dumping_internal_names(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    reference = rule_ref(artifacts)
    entry = requirement(reference, "已有可读要求。",
        scoring={"max_points": 20, "legacy_internal_name": {"another_internal_name": ["旧记录值", 0, False]}},
        scope={"private_scope_key": "局部适用"},
        evidence=[{"range_id": "opaque-range", "private_evidence_key": "原记录片段"}])
    data = requirements(entry)
    outline = {"chapters": [chapter("要求", [outline_target("technical", [reference])])]}

    result = export_authoring_document(artifacts, outline, [], requirements=data)

    row = word(artifacts, result).tables[0].rows[1]
    content = "\n".join(cell.text for cell in row.cells)
    assert "最高分：20" in content and "待核" in content
    assert "局部适用" in content and "旧记录值" in content and "原记录片段" in content
    assert not any(name in content for name in ["legacy_internal_name", "another_internal_name", "private_scope_key",
                                               "private_evidence_key", "range_id", "{", "}", "[", "]"])
    assert row.cells[2].text == ""
    assert any(g["code"] == "export_requirement_evidence_unverified" for g in result["gaps"])
    with ZipFile(BytesIO(artifacts.read_bytes(result["package_ref"]))) as archive:
        assert json.loads(archive.read("招标要求底稿.json")) == data


def test_original_json_code_examples_are_preserved_as_source_text(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    reference = rule_ref(artifacts)
    original = '接口示例：{"enabled": true, "ports": [80, 443]}'
    data = requirements(requirement(reference, original, evidence=[{"range_id": "example", "quote": original}]))
    outline = {"chapters": [chapter("接口要求", [outline_target("technical", [reference])])]}

    result = export_authoring_document(artifacts, outline, [], requirements=data)

    row = word(artifacts, result).tables[0].rows[1]
    assert original in row.cells[0].text
    assert row.cells[1].text == original
    assert row.cells[2].text == ""


def test_internal_classification_is_hidden_and_zero_false_values_remain_visible(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    reference = rule_ref(artifacts)
    data = requirements(requirement(reference, kind="scoring_rule", domains=["internal_domain_code"],
                                    scoring={"points": 0, "max_points": 0}, conditions=[False, 0]))
    outline = {"chapters": [chapter("要求", [outline_target("technical", [reference])])]}

    result = export_authoring_document(artifacts, outline, [], requirements=data)

    content = word(artifacts, result).tables[0].cell(1, 0).text
    assert "分值：0" in content and "最高分：0" in content
    assert "条件（待核）：否\n0" in content
    assert not any(value in content for value in ["scoring_rule", "internal_domain_code", "事项类型", "业务领域"])
    with ZipFile(BytesIO(artifacts.read_bytes(result["package_ref"]))) as archive:
        assert json.loads(archive.read("招标要求底稿.json")) == data
