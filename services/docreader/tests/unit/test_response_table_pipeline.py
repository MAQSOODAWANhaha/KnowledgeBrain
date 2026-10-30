"""Source selection, organization, writing, review and DOCX share the same response rows."""

from io import BytesIO
import json

from docx import Document
import pytest

from docreader.integrations.artifacts import ArtifactStore
from docreader.llm.budget import RequestBudget
from docreader.tender.authoring_excerpts import load_source_excerpt
from docreader.tender.authoring_pipeline import generate_document

from test_authoring_pipeline_boundaries import snapshot


SOURCES = ["系统必须支持双机热备。", "协同防御可提供定制开发承", "诺，须盖章并提交详细方案。"]
SUPPORTED = "产品支持双机热备和自动切换。"
SECOND_SUPPORTED = "支持按项目定制协同防御能力，可据此整理详细方案。"
INVENTED = "本产品已经取得本项目的定制开发盖章承诺。"


def request_payload(request):
    return json.loads(request["messages"][1]["content"][0]["text"])


def source_response(payload):
    assert payload.get("request_ref"), "This fixture must exercise real source WirePayload decoding."
    sources = payload["sources"]
    assert [source["text"] for source in sources] == SOURCES
    assert [source["range_id"] for source in sources] == ["r1", "r2", "r3"]
    return {"schema_version": 2, "chapters": [{"title": "技术响应", "children": [], "targets": [
        {"kind": "narrative", "purpose": "逐行对应本项目技术要求。", "range_ids": [], "rule_ids": [],
         "response_table": {"rows": [
             {"range_ids": [sources[0]["range_id"]]},
             {"range_ids": [source["range_id"] for source in sources[1:]]}]}}]}], "findings": []}


def assembly_response(payload, *, explicit):
    candidate = payload["chapters"][0]["targets"][0]
    assert len(candidate["response_table"]["rows"]) == 2
    sources = payload["source_excerpts"]
    assert [source["text"] for source in sources] == SOURCES
    assert [source["range_id"] for source in sources] == ["x1", "x2", "x3"]
    assert candidate["response_table"]["rows"][1]["range_ids"] == ["x2", "x3"]
    target = {"kind": "narrative", "purpose": "按完整来源逐项应答。", "target_ids": [candidate["target_id"]]}
    if explicit:
        target["response_table"] = {"rows": [
            {"range_ids": [sources[0]["range_id"]]}, {"range_ids": [source["range_id"] for source in sources[1:]]}]}
    return {"chapters": [{"title": "技术响应", "targets": [target], "children": []}]}


def raw_requirements(artifacts, outline):
    rows = outline["chapters"][0]["targets"][0]["response_table"]["rows"]
    return ["\n".join(load_source_excerpt(artifacts, ref)["text"] for ref in row["source_excerpts"]) for row in rows]


@pytest.mark.parametrize("explicit", [False, True])
async def test_no_knowledge_keeps_full_crosspage_requirements_without_write_or_content_review(tmp_path, explicit):
    artifacts = ArtifactStore(tmp_path)
    fixed = snapshot(artifacts, [{"heading": "技术要求", "text": text} for text in SOURCES])
    calls = []

    async def execute(operation, request, validate):
        payload = request_payload(request)
        stage = payload.get("operation", operation)
        calls.append(stage)
        if stage == "outline":
            return validate(source_response(payload))
        if stage == "assemble":
            return validate(assembly_response(payload, explicit=explicit))
        if stage == "structure":
            return validate({"findings": []})
        pytest.fail(f"No bidder material must not cause {stage}.")

    result = artifacts.read_json(await generate_document(
        artifacts, fixed, RequestBudget(model="fixture", context_tokens=131072, output_reserve=16384), execute))
    outline = artifacts.read_json(result["outline_ref"])
    assert raw_requirements(artifacts, outline) == [SOURCES[0], "\n".join(SOURCES[1:])]
    rows = outline["chapters"][0]["targets"][0]["response_table"]["rows"]
    assert len({row["row_id"] for row in rows}) == 2
    assert all(row["row_id"].startswith("chapter-1:target-1:row-") for row in rows)
    document = Document(BytesIO(artifacts.read_bytes(result["docx_ref"])))
    assert len(document.tables) == 1
    assert [[cell.text for cell in row.cells] for row in document.tables[0].rows[1:]] == [
        [SOURCES[0], ""], ["\n".join(SOURCES[1:]), ""]]
    assert not any("本部分内容待补" in paragraph.text for paragraph in document.paragraphs)
    assert calls == ["outline", "assemble", "structure"]
    assert not any(gap["code"] == "export_content_pending" for gap in result["gaps"])
    assert result["package_ref"]


@pytest.mark.parametrize("reject_second", [False, True])
async def test_row_only_content_is_reviewed_and_rejection_blanks_only_the_selected_answer(tmp_path, reject_second):
    artifacts = ArtifactStore(tmp_path)
    fixed = snapshot(artifacts, [{"heading": "技术要求", "text": text} for text in SOURCES])
    writes, reviews = [], []
    second_answer = INVENTED if reject_second else SECOND_SUPPORTED

    async def execute(operation, request, validate):
        payload = request_payload(request)
        stage = payload.get("operation", operation)
        if stage == "outline":
            return validate(source_response(payload))
        if stage == "assemble":
            return validate(assembly_response(payload, explicit=True))
        if stage == "structure":
            return validate({"findings": []})
        if operation == "authoring_basis":
            return validate({"shared_arrangements": []})
        if operation == "authoring_write":
            writes.append(payload)
            assert [row["row_id"] for row in payload["response_rows"]] == ["q1", "q2"]
            assert all(row["target_id"] == "t1" for row in payload["response_rows"])
            tender_supports = {support["support_id"]: support["text"] for support in payload["supports"]
                               if support["origin"] == "tender"}
            assert [tender_supports[identity] for identity in payload["response_rows"][1]["support_ids"]] == SOURCES[1:]
            return validate({"blocks": [], "row_fills": [
                {"row_id": "q1", "text": SUPPORTED, "support_ids": ["k1"]},
                {"row_id": "q2", "text": second_answer, "support_ids": ["k1"]}]})
        assert operation == "authoring_review"
        reviews.append(payload)
        unit = payload["units"][0]
        assert [row["row_id"] for row in unit["response_rows"]] == ["q1", "q2"]
        assert [fill["content_id"] for fill in unit["content"]["row_fills"]] == ["r1", "r2"]
        assert [fill["row_id"] for fill in unit["content"]["row_fills"]] == ["q1", "q2"]
        assert unit["content"]["status"] == "generated" and not unit["content"]["blocks"]
        findings = [{"code": "unsupported_content", "detail": "资料未支持本项目定制开发承诺已盖章。",
                     "task_ids": [unit["task_id"]], "content_ids": ["r2"], "can_revise": False}] if reject_second else []
        return validate({"findings": findings})

    result = artifacts.read_json(await generate_document(
        artifacts, fixed, RequestBudget(model="fixture", context_tokens=131072, output_reserve=16384), execute,
        materials=[{"material_id": "product-guide", "title": "产品说明",
                    "text": SUPPORTED if reject_second else SUPPORTED + SECOND_SUPPORTED}]))
    assert len(writes) == len(reviews) == 1
    outline = artifacts.read_json(result["outline_ref"])
    requirements = raw_requirements(artifacts, outline)
    assert requirements == [SOURCES[0], "\n".join(SOURCES[1:])]
    unit = artifacts.read_json(result["unit_refs"][0])
    assert unit["status"] == "generated" and not unit["blocks"] and not unit["cell_fills"]
    assert len(unit["row_fills"]) == (1 if reject_second else 2)
    assert unit["row_fills"][0]["material_ids"] == ["product-guide"]
    document = Document(BytesIO(artifacts.read_bytes(result["docx_ref"])))
    assert len(document.tables) == 1
    assert [[cell.text for cell in row.cells] for row in document.tables[0].rows[1:]] == [
        [requirements[0], SUPPORTED], [requirements[1], "" if reject_second else SECOND_SUPPORTED]]
    if reject_second:
        assert INVENTED not in document.element.xml
        assert any(fill["text"] == INVENTED for reference in result["superseded_unit_refs"]
                   for fill in artifacts.read_json(reference)["row_fills"])
