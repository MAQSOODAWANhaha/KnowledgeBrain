"""Grounded drafts survive runtime failures without reviving rejected content."""

from copy import deepcopy
from io import BytesIO
import json
from types import SimpleNamespace

from docx import Document
import pytest

from docreader.integrations.artifacts import ArtifactStore
from docreader.llm.budget import RequestBudget
from docreader.llm.request import ModelWorkUnitExhausted
from docreader.tender.authoring_pipeline import _adopt_reviewed, generate_document
from docreader.tests.submission_responses import assemble_independent_candidates, provided_template_ids

from test_authoring_pipeline_boundaries import snapshot as source_snapshot


def unit():
    return {'schema_version': 2, 'task_id': 'w1', 'chapter_id': 'chapter-1',
            'target_ids': ['target'], 'status': 'generated',
            'blocks': [{'kind': 'paragraph', 'text': '有据内容', 'target_ids': ['target'],
                        'material_ids': ['method'], 'source_refs': []}],
            'cell_fills': [], 'attachments': [], 'gaps': [], 'error_type': None}


def test_first_review_unavailable_preserves_the_referenced_draft(tmp_path):
    original = unit()
    value = _adopt_reviewed([original], [{'code': 'review_unavailable', 'task_ids': ['w1']}],
                            ArtifactStore(tmp_path), [])
    assert value == [original]


def test_rejected_attachment_does_not_remove_its_valid_sibling_or_prose(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    original = unit()
    original['attachments'] = [{'material_id': key, 'title': key,
                                'artifact_ref': artifacts.put_bytes(key.encode())} for key in ('first', 'second')]
    before = deepcopy(original)
    value = _adopt_reviewed([original], [{'code': 'unsupported_content', 'task_ids': ['w1'],
                                        'content_ids': ['a1']}], artifacts, [])[0]
    assert value['blocks'] == original['blocks']
    assert [a['material_id'] for a in value['attachments']] == ['second']
    assert original == before


@pytest.mark.parametrize('failure', ['revision_failed', 'revision_empty', 'rereview_failed'])
async def test_rejected_content_never_returns_after_a_failed_revision(tmp_path, failure):
    artifacts = ArtifactStore(tmp_path)
    calls = []

    async def execute(stage, request, validate):
        payload = json.loads(request['messages'][1]['content'][0]['text'])
        calls.append('structure' if payload.get('operation') == 'structure' else stage)
        if payload.get('operation') == 'assemble':
            return validate(assemble_independent_candidates(payload))
        if payload.get('operation') == 'structure':
            return validate({'findings': []})
        if stage == 'outline':
            return validate({'schema_version': 2, 'chapters': [{'title': '技术应答', 'children': [],
                'targets': [{'kind': 'narrative', 'purpose': '功能应答', 'rule_ids': [], 'range_ids': []}]}],
                'findings': []})
        if stage == 'authoring_basis':
            return validate({})
        if stage == 'authoring_write':
            if 'previous_content' in payload:
                if failure == 'revision_failed':
                    raise ModelWorkUnitExhausted(stage)
                if failure == 'revision_empty':
                    return validate({'blocks': []})
            return validate({'blocks': [{'kind': 'paragraph', 'text': text, 'target_ids': ['t1'],
                                          'support_ids': ['k1']} for text in ('保留的功能说明', '必须剔除的无据承诺')]})
        if calls.count('authoring_review') == 1:
            return validate({'findings': [{'code': 'unsupported_content', 'detail': '第二段超出资料。',
                'task_ids': [payload['units'][0]['task_id']], 'content_ids': ['b2'],
                'suggestion': '只保留第一段。', 'can_revise': True}]})
        raise ModelWorkUnitExhausted(stage)

    result = artifacts.read_json(await generate_document(artifacts,
        {'schema_version': 2, 'run_id': 'safe-adoption', 'source_version': 1, 'domain_revision': 1},
        RequestBudget(model='fixture', context_tokens=131072, output_reserve=16384), execute,
        materials=[{'material_id': 'method', 'title': '功能资料', 'text': '保留的功能说明。'}]))
    adopted = artifacts.read_json(result['unit_refs'][0])
    assert [b['text'] for b in adopted['blocks']] == ['保留的功能说明']
    assert result['docx_ref'] and result['package_ref']


async def test_tender_background_is_written_without_enterprise_knowledge(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    background = '本项目建设地点为甲园区。'
    snapshot = source_snapshot(artifacts, [{'text': background, 'heading': '项目背景'}])
    calls = []

    async def execute(stage, request, validate):
        value = json.loads(request['messages'][1]['content'][0]['text'])
        calls.append('structure' if value.get('operation') == 'structure' else stage)
        if value.get('operation') == 'assemble':
            return validate(assemble_independent_candidates(value))
        if value.get('operation') == 'structure':
            return validate({'findings': []})
        if stage == 'outline':
            return validate({'schema_version': 2, 'chapters': [{'title': '项目背景', 'children': [],
                'targets': [{'kind': 'narrative', 'purpose': '说明本项目建设地点', 'rule_ids': [],
                             'range_ids': [value['sources'][0]['range_id']], 'support_origins': ['tender']}]}],
                'findings': []})
        if stage == 'authoring_write':
            assert len(value['supports']) == 1 and value['supports'][0]['text'] == background
            return validate({'blocks': [{'kind': 'paragraph', 'text': background,
                'target_ids': ['t1'], 'support_ids': ['s1']}]})
        assert stage == 'authoring_review'
        assert value['units'][0]['content']['blocks'][0]['support_ids'] == ['s1']
        return validate({'findings': []})

    result = artifacts.read_json(await generate_document(artifacts, snapshot,
        RequestBudget(model='fixture', context_tokens=131072, output_reserve=16384), execute))
    adopted = artifacts.read_json(result['unit_refs'][0])
    assert calls == ['outline', 'outline', 'structure', 'authoring_write', 'authoring_review']
    assert adopted['blocks'][0]['text'] == background
    assert adopted['blocks'][0]['material_ids'] == [] and adopted['blocks'][0]['source_refs']


async def test_source_only_retrieval_and_semantic_empty_result_can_be_corrected(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    snapshot = source_snapshot(artifacts, [{'text': '应支持角色权限和异地容灾。'}])
    calls, searched = [], []

    class Knowledge:
        async def search(self, queries, _artifacts):
            searched.extend(queries)
            return {'materials': [{'material_id': 'permissions', 'title': '功能资料',
                'text': '本产品支持角色权限。', 'locator': {'target_ids': [queries[0]['target_id']]}}], 'issues': []}

    async def execute(stage, request, validate):
        value = json.loads(request['messages'][1]['content'][0]['text'])
        calls.append('structure' if value.get('operation') == 'structure' else stage)
        if value.get('operation') == 'assemble':
            return validate(assemble_independent_candidates(value))
        if value.get('operation') == 'structure':
            return validate({'findings': []})
        if stage == 'outline':
            reference = value['sources'][0]['range_id']
            return validate({'schema_version': 2, 'chapters': [{'title': '技术响应', 'children': [],
                'targets': [{'kind': 'narrative', 'purpose': '应答功能', 'rule_ids': [], 'range_ids': [reference]}]}],
                'findings': [], 'search_intents': [{'query': query, 'for_ids': [reference]}
                                                  for query in ('角色权限', '异地容灾')]})
        if stage == 'authoring_basis':
            return validate({})
        if stage == 'authoring_write':
            if 'previous_content' not in value:
                return validate({'blocks': [], 'gaps': ['初步认为资料不适用。']})
            return validate({'blocks': [{'kind': 'paragraph', 'text': '本产品支持角色权限。',
                'target_ids': ['t1'], 'support_ids': ['k1']}], 'gaps': ['异地容灾未取得适用资料，留空。']})
        if calls.count('authoring_review') == 1:
            assert value['units'][0]['content']['blocks'] == []
            return validate({'findings': [{'code': 'supported_response_missing', 'detail': '已有资料可支持角色权限。',
                'task_ids': [value['units'][0]['task_id']], 'suggestion': '仅补写角色权限，容灾留空。', 'can_revise': True}]})
        return validate({'findings': []})

    result = artifacts.read_json(await generate_document(artifacts, snapshot,
        RequestBudget(model='fixture', context_tokens=131072, output_reserve=16384), execute, knowledge=Knowledge()))
    adopted = artifacts.read_json(result['unit_refs'][0])
    assert [q['query'] for q in searched] == ['角色权限', '异地容灾']
    assert calls.count('authoring_write') == calls.count('authoring_review') == 2
    assert adopted['blocks'][0]['material_ids'] == ['permissions']
    assert any('容灾' in gap for gap in adopted['gaps'])


@pytest.mark.parametrize('text', ['', '冻结图片的可读OCR正文'])
def test_ocr_support_uses_actual_text_once_instead_of_the_status_dictionary(tmp_path, monkeypatch, text):
    from docreader.tender import authoring_pipeline as pipeline
    artifacts = ArtifactStore(tmp_path)
    ref = {'source_index_ref': artifacts.put_json({'fixture': 'image'}), 'unit_id': 'image1'}
    context = SimpleNamespace(payload={'rules': [], 'issues': []}, rules={}, sources={'range': [ref]}, input={'issues': []})
    task = {'targets': [{'target_id': 'target', 'kind': 'narrative', 'purpose': '项目背景', 'rule_refs': [],
                         'source_refs': [ref], 'support_origins': ['tender']}], 'title_path': ['背景']}
    monkeypatch.setattr(pipeline, 'load_outline_source', lambda *_args: {
        'document': {'role': 'main'}, 'unit': {'text': ''},
        'ocr': {'status': 'complete' if text else 'unresolved', 'text': text}})
    payload, _, _ = pipeline._load_task(artifacts, context, task, {'response': {'shared_arrangements': []}}, [])
    assert pipeline._has_allowed_support(payload) is bool(text)
    assert payload['supports'][0]['text'] == text
    assert 'text' not in payload['supports'][0]['ocr']
    if text:
        assert json.dumps(payload, ensure_ascii=False).count(text) == 1


async def test_basis_batches_never_resend_previous_batches_original_materials(tmp_path, monkeypatch):
    from docreader.tender import authoring_pipeline as pipeline
    artifacts = ArtifactStore(tmp_path)
    original_batches = pipeline._batches
    observed = []

    def batches(records, budget, stage, payload, **options):
        if stage == 'authoring_basis':
            yield from ([record] for record in records)
        else:
            yield from original_batches(records, budget, stage, payload, **options)
    monkeypatch.setattr(pipeline, '_batches', batches)

    async def execute(stage, request, validate):
        value = json.loads(request['messages'][1]['content'][0]['text'])
        if value.get('operation') == 'assemble':
            return validate(assemble_independent_candidates(value))
        if value.get('operation') == 'structure':
            return validate({'findings': []})
        if stage == 'outline':
            return validate({'schema_version': 2, 'chapters': [{'title': '技术响应', 'children': [],
                'targets': [{'kind': 'narrative', 'purpose': '应答功能', 'rule_ids': [], 'range_ids': []}]}], 'findings': []})
        if stage == 'authoring_basis':
            assert len(value['supports']) == 1 and 'previous' not in value
            support = value['supports'][0]
            observed.append(support['text'])
            return validate({'shared_arrangements': [{'topic': support['title'], 'text': support['text'],
                'nature': 'provided', 'support_ids': [support['support_id']]}]})
        if stage == 'authoring_write':
            return validate({'blocks': []})
        return validate({'findings': []})

    result = artifacts.read_json(await generate_document(artifacts,
        {'schema_version': 2, 'run_id': 'batch-basis', 'source_version': 1, 'domain_revision': 1},
        RequestBudget(model='fixture', context_tokens=131072, output_reserve=16384), execute,
        materials=[{'material_id': f'm{i}', 'title': f'功能{i}', 'text': f'第{i}份实际功能说明。'} for i in range(3)]))
    assert observed == [f'第{i}份实际功能说明。' for i in range(3)]
    basis = artifacts.read_json(result['basis_ref'])
    assert [a['material_ids'] for a in basis['response']['shared_arrangements']] == [['m0'], ['m1'], ['m2']]


async def test_mixed_nonprice_form_uses_source_and_company_support_without_copying(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    grid = {'row_count': 3, 'column_count': 2, 'cells': [
        {'row': row, 'column': col, 'text': label if col == 0 else '', 'row_span': 1, 'col_span': 1}
        for row, label in enumerate(('项目名称', '人员姓名', '证书编号')) for col in range(2)]}
    snapshot = source_snapshot(artifacts, [{'text': '本项目名称为甲项目。', 'grid': grid}])
    template_aliases = []

    async def execute(stage, request, validate):
        value = json.loads(request['messages'][1]['content'][0]['text'])
        if value.get('operation') == 'assemble':
            return validate(assemble_independent_candidates(value))
        if value.get('operation') == 'structure':
            return validate({'findings': []})
        if stage == 'outline':
            return validate({'schema_version': 2, 'chapters': [{'title': '项目响应表', 'children': [],
                'targets': [{'kind': 'form', 'purpose': '填写项目名称和企业人员，证书编号留空', 'rule_ids': [],
                             'range_ids': [s['range_id'] for s in value['sources']],
                             'template_part_ids': provided_template_ids(value),
                             'support_origins': ['tender', 'knowledge']}]}], 'findings': []})
        if stage == 'authoring_basis':
            return validate({})
        if stage == 'authoring_write':
            assert sum(bool(s.get('grid')) for s in value['supports']) == 1
            template_alias = next(s['support_id'] for s in value['supports'] if s.get('grid'))
            assert template_alias in value['targets'][0]['template_support_ids']
            template_aliases.append(template_alias)
            return validate({'blocks': [], 'cell_fills': [
                {'source_id': template_alias, 'row': row, 'column': 1, 'text': text,
                 'target_ids': ['t1'], 'support_ids': [support]}
                for row, text, support in ((0, '甲项目', template_alias), (1, '张三', 'k1'))]})
        assert [f['support_ids'] for f in value['units'][0]['content']['cell_fills']] == [[template_aliases[0]], ['k1']]
        return validate({'findings': []})

    result = artifacts.read_json(await generate_document(artifacts, snapshot,
        RequestBudget(model='fixture', context_tokens=131072, output_reserve=16384), execute,
        materials=[{'material_id': 'people', 'title': '人员资料', 'text': '可安排人员张三。'}]))
    adopted = artifacts.read_json(result['unit_refs'][0])
    assert len(adopted['cell_fills']) == 2
    assert adopted['cell_fills'][0]['source_refs'] and adopted['cell_fills'][0]['material_ids'] == []
    assert adopted['cell_fills'][1]['material_ids'] == ['people']
    document = Document(BytesIO(artifacts.read_bytes(result['docx_ref'])))
    tables = [table for table in document.tables if len(table.rows) == 3 and table.cell(0, 0).text == '项目名称']
    assert len(tables) == 1
    assert [tables[0].cell(row, 1).text for row in range(3)] == ['甲项目', '张三', '']


async def test_price_form_is_one_manual_instance_with_no_automatic_fills(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    grid = {'row_count': 2, 'column_count': 2, 'cells': [
        {'row': row, 'column': column, 'text': label if column == 0 else '', 'row_span': 1, 'col_span': 1}
        for row, label in enumerate(('人员姓名', '投标报价')) for column in range(2)]}
    snapshot = source_snapshot(artifacts, [{'grid': grid}])
    operations = []

    async def execute(stage, request, validate):
        value = json.loads(request['messages'][1]['content'][0]['text'])
        operations.append(value.get('operation', stage))
        if value.get('operation') == 'assemble':
            return validate(assemble_independent_candidates(value))
        if value.get('operation') == 'structure':
            selected = value['chapters'][0]['targets']
            assert len(selected) == 1 and selected[0]['manual_fill']
            return validate({'findings': []})
        assert stage == 'outline', 'Manual forms do not incur writing or content-review calls.'
        return validate({'schema_version': 2, 'chapters': [{'title': '报价表', 'children': [],
            'targets': [{'kind': 'form', 'purpose': '整表人工填写人员与最终报价', 'rule_ids': [],
                'range_ids': [], 'template_part_ids': provided_template_ids(value), 'manual_fill': True}]}],
            'findings': []})

    result = artifacts.read_json(await generate_document(artifacts, snapshot,
        RequestBudget(model='fixture', context_tokens=131072, output_reserve=16384), execute))
    assert operations == ['outline', 'assemble', 'structure']
    units = [artifacts.read_json(ref) for ref in result['unit_refs']]
    assert len(units) == 1 and units[0]['cell_fills'] == []
    document = Document(BytesIO(artifacts.read_bytes(result['docx_ref'])))
    assert len(document.tables) == 1
    assert [document.tables[0].cell(row, 1).text for row in range(2)] == ['', '']


async def test_explicit_rereview_changes_review_and_revision_but_not_initial_writing(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    runs = []

    async def run(review_request):
        calls = []
        runs.append(calls)
        async def execute(stage, request, validate):
            value = json.loads(request['messages'][1]['content'][0]['text'])
            calls.append((stage, value))
            if value.get('operation') == 'assemble':
                return validate(assemble_independent_candidates(value))
            if value.get('operation') == 'structure':
                return validate({'findings': []})
            if stage == 'outline':
                return validate({'schema_version': 2, 'chapters': [{'title': '技术响应', 'children': [],
                    'targets': [{'kind': 'narrative', 'purpose': '应答功能', 'rule_ids': [], 'range_ids': []}]}], 'findings': []})
            if stage == 'authoring_basis':
                return validate({})
            if stage == 'authoring_write':
                return validate({'blocks': [{'kind': 'paragraph', 'text': '有依据的功能说明。',
                                            'target_ids': ['t1'], 'support_ids': ['k1']}]})
            if sum(s == 'authoring_review' for s, _ in calls) == 1:
                return validate({'findings': [{'code': 'more_detail', 'detail': '已有资料可具体说明。',
                    'task_ids': [value['units'][0]['task_id']], 'suggestion': '根据资料细化。', 'can_revise': True}]})
            return validate({'findings': []})
        return await generate_document(artifacts,
            {'schema_version': 2, 'run_id': 'rereview', 'source_version': 1, 'domain_revision': 1},
            RequestBudget(model='fixture', context_tokens=131072, output_reserve=16384), execute,
            materials=[{'material_id': 'function', 'title': '功能资料', 'text': '有依据的功能说明。'}],
            review_request=review_request)

    await run(None)
    reason = '请重新核对权限范围与本项目要求是否对应。'
    await run({'target_ids': ['chapter-1:target-1'], 'reason': reason})
    initial = [[v for stage, v in calls if stage == 'authoring_write' and 'previous_content' not in v][0] for calls in runs]
    assert initial[0] == initial[1] and 'review_request' not in initial[1]
    review = next(v for stage, v in runs[1] if stage == 'authoring_review')
    assert review['review_request'] == {'reason': reason, 'targets': [{'task_id': 'w1', 'target_ids': ['t1']}]}
    correction = next(v for stage, v in runs[1] if stage == 'authoring_write' and 'previous_content' in v)
    assert correction['review_request'] == {'reason': reason, 'target_ids': ['t1']}
