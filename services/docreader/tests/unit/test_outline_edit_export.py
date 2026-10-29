"""Human outline edits reach the actual Word hierarchy, forms and stable anchors."""

from collections import Counter
from copy import deepcopy
from io import BytesIO
import json
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

from docreader.integrations.artifacts import ArtifactStore
from docreader.tender.authoring_pipeline import generate_document, plan_writing_tasks
from docreader.tender.outline_edit import apply_outline_edit
from docreader.tender.schema import AuthoringInput, AuthoringOutline
from docreader.tender.word_navigation import bookmark_name, read_docx_outline

from test_authoring_export import grid, source_ref
from test_authoring_pipeline_boundaries import budget, ordinary_response, payload


GOAL_CHAPTER = "manual-328f852d-3d53-4b7a-851d-fc0572420173"
EMPTY_CHAPTER = "manual-b78d5270-5938-48c2-acb7-e9538910fea7"


async def test_human_outline_projection_exports_its_actual_word_hierarchy_forms_and_bookmarks(tmp_path):
    artifacts = ArtifactStore(tmp_path / "artifacts")
    reference = source_ref(artifacts, grid=grid())
    descriptor = artifacts.read_json(reference["source_index_ref"])
    frozen_raw = deepcopy(artifacts.read_json(descriptor["raw_ref"]))
    frozen = artifacts.put_json({"schema_version": 2, "documents": [{
        **descriptor["document"], "role": "main", "raw_ref": descriptor["raw_ref"],
        "source_index_ref": reference["source_index_ref"]}]})
    snapshot = AuthoringInput(schema_version=2, run_id="outline-export", source_version=1,
                              domain_revision=1, frozen_ref=frozen).model_dump()
    form = {"target_id": "stable-form-target", "kind": "form", "purpose": "保留人员安排原表",
            "rule_refs": [], "source_refs": [reference],
            "template_refs": [{"source_ref": reference, "selectors": []}],
            "manual_fill": True, "support_origins": ["tender"]}
    narrative = {"target_id": "stable-narrative-target", "kind": "narrative", "purpose": "说明实施方法",
                 "rule_refs": [], "source_refs": [reference], "template_refs": [],
                 "related_target_ids": [form["target_id"]]}
    removed = {"target_id": "removed-target", "kind": "narrative", "purpose": "不再采用的响应位置",
               "rule_refs": [], "source_refs": []}
    base = AuthoringOutline(schema_version=2, authoring_input=snapshot, issues=[], chapters=[
        {"chapter_id": "old-root", "title": "原技术篇", "targets": [], "children": [
            {"chapter_id": "narrative", "title": "原实施方案", "targets": [narrative], "children": []},
            {"chapter_id": "removed", "title": "删除章节", "targets": [removed], "children": []}]},
        {"chapter_id": "form", "title": "原人员表", "targets": [form], "children": []},
        {"chapter_id": "container", "title": "原容器", "targets": [], "children": []},
    ]).model_dump(exclude_unset=True)
    before = deepcopy(base)
    goal = "依据已有企业资料补充交付安排，保留待确认事项。"
    edited = apply_outline_edit(base, [
        {"chapter_id": "container", "title": "商务响应篇", "children": [
            {"chapter_id": "form", "title": "规定人员格式", "children": []},
            {"chapter_id": "narrative", "title": "实施方案新版", "children": []}]},
        {"chapter_id": GOAL_CHAPTER, "title": "用户补充安排", "writing_goal": goal, "children": []},
        {"chapter_id": EMPTY_CHAPTER, "title": "留白新章", "children": []},
    ])
    assert base == before
    tasks = plan_writing_tasks(edited)
    goal_task = next(task for task in tasks if task["chapter_id"] == GOAL_CHAPTER)
    assert goal_task["targets"][0]["purpose"] == goal
    assert all(task["chapter_id"] not in {EMPTY_CHAPTER, "removed"} for task in tasks)
    observed = []

    async def execute(task, request, validate):
        assert task != "outline", "Human arrangement must survive without a new organization pass."
        value = payload(request)
        observed.append((task, value))
        if task == "authoring_write":
            title = value["title_path"][-1]
            text = "用户章节正文：按资料补充交付安排。" if title == "用户补充安排" else "保留章节正文：按资料实施。"
            return validate({"blocks": [{"kind": "paragraph", "text": text,
                "target_ids": [target["target_id"] for target in value["targets"]],
                "support_ids": [support["support_id"] for support in value["supports"]
                                if support["origin"] == "knowledge"]}]})
        return validate(ordinary_response(task, value))

    result = artifacts.read_json(await generate_document(artifacts, snapshot, budget(), execute,
        outline=edited, materials=[{"material_id": "provided-method", "title": "已给实施资料",
            "text": "按项目资料实施并补充交付安排，具体人员资料待确认。"}]))
    units = [artifacts.read_json(ref) for ref in result["unit_refs"]]
    goal_unit = next(unit for unit in units if unit["chapter_id"] == GOAL_CHAPTER)
    assert goal_unit["target_ids"] == ["manual-response-" + GOAL_CHAPTER]
    assert any(task == "authoring_write" and value["title_path"] == ["用户补充安排"]
               and value["targets"][0]["purpose"] == goal for task, value in observed)
    assert all(unit["chapter_id"] not in {EMPTY_CHAPTER, "removed"} for unit in units)

    data = artifacts.read_bytes(result["docx_ref"])
    document = Document(BytesIO(data))
    with ZipFile(BytesIO(data)) as archive:
        xml = archive.read("word/document.xml")
    root = etree.fromstring(xml)
    body = root.find(qn("w:body"))
    headings = []
    for paragraph in body.findall(qn("w:p")):
        style = paragraph.find(f"{qn('w:pPr')}/{qn('w:pStyle')}")
        if style is not None and style.get(qn("w:val")) in {"Heading1", "Heading2"}:
            text = "".join(part.text or "" for part in paragraph.iter(qn("w:t")))
            headings.append((text, style.get(qn("w:val"))))
    assert headings == [("1 商务响应篇", "Heading1"), ("1.1 规定人员格式", "Heading2"),
        ("1.2 实施方案新版", "Heading2"), ("2 用户补充安排", "Heading1"), ("3 留白新章", "Heading1")]
    all_text = "\n".join(part.text or "" for part in root.iter(qn("w:t")))
    assert "删除章节" not in all_text and "原技术篇" not in all_text
    assert "用户章节正文：按资料补充交付安排。" in all_text
    assert "保留章节正文：按资料实施。" in all_text
    assert "参见 1.1 规定人员格式" in all_text

    assert len(document.tables) == 1
    table = document.tables[0]
    assert table.cell(0, 0).text == "人员安排表"
    assert table.cell(0, 0)._tc is table.cell(0, 1)._tc
    assert table.cell(1, 0).text == "负责人" and table.cell(1, 1).text == ""
    assert table.cell(2, 0).text == "证书编号" and table.cell(2, 1).text == ""
    assert artifacts.read_json(descriptor["raw_ref"]) == frozen_raw

    bookmarks = Counter(anchor.get(qn("w:name")) for anchor in root.iter(qn("w:bookmarkStart")))
    for identity in ("container", "form", "narrative", GOAL_CHAPTER, EMPTY_CHAPTER):
        assert bookmarks[bookmark_name(identity)] == 1
    for identity in ("stable-form-target", "stable-narrative-target", "manual-response-" + GOAL_CHAPTER):
        assert bookmarks[bookmark_name(identity, "target")] == 1
    assert bookmarks[bookmark_name("removed")] == bookmarks[bookmark_name("old-root")] == 0
    assert bookmarks[bookmark_name("removed-target", "target")] == 0
    assert any(anchor.get(qn("w:name")) == bookmark_name("stable-form-target", "target")
               for anchor in root.find(f"{qn('w:body')}/{qn('w:tbl')}").iter(qn("w:bookmarkStart")))
    empty_heading = next(paragraph for paragraph in body.findall(qn("w:p"))
        if any(anchor.get(qn("w:name")) == bookmark_name(EMPTY_CHAPTER)
               for anchor in paragraph.iter(qn("w:bookmarkStart"))))
    assert all(node.tag == qn("w:sectPr") for node in list(body)[list(body).index(empty_heading) + 1:])
    navigation = read_docx_outline(data, edited["chapters"])
    assert navigation["missing_chapter_ids"] == []
    moved_form = navigation["chapters"][0]["children"][0]
    assert moved_form["chapter_id"] == "form"
    assert moved_form["targets"][0]["bookmark"] == bookmark_name("stable-form-target", "target")

    (tmp_path / "edited-outline.docx").write_bytes(data)
    (tmp_path / "word-document.xml").write_bytes(xml)
    (tmp_path / "edited-outline.json").write_text(json.dumps(edited, ensure_ascii=False, indent=2))
    (tmp_path / "export-summary.json").write_text(json.dumps({
        "headings": headings, "bookmarks": dict(bookmarks), "writing_tasks": tasks,
        "model_tasks": dict(Counter(task for task, _ in observed)),
        "gaps": result["gaps"], "navigation": navigation}, ensure_ascii=False, indent=2))
