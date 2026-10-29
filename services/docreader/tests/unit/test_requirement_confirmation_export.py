"""Unconfirmed analysis keeps its facts while exposing their review status."""

from copy import deepcopy
from io import BytesIO
import json
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
import pytest

from docreader.integrations.artifacts import ArtifactStore
from docreader.tender.authoring_export import export_authoring_document
from test_requirement_preservation_export import outline_target, requirement, requirements, rule_ref
from test_submission_export import chapter, grid, part, source, target, writing


@pytest.mark.parametrize("confirmation", ["absent", None, "confirmed"])
def test_confirmation_status_labels_main_volume_and_notes_without_changing_facts(tmp_path, confirmation):
    artifacts = ArtifactStore(tmp_path)
    reference = rule_ref(artifacts, decision="accepted-amendment", index=1)
    entry = requirement(reference, "有效补遗仅要求范围乙。", scope="范围乙",
                        conditions=["按当前有效条件"],
                        evidence=[{"range_id": "old", "quote": "原文曾要求范围甲。"}],
                        targets=("first", "second"))
    data = requirements(entry)
    if confirmation != "absent":
        data["authoring_input"]["confirmation_ref"] = (
            artifacts.put_json({"confirmed": True}) if confirmation == "confirmed" else None)
    selected = part(source(artifacts, "native-form", grid=grid()))
    targets = [outline_target(identity, [reference]) for identity in ("first", "second")]
    targets.append(target("form", templates=[selected]))
    outline = {"chapters": [chapter("技术", targets, volume=True)], "constraint_rule_refs": [reference]}
    units = [writing("shared", "技术", ["first", "second"], status="placeholder")]
    before = deepcopy((data, outline, units))

    result = export_authoring_document(artifacts, outline, units, requirements=data)

    pending = confirmation != "confirmed"
    with ZipFile(BytesIO(artifacts.read_bytes(result["package_ref"]))) as archive:
        assert json.loads(archive.read("招标要求底稿.json")) == data
        notes = archive.read("提交要求与待办.md").decode()
        assert ("分析整理（待确认）" in notes) is pending
        assert ("摘要及适用范围仍需结合原文和有效补遗核对" in notes) is pending
        assert "有效补遗仅要求范围乙。" in notes
        assert "原文曾要求范围甲。" in notes
        for content in (artifacts.read_bytes(result["docx_ref"]), archive.read("volumes/001.docx")):
            document = Document(BytesIO(content))
            basis, native = document.tables
            assert basis.cell(0, 0).text == ("分析整理（待确认）" if pending else "招标要求")
            assert basis.cell(1, 0).text == "有效补遗仅要求范围乙。\n适用范围：范围乙\n条件：按当前有效条件"
            assert basis.cell(1, 1).text == "原文曾要求范围甲。"
            assert basis.cell(1, 2).text == ""
            assert (any("摘要及适用范围仍需结合原文和有效补遗核对" in p.text
                        for p in document.paragraphs)) is pending
            assert native.cell(0, 0).text == "规定表头"
            assert native.cell(1, 1).text == native.cell(2, 1).text == ""
            bookmarks = [element.get(qn("w:name"))
                         for element in document.element.body.iter(qn("w:bookmarkStart"))]
            assert len(bookmarks) == len(set(bookmarks)) == 4
    assert (data, outline, units) == before


def test_source_only_export_without_requirement_snapshot_keeps_existing_notes(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    result = export_authoring_document(artifacts, {"chapters": [chapter("原格式")]}, [])
    with ZipFile(BytesIO(artifacts.read_bytes(result["package_ref"]))) as archive:
        assert "分析整理（待确认）" not in archive.read("提交要求与待办.md").decode()
