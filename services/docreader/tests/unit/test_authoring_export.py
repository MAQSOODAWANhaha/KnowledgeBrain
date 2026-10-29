"""Actual DOCX contents and packaged evidence remain usable when material is missing."""

from copy import deepcopy
from io import BytesIO
import json
from zipfile import ZipFile

from docx import Document
from PIL import Image

from docreader.integrations.artifacts import ArtifactStore
from docreader.tender.authoring_export import export_authoring_document
from docreader.tender.evidence.references import SourceIndex


def source_ref(artifacts, *, grid=None, text="", document_id="template"):
    unit = {"unit_id": "unit:0", "kind": "table_region" if grid else "section",
            "ordinal": 0, "text": text, "grid": grid,
            "locator": {"locator_kind": "document", "section_ordinal": 0}}
    raw = {"schema_version": 2, "units": [unit], "images": {}, "ocr": {},
           "source_contract": {"page_manifest": [], "units": [{
               "unit_id": "unit:0", "section_id": "section:0", "parent_section_id": None,
               "heading_path": "投标格式", "physical_locator": unit["locator"]}]}}
    document = {"document_id": document_id, "document_version": "1",
                "original_ref": artifacts.put_bytes(b"synthetic original")}
    index = SourceIndex(document, raw)
    ref = artifacts.put_json(index.descriptor(artifacts.put_json(index.raw)))
    return {"source_index_ref": ref, "unit_id": "unit:0"}


def target(identity, kind="narrative", refs=(), *, templates=()):
    return {"target_id": identity, "kind": kind, "purpose": "逐项应答",
            "source_refs": list(refs), "rule_refs": [],
            "template_refs": [{"source_ref": ref, "selectors": []} for ref in templates]}


def chapter(identity, title, targets=(), children=()):
    return {"chapter_id": identity, "title": title, "targets": list(targets), "children": list(children)}


def writing(identity, chapter_id, targets, **fields):
    return {"schema_version": 2, "task_id": identity, "chapter_id": chapter_id,
            "target_ids": targets, "status": "generated", "blocks": [],
            "cell_fills": [], "attachments": [], "gaps": [], **fields}


def document(artifacts, result):
    return Document(BytesIO(artifacts.read_bytes(result["docx_ref"])))


def grid():
    return {"row_count": 3, "column_count": 2, "widths_mm": [50, 100], "cells": [
        {"row": 0, "column": 0, "row_span": 1, "col_span": 2, "text": "人员安排表"},
        {"row": 1, "column": 0, "row_span": 1, "col_span": 1, "text": "负责人"},
        {"row": 1, "column": 1, "row_span": 1, "col_span": 1, "text": ""},
        {"row": 2, "column": 0, "row_span": 1, "col_span": 1, "text": "证书编号"},
        {"row": 2, "column": 1, "row_span": 1, "col_span": 1, "text": ""},
    ]}


def test_docx_keeps_nested_order_and_unfilled_required_positions(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    outline = {"chapters": [chapter("a", "技术文件", children=[
        chapter("b", "实施计划", [target("implementation")]),
        chapter("c", "人员", [target("staff")]),
    ])]}
    result = export_authoring_document(artifacts, outline, [writing(
        "task-1", "b", ["implementation"], blocks=[
            {"kind": "paragraph", "text": "按现有需求拟定实施步骤。"},
            {"kind": "table", "headers": ["阶段", "交付物"], "rows": [["部署", "配置记录"]]},
        ])])
    doc = document(artifacts, result)
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert headings == ["1 技术文件", "1.1 实施计划", "1.2 人员"]
    assert "按现有需求拟定实施步骤。" in [p.text for p in doc.paragraphs]
    assert [cell.text for cell in doc.tables[0].rows[1].cells] == ["部署", "配置记录"]
    assert any(g.get("target_id") == "staff" for g in result["gaps"])


def test_original_table_words_and_merges_are_preserved_and_only_empty_cells_filled(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    ref = source_ref(artifacts, grid=grid())
    original = artifacts.read_json(ref["source_index_ref"])
    original_raw = artifacts.read_json(original["raw_ref"])
    result = export_authoring_document(artifacts, {"chapters": [
        chapter("a", "人员表", [target("staff", "form", [ref], templates=[ref])])]}, [writing(
            "task-1", "a", ["staff"], cell_fills=[
                {"source_ref": ref, "row": 1, "column": 1, "text": "项目负责人（姓名待补）"},
                {"source_ref": ref, "row": 0, "column": 0, "text": "非法覆盖原表头"},
            ])])
    table = document(artifacts, result).tables[0]
    assert table.cell(0, 0).text == "人员安排表"
    assert table.cell(0, 0)._tc is table.cell(0, 1)._tc
    assert table.cell(1, 0).text == "负责人"
    assert table.cell(1, 1).text == "项目负责人（姓名待补）"
    assert table.cell(2, 1).text == ""
    assert any(g["code"] == "export_fixed_cell_preserved" for g in result["gaps"])
    assert artifacts.read_json(original["raw_ref"]) == original_raw


def test_multiple_tasks_fill_one_table_once_and_conflicts_stay_pending(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    ref = source_ref(artifacts, grid=grid())
    outline = {"chapters": [chapter("a", "人员表", [target("staff", "form", [ref], templates=[ref])])]}
    units = [writing("task-1", "a", ["staff"], cell_fills=[
        {"source_ref": ref, "row": 1, "column": 1, "text": "甲"}]),
        writing("task-2", "a", ["staff"], cell_fills=[
            {"source_ref": ref, "row": 1, "column": 1, "text": "乙"},
            {"source_ref": ref, "row": 2, "column": 1, "text": "真实证书编号"}])]
    result = export_authoring_document(artifacts, outline, units)
    tables = document(artifacts, result).tables
    assert len(tables) == 1
    assert tables[0].cell(1, 1).text == ""
    assert tables[0].cell(2, 1).text == "真实证书编号"
    assert any(g["code"] == "export_conflicting_cell_fills" for g in result["gaps"])


def test_broken_source_or_attachment_keeps_other_content_and_original_artifact(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    ref = source_ref(artifacts, text="固定声明：本文件所述内容真实。")
    broken = {**ref, "unit_id": "unknown-unit"}
    missing = artifacts.put_bytes(b"missing proof", "application/pdf")
    artifacts.path(missing).unlink()
    outline = {"chapters": [chapter("a", "承诺和证明", [target("form", "form", [broken], templates=[broken])]),
                             chapter("b", "服务方案", [target("service")])]}
    units = [writing("task-1", "a", ["form"], attachments=[
        {"material_id": "proof", "title": "缺失证明", "artifact_ref": missing}]),
        writing("task-2", "b", ["service"], blocks=[{"kind": "paragraph", "text": "具体服务流程。"}])]
    result = export_authoring_document(artifacts, outline, units)
    assert "具体服务流程。" in [p.text for p in document(artifacts, result).paragraphs]
    assert {"export_source_unavailable", "export_attachment_unavailable"} <= {g["code"] for g in result["gaps"]}
    assert result["attachments"] == []


def test_real_image_inserted_and_original_docx_attachment_packaged_without_recreation(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    png = BytesIO()
    Image.new("RGB", (24, 12), "navy").save(png, format="PNG")
    image_ref = artifacts.put_bytes(png.getvalue(), "image/png")
    original_docx = BytesIO()
    original = Document()
    original.add_paragraph("原件证明内容，不能改写。")
    original.save(original_docx)
    docx_ref = artifacts.put_bytes(original_docx.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    attachments = [{"material_id": "image", "title": "真实证明图", "artifact_ref": image_ref},
                   {"material_id": "docx", "title": "原始证明文件", "artifact_ref": docx_ref}]
    result = export_authoring_document(artifacts, {"chapters": [
        chapter("a", "证明材料", [target("evidence", "attachment")])]}, [writing(
            "task-1", "a", ["evidence"], attachments=attachments)])
    assert len(document(artifacts, result).inline_shapes) == 1
    assert result["attachments"] == attachments
    with ZipFile(BytesIO(artifacts.read_bytes(result["package_ref"]))) as archive:
        manifest = json.loads(archive.read("manifest.json"))
        for attachment in manifest["attachments"]:
            assert archive.read(attachment["package_path"]) == artifacts.read_bytes(attachment["artifact_ref"])
        assert archive.read("投标文件.docx") == artifacts.read_bytes(result["docx_ref"])
        assert json.loads(archive.read("待补事项.json")) == result["gaps"]
    assert any(g["code"] == "export_attachment_not_embedded" for g in result["gaps"])


def test_source_images_are_not_mistaken_for_selected_bid_evidence(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    ref = source_ref(artifacts, text="须提供检验报告。")
    result = export_authoring_document(artifacts, {"chapters": [
        chapter("a", "证明材料", [target("evidence", "attachment", [ref])])]}, [])
    assert result["attachments"] == []
    assert not document(artifacts, result).inline_shapes
    assert "须提供检验报告。" not in [p.text for p in document(artifacts, result).paragraphs]


def test_reexport_retains_prior_document_and_review_gaps_do_not_block(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    outline = {"chapters": [chapter("a", "技术方案", [target("solution")])]}
    units = [writing("task-1", "a", ["solution"], blocks=[{"kind": "paragraph", "text": "第一版草案"}])]
    original_input = deepcopy(units)
    first = export_authoring_document(artifacts, outline, units)
    old_bytes = artifacts.read_bytes(first["docx_ref"])
    units[0]["blocks"][0]["text"] = "第二版草案"
    second = export_authoring_document(artifacts, outline, units, [{
        "code": "response_unverified", "detail": "尚缺具体功能证明。", "chapter_id": "a"}])
    assert first["docx_ref"] != second["docx_ref"]
    assert artifacts.read_bytes(first["docx_ref"]) == old_bytes
    assert original_input[0]["blocks"][0]["text"] == "第一版草案"
    assert any(g["code"] == "response_unverified" for g in second["gaps"])


def test_only_explicit_volumes_produce_separate_documents_and_keep_subtree(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    outline = {"chapters": [
        {**chapter("a", "技术文件", children=[chapter("b", "实施", [target("implementation")])]), "is_volume": True},
        chapter("c", "补充说明", [target("notes")]),
    ]}
    units = [writing("task-1", "b", ["implementation"], blocks=[{"kind": "paragraph", "text": "实施正文"}]),
             writing("task-2", "c", ["notes"], blocks=[{"kind": "paragraph", "text": "说明正文"}])]
    result = export_authoring_document(artifacts, outline, units)
    with ZipFile(BytesIO(artifacts.read_bytes(result["package_ref"]))) as archive:
        manifest = json.loads(archive.read("manifest.json"))
        assert len(manifest["volumes"]) == 1
        volume = manifest["volumes"][0]
        assert volume["chapter_id"] == "a"
        paragraphs = [p.text for p in Document(BytesIO(archive.read(volume["package_path"]))).paragraphs]
        assert "实施正文" in paragraphs and "说明正文" not in paragraphs
    main_paragraphs = [p.text for p in document(artifacts, result).paragraphs]
    assert "实施正文" in main_paragraphs and "说明正文" in main_paragraphs


def test_real_pdf_pages_inserted_and_original_pdf_preserved(tmp_path):
    from pypdf import PdfWriter

    artifacts = ArtifactStore(tmp_path)
    output = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=210, height=297)
    writer.add_blank_page(width=297, height=210)
    writer.write(output)
    ref = artifacts.put_bytes(output.getvalue(), "application/pdf")
    result = export_authoring_document(artifacts, {"chapters": [
        chapter("a", "证明文件", [target("proof", "attachment")])]}, [writing(
            "task-1", "a", ["proof"], attachments=[
                {"material_id": "proof", "title": "两页证明", "artifact_ref": ref}])])
    assert len(document(artifacts, result).inline_shapes) == 2
    assert not any(g["code"] == "export_attachment_not_embedded" for g in result["gaps"])
    with ZipFile(BytesIO(artifacts.read_bytes(result["package_ref"]))) as archive:
        manifest = json.loads(archive.read("manifest.json"))
        assert archive.read(manifest["attachments"][0]["package_path"]) == output.getvalue()


def test_same_template_at_distinct_targets_keeps_each_target_values_and_order(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    ref = source_ref(artifacts, grid=grid())
    outline = {"chapters": [chapter("a", "两标段人员", [
        target("lot-1", "form", [ref], templates=[ref]), target("lot-2", "form", [ref], templates=[ref])])]}
    units = [writing("task-1", "a", ["lot-1"], cell_fills=[
        {"source_ref": ref, "row": 1, "column": 1, "text": "一标段负责人"}]),
        writing("task-2", "a", ["lot-2"], cell_fills=[
            {"source_ref": ref, "row": 1, "column": 1, "text": "二标段负责人"}])]
    result = export_authoring_document(artifacts, outline, units)
    tables = document(artifacts, result).tables
    assert len(tables) == 2
    assert [table.cell(1, 1).text for table in tables] == ["一标段负责人", "二标段负责人"]
    assert not any(g["code"] == "export_conflicting_cell_fills" for g in result["gaps"])


def test_equal_coordinates_in_two_templates_do_not_share_fill_values(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    ref_a = source_ref(artifacts, grid=grid(), document_id="first-template")
    ref_b = source_ref(artifacts, grid=grid(), document_id="second-template")
    outline = {"chapters": [chapter("a", "人员资料", [target(
        "staff", "form", [ref_a, ref_b], templates=[ref_a, ref_b])])]}
    result = export_authoring_document(artifacts, outline, [writing("task-1", "a", ["staff"], cell_fills=[
        {"source_ref": ref_a, "row": 1, "column": 1, "text": "主表姓名"},
        {"source_ref": ref_b, "row": 1, "column": 1, "text": "附件姓名"}])])
    assert [table.cell(1, 1).text for table in document(artifacts, result).tables] == ["主表姓名", "附件姓名"]


def test_empty_successful_unit_remains_a_visible_pending_position(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    result = export_authoring_document(artifacts, {"chapters": [
        chapter("a", "服务方案", [target("service")])]}, [writing("task-1", "a", ["service"])])
    assert any(g["code"] == "export_content_pending" for g in result["gaps"])
    assert any("待补" in p.text for p in document(artifacts, result).paragraphs)


def test_fill_cannot_write_into_another_tasks_template_even_in_same_chapter(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    ref_a = source_ref(artifacts, grid=grid(), document_id="template-a")
    ref_b = source_ref(artifacts, grid=grid(), document_id="template-b")
    outline = {"chapters": [chapter("a", "两个独立表格", [
        target("first", "form", [ref_a], templates=[ref_a]),
        target("second", "form", [ref_b], templates=[ref_b])])]}
    result = export_authoring_document(artifacts, outline, [writing("task-1", "a", ["first"], cell_fills=[
        {"source_ref": ref_b, "row": 1, "column": 1, "text": "不能串入另一个表格"}])])
    assert [table.cell(1, 1).text for table in document(artifacts, result).tables] == ["", ""]
    assert any(g["code"] == "export_cell_fill_unplaced" for g in result["gaps"])


def test_each_chapter_keeps_its_own_reuse_of_a_template(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    ref = source_ref(artifacts, grid=grid())
    outline = {"chapters": [chapter("a", "第一标段", [target("first", "form", [ref], templates=[ref])]),
                             chapter("b", "第二标段", [target("second", "form", [ref], templates=[ref])])]}
    result = export_authoring_document(artifacts, outline, [
        writing("task-1", "a", ["first"], cell_fills=[
            {"source_ref": ref, "row": 1, "column": 1, "text": "甲"}]),
        writing("task-2", "b", ["second"], cell_fills=[
            {"source_ref": ref, "row": 1, "column": 1, "text": "乙"}])])
    assert [table.cell(1, 1).text for table in document(artifacts, result).tables] == ["甲", "乙"]


def test_mixed_blocks_keep_table_introduction_and_misclassified_table_contents(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    result = export_authoring_document(artifacts, {"chapters": [chapter("a", "响应", [target("response")])]}, [
        writing("task-1", "a", ["response"], blocks=[
            {"kind": "table", "text": "以下安排仍待业务确认。", "headers": ["阶段", "措施"], "rows": [["部署", "联调"]]},
            {"kind": "paragraph", "text": "验收还包括以下检查。", "headers": ["项目"], "rows": [["功能测试"]]},
        ])])
    doc = document(artifacts, result)
    assert "以下安排仍待业务确认。" in [p.text for p in doc.paragraphs]
    assert "验收还包括以下检查。" in [p.text for p in doc.paragraphs]
    assert len(doc.tables) == 2 and doc.tables[1].cell(1, 0).text == "功能测试"
    assert any(g["code"] == "block_kind_mismatch" for g in result["gaps"])


def test_packaged_original_is_deduplicated_without_losing_material_occurrences(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    proof = artifacts.put_bytes(b"synthetic immutable proof", "text/plain")
    result = export_authoring_document(artifacts, {"chapters": [
        chapter("a", "证明材料", [target("proof-a", "attachment"), target("proof-b", "attachment")])]}, [
        writing("task-1", "a", ["proof-a"], attachments=[
            {"material_id": "first", "title": "资格证明", "artifact_ref": proof}]),
        writing("task-2", "a", ["proof-b"], attachments=[
            {"material_id": "second", "title": "评分证明", "artifact_ref": proof}]),
    ])
    with ZipFile(BytesIO(artifacts.read_bytes(result["package_ref"]))) as archive:
        manifest = json.loads(archive.read("manifest.json"))
        assert len(manifest["attachments"]) == 1
        assert manifest["attachments"][0]["occurrences"] == [
            {"material_id": "first", "title": "资格证明", "chapter_id": "a"},
            {"material_id": "second", "title": "评分证明", "chapter_id": "a"},
        ]


def test_manual_quote_blanks_preserve_source_numbers_and_same_chapter_technical_content(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    rows = [["招标要求", "预算30万元；最高限价28万元；价格评分10分"], ["投标总价", ""], ["折扣", ""]]
    ref = source_ref(artifacts, grid={"row_count": 3, "column_count": 2, "cells": [
        {"row": row, "column": column, "row_span": 1, "col_span": 1, "text": text}
        for row, cells in enumerate(rows) for column, text in enumerate(cells)]})
    outline = {"chapters": [chapter("a", "商务与技术响应", [
        {**target("quote", "form", [ref], templates=[ref]), "manual_fill": True}, target("technical")])]}
    result = export_authoring_document(artifacts, outline, [
        writing("task-quote", "a", ["quote"], status="placeholder", gaps=["最终报价及折扣由人工填写。"]),
        writing("task-technical", "a", ["technical"], blocks=[{"kind": "paragraph", "text": "技术实施按项目接口要求逐项联调。"}]),
    ])
    doc = document(artifacts, result)
    assert doc.tables[0].cell(0, 1).text == rows[0][1]
    assert doc.tables[0].cell(1, 1).text == "" and doc.tables[0].cell(2, 1).text == ""
    assert "技术实施按项目接口要求逐项联调。" in [paragraph.text for paragraph in doc.paragraphs]
    assert any("人工填写" in gap["detail"] for gap in result["gaps"])
