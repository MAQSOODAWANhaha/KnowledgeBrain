"""Complete native formats survive source selection, assembly and blank export."""
from copy import deepcopy
from io import BytesIO
import json
from zipfile import ZipFile

from docx import Document
import pytest

from docreader.integrations.artifacts import ArtifactStore
from docreader.llm.budget import RequestBudget
from docreader.llm.recovery import _fragments
from docreader.llm.request import ModelBudgetError, ModelContractError, ModelWorkUnitExhausted
from docreader.tender.authoring_export import export_authoring_document
from docreader.tender import authoring_pipeline as pipeline
from docreader.tender.authoring_templates import load_template_part
from test_authoring_templates import create_context, create_source
from test_structural_packing import grid


async def organize(tmp_path, units, select_source, select_assembly):
    artifacts, index_ref, native = create_source(tmp_path, units)
    context = create_context(artifacts, index_ref, native)
    calls = []

    async def execute(stage, request, validate):
        data = json.loads(request['messages'][1]['content'][0]['text'])
        calls.append(data)
        return validate(select_assembly(data) if data.get('operation') == 'assemble' else select_source(data))

    result = await pipeline._organize(artifacts, context, RequestBudget(
        model='fixture', context_tokens=131072), execute, [])
    return artifacts, result, calls


def source_result(targets):
    return {'schema_version': 2, 'chapters': [{'title': '正式填写', 'targets': targets}], 'findings': []}


def source_form(parts=(), **extra):
    return {'kind': 'form', 'purpose': '按完整格式填写', 'rule_ids': [], 'range_ids': [],
            'template_part_ids': list(parts), 'manual_fill': True, **extra}


def assembled(targets):
    return {'chapters': [{'title': '正式填写', 'targets': targets}]}


async def test_continuous_format_span_preserves_middle_pages_and_assembly_inherits_it(tmp_path):
    def select_source(data):
        parts = data['template_parts']
        return source_result([source_form(template_spans=[{
            'start_part_id': parts[0]['part_id'], 'end_part_id': parts[2]['part_id']}])])

    artifacts, outline, calls = await organize(tmp_path, [
        {'text': '函件第一条：____'}, {'text': '跨页固定承诺：完整保留。'},
        {'text': '函件签署栏：____'}, {'text': '另一份无关声明：____'}], select_source,
        lambda data: assembled([{'kind': 'form', 'purpose': '保留完整函件',
                                'target_ids': [data['chapters'][0]['targets'][0]['target_id']]}]))
    target = outline['chapters'][0]['targets'][0]
    text = '\n'.join(load_template_part(artifacts, ref)['text'] for ref in target['template_refs'])
    assert text == '函件第一条：____\n跨页固定承诺：完整保留。\n函件签署栏：____'
    assert len(calls) == 2
    exported = export_authoring_document(artifacts, outline, [])
    document = Document(BytesIO(artifacts.path(exported['docx_ref']).read_bytes()))
    assert '跨页固定承诺：完整保留。' in '\n'.join(p.text for p in document.paragraphs)
    assert '另一份无关声明' not in '\n'.join(p.text for p in document.paragraphs)


async def test_explicit_parts_omit_page_heading_without_deduplicating_substantive_text(tmp_path):
    def select_source(data):
        parts = data['template_parts']
        return source_result([source_form(template_spans=[{
            'start_part_id': parts[0]['part_id'], 'end_part_id': parts[-1]['part_id']}])])

    def select_assembly(data):
        parts = data['template_parts']
        return {'chapters': [{'title': '正式文件', 'is_volume': True, 'targets': [{
            'kind': 'form', 'purpose': '完整保留函件正文及签署栏', 'target_ids': ['t1'],
            'template_part_ids': [part['part_id'] for index, part in enumerate(parts) if index != 1]}]}]}

    texts = ['函件开头：____', '源文件固定页首文字', '应履行的实质条款。',
             '应履行的实质条款。', '函件签署栏：____']
    artifacts, outline, _ = await organize(tmp_path, [{'text': text} for text in texts],
                                          select_source, select_assembly)
    target = outline['chapters'][0]['targets'][0]
    assert [load_template_part(artifacts, ref)['text'] for ref in target['template_refs']] == [
        texts[0], *texts[2:]]
    exported = export_authoring_document(artifacts, outline, [])
    archive = ZipFile(BytesIO(artifacts.read_bytes(exported['package_ref'])))
    manifest = json.loads(archive.read('manifest.json'))
    assert len(manifest['volumes']) == 1
    for data in (artifacts.read_bytes(exported['docx_ref']),
                 archive.read(manifest['volumes'][0]['package_path'])):
        visible = '\n'.join(paragraph.text for paragraph in Document(BytesIO(data)).paragraphs)
        assert texts[1] not in visible
        assert visible.count(texts[2]) == 2
        assert texts[0] in visible and texts[-1] in visible


async def test_assembly_can_select_a_continuation_that_was_only_candidate_context(tmp_path):
    first, second = grid(rows=3), grid(rows=2)
    first['cells'][0]['text'] = '企业基本情况'
    second['cells'][0]['text'] = '设备续填与获奖情况'

    def select_source(data):
        parts = data['template_parts']
        return source_result([source_form([parts[0]['part_id']],
            range_ids=[s['range_id'] for s in data['sources']])])

    def select_assembly(data):
        parts = data['template_parts']
        assert len(parts) == 2  # The unselected but explicitly referenced continuation is reachable.
        return assembled([{'kind': 'form', 'purpose': '完整的两页企业资料', 'target_ids': ['t1'],
            'template_part_ids': [p['part_id'] for p in parts]}])

    artifacts, outline, _ = await organize(tmp_path, [{'grid': first}, {'grid': second}],
                                         select_source, select_assembly)
    target = outline['chapters'][0]['targets'][0]
    assert len(target['template_refs']) == 2
    exported = export_authoring_document(artifacts, outline, [])
    document = Document(BytesIO(artifacts.path(exported['docx_ref']).read_bytes()))
    assert len(document.tables) == 2
    assert '设备续填与获奖情况' in document.tables[1].cell(0, 0).text


async def test_merging_chapters_cannot_collapse_two_explicit_subject_instances(tmp_path):
    def select_source(data):
        identity = data['template_parts'][0]['part_id']
        return source_result([source_form([identity], subject='申请主体'),
                              source_form([identity], subject='生产主体')])

    artifacts, outline, _ = await organize(tmp_path, [{'grid': grid(rows=3)}], select_source,
        lambda data: assembled([{'kind': 'form', 'purpose': '分别填写财务情况', 'target_ids': ['t1', 't2']}]))
    targets = outline['chapters'][0]['targets']
    assert [t['subject'] for t in targets] == ['申请主体', '生产主体']
    assert len({t['target_id'] for t in targets}) == 2
    exported = export_authoring_document(artifacts, outline, [])
    document = Document(BytesIO(artifacts.path(exported['docx_ref']).read_bytes()))
    assert len(document.tables) == 2
    assert all(role in '\n'.join(p.text for p in document.paragraphs) for role in ['申请主体', '生产主体'])


async def test_external_operation_remains_a_note_and_cannot_create_a_fourth_volume(tmp_path):
    def select_source(data):
        return source_result([source_form([data['template_parts'][0]['part_id']])])

    def select_assembly(data):
        return {'chapters': [
            {'title': '正式文件', 'is_volume': True, 'targets': [
                {'kind': 'form', 'purpose': '提交正文', 'target_ids': ['t1']}]},
            {'title': '平台操作', 'presentation_role': 'submission_note', 'is_volume': True,
             'targets': [{'kind': 'narrative', 'purpose': '将现有文件的可编辑版本上传指定位置',
                          'target_ids': ['t1']}]}]}

    artifacts, outline, _ = await organize(tmp_path, [{'text': '完整函件：____'}],
                                         select_source, select_assembly)
    tasks = pipeline.plan_writing_tasks(outline)
    assert all(task['chapter_id'] != outline['chapters'][1]['chapter_id'] for task in tasks)
    exported = export_authoring_document(artifacts, outline, [])
    document = Document(BytesIO(artifacts.path(exported['docx_ref']).read_bytes()))
    assert '平台操作' not in '\n'.join(p.text for p in document.paragraphs)
    with ZipFile(artifacts.path(exported['package_ref'])) as archive:
        assert len([name for name in archive.namelist() if name.startswith('volumes/')]) == 1
        assert '将现有文件的可编辑版本上传指定位置' in archive.read('提交要求与待办.md').decode()


@pytest.mark.parametrize('field', ['template_part_ids', 'template_spans'])
async def test_assembly_explicit_empty_template_selection_keeps_original_source(tmp_path, field):
    def select_source(data):
        part = data['template_parts'][0]
        return source_result([source_form([part['part_id']], range_ids=part['range_ids'])])

    artifacts, outline, _ = await organize(tmp_path, [{'text': '这是解释依据，并非应当填写的原格式。'}],
        select_source, lambda data: assembled([{'kind': 'narrative', 'purpose': '按原要求说明',
                                                'target_ids': ['t1'], field: []}]))
    target = outline['chapters'][0]['targets'][0]
    assert target['template_refs'] == []
    assert [ref['unit_id'] for ref in target['source_refs']] == ['unit:0']
    assert [ref['source_ref']['unit_id'] for ref in target['source_excerpts']] == ['unit:0']
    exported = export_authoring_document(artifacts, outline, [])
    document = Document(BytesIO(artifacts.path(exported['docx_ref']).read_bytes()))
    assert '这是解释依据，并非应当填写的原格式。' not in '\n'.join(p.text for p in document.paragraphs)


async def test_subject_split_keeps_unknown_candidate_sources_and_templates_in_one_instance(tmp_path):
    def select_source(data):
        return source_result([source_form([part['part_id']], subject=subject, range_ids=part['range_ids'])
            for part, subject in zip(data['template_parts'], ['申请主体', '生产主体', None])])

    _, outline, _ = await organize(tmp_path, [
        {'text': '申请主体声明：____'}, {'text': '生产主体声明：____'}, {'text': '主体待核的独立声明：____'}],
        select_source, lambda data: assembled([{'kind': 'form', 'purpose': '各主体分别填写',
            'target_ids': ['t1', 't2', 't3'], 'subject': '申请主体'}]))
    targets = outline['chapters'][0]['targets']
    assert len(targets) == 3
    by_subject = {target.get('subject'): target for target in targets}
    assert set(by_subject) == {'申请主体', '生产主体', None}
    for subject, unit in [('申请主体', 'unit:0'), ('生产主体', 'unit:1'), (None, 'unit:2')]:
        assert [ref['unit_id'] for ref in by_subject[subject]['source_refs']] == [unit]
        assert [ref['source_ref']['unit_id'] for ref in by_subject[subject]['template_refs']] == [unit]


async def test_subject_split_does_not_hide_an_unknown_candidate_reference(tmp_path):
    raw = assembled([{'kind': 'form', 'purpose': '分别填写',
                      'target_ids': ['t1', 't2', 't999'], 'subject': '申请主体'}])

    def select_source(data):
        return source_result([source_form([data['template_parts'][0]['part_id']], subject='申请主体'),
                              source_form([data['template_parts'][0]['part_id']], subject='生产主体')])

    with pytest.raises(ModelContractError) as error:
        await organize(tmp_path, [{'text': '主体各自提交同一规定格式：____'}],
                       select_source, lambda data: deepcopy(raw))
    path = ['chapters', 0, 'targets', 0, 'target_ids', 2]
    assert error.value.errors[0]['path'] == path
    assert error.value.errors[0]['code'] == 'outline_unknown_reference'
    assert _fragments(raw, error.value.errors) == [{'path': path, 'value': 't999'}]


@pytest.mark.parametrize('unknown_count', [1, 2])
async def test_one_known_subject_does_not_absorb_independent_unknown_instances(tmp_path, unknown_count):
    def select_source(data):
        return source_result([source_form([part['part_id']], subject='甲主体' if number == 0 else None)
            for number, part in enumerate(data['template_parts'])])

    _, outline, _ = await organize(tmp_path, [{'text': '甲主体填写格式：____'},
        *({'text': f'主体待核的第{number}份独立格式：____'} for number in range(unknown_count))],
        select_source, lambda data: assembled([{'kind': 'form', 'purpose': '保留各填写实例',
            'target_ids': [target['target_id'] for target in data['chapters'][0]['targets']]}]))
    targets = outline['chapters'][0]['targets']
    assert len(targets) == unknown_count + 1
    assert [target.get('subject') for target in targets] == ['甲主体', *([None] * unknown_count)]
    for number, target in enumerate(targets):
        assert [ref['source_ref']['unit_id'] for ref in target['template_refs']] == [f'unit:{number}']


@pytest.mark.parametrize('selection,expected_instances', [
    ({'subject': '申请主体'}, 1), ({}, 2), ({'subject': None}, 2), ({'subject': '其他主体'}, 2)])
async def test_explicit_matching_subject_resolves_unknown_candidate_without_losing_sources(
        tmp_path, selection, expected_instances):
    def select_source(data):
        part = data['template_parts'][0]['part_id']
        return source_result([
            source_form([part], subject='申请主体', range_ids=data['template_parts'][0]['range_ids']),
            source_form([part], range_ids=data['template_parts'][-1]['range_ids'])])

    artifacts, outline, _ = await organize(tmp_path, [
        {'grid': grid(rows=3)}, {'text': '本格式的共同填写条件须保留。'}], select_source,
        lambda data: assembled([{'kind': 'form', 'purpose': '确认是同一主体的一份填写格式',
                                'target_ids': ['t1', 't2'], **selection}]))
    targets = outline['chapters'][0]['targets']
    assert len(targets) == expected_instances
    assert {ref['unit_id'] for target in targets for ref in target['source_refs']} == {'unit:0', 'unit:1'}
    if expected_instances == 1:
        assert targets[0]['subject'] == '申请主体'
        assert len(targets[0]['template_refs']) == 1
        assert {ref['source_ref']['unit_id'] for ref in targets[0]['source_excerpts']} == {'unit:0', 'unit:1'}
    else:
        assert [target.get('subject') for target in targets] == ['申请主体', None]
    exported = export_authoring_document(artifacts, outline, [])
    document = Document(BytesIO(artifacts.read_bytes(exported['docx_ref'])))
    assert len(document.tables) == expected_instances


@pytest.mark.parametrize('old_chapter_present,mixed_unit', [(True, False), (False, False), (True, True)])
def test_historical_body_content_moved_to_submission_note_does_not_flow_back(tmp_path, old_chapter_present, mixed_unit):
    artifacts = ArtifactStore(tmp_path)
    active_id = 'old-body' if old_chapter_present else 'new-body'
    body_target = {'target_id': 'body-target', 'kind': 'narrative', 'purpose': '保留正式正文',
                   'rule_refs': [], 'source_refs': []}
    note_target = {'target_id': 'moved-target', 'kind': 'narrative', 'purpose': '平台上传动作',
                   'rule_refs': [], 'source_refs': []}
    outline = {'chapters': [
        {'chapter_id': active_id, 'title': '正式正文', 'targets': [body_target], 'children': []},
        {'chapter_id': 'note', 'title': '外部操作', 'presentation_role': 'submission_note',
         'is_volume': True, 'targets': [note_target], 'children': []}]}
    historical = {'task_id': 'historical', 'chapter_id': 'old-body', 'target_ids': ['moved-target'],
        'status': 'generated', 'blocks': [{'kind': 'paragraph', 'text': '历史平台动作不得回流正文',
                                         'target_ids': ['moved-target']}],
        'cell_fills': [], 'row_fills': [], 'attachments': [], 'gaps': []}
    current = {'task_id': 'current', 'chapter_id': active_id, 'target_ids': ['body-target'],
        'status': 'generated', 'blocks': [{'kind': 'paragraph', 'text': '正确业务正文必须保留',
                                         'target_ids': ['body-target']}],
        'cell_fills': [], 'row_fills': [], 'attachments': [], 'gaps': []}
    if mixed_unit:
        historical['target_ids'].append('body-target')
        historical['blocks'].extend(current['blocks'])
        units = [historical]
    else:
        units = [historical, current]
    before = deepcopy(units)
    exported = export_authoring_document(artifacts, outline, units)
    document = Document(BytesIO(artifacts.read_bytes(exported['docx_ref'])))
    text = '\n'.join(paragraph.text for paragraph in document.paragraphs)
    assert '历史平台动作不得回流正文' not in text
    assert '正确业务正文必须保留' in text
    assert '补充内容' not in text
    assert units == before
    with ZipFile(artifacts.path(exported['package_ref'])) as archive:
        assert not any(name.startswith('volumes/') for name in archive.namelist())
        assert '平台上传动作' in archive.read('提交要求与待办.md').decode()


@pytest.mark.parametrize('assembly_outcome', ['omitted', 'exhausted', 'over_budget'])
async def test_source_note_subtree_remains_a_todo_when_assembly_cannot_place_it(tmp_path, assembly_outcome):
    source_text = '文件提交后，须在平台维护上传的可编辑版本。'
    purpose = '在平台维护既有文件的可编辑版本'

    def select_source(data):
        return {'schema_version': 2, 'findings': [], 'chapters': [
            {'title': '平台操作', 'presentation_role': 'submission_note', 'targets': [], 'children': [
                {'title': '维护动作', 'is_volume': True, 'targets': [
                    {'kind': 'narrative', 'purpose': purpose, 'rule_ids': [],
                     'range_ids': [data['sources'][0]['range_id']]}]}]}]}

    def select_assembly(data):
        if assembly_outcome == 'exhausted':
            raise ModelWorkUnitExhausted('outline')
        if assembly_outcome == 'over_budget':
            raise ModelBudgetError('synthetic assembly capacity exceeded')
        return {'chapters': []}

    artifacts, outline, _ = await organize(tmp_path, [{'text': source_text}], select_source, select_assembly)
    assert pipeline.plan_writing_tasks(outline) == []
    targets = [target for chapter in pipeline._flatten(outline['chapters']) for target in chapter['targets']]
    assert len(targets) == 1 and targets[0]['purpose'] == purpose
    assert [ref['unit_id'] for ref in targets[0]['source_refs']] == ['unit:0']
    exported = export_authoring_document(artifacts, outline, [])
    document = Document(BytesIO(artifacts.path(exported['docx_ref']).read_bytes()))
    body_text = '\n'.join(p.text for p in document.paragraphs)
    assert purpose not in body_text and source_text not in body_text
    with ZipFile(artifacts.path(exported['package_ref'])) as archive:
        assert not any(name.startswith('volumes/') for name in archive.namelist())
        notes = archive.read('提交要求与待办.md').decode()
        assert purpose in notes
        assert source_text in notes
