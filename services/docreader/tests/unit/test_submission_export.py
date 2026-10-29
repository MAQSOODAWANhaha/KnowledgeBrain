"""Physical submission contents follow selected form parts and business instances."""

from copy import deepcopy
from io import BytesIO
import json
from zipfile import ZipFile

from docx import Document
import pytest

from docreader.integrations.artifacts import ArtifactStore
from docreader.tender import authoring_export
from docreader.tender.evidence.references import EvidenceRepository, SourceIndex, fragment, selector


def source(artifacts, identity, *, text="", grid=None):
    unit = {"unit_id": identity, "kind": "table_region" if grid else "section",
            "ordinal": 0, "text": text, "grid": grid,
            "locator": {"locator_kind": "document", "section_ordinal": 0}}
    raw = {"schema_version": 2, "units": [unit], "images": {}, "ocr": {},
           "source_contract": {"page_manifest": [], "units": [{
               "unit_id": identity, "section_id": "section:0", "parent_section_id": None,
               "heading_path": "原格式", "physical_locator": unit["locator"]}]}}
    document = {"document_id": identity, "document_version": "1",
                "original_ref": artifacts.put_bytes(b"fixture original")}
    index = SourceIndex(document, raw)
    reference = artifacts.put_json(index.descriptor(artifacts.put_json(index.raw)))
    return {"source_index_ref": reference, "unit_id": identity}


def part(reference, identity=None):
    return {"source_ref": reference, "selectors": [] if identity is None else [{
        "range_id": identity, "source_range_id": identity, "start": 0, "end": 1}]}


def target(identity, *, refs=(), templates=(), kind="form", manual=False):
    return {"target_id": identity, "kind": kind, "purpose": "填写规定格式",
            "source_refs": list(refs), "template_refs": list(templates),
            "rule_refs": [], "manual_fill": manual}


def chapter(identity, targets=(), children=(), *, volume=False):
    return {"chapter_id": identity, "title": identity, "targets": list(targets),
            "children": list(children), "is_volume": volume}


def writing(identity, chapter_id, targets, **fields):
    return {"task_id": identity, "chapter_id": chapter_id, "target_ids": list(targets),
            "status": "generated", "blocks": [], "cell_fills": [], "attachments": [],
            "gaps": [], **fields}


def grid():
    return {"row_count": 3, "column_count": 2, "cells": [
        {"row": 0, "column": 0, "row_span": 1, "col_span": 2, "text": "规定表头"},
        {"row": 1, "column": 0, "row_span": 1, "col_span": 1, "text": "名称"},
        {"row": 1, "column": 1, "row_span": 1, "col_span": 1, "text": ""},
        {"row": 2, "column": 0, "row_span": 1, "col_span": 1, "text": "金额"},
        {"row": 2, "column": 1, "row_span": 1, "col_span": 1, "text": ""},
    ]}


@pytest.fixture
def parts(monkeypatch):
    """Isolate export from the separately tested frozen-part selection loader."""
    selected = {}

    def key(value):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)

    def load(artifacts, value, *, repository=None):
        if key(value) not in selected:
            raise ValueError("selected template part is unavailable")
        return deepcopy(selected[key(value)])

    def register(value, *, text="", grid=None):
        selected[key(value)] = {"document": {}, "unit": {"text": text, "grid": grid}}
        return value

    monkeypatch.setattr(authoring_export, "load_template_part", load, raising=False)
    monkeypatch.setattr(authoring_export, "template_part_key", key, raising=False)
    return register


def exported(artifacts, chapters, units=()):
    result = authoring_export.export_authoring_document(artifacts, {"chapters": chapters}, list(units))
    return result, Document(BytesIO(artifacts.read_bytes(result["docx_ref"])))


def test_context_reference_is_not_a_form_and_missing_template_stays_blank(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    context = source(artifacts, "evaluation", text="评审说明不得进入投标文件", grid=grid())
    result, word = exported(artifacts, [chapter("商务", [target("declaration", refs=[context])])])
    assert not word.tables
    assert "评审说明不得进入投标文件" not in [p.text for p in word.paragraphs]
    assert "（原格式待补）" in [p.text for p in word.paragraphs]
    assert any(g["code"] == "export_template_missing" for g in result["gaps"])


def test_selected_text_part_excludes_other_forms_in_same_unit(tmp_path, parts):
    artifacts = ArtifactStore(tmp_path)
    reference = source(artifacts, "two-forms", text="前一声明尾部\n本声明固定文字\n后一声明开头")
    selected = parts(part(reference, "this-form"), text="本声明固定文字")
    _, word = exported(artifacts, [chapter("声明", [target("declaration", refs=[reference], templates=[selected])])])
    values = [p.text for p in word.paragraphs]
    assert values.count("本声明固定文字") == 1
    assert not any("前一声明" in value or "后一声明" in value for value in values)


def test_grid_part_does_not_render_unit_summary_or_context_grid(tmp_path, parts):
    artifacts = ArtifactStore(tmp_path)
    reference = source(artifacts, "native-grid", text="整个来源unit的汇总文字", grid=grid())
    context = source(artifacts, "context", text="评审表", grid=grid())
    selected = parts(part(reference), grid=grid())
    _, word = exported(artifacts, [chapter("表格", [target("form", refs=[context, reference], templates=[selected])])])
    assert len(word.tables) == 1
    assert not any("来源unit" in p.text or p.text == "评审表" for p in word.paragraphs)


def test_exact_parts_deduplicate_only_within_one_instance_and_keep_order(tmp_path, parts):
    artifacts = ArtifactStore(tmp_path)
    reference = source(artifacts, "text", text="整段原文")
    first = parts(part(reference, "first"), text="第一部分")
    second = parts(part(reference, "second"), text="第二部分")
    _, word = exported(artifacts, [chapter("声明", [target(
        "one-instance", refs=[reference], templates=[second, first, deepcopy(second)])])])
    values = [p.text for p in word.paragraphs]
    assert values.count("第一部分") == values.count("第二部分") == 1
    assert values.index("第二部分") < values.index("第一部分")


def test_same_template_at_distinct_subjects_keeps_separate_fills(tmp_path, parts):
    artifacts = ArtifactStore(tmp_path)
    reference = source(artifacts, "financial-form", grid=grid())
    selected = parts(part(reference), grid=grid())
    targets = [target("bidder", templates=[selected]), target("manufacturer", templates=[selected])]
    units = [writing("w1", "财务", ["bidder"], cell_fills=[
        {"source_ref": reference, "row": 1, "column": 1, "text": "投标人资料"}]),
        writing("w2", "财务", ["manufacturer"], cell_fills=[
            {"source_ref": reference, "row": 1, "column": 1, "text": "制造商资料"}])]
    _, word = exported(artifacts, [chapter("财务", targets)], units)
    assert len(word.tables) == 2
    assert [t.cell(1, 1).text for t in word.tables] == ["投标人资料", "制造商资料"]


def test_fill_must_belong_to_selected_template_not_context_reference(tmp_path, parts):
    artifacts = ArtifactStore(tmp_path)
    reference = source(artifacts, "selected", grid=grid())
    context = source(artifacts, "not-selected", grid=grid())
    selected = parts(part(reference), grid=grid())
    result, word = exported(artifacts, [chapter("表格", [target(
        "form", refs=[context], templates=[selected])])], [writing("w", "表格", ["form"], cell_fills=[
            {"source_ref": reference, "row": 1, "column": 1, "text": "有据填写"},
            {"source_ref": context, "row": 2, "column": 1, "text": "不能串入"}])])
    assert len(word.tables) == 1
    assert word.tables[0].cell(1, 1).text == "有据填写"
    assert word.tables[0].cell(2, 1).text == ""
    assert any(g["code"] == "export_cell_fill_unplaced" for g in result["gaps"])


def test_manual_form_keeps_fixed_grid_and_empty_values_even_if_unit_has_fills(tmp_path, parts):
    artifacts = ArtifactStore(tmp_path)
    reference = source(artifacts, "price", grid=grid())
    selected = parts(part(reference), grid=grid())
    result, word = exported(artifacts, [chapter("报价", [target(
        "price", templates=[selected], manual=True)])], [writing("w", "报价", ["price"],
            blocks=[{"kind": "paragraph", "text": "不得自动采用的报价", "target_ids": ["price"]}],
            cell_fills=[{"source_ref": reference, "row": 2, "column": 1, "text": "999"}])])
    assert len(word.tables) == 1
    assert word.tables[0].cell(0, 0).text == "规定表头"
    assert word.tables[0].cell(0, 0)._tc is word.tables[0].cell(0, 1)._tc
    assert word.tables[0].cell(2, 1).text == ""
    assert "不得自动采用的报价" not in [p.text for p in word.paragraphs]
    assert any(g["code"] == "export_manual_content_omitted" for g in result["gaps"])


def test_unavailable_part_does_not_fall_back_to_whole_unit_or_generated_declaration(tmp_path, parts):
    artifacts = ArtifactStore(tmp_path)
    reference = source(artifacts, "source", text="不得兜底输出的整段声明")
    unavailable = part(reference, "missing-part")
    result, word = exported(artifacts, [chapter("声明", [target(
        "form", refs=[reference], templates=[unavailable])])], [writing("w", "声明", ["form"],
            blocks=[{"kind": "paragraph", "text": "不得另写完整声明"}])])
    values = [p.text for p in word.paragraphs]
    assert "不得兜底输出的整段声明" not in values
    assert "不得另写完整声明" not in values
    assert "（原格式待补）" in values
    assert any(g["code"] == "export_source_unavailable" for g in result["gaps"])


def test_non_grid_template_does_not_duplicate_generated_form_blocks_or_hide_narrative(tmp_path, parts):
    artifacts = ArtifactStore(tmp_path)
    reference = source(artifacts, "letter", text="原规定声明")
    selected = parts(part(reference), text="原规定声明")
    targets = [target("form", templates=[selected]), target("explanation", kind="narrative")]
    units = [writing("mixed", "声明", ["form", "explanation"], blocks=[
        {"kind": "paragraph", "text": "重复重写的完整声明", "target_ids": ["form"]},
        {"kind": "paragraph", "text": "独立说明内容", "target_ids": ["explanation"]}])]
    result, word = exported(artifacts, [chapter("声明", targets)], units)
    values = [p.text for p in word.paragraphs]
    assert values.count("原规定声明") == 1
    assert "重复重写的完整声明" not in values
    assert values.count("独立说明内容") == 1
    assert any(g["code"] == "export_form_content_unplaced" for g in result["gaps"])


def test_main_and_volume_render_same_selected_instance_and_preserve_other_volume(tmp_path, parts):
    artifacts = ArtifactStore(tmp_path)
    reference = source(artifacts, "price-grid", grid=grid())
    selected = parts(part(reference), grid=grid())
    outline = [chapter("商务", [target("business", kind="narrative")], volume=True),
               chapter("报价", children=[chapter("价格表", [target("price", templates=[selected], manual=True)])], volume=True)]
    result, word = exported(artifacts, outline, [writing("w", "商务", ["business"], blocks=[
        {"kind": "paragraph", "text": "商务说明"}])])
    assert len(word.tables) == 1
    with ZipFile(BytesIO(artifacts.read_bytes(result["package_ref"]))) as archive:
        manifest = json.loads(archive.read("manifest.json"))
        volumes = {v["chapter_id"]: Document(BytesIO(archive.read(v["package_path"]))) for v in manifest["volumes"]}
        assert len(volumes["商务"].tables) == 0
        assert len(volumes["报价"].tables) == 1
        assert [c.text for row in volumes["报价"].tables[0].rows for c in row.cells] == [
            c.text for row in word.tables[0].rows for c in row.cells]
        assert "商务说明" not in [p.text for p in volumes["报价"].paragraphs]


def test_real_frozen_selector_and_export_keep_only_the_selected_declaration(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    original_text = "声明甲🙂：____\n声明乙：____\n下一份格式"
    reference = source(artifacts, "declarations", text=original_text)
    native = EvidenceRepository(artifacts).source_index(reference["source_index_ref"])
    base = native.ranges[0]
    start = original_text.index("声明乙")
    selected = {"source_ref": reference, "selectors": [selector(fragment(
        base, start, start + len("声明乙：____")))]}
    _, word = exported(artifacts, [chapter("商务", [target(
        "declaration", refs=[reference], templates=[selected])])])
    values = [p.text for p in word.paragraphs]
    assert values.count("声明乙：____") == 1
    assert not any("声明甲" in value or "下一份格式" in value for value in values)
    assert EvidenceRepository(artifacts).source_index(reference["source_index_ref"]).raw["units"][0]["text"] == original_text


def test_real_whole_grid_selection_keeps_merges_and_fills_without_flattened_copy(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    reference = source(artifacts, "grid", text="不得重复输出的表格文本副本", grid=grid())
    result, word = exported(artifacts, [chapter("商务", [target(
        "form", refs=[reference], templates=[part(reference)])])], [writing("w", "商务", ["form"],
            cell_fills=[{"source_ref": reference, "row": 1, "column": 1, "text": "有据主体名称"},
                        {"source_ref": reference, "row": 0, "column": 0, "text": "不可覆盖表头"}])])
    assert len(word.tables) == 1
    table = word.tables[0]
    assert table.cell(0, 0).text == "规定表头"
    assert table.cell(0, 0)._tc is table.cell(0, 1)._tc
    assert table.cell(1, 1).text == "有据主体名称"
    assert table.cell(2, 1).text == ""
    assert "不得重复输出的表格文本副本" not in [p.text for p in word.paragraphs]
    assert any(g["code"] == "export_fixed_cell_preserved" for g in result["gaps"])


def test_distinct_template_parts_share_one_native_index_for_the_export(tmp_path, monkeypatch):
    from docreader.tender.evidence.references import EvidenceRepository, SourceIndex, fragment, selector

    artifacts = ArtifactStore(tmp_path)
    text = "声明甲：____\n声明乙：____"
    reference = source(artifacts, "declarations", text=text)
    base = EvidenceRepository(artifacts).source_index(reference["source_index_ref"]).ranges[0]
    selections = [{"source_ref": reference, "selectors": [selector(fragment(base, start, end))]}
                  for start, end in ((0, len("声明甲：____")), (len("声明甲：____\n"), len(text)))]
    calls = []
    original = SourceIndex.__init__

    def build(index, document, raw):
        calls.append(document["document_id"])
        original(index, document, raw)

    monkeypatch.setattr(SourceIndex, "__init__", build)
    _, word = exported(artifacts, [chapter("商务", [
        target("first", templates=[selections[0]]), target("second", templates=[selections[1]])])])
    assert len(calls) == 1
    assert [paragraph.text for paragraph in word.paragraphs].count("声明甲：____") == 1
    assert [paragraph.text for paragraph in word.paragraphs].count("声明乙：____") == 1
