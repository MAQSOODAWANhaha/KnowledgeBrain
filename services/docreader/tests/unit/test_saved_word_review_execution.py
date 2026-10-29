"""Saved-file reviews retain their real input, receipts and publication fences."""
import asyncio
from copy import deepcopy
import hashlib
from io import BytesIO
import json
import threading
from types import SimpleNamespace
from unittest.mock import AsyncMock

from docx import Document
import pytest

from docreader.integrations.artifacts import ArtifactIntegrityError, ArtifactStore
from docreader.integrations.store import Cancelled, Conflict, StaleExecution
from docreader.llm.budget import RequestBudget
from docreader.llm.request import ModelContractError, ModelWorkUnitExhausted
from docreader.tender import document_execution, saved_word_review
from docreader.tender.schema import AuthoringInput, AuthoringOutline
from docreader.tests.unit.test_saved_word_review import saved_bytes


@pytest.fixture
def runtime(tmp_path, monkeypatch):
    artifacts = ArtifactStore(tmp_path)
    snapshot = AuthoringInput(schema_version=2, run_id="saved-review-run", source_version=1,
                              domain_revision=1).model_dump()
    outline = AuthoringOutline(schema_version=2, authoring_input=snapshot, issues=[], chapters=[{
        "chapter_id": "chapter-original", "title": "旧标题", "targets": [], "children": [],
    }]).model_dump()
    document = {"outline_ref": artifacts.put_json(outline), "basis_ref": artifacts.put_json({"materials": []}),
                "unit_refs": [artifacts.put_json({"blocks": [{"text": "已过时的生成正文不得用于保存稿复核"}]})]}
    data, rows, calls = saved_bytes(), {}, []
    state = SimpleNamespace(hook=None, response=None, release_error=None)
    store = SimpleNamespace(setup=AsyncMock(), heartbeat=AsyncMock())

    async def get(execution_id, **_kwargs):
        return deepcopy(rows[execution_id])

    async def submit(execution_id, authoring_input, spec):
        if execution_id in rows:
            row = rows[execution_id]
            if row["authoring_input"] != authoring_input or row["execution_spec"] != spec:
                raise Conflict("execution_id already identifies another immutable review")
        else:
            rows[execution_id] = {"authoring_input": deepcopy(authoring_input), "execution_spec": deepcopy(spec),
                "status": "queued", "cancel_requested": False, "requests": [], "receipts": [], "epoch": 0}

    async def acquire(execution_id, **_kwargs):
        row = rows[execution_id]
        if row["status"] == "succeeded" or row["cancel_requested"]:
            return None
        row.update(status="running", epoch=row["epoch"] + 1)
        return SimpleNamespace(execution_id=execution_id, epoch=row["epoch"])

    async def release(lease, status, **kwargs):
        if status == "succeeded" and state.release_error:
            raise state.release_error
        rows[lease.execution_id].update(status=status, **kwargs)

    async def receipt(lease, stage, scope, **kwargs):
        return next((entry for entry in rows[lease.execution_id]["receipts"]
                     if entry["stage"] == stage and entry["scope"] == scope
                     and entry["input_digest"] == kwargs["input_digest"]), None)

    async def commit(lease, stage, artifact_ref, **kwargs):
        rows[lease.execution_id]["receipts"].append({"stage": stage, "artifact_ref": artifact_ref, **kwargs})

    store.get_execution, store.submit, store.acquire, store.release = get, submit, acquire, release
    store.receipt, store.commit_artifact = receipt, commit
    monkeypatch.setattr(saved_word_review, "AuthoringStore", lambda _dsn: store)
    active = AsyncMock()
    monkeypatch.setattr(saved_word_review, "assert_authoring_active", active)

    class Client:
        def __init__(self, config, *_args, **_kwargs):
            self.config = config

        async def __aenter__(self):
            return self

        async def __aexit__(self, *_args):
            pass

    class Model:
        def __init__(self, *_args):
            pass

        async def __call__(self, request, validate):
            payload = json.loads(request["messages"][1]["content"][0]["text"])
            calls.append(payload)
            if state.hook:
                await state.hook()
            return validate(deepcopy(state.response or {"findings": []}))

    monkeypatch.setattr(saved_word_review, "OpenAIModel", Client)
    monkeypatch.setattr(document_execution, "_DocumentModel", Model)
    options = {"execution_id": "saved-review", "snapshot": snapshot,
        "document_ref": artifacts.put_json(document),
        "saved_ref": artifacts.put_bytes(data, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        "version_id": "saved-version-one", "docx_sha256": hashlib.sha256(data).hexdigest(),
        "reason": "核对人工填写的工期", "heartbeat_seconds": 20}
    budget = RequestBudget(context_tokens=131072, model="fixture", prompts={})

    async def execute(**changes):
        return await saved_word_review.execute_saved_review(SimpleNamespace(dsn="fixture"), artifacts,
            SimpleNamespace(chat=SimpleNamespace(model="fixture", context_tokens=131072, base_url="http://fixture")),
            budget, **{**options, **changes})

    return SimpleNamespace(artifacts=artifacts, snapshot=snapshot, document=document, budget=budget,
        data=data, options=options, execute=execute, rows=rows, calls=calls, state=state, active=active, store=store)


async def test_review_uses_saved_human_text_and_tables_with_reason_and_no_generated_units(runtime, monkeypatch):
    original_read = runtime.artifacts.read_json

    def read(reference):
        assert reference not in runtime.document["unit_refs"], "Old generated WritingUnit is not review input."
        return original_read(reference)

    monkeypatch.setattr(runtime.artifacts, "read_json", read)
    report = original_read(await runtime.execute())
    actual = str([record["actual_content"] for call in runtime.calls for record in call["records"]])
    assert "人工修订后的实施内容" in actual and "人工填写：30天" in actual
    assert "此处也是实际保存内容" in actual and "已过时的生成正文" not in actual
    assert all(call["reason"] == runtime.options["reason"] for call in runtime.calls)
    assert report["version_id"] == runtime.options["version_id"]
    assert report["docx_sha256"] == runtime.options["docx_sha256"]
    assert report["record_count"] == report["reviewed_record_count"] > 0
    assert any(gap["code"] == "saved_word_unbound_content" for gap in report["gaps"])
    assert any(gap["code"] == "saved_word_review_scope" for gap in report["gaps"])


async def test_publication_failure_reuses_same_execution_accepted_response(runtime, monkeypatch):
    original_put = runtime.artifacts.put_json
    failed = False

    def put(value):
        nonlocal failed
        if "record_count" in value and not failed:
            failed = True
            raise OSError("synthetic report publication failure")
        return original_put(value)

    monkeypatch.setattr(runtime.artifacts, "put_json", put)
    with pytest.raises(OSError, match="publication failure"):
        await runtime.execute()
    assert runtime.rows["saved-review"]["status"] == "failed"
    assert len(runtime.rows["saved-review"]["receipts"]) == len(runtime.calls) == 1
    reference = await runtime.execute()
    assert runtime.artifacts.read_json(reference)["reviewed_record_count"] > 0
    assert len(runtime.calls) == 1
    assert await runtime.execute() == reference
    assert len(runtime.calls) == 1


@pytest.mark.parametrize("field,value", [
    ("version_id", "another-version"), ("docx_sha256", "0" * 64), ("reason", "另一个复核理由"),
    ("document_ref", {"changed": True}), ("saved_ref", {"changed": True}),
])
async def test_execution_identity_cannot_rebind_version_hash_reason_or_input(runtime, field, value):
    await runtime.execute()
    with pytest.raises(Conflict, match="immutable review"):
        await runtime.execute(**{field: value})
    assert len(runtime.calls) == 1


async def test_saved_hash_mismatch_fails_before_model_send(runtime):
    with pytest.raises(ArtifactIntegrityError, match="input changed"):
        await runtime.execute(docx_sha256="0" * 64)
    assert runtime.calls == [] and runtime.rows["saved-review"]["status"] == "failed"


@pytest.mark.parametrize("error", [Cancelled("parent cancelled"), StaleExecution("lease replaced")])
async def test_heartbeat_cancellation_or_lease_loss_cancels_work_and_prevents_success(runtime, error):
    started, cancelled = asyncio.Event(), asyncio.Event()

    async def block_model():
        started.set()
        try:
            await asyncio.Event().wait()
        finally:
            cancelled.set()

    async def fail_heartbeat(*_args, **_kwargs):
        await started.wait()
        raise error

    runtime.state.hook = block_model
    runtime.store.heartbeat.side_effect = fail_heartbeat
    with pytest.raises(type(error), match=str(error)):
        await asyncio.wait_for(runtime.execute(heartbeat_seconds=0.001), timeout=3)
    assert cancelled.is_set()
    assert runtime.rows["saved-review"]["status"] != "succeeded"
    assert "result_ref" not in runtime.rows["saved-review"]


async def test_parent_cancellation_at_publish_cannot_return_completed_report(runtime):
    runtime.state.release_error = Cancelled("parent cancelled before publish")
    with pytest.raises(Cancelled, match="before publish"):
        await runtime.execute()
    assert runtime.rows["saved-review"]["status"] != "succeeded"
    assert "result_ref" not in runtime.rows["saved-review"]


async def test_oversized_saved_paragraph_is_explicit_gap_without_truncation_or_model_send(runtime):
    document = Document()
    document.add_paragraph("人工保存的超长承诺内容。" * 2000)
    stream = BytesIO()
    document.save(stream)
    execute = AsyncMock()
    budget = RequestBudget(context_tokens=6000, model="fixture", prompts={"authoring_review": saved_word_review.PROMPT},
                           output_reserve=1024)
    report = await saved_word_review.review_saved_word(runtime.artifacts, runtime.document, stream.getvalue(), budget, execute)
    execute.assert_not_awaited()
    assert report["record_count"] == 1 and report["reviewed_record_count"] == 0
    assert report["findings"] == []
    assert any(gap["code"] == "saved_word_review_unavailable" and "saved-1" in gap["detail"] for gap in report["gaps"])


async def test_exhausted_model_work_unit_is_gap_without_claiming_reviewed_records(runtime):
    execute = AsyncMock(side_effect=ModelWorkUnitExhausted("authoring_review"))
    report = await saved_word_review.review_saved_word(runtime.artifacts, runtime.document, runtime.data,
                                                      runtime.budget, execute)
    execute.assert_awaited_once()
    assert report["record_count"] > 0 and report["reviewed_record_count"] == 0
    assert report["findings"] == []
    assert any(gap["code"] == "saved_word_review_unavailable" for gap in report["gaps"])


@pytest.mark.parametrize("finding", [
    {"task_ids": ["not-a-saved-record"], "content_ids": []},
    {"task_ids": ["saved-1"], "content_ids": ["old-generated-block"]},
    {"task_ids": [], "content_ids": []},
])
async def test_finding_cannot_cite_outside_saved_batch_or_generated_content(runtime, finding):
    runtime.state.response = {"findings": [{"code": "issue", "detail": "事实", **finding}]}
    with pytest.raises(ModelContractError):
        await runtime.execute()
    assert runtime.rows["saved-review"]["status"] == "failed"


async def test_saved_review_cannot_advertise_automatic_rewrite(runtime):
    runtime.state.response = {"findings": [{"code": "issue", "detail": "工期需要人工核实",
        "task_ids": ["saved-1"], "content_ids": [], "can_revise": True}]}
    report = runtime.artifacts.read_json(await runtime.execute())
    assert report["findings"][0]["can_revise"] is False


async def test_saved_file_preparation_leaves_event_loop_available_for_heartbeat(runtime, monkeypatch):
    loop, observations = asyncio.get_running_loop(), []
    original = saved_word_review.read_saved_word

    def slow_read(*args):
        progressed = threading.Event()
        loop.call_soon_threadsafe(progressed.set)
        observations.append(progressed.wait(timeout=0.25))
        return original(*args)

    async def execute(_stage, _request, validate):
        return validate({"findings": []})

    monkeypatch.setattr(saved_word_review, "read_saved_word", slow_read)
    await saved_word_review.review_saved_word(runtime.artifacts, runtime.document, runtime.data,
                                            runtime.budget, execute)
    assert observations == [True], "Slow saved Word parsing must let the heartbeat run."


def test_saved_word_preserves_hard_break_and_tab_separators():
    document = Document()
    run = document.add_paragraph().add_run("金额")
    run.add_tab()
    run.add_text("10元")
    run.add_break()
    run.add_text("工期30天")
    stream = BytesIO()
    document.save(stream)
    records, _ = saved_word_review.read_saved_word(stream.getvalue(), [])
    assert records[0]["actual_content"]["text"] == "金额\t10元\n工期30天"
