"""Shared writing content keeps every target navigable without duplicating prose."""

from collections import Counter
from io import BytesIO
import json
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
import pytest

from docreader.integrations.artifacts import ArtifactStore
from docreader.tender.authoring_export import export_authoring_document
from docreader.tender.word_navigation import bookmark_name
from test_submission_export import chapter, grid, part, source, target, writing


def files(artifacts, result):
    yield Document(BytesIO(artifacts.read_bytes(result["docx_ref"])))
    with ZipFile(BytesIO(artifacts.read_bytes(result["package_ref"]))) as archive:
        for volume in json.loads(archive.read("manifest.json"))["volumes"]:
            yield Document(BytesIO(archive.read(volume["package_path"])))


def target_paragraph(document, identity):
    name = bookmark_name(identity, "target")
    matches = [element for element in document.element.body.iter(qn("w:bookmarkStart"))
               if element.get(qn("w:name")) == name]
    assert len(matches) == 1, f"{identity} needs one stable anchor"
    anchor = matches[0]
    assert sum(element.get(qn("w:id")) == anchor.get(qn("w:id"))
               for element in document.element.body.iter(qn("w:bookmarkEnd"))) == 1
    return anchor.getparent()


def paragraph_text(paragraph):
    return "".join(element.text or "" for element in paragraph.iter(qn("w:t")))


@pytest.mark.parametrize("block_count", [0, 1, 2])
@pytest.mark.parametrize("volume", [False, True])
def test_shared_narrative_and_placeholder_have_all_anchors_in_main_and_volume(tmp_path, block_count, volume):
    artifacts = ArtifactStore(tmp_path)
    identities = ["first", "second", "third"]
    chapter_ = chapter("技术响应", [target(identity, kind="narrative") for identity in identities], volume=volume)
    blocks = [{"kind": "paragraph", "text": f"给定功能说明{number}"} for number in range(block_count)]
    unit = writing("shared", chapter_["chapter_id"], identities, blocks=blocks,
                   status="generated" if block_count else "placeholder")
    result = export_authoring_document(artifacts, {"chapters": [chapter_]}, [unit])
    documents = list(files(artifacts, result))
    assert len(documents) == (2 if volume else 1)
    for document in documents:
        anchors = [target_paragraph(document, identity) for identity in identities]
        assert all(anchor is anchors[0] for anchor in anchors)
        expected = blocks[0]["text"] if blocks else "（本部分内容待补）"
        assert paragraph_text(anchors[0]) == expected
        texts = Counter(paragraph.text for paragraph in document.paragraphs)
        assert all(texts[block["text"]] == 1 for block in blocks)
        assert texts["（本部分内容待补）"] == (0 if blocks else 1)
        # Only the existing page break is empty; no anchor-only visible paragraph.
        assert texts[""] == 1


def test_later_target_reuses_its_own_shared_unit_not_the_chapters_first_unit(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    chapter_ = chapter("技术响应", [target(identity, kind="narrative") for identity in ("both", "one", "two")])
    units = [writing("unit-1", chapter_["chapter_id"], ["both", "one"],
                     blocks=[{"kind": "paragraph", "text": "第一功能说明"}]),
             writing("unit-2", chapter_["chapter_id"], ["both", "two"],
                     blocks=[{"kind": "paragraph", "text": "第二功能说明"}])]
    result = export_authoring_document(artifacts, {"chapters": [chapter_]}, units)
    document = next(files(artifacts, result))
    assert paragraph_text(target_paragraph(document, "one")) == "第一功能说明"
    assert paragraph_text(target_paragraph(document, "two")) == "第二功能说明"


def test_own_template_and_cross_reference_take_priority_over_a_shared_unit_anchor(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    selected = part(source(artifacts, "original-form", grid=grid()))
    targets = [target("first", kind="narrative"), target("form", templates=[selected]),
               {**target("reference", kind="narrative"), "related_target_ids": ["first"]},
               target("shared-only", kind="narrative")]
    chapter_ = chapter("技术响应", targets, volume=True)
    unit = writing("shared", chapter_["chapter_id"], [target_["target_id"] for target_ in targets],
                   blocks=[{"kind": "paragraph", "text": "给定共享功能说明"}])
    result = export_authoring_document(artifacts, {"chapters": [chapter_]}, [unit])
    for document in files(artifacts, result):
        first = target_paragraph(document, "first")
        assert target_paragraph(document, "shared-only") is first
        assert paragraph_text(first) == "给定共享功能说明"
        assert len(document.tables) == 1
        assert target_paragraph(document, "form") is document.tables[0].cell(0, 0)._tc.find(qn("w:p"))
        assert paragraph_text(target_paragraph(document, "reference")) == "参见 1 技术响应"
        assert sum(paragraph.text == "给定共享功能说明" for paragraph in document.paragraphs) == 1
