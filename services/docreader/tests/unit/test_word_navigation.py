from io import BytesIO

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from zipfile import ZIP_DEFLATED, ZipFile

from docreader.integrations.artifacts import ArtifactStore
from docreader.tender.authoring_export import export_authoring_document
from docreader.tender.word_navigation import add_bookmark, bookmark_name, read_docx_outline


def saved(document):
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_export_has_stable_chapter_and_target_anchors_without_extra_text(tmp_path):
    store = ArtifactStore(tmp_path)
    chapter = {"chapter_id": "chapter-a", "title": "实施计划", "children": [], "targets": [
        {"target_id": "target-a", "kind": "narrative", "purpose": "实施", "rule_refs": [],
         "source_refs": [], "template_refs": []}]}
    result = export_authoring_document(store, {"chapters": [chapter]}, [])
    actual = read_docx_outline(store.read_bytes(result["docx_ref"]), [chapter])
    assert actual["chapters"][0]["chapter_id"] == "chapter-a"
    assert actual["chapters"][0]["bookmark"] == bookmark_name("chapter-a")
    assert bookmark_name("target-a", "target") in actual["bookmarks"]
    assert actual["missing_chapter_ids"] == []


def test_renamed_reordered_headings_bind_by_anchor_and_deleted_ones_remain_missing():
    known = [{"chapter_id": value, "title": "同名", "children": [], "targets": []} for value in ("a", "b", "deleted")]
    document = Document()
    add_bookmark(document.add_heading("已改名B", 1)._p, "b")
    add_bookmark(document.add_heading("同名", 1)._p, "a")
    document.add_heading("人工新增", 2)
    result = read_docx_outline(saved(document), known)
    assert [(item["chapter_id"], item["title"]) for item in result["chapters"]] == [("b", "已改名B"), ("a", "同名")]
    added = result["chapters"][1]["children"][0]
    assert added["manual"] is True and added["targets"] == [] and added["bookmark"] is None
    assert result["missing_chapter_ids"] == ["deleted"]


def test_duplicate_bookmarks_are_not_silently_bound():
    known = [{"chapter_id": "a", "title": "相同", "children": [], "targets": []}]
    document = Document()
    paragraph = document.add_heading("相同", 1)
    add_bookmark(paragraph._p, "a")
    from copy import deepcopy
    document.element.body.insert(1, deepcopy(paragraph._p))
    result = read_docx_outline(saved(document), known)
    assert all(chapter["manual"] and chapter["bookmark"] is None for chapter in result["chapters"])
    assert result["missing_chapter_ids"] == ["a"]


def test_direct_outline_level_takes_precedence_over_style_and_styles_are_optional():
    document = Document()
    heading = document.add_paragraph("普通样式的真实标题")
    hidden = document.add_heading("已经改为正文", 1)
    for paragraph, value in ((heading, "0"), (hidden, "9")):
        level = OxmlElement("w:outlineLvl")
        level.set(qn("w:val"), value)
        paragraph._p.get_or_add_pPr().append(level)
    original = saved(document)
    assert [item["title"] for item in read_docx_outline(original)["chapters"]] == ["普通样式的真实标题"]
    output = BytesIO()
    with ZipFile(BytesIO(original)) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for name in source.namelist():
            if name != "word/styles.xml":
                target.writestr(name, source.read(name))
    assert [item["title"] for item in read_docx_outline(output.getvalue())["chapters"]] == ["普通样式的真实标题"]


def test_cover_is_bound_without_adding_visible_heading(tmp_path):
    store = ArtifactStore(tmp_path)
    chapter = {"chapter_id": "cover", "title": "封面", "presentation_role": "cover", "children": [],
               "targets": [{"target_id": "cover-target", "kind": "narrative", "purpose": "项目名称",
                            "rule_refs": [], "source_refs": [], "template_refs": []}]}
    result = export_authoring_document(store, {"chapters": [chapter]}, [])
    data = store.read_bytes(result["docx_ref"])
    actual = read_docx_outline(data, [chapter])
    assert actual["chapters"][0]["chapter_id"] == "cover"
    assert actual["chapters"][0]["title"] == "封面"
    assert actual["chapters"][0]["targets"][0]["bookmark"] == bookmark_name("cover-target", "target")
    assert not any(paragraph.style.name.startswith("Heading") for paragraph in Document(BytesIO(data)).paragraphs)


def test_submission_note_subtree_is_not_a_missing_word_chapter_or_an_orphan_body(tmp_path):
    store = ArtifactStore(tmp_path)
    child = {"chapter_id": "note-child", "title": "外部子操作", "is_volume": True,
             "targets": [], "children": []}
    note = {"chapter_id": "note", "title": "上传与签章", "presentation_role": "submission_note",
            "is_volume": True, "children": [child], "targets": []}
    body = {"chapter_id": "body", "title": "正式正文", "children": [], "targets": []}
    deleted = {"chapter_id": "deleted", "title": "已删除正文", "children": [], "targets": []}
    historical = {"task_id": "historical", "chapter_id": child["chapter_id"], "target_ids": [],
                  "blocks": [{"kind": "paragraph", "text": "历史操作不能回流正文"}],
                  "cell_fills": [], "attachments": []}
    result = export_authoring_document(store, {"chapters": [note, body]}, [historical])
    data = store.read_bytes(result["docx_ref"])
    actual = read_docx_outline(data, [note, body, deleted])
    assert [chapter["chapter_id"] for chapter in actual["chapters"]] == ["body"]
    assert actual["missing_chapter_ids"] == ["deleted"]
    document = Document(BytesIO(data))
    text = '\n'.join(paragraph.text for paragraph in document.paragraphs)
    assert '1 正式正文' in text
    assert '历史操作不能回流正文' not in text and '外部子操作' not in text
    with ZipFile(store.path(result["package_ref"])) as archive:
        assert not any(name.startswith('volumes/') for name in archive.namelist())
