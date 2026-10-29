"""Assembly receives actual context; fixture choices do not test model semantics."""
from copy import deepcopy
from io import BytesIO
import json

from docx import Document
import pytest

from docreader.integrations.artifacts import ArtifactStore
from docreader.llm.budget import RequestBudget
from docreader.llm.request import ModelContractError
from docreader.tender import authoring_pipeline as pipeline
from docreader.tender.authoring import _OutlineContext
from docreader.tender.authoring_export import export_authoring_document
from docreader.tender.evidence.references import SourceIndex
from docreader.tender.schema import AuthoringOutline

from test_structural_packing import source


SUBMISSION_TEXT = (
    "商务、技术和报价文件分别提交；技术文件依次为技术响应、实施计划、售后服务。"
    "投标人与原厂为不同主体时分别提供财务资料，同一主体时按适用规定填写。"
)


def fixture_context(tmp_path, *, with_rule=False):
    artifacts = ArtifactStore(tmp_path)
    cells = [
        {"row": 0, "column": 0, "text": "8D财务情况表", "row_span": 1, "col_span": 2},
        {"row": 1, "column": 0, "text": "填写主体", "row_span": 1, "col_span": 1},
        {"row": 1, "column": 1, "text": "年度", "row_span": 1, "col_span": 1},
        {"row": 2, "column": 0, "text": "投标人或原厂：____", "row_span": 1, "col_span": 1},
        {"row": 2, "column": 1, "text": "", "row_span": 1, "col_span": 1},
        {"row": 3, "column": 0, "text": "不同主体分别填表；同一主体按适用规定填写。",
         "row_span": 1, "col_span": 2},
    ]
    native = source([
        {"text": SUBMISSION_TEXT, "heading": "投标人须知 > 提交组成与适用条件"},
        {"grid": {"row_count": 4, "column_count": 2, "cells": cells},
         "heading": "商务文件 > 资格审查资料 > 8D财务情况"},
    ])
    native = SourceIndex({**native.document, "original_ref": artifacts.put_bytes(b"synthetic tender")},
                         native.raw)
    index_ref = artifacts.put_json(native.descriptor(artifacts.put_json(native.raw)))
    snapshot = {"schema_version": 2, "run_id": "submission-followup", "source_version": 1,
        "domain_revision": 1, "frozen_ref": artifacts.put_json({"documents": [
            {**native.document, "role": "main", "source_index_ref": index_ref}]})}
    if with_rule:
        snapshot["extract_refs"] = [artifacts.put_json({"items": [{
            "item_id": "financial-form", "text": "按适用主体提供财务情况表。",
            "evidence": [{"range_id": native.ranges[-1]["range_id"],
                          "quote": native.ranges[-1]["text"]}]}]})]
    return artifacts, _OutlineContext(artifacts, snapshot), native


def local_response(payload):
    # These are explicit fixture references, not a semantic table classifier.
    constraint = next(entry["range_id"] for entry in payload["sources"]
                      if entry["text"] == SUBMISSION_TEXT)
    table = next(part for part in payload["template_parts"] if len(part["range_ids"]) == 4)
    return {"schema_version": 2, "chapters": [
        {"title": "商务文件", "is_volume": True, "targets": [], "children": [
            {"title": "资格审查资料", "targets": [], "children": [
                {"title": "8D财务情况", "children": [], "targets": [
                    {"kind": "form", "purpose": "按适用主体填写财务情况；主体关系未确认时保留待核。",
                     "rule_ids": [rule["rule_id"] for rule in payload["rules"]],
                     "range_ids": table["range_ids"],
                     "template_part_ids": [table["part_id"]]}]}]}]}],
        "constraint_range_ids": [constraint], "findings": []}


async def organize(tmp_path, assembly, *, with_rule=False, source_findings=()):
    artifacts, context, native = fixture_context(tmp_path, with_rule=with_rule)
    calls = []

    async def execute(task, request, validate):
        assert task == "outline"
        payload = json.loads(request["messages"][1]["content"][0]["text"])
        calls.append(deepcopy(payload))
        if payload.get("operation") == "assemble":
            return validate(assembly(payload, native))
        response = local_response(payload)
        response["findings"] = list(source_findings)
        return validate(response)

    result = await pipeline._organize(artifacts, context,
        RequestBudget(model="fixture", context_tokens=131072, output_reserve=16384), execute, [])
    assert len(calls) == 2
    return result, calls, native


def single_form(payload, *, parts=True):
    candidate = next(target for chapter in payload["chapters"] for target in chapter["targets"])
    return {"chapters": [{"title": "商务文件", "is_volume": True, "children": [], "targets": [
        {"kind": "form", "purpose": candidate["purpose"], "target_ids": [candidate["target_id"]],
         "template_part_ids": candidate["template_part_ids"] if parts else []}]}]}


async def test_assembly_preserves_submission_order_conditions_and_candidate_parent_path(tmp_path):
    def assemble(payload, native):
        assert any(entry["text"] == SUBMISSION_TEXT for entry in payload["submission_sources"])
        chapter = next(chapter for chapter in payload["chapters"] if chapter["targets"])
        assert chapter["title_path"] == ["商务文件", "资格审查资料", "8D财务情况"]
        assert chapter["targets"][0]["purpose"] == "按适用主体填写财务情况；主体关系未确认时保留待核。"
        return single_form(payload)

    await organize(tmp_path, assemble)


async def test_assembly_receives_complete_selected_grid_with_nonheader_subject_conditions(tmp_path):
    def assemble(payload, native):
        part = payload["template_parts"][0]
        assert part.get("grid") == native.raw["units"][1]["grid"]
        assert "heading_path" not in part
        headings = {section["section_id"]: section["heading_path"]
                    for section in payload["template_sections"]}
        assert headings[part["section_id"]] == "商务文件 > 资格审查资料 > 8D财务情况"
        assert len(payload["template_parts"]) == 1
        return single_form(payload)

    await organize(tmp_path, assemble)


async def test_assembly_can_bind_one_candidate_to_two_distinct_form_instances(tmp_path):
    def assemble(payload, native):
        result = single_form(payload)
        first = result["chapters"][0]["targets"][0]
        first["purpose"] = "投标人财务情况；依据适用主体填写。"
        second = deepcopy(first)
        second["purpose"] = "原厂财务情况；原厂与投标人为不同主体时填写。"
        result["chapters"][0]["targets"].append(second)
        return result

    result, _, _ = await organize(tmp_path, assemble)
    first, second = result["chapters"][0]["targets"]
    assert first["target_id"] != second["target_id"]
    assert first["template_refs"] == second["template_refs"] and first["template_refs"]
    assert first["source_refs"] == second["source_refs"]
    tasks = pipeline.plan_writing_tasks(result)
    assert len(tasks) == 2
    assert tasks[0]["task_id"] != tasks[1]["task_id"]
    assert [task["targets"][0]["target_id"] for task in tasks] == [first["target_id"], second["target_id"]]


async def test_assembly_omitting_template_never_adopts_candidate_template_implicitly(tmp_path):
    result, calls, _ = await organize(tmp_path, lambda payload, native: single_form(payload, parts=False))
    candidate = next(target for chapter in calls[-1]["chapters"] for target in chapter["targets"])
    assert candidate["template_part_ids"]
    selected = result["chapters"][0]["targets"][0]
    assert selected["source_refs"]
    assert selected["template_refs"] == []


async def test_explicit_candidate_exclusion_preserves_reason_without_rendering_or_false_unassigned_rule(tmp_path):
    reason = "已核对该候选属于另一适用主体，本次不作为提交表单。"

    def assemble(payload, native):
        candidate = next(target for chapter in payload["chapters"] for target in chapter["targets"])
        return {"chapters": [], "excluded_targets": [
            {"target_ids": [candidate["target_id"]], "reason": reason}]}

    result, _, _ = await organize(tmp_path, assemble, with_rule=True)
    assert not [target for chapter in pipeline._flatten(result["chapters"]) for target in chapter["targets"]]
    issue = next(issue for issue in result["issues"] if issue["code"] == "outline_target_excluded")
    assert reason in issue["detail"]
    assert not issue["detail"].startswith("来源读取阶段观察（最终情况以全局核对为准）：")
    assert "按适用主体填写财务情况；主体关系未确认时保留待核。" in issue["detail"]
    assert issue["rule_refs"][0]["item_ref"]["item_id"] == "financial-form"
    assert not any(issue["code"] == "outline_rule_unassigned" for issue in result["issues"])
    artifacts = ArtifactStore(tmp_path)
    exported = export_authoring_document(artifacts, result, [])
    word = Document(BytesIO(artifacts.read_bytes(exported["docx_ref"])))
    assert not word.tables
    assert "8D财务情况表" not in "\n".join(paragraph.text for paragraph in word.paragraphs)


async def test_undeclared_candidate_omission_still_keeps_template_in_pending_content(tmp_path):
    result, _, _ = await organize(tmp_path, lambda payload, native: {"chapters": []})
    assert [chapter["title"] for chapter in result["chapters"]] == ["待整理内容"]
    target = result["chapters"][0]["targets"][0]
    assert target["template_refs"]
    assert not any(issue["code"] == "outline_target_excluded" for issue in result["issues"])


@pytest.mark.parametrize("invalid, expected_code", [
    ("unknown", "outline_unknown_reference"),
    ("selected", "outline_conflicting_disposition"),
])
async def test_candidate_exclusion_rejects_unknown_or_already_selected_references(tmp_path, invalid, expected_code):
    def assemble(payload, native):
        response = single_form(payload) if invalid == "selected" else {"chapters": []}
        candidate = next(target for chapter in payload["chapters"] for target in chapter["targets"])
        response["excluded_targets"] = [{"target_ids": [
            candidate["target_id"] if invalid == "selected" else "unknown-candidate"],
            "reason": "排除理由必须针对本请求候选，且不能同时采用。"}]
        return response

    with pytest.raises(ModelContractError) as caught:
        await organize(tmp_path, assemble)
    assert any(error["code"] == expected_code for error in caught.value.errors)


async def test_assembly_receives_source_observations_but_not_local_unassigned_noise(tmp_path):
    observation = {"code": "financial_subject_context",
        "detail": "本批仅见主体条件，尚未确认投标人与原厂是否相同，归并时核对完整原文。",
        "rule_ids": []}

    def assemble(payload, native):
        assert {"code": observation["code"], "detail": observation["detail"]} in payload.get("findings", [])
        assert all(set(finding) == {"code", "detail"} for finding in payload["findings"])
        assert not any(finding["code"] == "outline_rule_unassigned" for finding in payload["findings"])
        return single_form(payload)

    await organize(tmp_path, assemble, source_findings=[observation,
        {"code": "outline_rule_unassigned", "detail": "仅本批未安排的规则不是全局漏项。", "rule_ids": []}])


async def test_supplemental_frozen_source_reaches_assembly_final_review_and_persisted_outline_without_extra_read(tmp_path):
    artifacts, context, _ = fixture_context(tmp_path)
    supplied = next(ref for refs in context.sources.values() for ref in refs if ref["unit_id"] == "unit:0")
    calls = []

    async def execute(task, request, validate):
        payload = json.loads(request["messages"][1]["content"][0]["text"])
        operation = payload.get("operation", "source")
        calls.append((task, operation))
        if operation in {"assemble", "structure"}:
            assert [source["text"] for source in payload["submission_sources"]] == [SUBMISSION_TEXT]
        if operation == "structure":
            return validate({"findings": []})
        if operation == "assemble":
            return validate(single_form(payload))
        assert task == "outline"
        response = local_response(payload)
        response["constraint_range_ids"] = []
        return validate(response)

    reference = await pipeline.generate_document(artifacts, context.input,
        RequestBudget(model="fixture", context_tokens=131072, output_reserve=16384), execute,
        submission_sources=[supplied, deepcopy(supplied)])
    document = artifacts.read_json(reference)
    outline = artifacts.read_json(document["outline_ref"])
    assert outline["submission_refs"] == [supplied]
    assert AuthoringOutline.model_validate(outline).model_dump()["submission_refs"] == [supplied]
    assert calls == [("outline", "source"), ("outline", "assemble"), ("authoring_review", "structure")]


async def test_supplemental_source_outside_current_frozen_units_is_a_gap_and_never_enters_assembly(tmp_path):
    artifacts, context, _ = fixture_context(tmp_path)
    known = next(ref for refs in context.sources.values() for ref in refs)
    foreign = {**known, "unit_id": "missing-from-current-snapshot"}
    gaps, calls = [], []

    async def execute(task, request, validate):
        payload = json.loads(request["messages"][1]["content"][0]["text"])
        calls.append(payload.get("operation", "source"))
        if payload.get("operation") == "assemble":
            assert payload["submission_sources"] == []
            return validate(single_form(payload))
        response = local_response(payload)
        response["constraint_range_ids"] = []
        return validate(response)

    outline = await pipeline._organize(artifacts, context,
        RequestBudget(model="fixture", context_tokens=131072, output_reserve=16384), execute, gaps,
        submission_sources=[foreign])
    assert outline["submission_refs"] == []
    assert any(gap["code"] == "outline_submission_source_unavailable" for gap in gaps)
    assert calls == ["source", "assemble"]


async def test_excluded_candidate_scoped_issues_are_retired_but_other_candidate_and_global_observations_remain(tmp_path):
    artifacts, context, _ = fixture_context(tmp_path)
    original_issues = []
    global_detail = "全局格式关系尚需核对，不能因排除一个候选而删除本观察。"

    async def execute(task, request, validate):
        payload = json.loads(request["messages"][1]["content"][0]["text"])
        if payload.get("operation") == "assemble":
            first, second = next(chapter["targets"] for chapter in payload["chapters"] if chapter["targets"])
            return validate({"chapters": [{"title": "商务文件", "children": [], "targets": [
                {"kind": "form", "purpose": second["purpose"], "target_ids": [second["target_id"]],
                 "template_part_ids": []}]}], "excluded_targets": [
                    {"target_ids": [first["target_id"]], "reason": "该候选不属于本次提交要求。"}]})
        response = local_response(payload)
        candidates = response["chapters"][0]["children"][0]["children"][0]["targets"]
        candidates[0]["template_part_ids"] = []
        candidates[0]["purpose"] = "仅作为局部观察保留的候选，归并时判断适用性。"
        second = deepcopy(candidates[0])
        second["purpose"] = "本次仍需保留填写位置的表单，其原格式待补。"
        candidates.append(second)
        response["findings"] = [{"code": "outline_template_missing", "detail": global_detail, "rule_ids": []}]
        result = validate(response)
        original_issues.extend(deepcopy(result["issues"]))
        return result

    outline = await pipeline._organize(artifacts, context,
        RequestBudget(model="fixture", context_tokens=131072, output_reserve=16384), execute, [])
    assert len([issue for issue in original_issues
                if issue["code"] == "outline_template_missing" and issue.get("target_id")]) == 2
    remaining = [issue for issue in outline["issues"] if issue["code"] == "outline_template_missing"]
    assert len(remaining) == 2
    observation = next(issue for issue in remaining if issue["detail"].endswith(global_detail))
    assert observation["detail"].startswith("来源读取阶段观察（最终情况以全局核对为准）：")
    # Both candidates have actual source references; missing a rule is not an ungrounded position.
    assert not any(issue["code"] == "outline_target_unbound" for issue in outline["issues"])
    assert len([issue for issue in outline["issues"] if issue["code"] == "outline_target_excluded"]) == 1
