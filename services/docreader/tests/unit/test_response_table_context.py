"""Shared source headers and conditions belong to the actual business table."""

from copy import deepcopy
from io import BytesIO
import json

from docx import Document
import pytest

from docreader.integrations.artifacts import ArtifactStore
from docreader.llm.budget import RequestBudget
from docreader.llm.request import ModelContractError
from docreader.llm.wire import WirePayload
from docreader.tender.authoring import validate_outline_references
from docreader.tender.authoring_export import export_authoring_document
from docreader.tender.authoring_excerpts import bind_source_excerpts, load_source_excerpt
from docreader.tender.authoring_fragments import split_writing_task
from docreader.tender import authoring_pipeline as pipeline
from docreader.tender.authoring_pipeline import generate_document
from docreader.tender.authoring_templates import prepare_template_parts
from docreader.tender.authoring_view import compact_assembly_payload, compact_structure_payload
from docreader.tender.schema import OutlineResponse

from test_authoring_pipeline_boundaries import snapshot
from test_authoring_templates import create_context, create_source
from test_authoring_compact_view import native_grid, payload as compact_payload
from test_response_table_export import excerpts, opened, outline, unit


HEADER = "指标要求（同等及以上配置）"
GUIDANCE = "以下条目共同适用：单位为Mbps，未提供资料时保留空白应答。"
REQUIREMENTS = ["甲项须支持双机热备。", "乙项须提供独立证明。"]


@pytest.mark.parametrize("manual", [False, True])
def test_shared_context_exports_once_before_the_table_without_an_answer_slot(tmp_path, manual):
    artifacts = ArtifactStore(tmp_path)
    selected = excerpts(artifacts, ["### " + HEADER, GUIDANCE, *REQUIREMENTS, "只供理解，不是成稿内容。"])
    rows = [{"row_id": f"row-{i}", "source_excerpts": [ref], "rule_refs": []}
            for i, ref in enumerate(selected[2:4])]
    tree = outline(rows, manual_fill=manual, source_excerpts=selected[4:])
    tree["chapters"][0]["targets"][0]["response_table"]["context_source_excerpts"] = selected[:2]
    before = deepcopy(tree)

    document = opened(artifacts, export_authoring_document(artifacts, tree, [unit()]))

    assert [p.text for p in document.paragraphs].count(HEADER + "\n" + GUIDANCE) == 1
    assert len(document.tables) == 1 and len(document.tables[0].rows) == 3
    assert [[cell.text for cell in row.cells] for row in document.tables[0].rows[1:]] == [
        [REQUIREMENTS[0], ""], [REQUIREMENTS[1], ""]]
    assert "只供理解" not in document.element.xml
    assert document.element.xml.index(HEADER) < document.element.xml.index("招标要求")
    assert tree == before
    assert [load_source_excerpt(artifacts, ref)["text"] for ref in selected[:2]] == ["### " + HEADER, GUIDANCE]


@pytest.mark.parametrize("explicit", [False, True])
@pytest.mark.parametrize("with_knowledge", [False, True])
async def test_source_assembly_review_and_revision_retain_shared_context(tmp_path, explicit, with_knowledge):
    artifacts = ArtifactStore(tmp_path)
    fixed = snapshot(artifacts, [{"heading": "业务要求", "text": text}
                                 for text in [HEADER, GUIDANCE, *REQUIREMENTS]])
    calls = []
    writes, reviews = [], []

    def check_context(target, supports):
        by_id = {support["support_id"]: support for support in supports}
        assert [by_id[i]["text"] for i in target["response_context_support_ids"]] == [HEADER, GUIDANCE]
        assert all(by_id[i]["origin"] == "tender" for i in target["response_context_support_ids"])

    async def execute(operation, request, validate):
        payload = json.loads(request["messages"][1]["content"][0]["text"])
        stage = payload.get("operation", operation)
        calls.append(stage)
        if stage == "outline":
            by_text = {source["text"]: source["range_id"] for source in payload["sources"]}
            assert all(identity.startswith("r") for identity in by_text.values())
            return validate({"schema_version": 2, "chapters": [{"title": "业务应答", "targets": [
                {"kind": "narrative", "purpose": "按完整要求逐项应答", "rule_ids": [], "range_ids": [],
                 "support_origins": ["tender", "knowledge"],
                 "response_table": {"context_range_ids": [by_text[HEADER], by_text[GUIDANCE]],
                    "rows": [{"range_ids": [by_text[text]]} for text in REQUIREMENTS]}}]}], "findings": []})
        if stage == "assemble":
            candidate = payload["chapters"][0]["targets"][0]
            by_text = {source["text"]: source["range_id"] for source in payload["source_excerpts"]}
            assert candidate["response_table"]["context_range_ids"] == [by_text[HEADER], by_text[GUIDANCE]]
            assert len(payload["source_excerpts"]) == 4
            target = {"kind": "narrative", "purpose": "逐项应答", "target_ids": [candidate["target_id"]]}
            if explicit:
                target["response_table"] = {"context_range_ids": [by_text[HEADER], by_text[HEADER]],
                    "rows": [{"range_ids": [by_text[text]]} for text in REQUIREMENTS]}
                unknown_context = deepcopy(target)
                unknown_context["response_table"]["context_range_ids"] = ["absent"]
                with pytest.raises(ModelContractError):
                    validate({"chapters": [{"title": "业务应答", "targets": [unknown_context]}]})
            return validate({"chapters": [{"title": "业务应答", "targets": [target]}]})
        if stage == "structure":
            target = payload["chapters"][0]["targets"][0]
            check_context(target, payload["supports"])
            assert len(target["response_rows"]) == 2
            return validate({"findings": []})
        if with_knowledge:
            if operation == "authoring_basis":
                return validate({"shared_arrangements": []})
            if operation == "authoring_write":
                writes.append(payload)
                check_context(payload["targets"][0], payload["supports"])
                assert len(payload["response_rows"]) == 2
                assert len([s for s in payload["supports"] if s.get("text") == HEADER]) == 1
                assert bool(payload.get("previous_content")) == (len(writes) == 2)
                return validate({"row_fills": [{"row_id": "q1", "text": "资料确认支持双机热备。" +
                                                ("支持自动切换。" if len(writes) == 2 else ""),
                                                "support_ids": ["k1"]}]})
            if operation == "authoring_review":
                reviews.append(payload)
                record = payload["units"][0]
                check_context(record["targets"][0], record["supports"])
                findings = [{"code": "response_needs_detail", "detail": "按已有资料补充表述。",
                    "task_ids": [record["task_id"]], "content_ids": ["r1"], "can_revise": True,
                    "reject_content": False}] if len(reviews) == 1 else []
                return validate({"findings": findings})
        pytest.fail(f"No knowledge must not cause {stage}.")

    result = artifacts.read_json(await generate_document(
        artifacts, fixed, RequestBudget(model="fixture", context_tokens=131072, output_reserve=16384), execute,
        materials=[{"material_id": "guide", "title": "产品资料", "text": "资料确认支持双机热备。支持自动切换。"}]
                  if with_knowledge else []))
    tree = artifacts.read_json(result["outline_ref"])
    table = tree["chapters"][0]["targets"][0]["response_table"]
    assert [load_source_excerpt(artifacts, ref)["text"] for ref in table["context_source_excerpts"]] == [HEADER, GUIDANCE]
    document = Document(BytesIO(artifacts.read_bytes(result["docx_ref"])))
    assert [p.text for p in document.paragraphs].count(HEADER + "\n" + GUIDANCE) == 1
    assert len(document.tables[0].rows) == 3
    assert calls == ["outline", "assemble", "structure"] + (["authoring_basis", "authoring_write",
        "authoring_review", "authoring_write", "authoring_review"] if with_knowledge else [])


def test_shared_context_uses_current_read_ranges_and_keeps_distinct_slices_of_one_unit(tmp_path):
    artifacts, index_ref, native = create_source(tmp_path, [{"text": HEADER + "\n" + REQUIREMENTS[0]}])
    context = create_context(artifacts, index_ref, native)
    records, _ = prepare_template_parts(artifacts, context)
    sources = [record["source"] for record in records]
    wire = WirePayload("outline", {"schema_version": 2, "sources": sources})
    ids = [entry["range_id"] for entry in wire.payload["sources"]]
    response = wire.decode({"schema_version": 2, "chapters": [{"title": "应答", "targets": [
        {"kind": "narrative", "purpose": "逐项应答", "range_ids": [], "rule_ids": [],
         "response_table": {"context_range_ids": [ids[0], ids[0]], "rows": [{"range_ids": [ids[1]]}]}}]}],
        "findings": []})
    parsed = OutlineResponse.model_validate(response).model_dump()
    actual_ids = [entry["range_id"] for entry in sources]
    validate_outline_references(parsed, rule_ids=[], source_ids=actual_ids, template_part_ids=[])
    target = context.validate(parsed)["chapters"][0]["targets"][0]
    shared = target["response_table"]["context_source_excerpts"]
    row = target["response_table"]["rows"][0]
    assert len(shared) == 1 and len(target["source_refs"]) == 1
    assert shared[0]["source_ref"] == row["source_excerpts"][0]["source_ref"]
    assert load_source_excerpt(artifacts, shared[0])["text"] == HEADER
    assert load_source_excerpt(artifacts, row["source_excerpts"][0])["text"] == REQUIREMENTS[0]
    with pytest.raises(ModelContractError):
        validate_outline_references(parsed, rule_ids=[], source_ids=actual_ids,
                                    constraint_source_ids=actual_ids[1:], template_part_ids=[])
    parsed["chapters"][0]["targets"][0]["response_table"]["context_range_ids"] = ["absent"]
    with pytest.raises(ModelContractError):
        context.validate(parsed)


def test_split_keeps_table_context_even_when_it_shares_the_omitted_rows_unit(tmp_path):
    artifacts, index_ref, native = create_source(tmp_path, [{"text": HEADER + "\n" + REQUIREMENTS[0]},
                                                        {"text": REQUIREMENTS[1]}])
    context = create_context(artifacts, index_ref, native)
    records, _ = prepare_template_parts(artifacts, context)
    selected = [bind_source_excerpts(context, [record["source"]["range_id"]])[0] for record in records]
    # The common header and first requirement occupy distinct ranges of one native unit.
    common = selected[0]
    rows = [{"row_id": str(i), "source_excerpts": [ref], "rule_refs": []}
            for i, ref in enumerate(selected[1:])]
    target = outline(rows)["chapters"][0]["targets"][0]
    target["response_table"]["context_source_excerpts"] = [common]
    target["source_refs"] = [ref["source_ref"] for ref in selected]
    task = {"task_id": "task", "chapter_id": "chapter", "title_path": ["应答"], "targets": [target]}
    before = deepcopy(task)
    parts = split_writing_task(task)
    assert len(parts) == 2 and task == before
    assert [row for part in parts for row in part["targets"][0]["response_table"]["rows"]] == rows
    for part in parts:
        narrowed = part["targets"][0]
        assert narrowed["response_table"]["context_source_excerpts"] == [common]
        assert common["source_ref"] in narrowed["source_refs"]
    parts[0]["targets"][0]["response_table"]["context_source_excerpts"][0]["selectors"].clear()
    assert task == before


def test_compact_views_preserve_table_context_references_and_text():
    original = compact_payload(native_grid())
    table = {"context_range_ids": ["x1", "x2"], "rows": [{"range_ids": ["x3"], "rule_ids": []}]}
    original["chapters"][0]["targets"][1]["response_table"] = table
    original["source_excerpts"] = [{"range_id": f"x{i}", "text": text}
        for i, text in enumerate([HEADER, GUIDANCE, REQUIREMENTS[0]], 1)]
    before = deepcopy(original)
    compact = compact_assembly_payload(original)
    assert original == before
    assert compact["chapters"][0]["targets"][1]["response_table"] == table
    assert compact["source_excerpts"] == {entry["range_id"]: {key: value for key, value in entry.items()
        if key != "range_id"} for entry in original["source_excerpts"]}
    structure = {"chapters": [{"targets": [{"response_context_support_ids": ["s1", "s2"],
        "response_rows": [{"row_id": "q1", "source_support_ids": ["s3"]}]}]}],
        "supports": [{"support_id": f"s{i}", "origin": "tender", "text": text}
            for i, text in enumerate([HEADER, GUIDANCE, REQUIREMENTS[0]], 1)], "submission_sources": []}
    assert compact_structure_payload(structure) == {**structure, "structure_view": "compact_v1"}


def test_unreadable_shared_context_keeps_a_located_placeholder_and_all_requirements(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    selected = excerpts(artifacts, REQUIREMENTS)
    rows = [{"row_id": str(i), "source_excerpts": [ref], "rule_refs": []} for i, ref in enumerate(selected)]
    tree = outline(rows)
    missing = deepcopy(selected[0])
    missing["source_ref"]["unit_id"] = "unavailable"
    tree["chapters"][0]["targets"][0]["response_table"]["context_source_excerpts"] = [missing]
    result = export_authoring_document(artifacts, tree, [])
    document = opened(artifacts, result)
    assert any("原表共同条件待核对" in p.text for p in document.paragraphs)
    assert [[cell.text for cell in row.cells] for row in document.tables[0].rows[1:]] == [[text, ""] for text in REQUIREMENTS]
    assert any(gap["code"] == "export_response_source_unavailable" and gap["target_id"] == "technical"
               for gap in result["gaps"])


async def test_structure_revision_keeps_context_added_by_first_assembly(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    fixed = snapshot(artifacts, [{"heading": "业务要求", "text": text}
                                 for text in [HEADER, GUIDANCE, REQUIREMENTS[0]]])
    calls = []
    structures = []

    async def execute(operation, request, validate):
        payload = json.loads(request["messages"][1]["content"][0]["text"])
        stage = payload.get("operation", operation)
        calls.append(stage)
        if stage == "outline":
            by_text = {source["text"]: source["range_id"] for source in payload["sources"]}
            return validate({"schema_version": 2, "chapters": [{"title": "业务应答", "targets": [
                {"kind": "narrative", "purpose": "逐项应答", "rule_ids": [], "range_ids": [by_text[GUIDANCE]],
                 "response_table": {"context_range_ids": [by_text[HEADER]],
                    "rows": [{"range_ids": [by_text[REQUIREMENTS[0]]]}]}}]}], "findings": []})
        if stage == "assemble":
            candidate = payload["chapters"][0]["targets"][0]
            target = {"kind": "narrative", "purpose": "按完整要求应答", "target_ids": [candidate["target_id"]]}
            if not payload.get("previous_structure"):
                by_text = {source["text"]: source["range_id"] for source in payload["source_excerpts"]}
                target["response_table"] = {"context_range_ids": [by_text[GUIDANCE]],
                    "rows": candidate["response_table"]["rows"]}
            else:
                target["purpose"] = "完善后的逐项应答"
            return validate({"chapters": [{"title": "业务应答", "targets": [target]}]})
        if stage == "structure":
            target = payload["chapters"][0]["targets"][0]
            by_id = {support["support_id"]: support["text"] for support in payload["supports"]}
            structures.append([by_id[identity] for identity in target["response_context_support_ids"]])
            findings = [{"code": "purpose_needs_detail", "detail": "完善用途文字，保留既有实际原文。",
                "can_revise": True}] if len(structures) == 1 else []
            return validate({"findings": findings})
        pytest.fail(f"No knowledge must not cause {stage}.")

    result = artifacts.read_json(await generate_document(
        artifacts, fixed, RequestBudget(model="fixture", context_tokens=131072, output_reserve=16384), execute))
    assert structures == [[HEADER, GUIDANCE], [HEADER, GUIDANCE]]
    assert calls == ["outline", "assemble", "structure", "assemble", "structure"]
    document = opened(artifacts, result)
    assert [p.text for p in document.paragraphs].count(HEADER + "\n" + GUIDANCE) == 1


async def test_structure_revision_keeps_rows_added_by_first_assembly(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    fixed = snapshot(artifacts, [{"heading": "业务要求", "text": text} for text in REQUIREMENTS])
    reviewed_rows = []

    async def execute(operation, request, validate):
        payload = json.loads(request["messages"][1]["content"][0]["text"])
        stage = payload.get("operation", operation)
        if stage == "outline":
            by_text = {source["text"]: source["range_id"] for source in payload["sources"]}
            return validate({"schema_version": 2, "chapters": [{"title": "业务应答", "targets": [
                {"kind": "narrative", "purpose": "逐项应答", "rule_ids": [], "range_ids": [],
                 "response_table": {"rows": [{"range_ids": [by_text[REQUIREMENTS[0]]]}]}},
                {"kind": "narrative", "purpose": "可选的补充要求", "rule_ids": [],
                 "range_ids": [by_text[REQUIREMENTS[1]]]}]}], "findings": []})
        if stage == "assemble":
            target = {"kind": "narrative", "purpose": "已改好的用途" if payload.get("previous_structure") else "逐项应答",
                      "target_ids": [payload["chapters"][0]["targets"][0]["target_id"]]}
            if not payload.get("previous_structure"):
                target["response_table"] = {"rows": [{"range_ids": [source["range_id"]]}
                    for source in payload["source_excerpts"]]}
            return validate({"chapters": [{"title": "业务应答", "targets": [target]}],
                "excluded_targets": [{"target_ids": [payload["chapters"][0]["targets"][1]["target_id"]],
                                      "reason": "补充要求已选入业务表，不另成一份交付。"}]})
        if stage == "structure":
            supports = {support["support_id"]: support["text"] for support in payload["supports"]}
            reviewed_rows.append(["\n".join(supports[identity] for identity in row["source_support_ids"])
                for row in payload["chapters"][0]["targets"][0]["response_rows"]])
            findings = [{"code": "purpose_needs_detail", "detail": "只调整用途文字，保留已安排的完整要求。"}]
            return validate({"findings": findings if len(reviewed_rows) == 1 else []})
        pytest.fail(f"No knowledge must not cause {stage}.")

    result = artifacts.read_json(await generate_document(
        artifacts, fixed, RequestBudget(model="fixture", context_tokens=131072, output_reserve=16384), execute))
    assert len(reviewed_rows) == 2
    assert sorted(reviewed_rows[0]) == sorted(reviewed_rows[1]) == sorted(REQUIREMENTS)
    tree = artifacts.read_json(result["outline_ref"])
    assert {ref["unit_id"] for ref in tree["chapters"][0]["targets"][0]["source_refs"]} == {"unit:0", "unit:1"}
    document = opened(artifacts, result)
    assert sorted(row.cells[0].text for row in document.tables[0].rows[1:]) == sorted(REQUIREMENTS)


@pytest.mark.parametrize("selection", ["omitted", "explicit_complete", "explicit_incomplete",
                                      "implicit_context", "incomplete_rows", "revision_after_explicit_complete"])
async def test_distinct_table_conditions_keep_their_scope_when_assembly_merges_candidates(tmp_path, selection):
    artifacts = ArtifactStore(tmp_path)
    conditions = ["仅甲项适用：使用甲环境。", "仅乙项适用：使用乙环境。"]
    fixed = snapshot(artifacts, [{"heading": "业务要求", "text": text}
                                 for text in [conditions[0], REQUIREMENTS[0], conditions[1], REQUIREMENTS[1]]])
    calls = []
    structures = []

    async def execute(operation, request, validate):
        payload = json.loads(request["messages"][1]["content"][0]["text"])
        stage = payload.get("operation", operation)
        calls.append(stage)
        if stage == "outline":
            by_text = {source["text"]: source["range_id"] for source in payload["sources"]}
            return validate({"schema_version": 2, "findings": [], "chapters": [{"title": "业务应答", "targets": [
                {"kind": "narrative", "purpose": f"第{number}份原业务表", "range_ids": [], "rule_ids": [],
                 "response_table": {"context_range_ids": [by_text[condition]],
                    "rows": [{"range_ids": [by_text[requirement]]}]}}
                for number, (condition, requirement) in enumerate(zip(conditions, REQUIREMENTS))]}]})
        if stage == "assemble":
            target = {"kind": "narrative", "purpose": "完善用途后合并组织业务应答" if payload.get("previous_structure") else "合并组织业务应答",
                      "target_ids": [target["target_id"] for target in payload["chapters"][0]["targets"]]}
            if selection != "omitted" and not payload.get("previous_structure"):
                by_text = {source["text"]: source["range_id"] for source in payload["source_excerpts"]}
                target["response_table"] = {"context_range_ids": [], "rows": [
                    {"range_ids": ([by_text[condition]] if selection != "explicit_incomplete" else []) +
                                  [by_text[requirement]]} for condition, requirement in zip(conditions, REQUIREMENTS)]}
                if selection == "implicit_context":
                    target["response_table"].pop("context_range_ids")
                if selection == "incomplete_rows":
                    target["response_table"]["rows"][1]["range_ids"].pop()
            return validate({"chapters": [{"title": "业务应答", "targets": [target]}]})
        assert stage == "structure"
        structures.append(payload)
        return validate({"findings": [{"code": "purpose_needs_detail", "detail": "只修改用途说明。"}]
            if selection == "revision_after_explicit_complete" and len(structures) == 1 else []})

    result = artifacts.read_json(await generate_document(
        artifacts, fixed, RequestBudget(model="fixture", context_tokens=131072, output_reserve=16384), execute))
    tree = artifacts.read_json(result["outline_ref"])
    targets = tree["chapters"][0]["targets"]
    document = opened(artifacts, result)
    if selection in {"explicit_complete", "revision_after_explicit_complete"}:
        assert len(targets) == len(document.tables) == 1
        assert targets[0]["response_table"].get("context_source_excerpts", []) == []
        assert [row.cells[0].text for row in document.tables[0].rows[1:]] == [
            condition + "\n" + requirement for condition, requirement in zip(conditions, REQUIREMENTS)]
    else:
        assert len(targets) == len(document.tables) == 2
        for target, condition, requirement, table in zip(targets, conditions, REQUIREMENTS, document.tables):
            assert [load_source_excerpt(artifacts, ref)["text"] for ref in
                    target["response_table"]["context_source_excerpts"]] == [condition]
            assert [row.cells[0].text for row in table.rows[1:]] == [requirement]
        assert any(issue["code"] == "outline_response_context_unassigned" for issue in tree["issues"])
    assert calls == ["outline", "assemble", "structure"] + (
        ["assemble", "structure"] if selection == "revision_after_explicit_complete" else [])


async def test_context_fallback_keeps_additional_sources_rules_and_reference_lineage(tmp_path):
    conditions = ["仅甲项适用。", "仅乙项适用。"]
    additions = ["新增的理解条件，归属待核。", "新增的理解条目，归属待核。"]
    artifacts, index_ref, native = create_source(tmp_path, [
        {"text": text} for text in [conditions[0], REQUIREMENTS[0], conditions[1], REQUIREMENTS[1], *additions]])
    context = create_context(artifacts, index_ref, native)
    rule_artifact = artifacts.put_json({"items": [{"item_id": f"extra-{i}"} for i in range(2)]})
    references = [{"item_ref": {"artifact_ref": rule_artifact, "item_id": f"extra-{i}"}} for i in range(2)]
    context.rules = {f"rule-{i}": ref for i, ref in enumerate(references)}
    context.payload["rules"] = [{"rule_id": identity, "item": {"text": "新增规则归属待核。", "evidence": []}}
                                for identity in context.rules]

    async def execute(operation, request, validate):
        payload = json.loads(request["messages"][1]["content"][0]["text"])
        if payload.get("operation") != "assemble":
            by_text = {source["text"]: source["range_id"] for source in payload["sources"]}
            return validate({"schema_version": 2, "findings": [], "chapters": [{"title": "原业务应答", "targets": [
                *[{"kind": "narrative", "purpose": f"第{i}份原业务表", "range_ids": [], "rule_ids": [],
                   "support_origins": ["tender"], "response_table": {"context_range_ids": [by_text[condition]],
                       "rows": [{"range_ids": [by_text[requirement]]}]}}
                  for i, (condition, requirement) in enumerate(zip(conditions, REQUIREMENTS))],
                {"kind": "narrative", "purpose": "可供理解的其他来源", "rule_ids": [],
                 "range_ids": [by_text[text] for text in additions]}]}]})
        by_text = {source["text"]: source["range_id"] for source in payload["source_excerpts"]}
        rule_ids = [entry["rule_id"] for entry in payload["unassigned_rules"]]
        response = {"chapters": [{"title": "归并应答", "targets": [
            {"kind": "narrative", "purpose": "试图合并两份业务表", "target_ids": ["t1", "t2"],
             "rule_ids": rule_ids[:1], "support_origins": ["tender", "knowledge"],
             "response_table": {"context_range_ids": [by_text[additions[0]]], "rows": [
                 {"range_ids": [by_text[REQUIREMENTS[0]], by_text[additions[1]]], "rule_ids": rule_ids[1:]},
                 {"range_ids": [by_text[REQUIREMENTS[1]]]}]}},
            {"kind": "narrative", "purpose": "见第一份原业务表", "related_target_ids": ["t1"]}]}],
            "excluded_targets": [{"target_ids": ["t3"], "reason": "仅作为理解来源，不单独交付。"}]}
        for field in ("range_ids", "rule_ids"):
            invalid = deepcopy(response)
            invalid["chapters"][0]["targets"][0]["response_table"]["rows"][0][field] = ["absent"]
            with pytest.raises(ModelContractError):
                validate(invalid)
        return validate(response)

    tree = await pipeline._organize(artifacts, context, RequestBudget(
        model="fixture", context_tokens=131072, output_reserve=16384), execute, [])
    first, second, cross_reference = tree["chapters"][0]["targets"]
    assert cross_reference["related_target_ids"] == [first["target_id"]]
    assert tree["submission_refs"] == [] and tree["constraint_rule_refs"] == []
    for target, condition, requirement in zip([first, second], conditions, REQUIREMENTS):
        assert target["support_origins"] == ["tender"] and target["rule_refs"] == []
        assert [load_source_excerpt(artifacts, ref)["text"] for ref in
                target["response_table"]["context_source_excerpts"]] == [condition]
        assert [[load_source_excerpt(artifacts, ref)["text"] for ref in row["source_excerpts"]]
                for row in target["response_table"]["rows"]] == [[requirement]]
    retained_sources = {load_source_excerpt(artifacts, ref)["text"]
                        for target in [first, second] for ref in target.get("source_excerpts", [])}
    assert retained_sources == set(additions)
    assert all(any(ref["unit_id"] == f"unit:{number}" for target in [first, second]
                   for ref in target["source_refs"]) for number in (4, 5))
    assert {json.dumps(ref, sort_keys=True) for issue in tree["issues"]
            if issue["code"] == "outline_rule_unassigned" for ref in issue["rule_refs"]} == {
                json.dumps(ref, sort_keys=True) for ref in references}
    issue = next(issue for issue in tree["issues"] if issue["code"] == "outline_response_context_unassigned")
    assert all(text in issue["detail"] for text in ["归并应答", "第0份原业务表", "第1份原业务表", "归属待核"])
