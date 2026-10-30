"""Business response rows preserve complete source text and leave unknown answers blank."""

from copy import deepcopy
from io import BytesIO

from docx import Document
import pytest

from docreader.integrations.artifacts import ArtifactStore
from docreader.tender.authoring_export import export_authoring_document
from docreader.tender.authoring_excerpts import load_source_excerpt
from docreader.tender.evidence.references import EvidenceRepository, SourceIndex, fragment, selector


def excerpts(artifacts, texts):
    units = [{"unit_id": f"page:{ordinal}", "kind": "section", "ordinal": ordinal,
              "text": text, "locator": {"locator_kind": "document", "section_ordinal": ordinal}}
             for ordinal, text in enumerate(texts)]
    raw = {"schema_version": 2, "units": units, "images": {}, "ocr": {},
           "source_contract": {"page_manifest": [], "units": [
               {"unit_id": unit["unit_id"], "section_id": "technical", "parent_section_id": None,
                "heading_path": "项目要求", "physical_locator": unit["locator"]} for unit in units]}}
    original = {"document_id": "source", "document_version": "1",
                "original_ref": artifacts.put_bytes(b"synthetic source")}
    index = SourceIndex(original, raw)
    index_ref = artifacts.put_json(index.descriptor(artifacts.put_json(raw)))
    result = []
    for entry in index.ranges:
        selected = {**selector(entry), "source_range_id": entry["range_id"]}
        result.append({"source_ref": {"source_index_ref": index_ref, "unit_id": entry["unit_id"]},
                       "selectors": [selected]})
    return result


def outline(rows, **fields):
    target = {"target_id": "technical", "kind": "narrative", "purpose": "技术逐项应答",
              "rule_refs": [], "source_refs": [], "response_table": {"rows": rows}, **fields}
    return {"chapters": [{"chapter_id": "chapter", "title": "技术响应", "targets": [target], "children": []}]}


def unit(**fields):
    return {"schema_version": 2, "task_id": "task", "chapter_id": "chapter", "target_ids": ["technical"],
            "status": "placeholder", "blocks": [], "cell_fills": [], "row_fills": [],
            "attachments": [], "gaps": [], **fields}


def opened(artifacts, result):
    return Document(BytesIO(artifacts.read_bytes(result["docx_ref"])))


def test_complete_cross_page_rows_are_exported_without_knowledge_or_generic_placeholder(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    selected = excerpts(artifacts, ["系统应支持协同防御，或提供定制开发承", "诺，盖章并附详细方案。", "另须提供独立功能证明。"])
    rows = [{"row_id": "cooperation", "source_excerpts": selected[:2], "rule_refs": []},
            {"row_id": "proof", "source_excerpts": selected[2:], "rule_refs": []}]
    result = export_authoring_document(artifacts, outline(rows), [unit()])
    doc = opened(artifacts, result)
    assert len(doc.tables) == 1
    assert [cell.text for cell in doc.tables[0].rows[0].cells] == ["招标要求", "投标响应"]
    assert [cell.text for cell in doc.tables[0].rows[1].cells] == ["系统应支持协同防御，或提供定制开发承\n诺，盖章并附详细方案。", ""]
    assert [cell.text for cell in doc.tables[0].rows[2].cells] == ["另须提供独立功能证明。", ""]
    assert not any("本部分内容待补" in p.text for p in doc.paragraphs)
    assert not any(gap["code"] == "export_content_pending" for gap in result["gaps"])


@pytest.mark.parametrize('manual_fill', [False, True])
def test_form_response_rows_are_a_real_table_without_native_template_missing_warning(tmp_path, manual_fill):
    artifacts = ArtifactStore(tmp_path)
    selected = excerpts(artifacts, ['完整技术要求及其适用条件。'])
    rows = [{'row_id': 'requirement', 'source_excerpts': selected, 'rule_refs': []}]
    result = export_authoring_document(artifacts,
        outline(rows, kind='form', manual_fill=manual_fill), [unit()])
    doc = opened(artifacts, result)
    assert len(doc.tables) == 1
    assert [cell.text for cell in doc.tables[0].rows[1].cells] == ['完整技术要求及其适用条件。', '']
    assert '（原格式待补）' not in [p.text for p in doc.paragraphs]
    assert not any(gap['code'] == 'export_template_missing' for gap in result['gaps'])


def test_response_requirements_hide_parser_headings_without_changing_cross_page_sources(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    texts = ["#### （1）协同能力至少支持30Mbps。", "### 续页要求：提供接口 #1，不以型号代替功能。"]
    selected = excerpts(artifacts, texts)
    frozen_refs = deepcopy(selected)
    frozen_indexes = [artifacts.read_bytes(ref["source_ref"]["source_index_ref"]) for ref in selected]
    rows = [{"row_id": "cross-page", "source_excerpts": selected, "rule_refs": []}]

    doc = opened(artifacts, export_authoring_document(artifacts, outline(rows), [unit()]))

    assert [cell.text for cell in doc.tables[0].rows[1].cells] == [
        "（1）协同能力至少支持30Mbps。\n续页要求：提供接口 #1，不以型号代替功能。", ""]
    assert selected == frozen_refs
    assert [load_source_excerpt(artifacts, ref)["text"] for ref in selected] == texts
    assert [artifacts.read_bytes(ref["source_ref"]["source_index_ref"]) for ref in selected] == frozen_indexes


def test_row_answers_join_by_frozen_identity_across_units_and_conflicts_only_blank_their_row(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    selected = excerpts(artifacts, ["要求甲。", "要求乙。", "要求丙。"])
    rows = [{"row_id": identity, "source_excerpts": [excerpt], "rule_refs": []}
            for identity, excerpt in zip(("a", "b", "c"), selected)]
    first = unit(status="generated", row_fills=[{"row_id": "b", "text": "有据乙。", "material_ids": ["knowledge"]},
                                               {"row_id": "a", "text": "甲答案一。", "material_ids": ["knowledge"]}])
    second = unit(task_id="task-two", status="generated", row_fills=[
        {"row_id": "a", "text": "甲答案二。", "material_ids": ["knowledge"]},
        {"row_id": "unknown", "text": "不得串入正文。", "material_ids": ["knowledge"]}])
    result = export_authoring_document(artifacts, outline(rows), [first, second])
    doc = opened(artifacts, result)
    assert len(doc.tables) == 1
    assert [[cell.text for cell in row.cells] for row in doc.tables[0].rows[1:]] == [
        ["要求甲。", ""], ["要求乙。", "有据乙。"], ["要求丙。", ""]]
    assert "不得串入正文。" not in doc.element.xml
    assert {gap["code"] for gap in result["gaps"]} >= {"export_conflicting_row_fills", "export_row_fill_unplaced"}


def test_manual_response_table_keeps_requirements_and_ignores_answers(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    selected = excerpts(artifacts, ["本次报价上限及要求不得丢失。"])
    rows = [{"row_id": "price", "source_excerpts": selected, "rule_refs": []}]
    result = export_authoring_document(artifacts, outline(rows, manual_fill=True), [unit(
        status="generated", row_fills=[{"row_id": "price", "text": "123元", "material_ids": ["historical-price"]}])])
    doc = opened(artifacts, result)
    assert doc.tables[0].cell(1, 0).text == "本次报价上限及要求不得丢失。"
    assert doc.tables[0].cell(1, 1).text == ""
    assert "123元" not in doc.element.xml


def test_row_table_uses_selected_excerpt_not_entire_native_unit(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    selected = excerpts(artifacts, ["前文背景。必须支持双机热备。后文无关说明。"])
    source = EvidenceRepository(artifacts).source_index(selected[0]["source_ref"]["source_index_ref"])
    selected[0]["selectors"] = [selector(fragment(source.ranges[0], 5, 14))]
    rows = [{"row_id": "ha", "source_excerpts": selected, "rule_refs": []}]
    doc = opened(artifacts, export_authoring_document(artifacts, outline(rows), []))
    assert doc.tables[0].cell(1, 0).text == "必须支持双机热备。"
    assert "前文背景" not in doc.element.xml and "后文无关" not in doc.element.xml
