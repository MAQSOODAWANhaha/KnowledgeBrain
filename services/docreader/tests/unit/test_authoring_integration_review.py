"""Regressions found by reviewing the exact-source and row-response integration."""
from io import BytesIO
import json
from types import SimpleNamespace

from docx import Document
import pytest

from docreader.integrations.artifacts import ArtifactStore
from docreader.tender import authoring_pipeline as pipeline
from docreader.tender.authoring import _OutlineContext
from test_authoring_pipeline_boundaries import budget, payload, snapshot


def test_partial_source_routing_includes_parent_but_not_unselected_sibling_fragment():
    relations = [{'relation_id': identity, 'from_item_id': 'a', 'to_item_id': 'b',
                  'evidence': [{'range_id': selected}]} for identity, selected in (
                      ('selected', 'fragment-a'), ('parent', 'native'), ('sibling', 'fragment-b'))]
    context = SimpleNamespace(range_roots={'native': 'native', 'fragment-a': 'native', 'fragment-b': 'native'},
        payload={'rules': [], 'relations': relations, 'issues': []})
    precise = pipeline._analysis_context(context, [], ['fragment-a'])
    assert {relation['relation_id'] for relation in precise['relations']} == {'selected', 'parent'}
    complete = pipeline._analysis_context(context, [], ['native'])
    assert {relation['relation_id'] for relation in complete['relations']} == {'selected', 'parent', 'sibling'}


def test_selected_text_template_supports_keep_their_exact_provenance(tmp_path):
    from docreader.tender.authoring_templates import prepare_template_parts
    from docreader.tender.authoring_excerpts import load_source_excerpt
    from test_authoring_templates import create_context, create_source

    artifacts, index_ref, native = create_source(tmp_path, [{'text': '原格式甲：____\n原格式乙：____'}])
    context = create_context(artifacts, index_ref, native)
    _, views = prepare_template_parts(artifacts, context)
    task = {'task_id': 'task', 'chapter_id': 'chapter', 'title_path': ['采用原格式说明'], 'targets': [
        {'target_id': 'target', 'kind': 'narrative', 'purpose': '依据实际格式说明', 'rule_refs': [],
         'source_refs': [], 'support_origins': ['tender'],
         'template_refs': [context.template_parts[part['part_id']] for part in views]}]}
    value, bindings, _ = pipeline._load_task(artifacts, context, task, {'response': {'shared_arrangements': []}}, [])
    ids = value['targets'][0]['template_support_ids']
    assert len(ids) == 2
    assert [load_source_excerpt(artifacts, bindings[identity])['text'] for identity in ids] == [
        '原格式甲：____', '原格式乙：____']
    support_bindings = pipeline._support_bindings(bindings, [])
    assert support_bindings[ids[0]]['source_excerpt'] != support_bindings[ids[1]]['source_excerpt']


@pytest.mark.parametrize('reject_answer', [False, True])
async def test_derived_form_rows_are_written_and_review_can_reject_only_the_answer(tmp_path, reject_answer):
    artifacts = ArtifactStore(tmp_path)
    requirement = '系统须提供双机热备能力，维护时仍须保持现有业务。'
    fixed = snapshot(artifacts, [{'text': requirement}])
    context = _OutlineContext(artifacts, fixed)
    ranges = list(context.excerpts)
    outline = context.validate({'schema_version': 2, 'chapters': [{'title': '技术响应', 'children': [],
        'targets': [{'kind': 'form', 'purpose': '逐项填写功能响应', 'rule_ids': [], 'range_ids': [],
                     'response_table': {'rows': [{'range_ids': ranges}]}}]}], 'findings': []})
    calls = []

    async def execute(stage, request, validate):
        value = payload(request)
        calls.append(stage)
        if value.get('operation') == 'structure':
            return validate({'findings': []})
        if stage == 'authoring_basis':
            return validate({})
        if stage == 'authoring_write':
            assert value['response_rows'][0]['row_id'] == 'q1'
            assert not any(support.get('grid') for support in value['supports'])
            return validate({'blocks': [], 'row_fills': [
                {'row_id': 'q1', 'text': '支持双机热备。', 'support_ids': ['k1']}]})
        assert stage == 'authoring_review'
        record = value['units'][0]
        assert record['content']['row_fills'][0]['content_id'] == 'r1'
        return validate({'findings': [{'code': 'unsupported_content', 'detail': '该句的维护能力尚无支持。',
            'task_ids': [record['task_id']], 'content_ids': ['r1']}] if reject_answer else []})

    result = artifacts.read_json(await pipeline.generate_document(artifacts, fixed, budget(), execute,
        outline=outline, materials=[{'material_id': 'function', 'title': '已提供功能说明', 'text': '支持双机热备。'}]))
    assert calls.count('authoring_write') == 1 and calls.count('authoring_review') == 1
    unit = artifacts.read_json(result['unit_refs'][0])
    assert bool(unit['row_fills']) is not reject_answer
    document = Document(BytesIO(artifacts.read_bytes(result['docx_ref'])))
    table = next(table for table in document.tables if table.cell(0, 0).text == '招标要求')
    assert table.cell(1, 0).text == requirement
    assert table.cell(1, 1).text == ('' if reject_answer else '支持双机热备。')


async def test_assembly_preserves_known_row_missing_from_explicit_regrouping(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    requirements = ['须提供双机热备能力。', '故障恢复须保留全部审计日志。']
    fixed = snapshot(artifacts, [{'text': text} for text in requirements])

    async def execute(stage, request, validate):
        value = payload(request)
        if value.get('operation') == 'structure':
            return validate({'findings': []})
        assert stage == 'outline'
        if value.get('operation') == 'assemble':
            candidate = next(target for chapter in value['chapters'] for target in chapter['targets'])
            # A complete candidate was selected, but one row accidentally vanished
            # while its table grouping was rewritten. Do not hide it in rule_refs.
            return validate({'chapters': [{'title': '技术响应', 'children': [], 'targets': [
                {'kind': 'narrative', 'purpose': '逐项响应系统功能', 'target_ids': [candidate['target_id']],
                 'response_table': {'rows': [candidate['response_table']['rows'][0]]}}]}]})
        return validate({'schema_version': 2, 'chapters': [{'title': '技术响应', 'children': [],
            'targets': [{'kind': 'narrative', 'purpose': '逐项响应系统功能', 'rule_ids': [], 'range_ids': [],
                'response_table': {'rows': [{'range_ids': [source['range_id']]} for source in value['sources']]}}]}],
            'findings': []})

    saved = artifacts.read_json(await pipeline.generate_document(
        artifacts, fixed, budget(), execute, outline_only=True))
    rows = [row for chapter in pipeline._flatten(saved['chapters']) for target in chapter['targets']
            for row in (target.get('response_table') or {}).get('rows', [])]
    from docreader.tender.authoring_excerpts import load_source_excerpt
    assert [load_source_excerpt(artifacts, part)['text'] for row in rows for part in row['source_excerpts']] == requirements
    assert any(issue['code'] == 'outline_response_row_unassigned' for issue in saved['issues'])


async def test_assembly_cannot_exclude_a_rule_assigned_only_through_a_response_row(tmp_path):
    from docreader.llm.request import ModelContractError
    from test_authoring_templates import create_source
    from test_outline_source_binding import outline_context

    artifacts, _, native = create_source(tmp_path, [{'text': '须保留完整审计日志。'}])
    context, _ = outline_context(artifacts, native)

    async def execute(stage, request, validate):
        value = payload(request)
        if value.get('operation') == 'assemble':
            candidate = next(target for chapter in value['chapters'] for target in chapter['targets'])
            missing = value['unassigned_rules'][0]['rule_id']
            row = {**candidate['response_table']['rows'][0], 'rule_ids': [missing]}
            return validate({'chapters': [{'title': '技术响应', 'children': [], 'targets': [
                {'kind': 'narrative', 'purpose': '逐项响应', 'target_ids': [candidate['target_id']],
                 'response_table': {'rows': [row]}}]}],
                'excluded_rules': [{'rule_ids': [missing], 'reason': '同时错误声明本次不适用。'}]})
        assert stage == 'outline'
        return validate({'schema_version': 2, 'chapters': [{'title': '技术响应', 'children': [], 'targets': [
            {'kind': 'narrative', 'purpose': '逐项响应', 'rule_ids': [], 'range_ids': [],
             'response_table': {'rows': [{'range_ids': [value['sources'][0]['range_id']]}]}}]}], 'findings': []})

    with pytest.raises(ModelContractError) as caught:
        await pipeline._organize(artifacts, context, budget(), execute, [])
    assert caught.value.errors[0]['code'] == 'outline_conflicting_disposition'


async def test_accepted_structure_revision_keeps_new_constraints_when_tree_is_unchanged(tmp_path):
    from test_authoring_templates import create_source
    from test_outline_source_binding import outline_context

    artifacts, _, native = create_source(tmp_path, [{'text': '平台提交时须附可编辑版本。'}])
    context, _ = outline_context(artifacts, native)
    calls = []

    async def execute(stage, request, validate):
        value = payload(request)
        calls.append(value.get('operation', 'source'))
        if value.get('operation') == 'structure':
            return validate({'findings': [{'code': 'submission_constraint_missing',
                'detail': '提交格式要求尚未进入交付约束。', 'suggestion': '保留为提交待办。', 'can_revise': True}]})
        if value.get('operation') == 'assemble':
            candidate = next(target for chapter in value['chapters'] for target in chapter['targets'])
            return validate({'chapters': [{'title': '正文位置', 'children': [], 'targets': [
                {'kind': 'narrative', 'purpose': '保持已有正文位置', 'target_ids': [candidate['target_id']]}]}],
                'constraint_rule_ids': [value['unassigned_rules'][0]['rule_id']] if value.get('previous_structure') else []})
        assert stage == 'outline'
        return validate({'schema_version': 2, 'chapters': [{'title': '正文位置', 'children': [], 'targets': [
            {'kind': 'narrative', 'purpose': '保持已有正文位置', 'rule_ids': [],
             'range_ids': [value['sources'][0]['range_id']]}]}], 'findings': []})

    result = await pipeline._organize(artifacts, context, budget(), execute, [], structure_results={})
    assert result['constraint_rule_refs'] == list(context.rules.values())
    assert not any(issue['code'] == 'outline_rule_unassigned' for issue in result['issues'])
    assert calls == ['source', 'assemble', 'structure', 'assemble']
