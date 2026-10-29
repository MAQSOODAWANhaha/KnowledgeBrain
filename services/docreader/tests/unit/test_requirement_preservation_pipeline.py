"""Original tender requirements survive absent or rejected bidder responses."""

from copy import deepcopy
from io import BytesIO
import json
from zipfile import ZipFile

from docx import Document
import pytest

from docreader.integrations.artifacts import ArtifactStore
from docreader.llm.budget import RequestBudget
from docreader.llm.request import ModelWorkUnitExhausted
from docreader.tender.authoring import _OutlineContext
from docreader.tender.authoring_pipeline import generate_document
from docreader.tender.schema import AuthoringRequirements
from docreader.tests.submission_responses import assemble_independent_candidates

from test_analysis_checks import Model, run, scenario
from test_authoring_pipeline_boundaries import snapshot as source_snapshot
from test_outline_organization import authoring_input


REQUIREMENT = "在内网部署时，应支持独立角色权限配置；灾备演练环境除外。"
UNASSIGNED = "运维服务应每季度提交一次巡检报告。"
CONDITION = "仅内网部署场景适用。"
EXCEPTION = "灾备演练环境不适用。"
INVENTED = "本产品已建成三地容灾中心。"
QUOTE_REQUIREMENT = "投标报价不得超过最高限价60万元，最终报价由人工填写。"


async def frozen_requirements(tmp_path, *, requirement=REQUIREMENT, retained=None):
    data = await scenario(tmp_path, documents=[("main", "main", [
        (requirement, "obligation"), (UNASSIGNED, "obligation")])])
    artifacts, plan, items, _ = data
    extracted = artifacts.read_json(items["packs"][0]["artifact_ref"])
    if requirement == REQUIREMENT:
        extracted["items"][0].update(scope="内网部署", conditions=[CONDITION], exceptions=[EXCEPTION])
    items["packs"][0]["artifact_ref"] = artifacts.put_json(extracted)
    verification = await run(data, Model())
    if retained is not None:
        decision = {"decision_id": "accepted-partial", "amendment": {
            "target_item_id": extracted["items"][0]["item_id"], "retained_rules": retained,
            "replacement_item_ids": []}}
        verification["decisions_ref"] = artifacts.put_json([decision])
        verification["accepted_decision_ids"] = [decision["decision_id"]]
    fixed = authoring_input(artifacts, plan, items, verification)
    context = _OutlineContext(artifacts, fixed)
    expected = [deepcopy(entry.get("item", entry.get("effective_rule"))) for entry in context.payload["rules"]]
    return artifacts, fixed, expected


def budget():
    return RequestBudget(model="fixture", context_tokens=131072, output_reserve=16384)


def payload(request):
    return json.loads(request["messages"][1]["content"][0]["text"])


def document_text(artifacts, result):
    document = Document(BytesIO(artifacts.read_bytes(result["docx_ref"])))
    return "\n".join([paragraph.text for paragraph in document.paragraphs] +
                     [cell.text for table in document.tables for row in table.rows for cell in row.cells])


def submission_notes(artifacts, result):
    with ZipFile(BytesIO(artifacts.read_bytes(result["package_ref"]))) as package:
        return package.read("提交要求与待办.md").decode()


def executor(mode, calls, *, manual=False, source_only=False):
    async def execute(stage, request, validate):
        value = payload(request)
        operation = value.get("operation", stage)
        calls.append((operation, "previous_content" in value))
        if operation == "assemble":
            return validate(assemble_independent_candidates(value))
        if operation == "structure":
            return validate({"findings": []})
        if stage == "outline":
            if mode == "outline_exhausted":
                raise ModelWorkUnitExhausted(stage)
            selected = [value["rules"][0]["rule_id"]] if value["rules"] else []
            return validate({"schema_version": 2, "chapters": [{
                "title": "投标报价" if manual else "技术响应", "children": [], "targets": [{
                    "kind": "narrative", "purpose": "最终报价由人工补充。" if manual else "说明产品功能应答。",
                    "rule_ids": selected, "range_ids": [entry["range_id"] for entry in value["sources"]]
                    if source_only else [], "manual_fill": manual}]}], "findings": []})
        if stage == "authoring_basis":
            return validate({"shared_arrangements": []})
        if stage == "authoring_write":
            if mode == "write_exhausted" or "previous_content" in value:
                raise ModelWorkUnitExhausted(stage)
            support = next(entry["support_id"] for entry in value["supports"] if entry["origin"] == "knowledge")
            return validate({"blocks": [{"kind": "paragraph", "text": INVENTED,
                "target_ids": [value["targets"][0]["target_id"]], "support_ids": [support]}]})
        if stage == "authoring_review":
            assert mode in {"rejected", "revision_exhausted"}
            return validate({"findings": [{"code": "unsupported_content",
                "detail": "现有资料未支持已建成三地容灾中心。",
                "task_ids": [value["units"][0]["task_id"]],
                "can_revise": mode == "revision_exhausted"}]})
        pytest.fail(f"Unexpected model operation: {operation}")

    return execute


def preserved(artifacts, result):
    assert result.get("requirements_ref"), "Tender requirements need a snapshot independent of WritingUnit."
    raw = artifacts.read_json(result["requirements_ref"])
    with ZipFile(BytesIO(artifacts.read_bytes(result["package_ref"]))) as package:
        entry = json.loads(package.read("manifest.json"))["requirements"]
        assert entry["artifact_ref"] == result["requirements_ref"]
        assert json.loads(package.read(entry["package_path"])) == raw
    return AuthoringRequirements.model_validate(raw).model_dump(exclude_unset=True)


@pytest.mark.parametrize("mode", ["no_material", "write_exhausted", "rejected", "revision_exhausted"])
async def test_requirement_snapshot_and_docx_survive_response_failure_or_rejection(tmp_path, mode):
    artifacts, fixed, expected = await frozen_requirements(tmp_path)
    calls = []
    materials = [] if mode == "no_material" else [{"material_id": "permissions", "title": "功能说明",
        "text": "本产品支持角色权限管理。"}]
    result = artifacts.read_json(await generate_document(artifacts, fixed, budget(), executor(mode, calls),
                                                         materials=materials))
    requirements = preserved(artifacts, result)
    assert requirements["authoring_input"] == fixed
    assert requirements["rule_origin"] == "verified_projection"
    assert [record["rule"] for record in requirements["rules"]] == expected
    selected = next(record for record in requirements["rules"] if record["rule"]["text"] == REQUIREMENT)
    assert selected["target_ids"] == ["chapter-1:target-1"]
    assert selected["rule"]["conditions"] == [CONDITION] and selected["rule"]["exceptions"] == [EXCEPTION]
    units = [artifacts.read_json(ref) for ref in result["unit_refs"]]
    assert all(unit["status"] == "placeholder" and unit["blocks"] == [] for unit in units)
    text = document_text(artifacts, result)
    assert all(value not in text for value in (REQUIREMENT, CONDITION, EXCEPTION,
        "招标要求依据（企业应答另行填写）", "分析整理（待确认）"))
    notes = submission_notes(artifacts, result)
    assert REQUIREMENT in notes and CONDITION in notes and EXCEPTION in notes
    assert INVENTED not in text and "待补" in text and result["package_ref"]
    if mode == "no_material":
        assert calls == [("outline", False), ("assemble", False), ("structure", False)]
    else:
        assert sum(operation == "authoring_write" for operation, _ in calls) == (
            2 if mode == "revision_exhausted" else 1)
    if mode in {"rejected", "revision_exhausted"}:
        old = [artifacts.read_json(ref) for ref in result["superseded_unit_refs"]]
        assert any(block["text"] == INVENTED for unit in old for block in unit["blocks"])


async def test_unassigned_rule_remains_in_requirement_snapshot_without_false_target_mapping(tmp_path):
    artifacts, fixed, expected = await frozen_requirements(tmp_path)
    result = artifacts.read_json(await generate_document(artifacts, fixed, budget(), executor("no_material", [])))
    records = preserved(artifacts, result)["rules"]
    pending = next(record for record in records if record["rule"]["text"] == UNASSIGNED)
    assert pending["rule"] == expected[1] and pending["target_ids"] == []
    assert UNASSIGNED not in document_text(artifacts, result), "The host must not assign unknown semantic placement."
    assert UNASSIGNED in submission_notes(artifacts, result)
    assert any(gap["code"] == "export_requirement_unassigned" for gap in result["gaps"])
    outline = artifacts.read_json(result["outline_ref"])
    assert any(issue["code"] == "outline_rule_unassigned" for issue in outline["issues"])


async def test_exhausted_organization_still_delivers_all_requirements_in_the_independent_snapshot(tmp_path):
    artifacts, fixed, expected = await frozen_requirements(tmp_path)
    calls = []
    result = artifacts.read_json(await generate_document(artifacts, fixed, budget(), executor("outline_exhausted", calls)))
    records = preserved(artifacts, result)["rules"]
    assert [record["rule"] for record in records] == expected
    assert all(record["target_ids"] == [] for record in records)
    assert any(gap["code"] == "outline_batch_unavailable" for gap in result["gaps"])
    assert all(operation not in {"authoring_write", "authoring_basis"} for operation, _ in calls)
    assert INVENTED not in document_text(artifacts, result) and result["package_ref"]
    notes = submission_notes(artifacts, result)
    assert all(value in notes for value in (REQUIREMENT, CONDITION, EXCEPTION, UNASSIGNED))


async def test_manual_position_keeps_requirements_without_automatic_bidder_response(tmp_path):
    artifacts, fixed, expected = await frozen_requirements(tmp_path, requirement=QUOTE_REQUIREMENT)
    calls = []
    result = artifacts.read_json(await generate_document(artifacts, fixed, budget(),
                                                         executor("no_material", calls, manual=True)))
    assert [record["rule"] for record in preserved(artifacts, result)["rules"]] == expected
    text = document_text(artifacts, result)
    assert QUOTE_REQUIREMENT not in text and "人工填写" in text and "________" in text
    assert "招标要求依据（企业应答另行填写）" not in text and "分析整理（待确认）" not in text
    assert QUOTE_REQUIREMENT in submission_notes(artifacts, result)
    assert calls == [("outline", False), ("assemble", False), ("structure", False)]
    assert all(not artifacts.read_json(ref)["cell_fills"] for ref in result["unit_refs"])


async def test_effective_partial_rule_does_not_restore_old_conditions_from_original_item(tmp_path):
    _, _, original = await frozen_requirements(tmp_path / "original")
    retained = [{"text": "一标段应支持独立权限配置。", "scope": "仅一标段",
                 "evidence": original[0]["evidence"]}]
    artifacts, fixed, expected = await frozen_requirements(tmp_path / "retained", retained=retained)
    result = artifacts.read_json(await generate_document(artifacts, fixed, budget(), executor("no_material", [])))
    records = preserved(artifacts, result)["rules"]
    assert [record["rule"] for record in records] == expected
    active = next(record for record in records if record["rule"]["text"] == retained[0]["text"])
    assert active["rule"] == retained[0] and active["rule_ref"]["decision_id"] == "accepted-partial"
    assert "conditions" not in active["rule"] and "exceptions" not in active["rule"]
    text = document_text(artifacts, result)
    assert retained[0]["text"] not in text and CONDITION not in text and EXCEPTION not in text
    notes = submission_notes(artifacts, result)
    assert retained[0]["text"] in notes and CONDITION not in notes and EXCEPTION not in notes
    assert all(entry["quote"] in notes for entry in retained[0]["evidence"])


async def test_source_only_context_is_retained_as_reference_without_becoming_response_text(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    context_text = "仅供评委内部使用的评审提示，不属于应答正文。"
    fixed = source_snapshot(artifacts, [{"text": context_text}])
    calls = []
    result = artifacts.read_json(await generate_document(artifacts, fixed, budget(),
                                                         executor("no_material", calls, source_only=True)))
    requirements = preserved(artifacts, result)
    assert requirements["rule_origin"] == "source_only" and requirements["rules"] == []
    assert context_text not in document_text(artifacts, result)
    outline = artifacts.read_json(result["outline_ref"])
    assert outline["chapters"][0]["targets"][0]["source_refs"]
    assert result["package_ref"]
    assert not any(issue["code"] == "outline_target_unbound" for issue in outline["issues"])
    assert all(not artifacts.read_json(ref)["blocks"] for ref in result["unit_refs"])
    assert all(operation not in {"authoring_basis", "authoring_write"} for operation, _ in calls)
