"""Outline source choices resolve to loadable native templates, including XLSX grids."""

from copy import deepcopy
import io

import pytest

from docreader.integrations.artifacts import ArtifactStore
from docreader.llm.request import ModelContractError
from docreader.llm.wire import WirePayload
from docreader.tender.authoring import _OutlineContext, load_outline_source
from docreader.tender.evidence import freeze_sources
from docreader.tender.evidence.references import EvidenceRepository, SourceIndex, fragment, selector

from test_structural_packing import source


def outline_context(artifacts, index, *, body=None, evidence_source=None):
    raw_ref = artifacts.put_json(index.raw)
    source_ref = artifacts.put_json(index.descriptor(raw_ref))
    identity = {"run_id": "outline-native", "source_version": 1, "domain_revision": 1}
    frozen_ref = artifacts.put_json({"schema_version": 2, "documents": [
        {**index.document, "role": "main", "source_index_ref": source_ref, "raw_ref": raw_ref}]})
    pack = {"schema_version": 2, "pack_id": "native-pack", **index.document,
            "source_index_ref": source_ref,
            "body": [selector(row) for row in (index.ranges if body is None else body)],
            "supporting": []}
    pack_ref = artifacts.put_json(pack)
    plan_ref = artifacts.put_json({"schema_version": 2, "run_identity": identity,
        "frozen_ref": frozen_ref, "packs": [{"pack_id": "native-pack", "artifact_ref": pack_ref}]})
    evidence_source = evidence_source or next(row for row in index.ranges if row["text"])
    item = {"item_id": "native-item", "kind": "submission_requirement", "text": evidence_source["text"],
            "evidence": [{"range_id": evidence_source["range_id"], "quote": evidence_source["text"]}]}
    extraction_ref = artifacts.put_json({"schema_version": 2, "items": [item]})
    items_ref = artifacts.put_json({"schema_version": 2,
        "packs": [{"pack_id": "native-pack", "artifact_ref": extraction_ref}]})
    verification_ref = artifacts.put_json({"schema_version": 2, "items_ref": items_ref,
        "item_refs": [{"artifact_ref": extraction_ref, "item_id": item["item_id"]}],
        "relations_ref": artifacts.put_json([]), "unverified_relation_ids": [], "issues": [],
        "decisions_ref": artifacts.put_json([]), "accepted_decision_ids": []})
    value = {"schema_version": 2, **identity, "checklist_digest": "a" * 64,
        "frozen_ref": frozen_ref, "pack_plan_ref": plan_ref, "items_ref": items_ref,
        "verification_ref": verification_ref, "confirmation_ref": artifacts.put_json({
            "source_version": 1, "domain_revision": 1, "checklist_digest": "a" * 64})}
    return _OutlineContext(artifacts, value), source_ref


def form_response(ctx, range_ids):
    return {"schema_version": 2, "chapters": [{"title": "规定响应格式", "children": [],
        "targets": [{"kind": "form", "purpose": "保留原始格式填写人员及响应说明。",
                     "rule_ids": [next(iter(ctx.rules))], "range_ids": range_ids}]}], "findings": []}


def spreadsheet(artifacts, table_count):
    """Rows reference shared cells; only grid units retain the complete table."""
    tables, rows, cells = [], [], {}
    for table_number in range(table_count):
        unit_id = f"grid:{table_number}"
        start_column = 1 + table_number * 3
        grid_cells = [{"row": row, "column": column, "row_span": 1, "col_span": 1,
                       "text": f"表{table_number}行{row}列{column}"}
                      for row in range(2) for column in range(2)]
        tables.append({"unit_id": unit_id, "kind": "table", "text": "", "ordinal": table_number,
            "grid": {"row_count": 2, "column_count": 2, "cells": grid_cells},
            "header_cells": [{"row": 0, "column": column} for column in range(2)],
            "locator": {"locator_kind": "spreadsheet", "sheet_ordinal": 0,
                "region": {"start_row": 1, "end_row": 2, "start_column": start_column,
                           "end_column": start_column + 1}}})
        for cell in grid_cells:
            row, column = cell["row"] + 1, cell["column"] + start_column
            cells[f"r{row}:c{column}"] = {"row": row, "column": column,
                "grid_cell_ref": {"unit_id": unit_id, "row": cell["row"], "column": cell["column"]},
                "grid_value_fields": ["text"]}
    for row in range(1, 3):
        rows.append({"unit_id": f"row:{row}", "kind": "table_row", "text": "", "ordinal": row,
            "locator": {"locator_kind": "spreadsheet", "sheet_ordinal": 0,
                "region": {"start_row": row, "end_row": row, "start_column": 1,
                           "end_column": 2 + (table_count - 1) * 3},
                "cell_refs": [key for key, cell in cells.items() if cell["row"] == row]}})
    units = tables + rows
    raw = {"schema_version": 2, "units": units, "spreadsheet_cells": cells,
        "images": {}, "ocr": {}, "source_contract": {"page_manifest": [], "units": [
            {"unit_id": unit["unit_id"], "section_id": "sheet:0", "parent_section_id": None,
             "heading_path": "人员及偏离表", "physical_locator": unit["locator"]} for unit in units]}}
    document = {"document_id": "spreadsheet", "document_version": "1",
                "original_ref": artifacts.put_bytes(b"synthetic spreadsheet original")}
    return SourceIndex(document, raw), tables


def test_virtual_blank_pages_are_neither_catalog_templates_nor_selectable_sources(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    native = source([{"text": "投标文件须提供人员表。"}])
    raw = deepcopy(native.raw)
    raw["source_contract"]["page_manifest"] = [{"page_ordinal": 1, "classification": "blank", "unit_ids": []}]
    document = {**native.document, "original_ref": artifacts.put_bytes(b"synthetic document original")}
    index = SourceIndex(document, raw)
    blank = next(row for row in index.ranges if row["unit_id"] == "blank:page:1")
    ctx, _ = outline_context(artifacts, index)
    assert blank["range_id"] not in {row["range_id"] for row in ctx.payload["source_catalog"]}
    with pytest.raises(ModelContractError) as caught:
        ctx.validate(form_response(ctx, [blank["range_id"]]))
    assert caught.value.errors[0]["code"] == "outline_unknown_reference"
    assert caught.value.errors[0]["path"][-1] == "range_ids"
    for entry in ctx.payload["source_catalog"]:
        outlined = ctx.validate(form_response(ctx, [entry["range_id"]]))
        assert all(load_outline_source(artifacts, ref)["unit"]["unit_id"] == ref["unit_id"]
                   for ref in outlined["chapters"][0]["targets"][0]["source_refs"])


@pytest.mark.parametrize("table_count", [1, 2])
def test_spreadsheet_row_choices_bind_all_complete_grid_units_and_read_full_tables(tmp_path, table_count):
    artifacts = ArtifactStore(tmp_path)
    index, tables = spreadsheet(artifacts, table_count)
    assert len(index.ranges) == 2 and all(row["unit_id"].startswith("row:") for row in index.ranges)
    assert all("cell_refs" in unit["locator"] and "grid" not in unit
               for unit in index.raw["units"] if unit["kind"] == "table_row")
    ctx, source_ref = outline_context(artifacts, index)
    outlined = ctx.validate(form_response(ctx, [index.ranges[0]["range_id"]]))
    target = outlined["chapters"][0]["targets"][0]
    assert target["source_refs"] == [
        {"source_index_ref": source_ref, "unit_id": table["unit_id"]} for table in tables]
    assert [load_outline_source(artifacts, ref)["unit"] for ref in target["source_refs"]] == tables
    read_ids = {row["range_id"] for view in ctx.payload["source_views"]
                for role in ("body", "supporting") for row in view.get(role, [])}
    assert {row["range_id"] for row in index.ranges} <= read_ids
    assert not any(issue["code"] in {"outline_source_not_read", "outline_template_missing"}
                   for issue in outlined["issues"])


def test_multiple_spreadsheet_rows_do_not_duplicate_shared_template_grids(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    index, tables = spreadsheet(artifacts, 2)
    ctx, source_ref = outline_context(artifacts, index)
    selected = [row["range_id"] for row in index.ranges]
    outlined = ctx.validate(form_response(ctx, [*selected, selected[0]]))
    target = outlined["chapters"][0]["targets"][0]
    assert target["source_refs"] == [
        {"source_index_ref": source_ref, "unit_id": table["unit_id"]} for table in tables]


def test_fragment_evidence_uses_text_free_alias_and_short_selection_preserves_full_native_source(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    text = "项目背景：服务于本项目。\n服务响应须在2小时内完成，并提交人员表。"
    native = source([{"text": text, "heading": "服务响应", "section": "service"}])
    index = SourceIndex({**native.document,
                         "original_ref": artifacts.put_bytes(b"synthetic fragment original")}, native.raw)
    original = index.ranges[0]
    selected = fragment(original, text.index("服务响应须"), original["end"])
    ctx, source_ref = outline_context(artifacts, index, body=[selected], evidence_source=selected)
    plan = artifacts.read_json(ctx.input["pack_plan_ref"])
    pack = EvidenceRepository(artifacts).load_pack(plan["packs"][0])
    assert pack["body"][0]["range_id"] == selected["range_id"]
    assert pack["body"][0]["text"] == selected["text"]
    assert ctx.payload["rules"][0]["item"]["evidence"][0]["range_id"] == selected["range_id"]

    view = ctx.payload["source_views"][0]
    assert [(entry["range_id"], entry["text"]) for entry in view["body"]] == [(original["range_id"], text)]
    assert view["supporting"] == []
    assert view["source_aliases"] == [{"range_id": selected["range_id"], "segments": [
        {"range_id": original["range_id"], "start": selected["start"], "end": selected["end"]}]}]

    wire = WirePayload("outline", ctx.payload)
    assert wire.enabled
    raw = form_response(ctx, [wire.mapping["r"][selected["range_id"]]])
    raw["chapters"][0]["targets"][0]["rule_ids"] = [wire.mapping["e"][next(iter(ctx.rules))]]
    decoded = wire.decode(raw)
    assert decoded["chapters"][0]["targets"][0]["range_ids"] == [selected["range_id"]]
    saved = artifacts.read_json(artifacts.put_json(ctx.validate(decoded)))
    references = saved["chapters"][0]["targets"][0]["source_refs"]
    assert references == [{"source_index_ref": source_ref, "unit_id": original["unit_id"]}]
    loaded = load_outline_source(artifacts, references[0])
    assert loaded["unit"] == index.raw["units"][0]
    assert loaded["unit"]["text"] == text
    assert not any(issue["code"] == "outline_source_not_read" for issue in saved["issues"])


@pytest.mark.parametrize("ocr", [
    {"text": "投标人须逐项填写技术偏离表并盖章。", "status": "complete", "issues": []},
    {"text": "部分可识别内容", "status": "unresolved", "issues": ["ocr_incomplete"]},
    {"text": "", "status": "complete", "issues": [],
     "disposition": {"status": "reviewed", "classification": "decorative", "reason_code": "logo"}},
])
def test_image_source_roundtrip_preserves_frozen_ocr_image_and_anchor(tmp_path, ocr):
    artifacts = ArtifactStore(tmp_path)
    native = source([{"text": "须按扫描原表填写。"}, {"text": ""}])
    raw = deepcopy(native.raw)
    unit = raw["units"][1]
    unit.update(kind="image_region", locator={"locator_kind": "image", "image_key": "scan.png"})
    raw["source_contract"]["units"][1]["physical_locator"] = unit["locator"]
    image_ref = artifacts.put_bytes(b"synthetic frozen image", "image/png")
    raw["images"] = {"scan.png": image_ref, "unrelated.png": image_ref}
    raw["ocr"][unit["unit_id"]] = deepcopy(ocr)
    anchor = {"sheet_ordinal": 0, "from": {"row": 3, "column": 2, "row_offset_emu": 125}}
    if ocr["status"] == "complete":
        raw["image_anchors"] = {unit["unit_id"]: anchor}
    document = {**native.document, "original_ref": artifacts.put_bytes(b"synthetic scan original")}
    index = SourceIndex(document, raw)
    frozen = deepcopy(raw)
    ctx, source_ref = outline_context(artifacts, index)
    chosen = next(row for row in index.ranges if row["unit_id"] == unit["unit_id"])
    outlined = artifacts.read_json(artifacts.put_json(ctx.validate(form_response(ctx, [chosen["range_id"]]))))
    references = outlined["chapters"][0]["targets"][0]["source_refs"]
    assert references == [{"source_index_ref": source_ref, "unit_id": unit["unit_id"]}]
    before = {p: p.read_bytes() for p in artifacts.root.rglob("*") if p.is_file()}
    loaded = load_outline_source(artifacts, references[0])
    assert loaded["unit"] == unit and loaded["unit"]["text"] == ""
    assert loaded["document"] == document
    assert loaded["ocr"] == ocr
    assert loaded["image_ref"] == image_ref
    assert artifacts.read_bytes(loaded["image_ref"]) == b"synthetic frozen image"
    if "image_anchors" in raw:
        assert loaded["image_anchor"] == anchor
        loaded["image_anchor"]["from"]["row"] = 99
        assert load_outline_source(artifacts, references[0])["image_anchor"] == anchor
    else:
        assert "image_anchor" not in loaded
    loaded["ocr"]["issues"].append("caller mutation")
    assert load_outline_source(artifacts, references[0])["ocr"] == ocr
    assert index.raw == frozen
    assert {p: p.read_bytes() for p in artifacts.root.rglob("*") if p.is_file()} == before


@pytest.mark.parametrize("table_count", [1, 2])
@pytest.mark.parametrize("use_fragment", [False, True])
def test_spreadsheet_row_preserves_cells_outside_defined_tables(tmp_path, table_count, use_fragment):
    artifacts = ArtifactStore(tmp_path)
    native, tables = spreadsheet(artifacts, table_count)
    raw = deepcopy(native.raw)
    row = next(unit for unit in raw["units"] if unit["unit_id"] == "row:1")
    column = table_count * 3
    key = f"r1:c{column}"
    outside = {"row": 1, "column": column, "text": "必须另附盖章证明。", "raw_value": "必须另附盖章证明。"}
    raw["spreadsheet_cells"][key] = outside
    row["locator"]["cell_refs"].append(key)
    row["locator"]["region"]["end_column"] = column
    index = SourceIndex(native.document, raw)
    original = index.ranges[0]
    assert outside["text"] in original["text"]
    chosen = fragment(original, 0, original["text"].index(" | ")) if use_fragment else original
    ctx, source_ref = outline_context(artifacts, index, body=[chosen], evidence_source=chosen)
    outlined = artifacts.read_json(artifacts.put_json(ctx.validate(form_response(ctx, [chosen["range_id"]]))))
    references = outlined["chapters"][0]["targets"][0]["source_refs"]
    assert references == [{"source_index_ref": source_ref, "unit_id": table["unit_id"]} for table in tables] + [
        {"source_index_ref": source_ref, "unit_id": row["unit_id"]}]
    loaded = [load_outline_source(artifacts, ref) for ref in references]
    assert [entry["unit"] for entry in loaded] == [*tables, row]
    assert loaded[-1]["spreadsheet_cells"][key] == outside
    assert len({ref["unit_id"] for ref in references}) == len(references)
    assert not any(issue["code"] == "outline_source_not_read" for issue in outlined["issues"])


@pytest.mark.parametrize("unit_id", ["grid:0", "row:2"])
def test_spreadsheet_loader_resolves_shared_cell_values_and_preserves_formula_metadata(tmp_path, unit_id):
    artifacts = ArtifactStore(tmp_path)
    native, _ = spreadsheet(artifacts, 1)
    raw = deepcopy(native.raw)
    cell = raw["spreadsheet_cells"]["r2:c2"]
    inherited = ["text", "raw_value", "display_text", "cached_value"]
    cell.update(grid_value_fields=inherited, formula="=A2*2", formula_references=["A2"],
                value_type="formula", cached_value_type="numeric", display_complete=False,
                display_incomplete_reason="number_format_unavailable")
    grid = raw["units"][0]["grid"]
    grid["cells"][-1]["text"] = "6"
    raw["units"][0]["locator"]["cell_refs"] = list(raw["spreadsheet_cells"])
    index = SourceIndex(native.document, raw)
    raw_ref = artifacts.put_json(raw)
    source_ref = artifacts.put_json(index.descriptor(raw_ref))
    reference = {"source_index_ref": source_ref, "unit_id": unit_id}
    expected_unit = next(unit for unit in raw["units"] if unit["unit_id"] == unit_id)
    before = deepcopy(raw)
    loaded = load_outline_source(artifacts, reference)
    assert loaded["unit"] == expected_unit
    assert set(loaded["spreadsheet_cells"]) == set(expected_unit["locator"]["cell_refs"])
    assert loaded["spreadsheet_cells"]["r2:c2"] == {**cell, **dict.fromkeys(inherited, "6")}
    loaded["spreadsheet_cells"]["r2:c2"]["formula_references"].append("Z99")
    loaded["unit"]["locator"]["cell_refs"].clear()
    reloaded = load_outline_source(artifacts, reference)
    assert reloaded["unit"] == expected_unit
    assert reloaded["spreadsheet_cells"]["r2:c2"]["formula_references"] == ["A2"]
    assert index.raw == before and artifacts.read_json(raw_ref) == before


async def test_native_workbook_freeze_to_outline_preserves_outside_cells_formula_and_image(tmp_path):
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as ExcelImage
    from openpyxl.worksheet.table import Table
    from PIL import Image

    artifacts = ArtifactStore(tmp_path / "artifacts")
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["必须另附盖章证明。", "项目名称", "合同金额"])
    sheet.append(["表外说明", "服务项目", "=1+2"])
    sheet.add_table(Table(displayName="RequiredForm", ref="B1:C2"))
    pixels, output = io.BytesIO(), io.BytesIO()
    Image.new("RGB", (12, 8), "blue").save(pixels, format="PNG")
    pixels.seek(0)
    sheet.add_image(ExcelImage(pixels), "E3")
    workbook.save(output)
    original_ref = artifacts.put_bytes(output.getvalue())
    frozen_ref = await freeze_sources({"documents": [{"document_id": "native-workbook", "document_version": "1",
        "role": "main", "ordinal": 0, "file_name": "form.xlsx", "original_ref": original_ref}]}, artifacts,
        ocr=lambda image_ref, context: {"text": "扫描原表不得更改。", "status": "complete", "issues": []}, plan=False)
    frozen = artifacts.read_json(frozen_ref)
    index = EvidenceRepository(artifacts).source_index(frozen["documents"][0]["source_index_ref"])
    chosen = next(row for row in index.ranges if "必须另附盖章证明。" in row["text"])
    image = next(row for row in index.ranges if row["kind"] == "image_region")
    ctx, _ = outline_context(artifacts, index, evidence_source=chosen)
    outlined = ctx.validate(form_response(ctx, [chosen["range_id"], image["range_id"]]))
    loaded = [load_outline_source(artifacts, ref) for ref in outlined["chapters"][0]["targets"][0]["source_refs"]]
    grid = next(entry for entry in loaded if entry["unit"].get("grid"))
    row = next(entry for entry in loaded if entry["unit"]["kind"] == "table_row")
    scan = next(entry for entry in loaded if entry["unit"]["kind"] == "image_region")
    assert any(cell["text"] == "必须另附盖章证明。" for cell in row["spreadsheet_cells"].values())
    assert any(cell.get("formula") == "=1+2" for cell in grid["spreadsheet_cells"].values())
    assert scan["ocr"] == index.raw["ocr"][image["unit_id"]]
    assert scan["image_anchor"] == index.raw["image_anchors"][image["unit_id"]]
    assert artifacts.read_bytes(scan["image_ref"])


@pytest.mark.parametrize("merge", ["A1:A2", "A4:C4"])
async def test_outside_merge_survives_freeze_and_outline_source_loading(tmp_path, merge):
    from openpyxl import Workbook
    from openpyxl.worksheet.table import Table

    artifacts = ArtifactStore(tmp_path / "artifacts")
    workbook = Workbook()
    sheet = workbook.active
    for address, value in {"D1": "项目", "E1": "响应", "D2": "技术", "E2": "提供"}.items():
        sheet[address] = value
    sheet.add_table(Table(displayName="DefinedForm", ref="D1:E2"))
    sheet[merge.split(":")[0]] = "必须保留表外合并证明说明。"
    sheet.merge_cells(merge)
    output = io.BytesIO()
    workbook.save(output)
    original_ref = artifacts.put_bytes(output.getvalue())
    frozen_ref = await freeze_sources({"documents": [{"document_id": "merge-workbook", "document_version": "1",
        "role": "main", "ordinal": 0, "file_name": "merge.xlsx", "original_ref": original_ref}]}, artifacts, plan=False)
    frozen = artifacts.read_json(frozen_ref)
    index = EvidenceRepository(artifacts).source_index(frozen["documents"][0]["source_index_ref"])
    chosen = next(row for row in index.ranges if "必须保留表外合并证明说明。" in row["text"])
    ctx, source_ref = outline_context(artifacts, index, evidence_source=chosen)
    outlined = artifacts.read_json(artifacts.put_json(ctx.validate(form_response(ctx, [chosen["range_id"]]))))
    references = outlined["chapters"][0]["targets"][0]["source_refs"]
    row_ref = {"source_index_ref": source_ref, "unit_id": chosen["unit_id"]}
    assert row_ref in references
    native = next(unit for unit in index.raw["units"] if unit["unit_id"] == chosen["unit_id"])
    loaded = load_outline_source(artifacts, row_ref)
    assert loaded["unit"] == native
    assert [region["a1_range"] for region in loaded["unit"]["locator"]["merged_ranges"]] == [merge]
    assert artifacts.read_bytes(loaded["document"]["original_ref"]) == output.getvalue()
