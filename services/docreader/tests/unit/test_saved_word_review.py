"""Review input must come from the immutable saved Word, including its tables."""
from io import BytesIO

from docx import Document

from docreader.tender.word_navigation import add_bookmark


def saved_bytes():
    document = Document()
    heading = document.add_heading("人工改名章节", level=1)
    add_bookmark(heading._p, "chapter-original")
    document.add_paragraph("人工修订后的实施内容，旧稿中不存在。")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "工期"
    table.cell(0, 1).text = "人工填写：30天"
    document.add_heading("人工新增章节", level=1)
    document.add_paragraph("此处也是实际保存内容。")
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_saved_word_records_include_human_text_tables_and_unbound_content():
    from docreader.tender.saved_word_review import read_saved_word

    records, gaps = read_saved_word(saved_bytes(), [{"chapter_id": "chapter-original", "title": "旧标题", "children": [], "targets": []}])
    text = str(records)
    assert "人工修订后的实施内容" in text and "人工填写：30天" in text and "此处也是实际保存内容" in text
    assert records[0]["chapter_id"] == "chapter-original"
    assert any(record["chapter_id"] is None for record in records)
    assert any(gap["code"] == "saved_word_unbound_content" for gap in gaps)


def test_unanchored_document_is_reviewed_without_guessing_chapter_from_title():
    from docreader.tender.saved_word_review import read_saved_word

    document = Document()
    document.add_heading("旧标题", level=1)
    document.add_paragraph("同名标题不构成原章关系。")
    stream = BytesIO()
    document.save(stream)
    records, gaps = read_saved_word(stream.getvalue(), [{"chapter_id": "original", "title": "旧标题", "children": [], "targets": []}])
    assert all(record["chapter_id"] is None for record in records)
    assert any(gap["code"] == "saved_word_missing_chapter" for gap in gaps)


async def test_saved_word_review_only_reports_without_adopting_generated_content(tmp_path, monkeypatch):
    import json

    from docreader.integrations.artifacts import ArtifactStore
    from docreader.llm.budget import RequestBudget
    from docreader.tender import saved_word_review

    record = {"task_id": "saved-1", "chapter_id": "chapter-original", "title": "人工保存章节",
              "actual_content": {"kind": "paragraph", "text": "人工保存的当前正文。"}}
    monkeypatch.setattr(saved_word_review, "_prepare_review", lambda *_args: (
        [record], [], {"chapter-original": [{"rules": [], "supports": []}]}, {}))
    calls = []

    async def execute(stage, request, validate):
        payload = json.loads(request["messages"][1]["content"][0]["text"])
        calls.append(payload)
        assert stage == "authoring_review" and payload["operation"] == "review_saved_word"
        assert payload["records"][0]["actual_content"] == record["actual_content"]
        return validate({"findings": [{"code": "source_meaning_changed", "detail": "人工稿条件需要核对。",
            "task_ids": ["saved-1"], "can_revise": True, "reject_content": True}]})

    report = await saved_word_review.review_saved_word(ArtifactStore(tmp_path), {}, saved_bytes(),
        RequestBudget(model="fixture", context_tokens=131072, prompts={}), execute)
    assert len(calls) == report["reviewed_record_count"] == 1
    finding, = report["findings"]
    assert finding["code"] == "source_meaning_changed" and finding["content_ids"] == []
    assert finding["can_revise"] is False
    assert finding["reject_content"] is False, "Saved Word review reports opinions without content adoption."
