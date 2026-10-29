"""Display parser heading decoration without altering frozen template evidence."""

from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile

from docx import Document
import pytest

from docreader.integrations.artifacts import ArtifactStore
from docreader.tender.authoring_export import export_authoring_document
from docreader.tender.authoring_templates import load_template_part
from docreader.tender.evidence.references import EvidenceRepository, fragment, selector

from test_submission_export import chapter, grid, part, source, target, writing


@pytest.mark.parametrize("original,display", [
    ("### 6. 我方理解：____", "6. 我方理解：____"),
    ("# 固定声明\n正文：____\n## 1. 承诺内容", "固定声明\n正文：____\n1. 承诺内容"),
    ("   ######\t签章：____", "   签章：____"),
])
def test_template_heading_decoration_is_display_only(tmp_path, original, display):
    artifacts = ArtifactStore(tmp_path)
    reference = source(artifacts, "declaration", text=original)
    selected = part(reference)
    outline = {"chapters": [chapter("声明", [target("declaration", templates=[selected])], volume=True)]}
    before = deepcopy(outline)
    source_bytes = artifacts.read_bytes(reference["source_index_ref"])

    exported = export_authoring_document(artifacts, outline, [])

    with ZipFile(BytesIO(artifacts.read_bytes(exported["package_ref"]))) as archive:
        for content in [artifacts.read_bytes(exported["docx_ref"]), archive.read("volumes/001.docx")]:
            word = Document(BytesIO(content))
            assert display in [paragraph.text for paragraph in word.paragraphs]
            assert original not in [paragraph.text for paragraph in word.paragraphs]
    assert outline == before
    assert artifacts.read_bytes(reference["source_index_ref"]) == source_bytes
    assert load_template_part(artifacts, selected)["unit"]["text"] == original


@pytest.mark.parametrize("text", [
    "规格#型号；使用C#；编号A#3；签章：____",
    "正文含 # 号和 ### 标识，不是标题。",
    "###没有分隔空白，不作为标题",
    "####### 七个标记不作为标准标题",
    "    ### 缩进代码片段保留",
])
def test_template_non_heading_hashes_are_preserved(tmp_path, text):
    artifacts = ArtifactStore(tmp_path)
    reference = source(artifacts, "specification", text=text)
    outline = {"chapters": [chapter("格式", [target("specification", templates=[part(reference)])])]}

    exported = export_authoring_document(artifacts, outline, [])

    word = Document(BytesIO(artifacts.read_bytes(exported["docx_ref"])))
    assert text in [paragraph.text for paragraph in word.paragraphs]


def test_selected_template_display_retains_exact_frozen_selector(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    selected_text = "### 6. 固定声明：____"
    original = f"其他格式前文\n{selected_text}\n另一声明"
    reference = source(artifacts, "separate-declarations", text=original)
    base = EvidenceRepository(artifacts).source_index(reference["source_index_ref"]).ranges[0]
    start = original.index(selected_text)
    selected = {"source_ref": reference, "selectors": [selector(fragment(base, start, start + len(selected_text)))]}
    outline = {"chapters": [chapter("声明", [target("declaration", templates=[selected])])]}
    before = deepcopy(outline)

    exported = export_authoring_document(artifacts, outline, [])

    paragraphs = [paragraph.text for paragraph in Document(BytesIO(artifacts.read_bytes(exported["docx_ref"]))).paragraphs]
    assert "6. 固定声明：____" in paragraphs
    assert not any("其他格式前文" in value or "另一声明" in value for value in paragraphs)
    assert outline == before
    assert load_template_part(artifacts, selected)["text"] == selected_text


def test_shared_form_instances_display_each_existing_purpose_before_its_template(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    reference = source(artifacts, "financial", text="8D 财务格式", grid=grid())
    selected = part(reference)
    targets = [target("bidder", templates=[selected]), target("manufacturer", templates=[selected])]
    purposes = ["投标人财务情况；与原厂同主体时按适用要求处理", "原厂财务情况；与投标人不同主体时填写"]
    for item, purpose in zip(targets, purposes):
        item["purpose"] = purpose
    outline = {"chapters": [chapter("财务", targets, volume=True)]}
    units = [writing("w1", "财务", ["bidder"], status="placeholder"),
             writing("w2", "财务", ["manufacturer"], status="placeholder")]
    before = deepcopy(outline)

    exported = export_authoring_document(artifacts, outline, units)

    with ZipFile(BytesIO(artifacts.read_bytes(exported["package_ref"]))) as archive:
        for content in [artifacts.read_bytes(exported["docx_ref"]), archive.read("volumes/001.docx")]:
            word = Document(BytesIO(content))
            paragraphs = [paragraph.text for paragraph in word.paragraphs]
            assert all(paragraphs.count(purpose) == 1 for purpose in purposes)
            body = list(word.element.body)
            captions = [next(p for p in word.paragraphs if p.text == purpose) for purpose in purposes]
            assert body.index(captions[0]._p) < body.index(word.tables[0]._tbl)
            assert body.index(word.tables[0]._tbl) < body.index(captions[1]._p) < body.index(word.tables[1]._tbl)
            assert [table.cell(1, 1).text for table in word.tables] == ["", ""]
    assert outline == before
    assert load_template_part(artifacts, selected)["grid"] == grid()


def test_duplicate_part_within_one_instance_does_not_add_a_purpose_caption(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    reference = source(artifacts, "single", grid=grid())
    selected = part(reference)
    item = target("one", templates=[selected, deepcopy(selected)])
    item["purpose"] = "只填写一份；重复片段不是第二个实例"
    outline = {"chapters": [chapter("财务", [item])]}

    exported = export_authoring_document(artifacts, outline, [writing("w", "财务", ["one"], status="placeholder")])

    word = Document(BytesIO(artifacts.read_bytes(exported["docx_ref"])))
    assert len(word.tables) == 1
    assert item["purpose"] not in [paragraph.text for paragraph in word.paragraphs]


def test_same_template_in_separate_chapters_does_not_add_purpose_captions(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    reference = source(artifacts, "shared-source", grid=grid())
    first = target("one", templates=[part(reference)])
    second = target("two", templates=[part(reference)])
    first["purpose"], second["purpose"] = "甲章节填写要求", "乙章节填写要求"
    outline = {"chapters": [chapter("甲", [first]), chapter("乙", [second])]}
    units = [writing("w1", "甲", ["one"], status="placeholder"), writing("w2", "乙", ["two"], status="placeholder")]

    exported = export_authoring_document(artifacts, outline, units)

    word = Document(BytesIO(artifacts.read_bytes(exported["docx_ref"])))
    assert len(word.tables) == 2
    assert not {first["purpose"], second["purpose"]} & {paragraph.text for paragraph in word.paragraphs}
