"""Explicit layout semantics survive into the editable submission artifacts."""

from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Mm
import pytest

from docreader.integrations.artifacts import ArtifactStore
from docreader.tender.authoring_export import export_authoring_document
from docreader.tender.evidence.references import SourceIndex


def chapter(identity, *, targets=(), children=(), role="body", volume=False, title=None):
    return {"chapter_id": identity, "title": title or identity, "targets": list(targets),
            "children": list(children), "presentation_role": role, "is_volume": volume}


def target(identity, *, templates=(), related=(), rules=(), manual=False):
    return {"target_id": identity, "kind": "form" if templates else "narrative",
            "purpose": "填写" + identity, "template_refs": list(templates), "source_refs": [],
            "rule_refs": list(rules), "related_target_ids": list(related), "manual_fill": manual}


def writing(identity, chapter_id, target_ids, text):
    return {"task_id": identity, "chapter_id": chapter_id, "target_ids": list(target_ids),
            "status": "generated", "blocks": [{"kind": "paragraph", "text": text}],
            "cell_fills": [], "attachments": [], "gaps": []}


def word(artifacts, result, key="docx_ref"):
    return Document(BytesIO(artifacts.read_bytes(result[key])))


def native_part(artifacts, *, headers=(), count=None):
    grid = {"row_count": 3, "column_count": 2, "widths_mm": [60, 120], "cells": [
        {"row": 0, "column": 0, "row_span": 1, "col_span": 2, "text": "规定原表"},
        {"row": 1, "column": 0, "row_span": 1, "col_span": 1, "text": "功能要求"},
        {"row": 1, "column": 1, "row_span": 1, "col_span": 1, "text": ""},
        {"row": 2, "column": 0, "row_span": 1, "col_span": 1, "text": "最终报价"},
        {"row": 2, "column": 1, "row_span": 1, "col_span": 1, "text": ""}]}
    unit = {"unit_id": "table", "kind": "table_region", "ordinal": 0, "text": "", "grid": grid,
            "header_cells": [{"row": row, "column": 0, "role": "column_header"} for row in headers],
            "locator": {"locator_kind": "document", "section_ordinal": 0}}
    raw = {"schema_version": 2, "units": [unit], "images": {}, "ocr": {}, "source_contract": {
        "page_manifest": [], "units": [{"unit_id": "table", "section_id": "section:0",
            "parent_section_id": None, "heading_path": "原格式", "physical_locator": unit["locator"]}]}}
    document = {"document_id": "template", "document_version": "1",
                "original_ref": artifacts.put_bytes(b"synthetic original")}
    index = SourceIndex(document, raw)
    index_ref = artifacts.put_json(index.descriptor(artifacts.put_json(index.raw)))
    part = {"source_ref": {"source_index_ref": index_ref, "unit_id": "table"}, "selectors": []}
    if count is not None:
        part["header_row_count"] = count
    return part, raw


def header_rows(table):
    return [index for index, row in enumerate(table.rows)
            if row._tr.find(qn("w:trPr")) is not None
            and row._tr.find(qn("w:trPr")).find(qn("w:tblHeader")) is not None]


def test_explicit_cover_precedes_directory_without_chapter_number_or_title_guessing(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    outline = {"chapters": [chapter("ordinary", title="封面"), chapter("actual", role="cover"),
                            chapter("response", title="技术应答")]}
    original = deepcopy(outline)
    units = [writing("cover", "actual", ["cover-target"], "实际封面内容，企业名称____。")]
    result = export_authoring_document(artifacts, outline, units)
    document = word(artifacts, result)
    text = [p.text for p in document.paragraphs]
    assert text.index("实际封面内容，企业名称____。") < text.index("目录")
    headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]
    assert headings == ["1 封面", "2 技术应答"]
    assert text.count("实际封面内容，企业名称____。") == 1
    assert outline == original


def test_each_volume_keeps_its_own_cover_and_formal_cross_references(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    outline = {"chapters": [chapter("root-cover", role="cover"),
        chapter("commercial", title="商务文件", volume=True, children=[
            chapter("commercial-cover", role="cover"),
            chapter("declaration", title="投标函", targets=[target("letter", related=["price", "price"])])]),
        chapter("pricing", title="报价文件", volume=True, children=[
            chapter("pricing-cover", role="cover"),
            chapter("prices", title="分项报价", targets=[target("price", manual=True)])])]}
    units = [writing(identity, identity, [], text) for identity, text in [
        ("root-cover", "整稿原封面"), ("commercial-cover", "商务原封面"), ("pricing-cover", "报价原封面")]]
    result = export_authoring_document(artifacts, outline, units)
    document = word(artifacts, result)
    text = [p.text for p in document.paragraphs]
    assert text.index("整稿原封面") < text.index("目录") < text.index("商务原封面")
    assert text.index("商务原封面") < text.index("1 商务文件") < text.index("报价原封面")
    assert "参见 2.1 分项报价" in text
    assert text.count("参见 2.1 分项报价") == 1
    with ZipFile(BytesIO(artifacts.read_bytes(result["package_ref"]))) as package:
        commercial = Document(BytesIO(package.read("volumes/001.docx")))
        pricing = Document(BytesIO(package.read("volumes/002.docx")))
    commercial_text, pricing_text = ([p.text for p in doc.paragraphs] for doc in (commercial, pricing))
    assert commercial_text.index("商务原封面") < commercial_text.index("目录")
    assert pricing_text.index("报价原封面") < pricing_text.index("目录")
    assert "报价原封面" not in commercial_text and "商务原封面" not in pricing_text
    assert "参见 2.1 分项报价" in commercial_text


def test_directory_has_updatable_toc_field_and_existing_page_field(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    result = export_authoring_document(artifacts, {"chapters": [chapter("body", title="正文")]}, [])
    with ZipFile(BytesIO(artifacts.read_bytes(result["docx_ref"]))) as archive:
        from lxml import etree

        xml = etree.fromstring(archive.read("word/document.xml"))
        settings = etree.fromstring(archive.read("word/settings.xml"))
        namespaces = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        instructions = xml.xpath("//w:instrText/text() | //w:fldSimple/@w:instr", namespaces=namespaces)
        assert any('TOC' in instruction and '\\o "1-9"' in instruction for instruction in instructions)
        update = settings.find(qn("w:updateFields"))
        assert update is not None and update.get(qn("w:val")) in {"true", "1"}
        footer = archive.read("word/footer1.xml").decode()
        assert 'w:instr="PAGE"' in footer
    assert [p.text for p in word(artifacts, result).paragraphs].count("1 正文") == 1


def test_cover_marked_as_volume_does_not_silently_drop_its_source_content(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    outline = {"chapters": [chapter("cover-volume", role="cover", volume=True,
        children=[chapter("body", title="响应正文")])]}
    result = export_authoring_document(artifacts, outline,
        [writing("cover-content", "cover-volume", [], "必须保留的原封面。")])
    text = [paragraph.text for paragraph in word(artifacts, result).paragraphs]
    assert text.count("必须保留的原封面。") == 1
    assert text.index("必须保留的原封面。") < text.index("目录") < text.index("1 响应正文")


@pytest.mark.parametrize(("native_headers", "explicit_count", "expected"), [
    ([], 1, [0]), ([0], 0, []), ([0], None, [0]), ([1], None, [])])
def test_native_table_layout_uses_explicit_or_native_headers_and_keeps_original_cells(
        tmp_path, native_headers, explicit_count, expected):
    artifacts = ArtifactStore(tmp_path)
    part, raw = native_part(artifacts, headers=native_headers, count=explicit_count)
    original = deepcopy(raw)
    outline = {"chapters": [chapter("table", targets=[target("form", templates=[part], manual=True)])]}
    result = export_authoring_document(artifacts, outline, [])
    document = word(artifacts, result)
    table = document.tables[0]
    assert header_rows(table) == expected
    assert table.cell(0, 0)._tc is table.cell(0, 1)._tc
    assert table.cell(0, 0).text == "规定原表"
    assert table.cell(1, 0).text == "功能要求" and table.cell(1, 1).text == ""
    assert table.cell(2, 0).text == "最终报价" and table.cell(2, 1).text == ""
    assert all(cell._tc.get_or_add_tcPr().find(qn("w:vAlign")).get(qn("w:val")) == "top"
               for row in table.rows for cell in row.cells)
    assert sum(column.width for column in table.columns) <= Mm(160) + Mm(0.1)
    assert table.columns[1].width / table.columns[0].width == pytest.approx(2, rel=.01)
    assert raw == original


def test_requirement_table_reserves_readable_response_width_and_repeated_headers(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    reference = {"item_ref": {"item_id": "rule", "artifact_ref": artifacts.put_json({"rule": True})}}
    requirements = {"rules": [{"rule_ref": reference, "target_ids": ["response"], "rule": {
        "text": "至少提供三年售后服务。", "evidence": [{"range_id": "range", "quote": "至少提供三年售后服务。"}]}}]}
    result = export_authoring_document(artifacts, {"chapters": [chapter("service", targets=[
        target("response", rules=[reference])])]}, [], requirements=requirements)
    table = word(artifacts, result).tables[0]
    assert header_rows(table) == [0]
    widths = [column.width for column in table.columns]
    assert [value / sum(widths) for value in widths] == pytest.approx([.4, .35, .25], rel=.01)
    assert table.cell(1, 2).text == ""


def test_package_submission_notes_keep_constraints_unassigned_rules_and_review_opinions(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    refs = [{"item_ref": {"item_id": value, "artifact_ref": artifacts.put_json({value: True})}}
            for value in ("upload", "pending", "body")]
    texts = ["须上传可编辑文件压缩包。", "是否适用需核对的未安排要求。", "已有正文位置要求。"]
    requirements = {"rules": [{"rule_ref": ref, "target_ids": [], "rule": {"text": text,
        "evidence": [{"range_id": str(index), "quote": text}]}}
        for index, (ref, text) in enumerate(zip(refs, texts))]}
    outline = {"chapters": [chapter("response", targets=[target("body", rules=[refs[2]])])],
               "constraint_rule_refs": [refs[0]]}
    result = export_authoring_document(artifacts, outline, [], [{"code": "cross_volume_reference",
        "detail": "报价表跨卷引用待核对，并非缺少企业资料。", "chapter_id": "response"}], requirements=requirements)
    with ZipFile(BytesIO(artifacts.read_bytes(result["package_ref"]))) as package:
        notes = package.read("提交要求与待办.md").decode()
    assert "已知提交与编制约束" in notes and texts[0] in notes
    assert "尚未安排的要求" in notes and texts[1] in notes
    assert texts[2] not in notes.split("尚未安排的要求")[1].split("当前核对意见")[0]
    assert "当前核对意见" in notes and "并非缺少企业资料" in notes


def test_explicit_rule_exclusion_is_not_reported_as_unassigned_in_document_or_notes(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    reference = {"item_ref": {"item_id": "excluded", "artifact_ref": artifacts.put_json({"excluded": True})}}
    requirements = {"rules": [{"rule_ref": reference, "target_ids": [], "rule": {
        "text": "仅另一标段需要此证明。", "evidence": [{"range_id": "other", "quote": "仅另一标段需要此证明。"}]}}]}
    outline = {"chapters": [], "issues": [{"code": "outline_rule_excluded", "detail": "原文明示本标段不适用。",
                                         "rule_refs": [reference]}]}
    result = export_authoring_document(artifacts, outline, [], requirements=requirements)
    assert not any(gap["code"] == "export_requirement_unassigned" for gap in result["gaps"])
    with ZipFile(BytesIO(artifacts.read_bytes(result["package_ref"]))) as package:
        notes = package.read("提交要求与待办.md").decode()
    assert "仅另一标段需要此证明。" not in notes.split("尚未安排的要求")[1].split("人工填写位置")[0]
