"""Real-parser contract tests; OCR text is deterministic and not semantic acceptance."""
from __future__ import annotations

import io
from dataclasses import replace

import pytest
from docx import Document
from docx.oxml import OxmlElement
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from PIL import Image
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject, NumberObject

from docreader.integrations.artifacts import ArtifactStore
from docreader.tender.evidence import EvidenceRepository, freeze_sources, register_manifest


def save_office(document):
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def png():
    output = io.BytesIO()
    Image.new("RGB", (12, 8), "blue").save(output, format="PNG")
    return output.getvalue()


def pdf(*, tiny_image=False):
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    font = writer._add_object(DictionaryObject({NameObject('/Type'): NameObject('/Font'),
        NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')}))
    resources = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
    commands = b'BT /F1 12 Tf 10 270 Td (Bidder must submit proof within 30 days.) Tj ET'
    if tiny_image:
        pixels = Image.new('RGB', (12, 8), 'blue')
        image = DecodedStreamObject()
        image.set_data(pixels.tobytes())
        image.update({NameObject('/Type'): NameObject('/XObject'), NameObject('/Subtype'): NameObject('/Image'),
                      NameObject('/Width'): NumberObject(12), NameObject('/Height'): NumberObject(8),
                      NameObject('/ColorSpace'): NameObject('/DeviceRGB'), NameObject('/BitsPerComponent'): NumberObject(8)})
        resources[NameObject('/XObject')] = DictionaryObject({NameObject('/Im0'): writer._add_object(image.flate_encode())})
        # Two actual occurrences sharing one payload; both must have source identity.
        commands += b' q 12 0 0 8 10 10 cm /Im0 Do Q q 12 0 0 8 40 10 cm /Im0 Do Q'
    page[NameObject('/Resources')] = resources
    stream = DecodedStreamObject()
    stream.set_data(commands)
    page[NameObject('/Contents')] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def registration(tmp_path, artifacts, documents):
    entries = []
    for name, content in documents:
        path = tmp_path / name
        path.write_bytes(content)
        entries.append({"document_id": name, "version": "1", "role": "main", "path": str(path)})
    return register_manifest({"documents": entries}, artifacts)


@pytest.fixture
def artifacts(tmp_path):
    return ArtifactStore(tmp_path / "artifacts")


@pytest.fixture(autouse=True)
def single_pdf_renderer(monkeypatch):
    from docreader.config import CONFIG
    monkeypatch.setattr("docreader.parser.pdf_parser.CONFIG", replace(CONFIG, pdf_render_parallelism=1,
                       pdf_render_dpi=72, pdf_render_max_edge=300))


def test_registration_copies_inputs_and_reports_all_document_problems(tmp_path, artifacts):
    first = tmp_path / 'main.pdf'
    first.write_bytes(pdf())
    record = register_manifest({"documents": [
        {"document_id": "main", "version": 1, "path": str(first)},
        {"document_id": "copy", "version": 1, "path": str(first)},
        {"document_id": "main", "version": 1, "path": str(tmp_path / 'missing')},
        {"document_id": "directory", "version": 1, "path": str(tmp_path)},
    ], "expected_document_ids": ["amendment"]}, artifacts)
    codes = {entry["code"] for entry in record["issues"]}
    assert {"duplicate_document_content", "duplicate_document_id", "missing_document",
            "unreadable_document", "missing_expected_document"} <= codes
    first.unlink()
    assert artifacts.read_bytes(record["documents"][0]["original_ref"]) == pdf()
    assert all('path' not in document for document in record['documents'])
    replay = register_manifest(artifacts.read_json(artifacts.put_json(record)), artifacts)
    assert replay == record


def test_registration_accepts_submit_ref_aliases(tmp_path, artifacts):
    ref = artifacts.put_bytes(pdf())
    output = register_manifest({"documents": [{"document_id": "a", "version": "v1", "role": "attachment",
                                                "file_ref": ref, "filename": "a.pdf"}]}, artifacts)
    assert output['documents'][0]['original_ref'] == ref
    assert output['documents'][0]['file_name'] == 'a.pdf'
    assert output['issues'] == []


async def test_text_pdf_never_renders_or_calls_ocr(tmp_path, artifacts, monkeypatch):
    def forbidden(*args, **kwargs):
        raise AssertionError('native text page must not render')
    monkeypatch.setattr('docreader.parser.pdf_parser._render_scanned_pages', forbidden)
    monkeypatch.setattr('docreader.parser.pdf_parser.FORCE_SCANNED_PDF', True)
    registered = registration(tmp_path, artifacts, [('native.pdf', pdf())])
    ref = await freeze_sources(registered, artifacts, forbidden)
    index = artifacts.read_json(ref)
    assert index['issues'] == []
    assert index['stats']['ocr_request_count'] == 0
    assert '30 days' in str(list(EvidenceRepository(artifacts).iter_packs(ref)))


async def test_pdf_tiny_shared_images_keep_each_occurrence_without_page_render(tmp_path, artifacts, monkeypatch):
    calls = []
    async def ocr(ref, source):
        calls.append((ref, source))
        assert artifacts.read_bytes(ref)
        return {'text': 'Stamped proof required', 'status': 'complete', 'issues': []}
    def forbidden(*args, **kwargs):
        raise AssertionError('text page must not render')
    monkeypatch.setattr('docreader.parser.pdf_parser._render_scanned_pages', forbidden)
    registered = registration(tmp_path, artifacts, [('native.pdf', pdf(tiny_image=True))])
    ref = await freeze_sources(registered, artifacts, ocr)
    index = artifacts.read_json(ref)
    assert index['issues'] == []
    assert len(calls) == 2
    assert len({source['unit_id'] for _, source in calls}) == 2
    assert len({item['sha256'] for item, _ in calls}) == 1
    assert all(source['locator']['left'] is not None for _, source in calls)


async def test_mixed_native_pdf_only_scanned_page_uses_ocr(tmp_path, artifacts):
    from docreader.scripts.generate_source_contract_fixtures import mixed_pdf
    calls = []
    async def ocr(ref, source):
        calls.append(source['locator']['page_ordinal'])
        return {'text': 'Scanned evidence 10 / 20', 'status': 'complete', 'issues': []}
    ref = await freeze_sources(registration(tmp_path, artifacts, [('mixed.pdf', mixed_pdf())]), artifacts, ocr)
    assert calls == [2]
    packs = list(EvidenceRepository(artifacts).iter_packs(ref))
    assert any(r.get('non_requirement_reason') == 'blank_page' for pack in packs for r in pack['body'])
    assert artifacts.read_json(ref)['issues'] == []


async def test_xlsx_images_anchors_blank_rows_and_partial_fields_survive(tmp_path, artifacts):
    workbook = Workbook()
    for sheet in [workbook.active, workbook.create_sheet('Other')]:
        sheet.append(['Field', 'Response'])
        sheet.append(['Name', None])
        sheet.append([None, None])
        sheet.append(['Value', 'required'])
        sheet.add_image(ExcelImage(io.BytesIO(png())), 'C2')
    calls = []
    async def ocr(ref, source):
        calls.append(source)
        return {'text': 'Sign here', 'status': 'complete', 'issues': []}
    frozen = await freeze_sources(registration(tmp_path, artifacts, [('forms.xlsx', save_office(workbook))]), artifacts, ocr)
    index = artifacts.read_json(frozen)
    assert index['issues'] == []
    assert len(calls) == 2
    assert {source['image_anchor']['sheet_ordinal'] for source in calls} == {0, 1}
    assert all(source['image_anchor']['from']['row'] == 2 and source['image_anchor']['from']['column'] == 3 for source in calls)
    packs = list(EvidenceRepository(artifacts).iter_packs(frozen))
    rows = [entry for pack in packs for entry in pack['body']]
    names = [entry for entry in rows if entry['text'].startswith('Name')]
    assert len(names) == 2 and all(entry['text'] == 'Name | ' for entry in names)
    blank = [entry for entry in rows if entry.get('empty')]
    assert len(blank) == 2 and all(entry['text'] == '' for entry in blank)
    assert sum(len(pack['compacted_blank_ranges']) for pack in packs) == 2
    raw = artifacts.read_json(index['documents'][0]['raw_ref'])
    grids = [entry['grid'] for entry in raw['units'] if entry.get('grid')]
    assert len(grids) == 2 and all(grid['row_count'] == 4 for grid in grids)
    assert len({entry['range_id'] for entry in rows}) == len(rows)


async def test_docx_raw_grid_merge_and_cross_pack_context(tmp_path, artifacts):
    document = Document()
    document.add_heading('第一章 条件', 1)
    document.add_paragraph('本章适用于联合体，必须同时提交全部成员资料。')
    document.add_heading('第一节 响应', 2)
    document.add_paragraph('以下内容属于联合体响应。')
    for i in range(8):
        document.add_paragraph(f'必须提供第{i}项相关证明。' * 10)
    table = document.add_table(rows=8, cols=2)
    table.cell(0, 0).text = '证明字段'
    table.cell(0, 1).text = '响应内容'
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    table.cell(1, 0).merge(table.cell(1, 1)).text = '合并说明'
    for row in range(2, 8):
        table.cell(row, 0).text = f'材料{row}'
        table.cell(row, 1).text = '必须盖章'
    document.add_heading('第二章 结束', 1)
    document.add_paragraph('必须提交独立附件。')
    ref = await freeze_sources(registration(tmp_path, artifacts, [('source.docx', save_office(document))]),
                               artifacts, range_chars=180, pack_chars=700)
    index = artifacts.read_json(ref)
    packs = list(EvidenceRepository(artifacts).iter_packs(ref))
    assert len(packs) > 2
    body_ids = [entry['range_id'] for pack in packs for entry in pack['body']]
    assert len(body_ids) == len(set(body_ids)) == index['range_count']
    assert any('本章适用于联合体' in entry['text'] for pack in packs[1:] for entry in pack['supporting'])
    assert any(entry.get('reason') == 'header' for pack in packs for entry in pack['supporting'])
    assert all(not ({entry['range_id'] for entry in pack['body']} &
                    {entry['range_id'] for entry in pack['supporting']}) for pack in packs)
    raw = artifacts.read_json(index['documents'][0]['raw_ref'])
    grid = next(entry['grid'] for entry in raw['units'] if entry.get('grid'))
    assert any(cell['col_span'] == 2 for cell in grid['cells'])
    assert all(entry['physical_path'] for entry in raw['units'])


async def test_failed_document_and_missing_ocr_do_not_discard_successful_ranges(tmp_path, artifacts):
    registered = registration(tmp_path, artifacts, [('good.pdf', pdf()), ('broken.docx', b'broken'), ('scan.png', png()), ('unknown.bin', b'unknown')])
    ref = await freeze_sources(registered, artifacts)
    index = artifacts.read_json(ref)
    assert {'document_parse_failed', 'ocr_provider_unavailable', 'unsupported_document'} <= {problem['code'] for problem in index['issues']}
    assert any(document['document_id'] == 'good.pdf' and document['status'] == 'complete' for document in index['documents'])
    assert any('30 days' in entry['text'] for pack in EvidenceRepository(artifacts).iter_packs(ref) for entry in pack['body'])


async def test_repository_reuses_frozen_refs_without_native_parse(tmp_path, artifacts, monkeypatch):
    ref = await freeze_sources(registration(tmp_path, artifacts, [('native.pdf', pdf())]), artifacts)
    monkeypatch.setattr('docreader.tender.evidence.freeze.parse_native', lambda *args: pytest.fail('must reuse frozen source'))
    repository = EvidenceRepository(artifacts)
    first = list(repository.iter_packs(ref))
    assert first == list(repository.iter_packs(ref))
    index = artifacts.read_json(ref)
    assert repository.load_pack(index['packs'][0]) == first[0]


async def test_small_sections_merge_while_ranges_keep_their_identities(tmp_path, artifacts):
    document = Document()
    for i in range(3):
        document.add_heading(f'Chapter {i}', 1)
        document.add_paragraph(f'Submit document {i}.')
    ref = await freeze_sources(registration(tmp_path, artifacts, [('small.docx', save_office(document))]), artifacts)
    packs = list(EvidenceRepository(artifacts).iter_packs(ref))
    assert len(packs) == 1
    assert len({entry['section_id'] for entry in packs[0]['body']}) == 3


async def test_formula_display_incomplete_is_not_silent_success(tmp_path, artifacts):
    workbook = Workbook()
    workbook.active.append(['Total', '=SUM(B2:B3)'])
    frozen = await freeze_sources(registration(tmp_path, artifacts, [('formula.xlsx', save_office(workbook))]), artifacts)
    index = artifacts.read_json(frozen)
    assert any(problem['code'] == 'spreadsheet_display_incomplete' and problem['detail'] == 'formula_cached_value_missing'
               for problem in index['issues'])
    assert index['documents'][0]['status'] == 'partial'


async def test_embedded_pdf_attachment_is_explicitly_unresolved(tmp_path, artifacts):
    from pypdf import PdfReader
    writer = PdfWriter()
    writer.append(PdfReader(io.BytesIO(pdf())))
    writer.add_attachment('terms.txt', b'Additional conditions apply.')
    output = io.BytesIO()
    writer.write(output)
    frozen = await freeze_sources(registration(tmp_path, artifacts, [('attachment.pdf', output.getvalue())]), artifacts)
    index = artifacts.read_json(frozen)
    assert any(problem['code'] == 'attachment_content_not_extracted' for problem in index['issues'])
    raw = artifacts.read_json(index['documents'][0]['raw_ref'])
    assert any(unit['locator'].get('part_name') == 'terms.txt' for unit in raw['units'])


async def test_storage_corruption_remains_technical_failure(tmp_path, artifacts):
    from docreader.integrations.artifacts import ArtifactIntegrityError
    registered = registration(tmp_path, artifacts, [('source.pdf', pdf())])
    path = artifacts.path(registered['documents'][0]['original_ref'])
    path.chmod(0o644)
    path.write_bytes(b'corrupt')
    with pytest.raises(ArtifactIntegrityError):
        await freeze_sources(registered, artifacts)


async def test_changed_original_gets_new_range_identity_even_when_manifest_version_is_reused(tmp_path, artifacts):
    initial = registration(tmp_path, artifacts, [('source.pdf', pdf())])
    frozen = await freeze_sources(initial, artifacts)
    changed = registration(tmp_path, artifacts, [('source.pdf', pdf().replace(b'30 days', b'60 days'))])
    revised = await freeze_sources(changed, artifacts)
    repository = EvidenceRepository(artifacts)
    original_ranges = [entry for pack in repository.iter_packs(frozen) for entry in pack['body']]
    revised_ranges = [entry for pack in repository.iter_packs(revised) for entry in pack['body']]
    assert not {entry['range_id'] for entry in original_ranges} & {entry['range_id'] for entry in revised_ranges}
    assert any('30 days' in entry['text'] for entry in original_ranges)
    assert any('60 days' in entry['text'] for entry in revised_ranges)


async def test_offset_adjacent_excel_tables_each_keep_their_own_headers(tmp_path, artifacts):
    from openpyxl.worksheet.table import Table
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(['First field', 'First value'])
    for row in range(2, 8):
        sheet.cell(row, 1).value = f'First {row}'
        sheet.cell(row, 2).value = 'Required'
    sheet['D3'], sheet['E3'] = 'Second field', 'Second value'
    for row in range(4, 8):
        sheet.cell(row, 4).value = f'Second {row}'
        sheet.cell(row, 5).value = 'Required'
    sheet.add_table(Table(displayName='FirstTable', ref='A1:B7'))
    sheet.add_table(Table(displayName='SecondTable', ref='D3:E7'))
    frozen = await freeze_sources(registration(tmp_path, artifacts, [('tables.xlsx', save_office(workbook))]),
                                   artifacts, pack_chars=350)
    packs = list(EvidenceRepository(artifacts).iter_packs(frozen))
    later = next(pack for pack in packs if any('Second 7' in entry['text'] for entry in pack['body']))
    assert any('Second field' in entry['text'] and entry['reason'] == 'header' for entry in later['supporting'])
    assert any('First field' in entry['text'] and entry['reason'] == 'header' for entry in later['supporting'])
