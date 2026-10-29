"""Regression boundaries found while reviewing the document-authoring path."""

from copy import deepcopy
import json

import pytest

from docreader.integrations.artifacts import ArtifactStore
from docreader.llm.budget import RequestBudget
from docreader.llm.request import ModelContractError
from docreader.tender.authoring_pipeline import bind_writing_response, generate_document
from docreader.tender.evidence.references import SourceIndex
from docreader.tests.submission_responses import assemble_independent_candidates, provided_template_ids

from test_structural_packing import source


def budget():
    return RequestBudget(model="fixture", context_tokens=131072, prompts={}, output_reserve=16384)


def snapshot(artifacts, specs=()):
    value = {"schema_version": 2, "run_id": "boundary-run", "source_version": 1, "domain_revision": 1}
    if specs:
        native = source(specs)
        native = SourceIndex({**native.document, "original_ref": artifacts.put_bytes(b"synthetic original")}, native.raw)
        raw_ref = artifacts.put_json(native.raw)
        index_ref = artifacts.put_json(native.descriptor(raw_ref))
        value["frozen_ref"] = artifacts.put_json({"schema_version": 2, "documents": [
            {**native.document, "role": "main", "raw_ref": raw_ref, "source_index_ref": index_ref}]})
    return value


def payload(request):
    return json.loads(request["messages"][1]["content"][0]["text"])


def outline(value, *, kind="narrative", chapters=1):
    return {"schema_version": 2, "chapters": [
        {"title": f"响应部分{number + 1}", "targets": [
            {"kind": kind, "purpose": "根据给定原文填写本响应位置", "rule_ids": [r["rule_id"] for r in value["rules"]],
             "range_ids": [s["range_id"] for s in value.get("sources", [])],
             "template_part_ids": provided_template_ids(value) if kind == "form" else []}], "children": []}
        for number in range(chapters)], "findings": []}


def ordinary_response(task, value):
    if task == "outline":
        if value.get("operation") == "assemble":
            return assemble_independent_candidates(value)
        return outline(value)
    if task == "authoring_basis":
        return {"approach": "根据可用资料起草；未知事实保留待补。"}
    if task == "authoring_write":
        return {"blocks": [{"kind": "paragraph", "text": "待补充实际资料。",
            "target_ids": [value["targets"][0]["target_id"]],
            "support_ids": [entry["support_id"] for entry in value["supports"] if entry["origin"] == "knowledge"]}]}
    return {"findings": []}


async def test_empty_outline_retains_a_located_placeholder_instead_of_a_cover_only_document(tmp_path):
    artifacts = ArtifactStore(tmp_path)

    async def execute(task, request, validate):
        if task == "outline":
            value = payload(request)
            if value.get("operation") == "assemble":
                return validate(assemble_independent_candidates(value))
            return validate({"schema_version": 2, "chapters": [], "findings": []})
        if task == "authoring_write":
            return validate({"blocks": []})
        return validate(ordinary_response(task, payload(request)))

    result = artifacts.read_json(await generate_document(artifacts, snapshot(artifacts), budget(), execute))
    assert result["unit_refs"], "An empty model outline must retain an explicit writing position."
    assert all(artifacts.read_json(ref)["status"] == "placeholder" for ref in result["unit_refs"])
    assert result["gaps"]
    assert artifacts.read_bytes(result["docx_ref"]).startswith(b"PK")


def test_whitespace_only_table_headers_are_not_generated_content():
    task = {"task_id": "w1", "chapter_id": "chapter-1", "title_path": ["技术方案"], "targets": [
        {"target_id": "target-1", "kind": "narrative", "purpose": "技术说明", "rule_refs": [], "source_refs": []}]}
    unit = bind_writing_response(task, {"blocks": [
        {"kind": "table", "headers": ["", "  "], "rows": [["", "\n"]]}]}, {}, [])
    assert unit["status"] == "placeholder" and unit["gaps"]


async def test_attachment_selection_cannot_take_proof_from_another_target(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    original = artifacts.put_bytes(b"synthetic second-lot certificate", "text/plain")
    materials = [{"material_id": "other-lot", "title": "二标段证明", "text": "仅供二标段使用的证明",
                  "original_ref": original, "attachment": True, "usage": "proof",
                  "locator": {"target_ids": ["chapter-2:target-1"]}}]
    visible_materials = []

    async def execute(task, request, validate):
        value = payload(request)
        if task == "outline":
            if value.get("operation") == "assemble":
                return validate(assemble_independent_candidates(value))
            organized = outline(value, kind="attachment", chapters=2)
            # Narrative may select proof too; it must not borrow another target's
            # original, even when a model returns that material's identity.
            organized["chapters"][0]["targets"][0]["kind"] = "narrative"
            return validate(organized)
        if task == "authoring_write":
            visible_materials.append([m["support_id"] for m in value["supports"] if m["origin"] == "knowledge"])
            try:
                return validate({"blocks": [], "attachment_ids": ["k1"]})
            except ModelContractError:
                # A durable executor would supply this feedback on its correction attempt.
                return validate({"blocks": [], "gaps": ["本位置证明尚未提供。"]})
        return validate(ordinary_response(task, value))

    result = artifacts.read_json(await generate_document(artifacts, snapshot(artifacts), budget(), execute,
                                                         materials=materials))
    first, second = [artifacts.read_json(ref) for ref in result["unit_refs"]]
    assert visible_materials == [["k1"]]
    assert first["attachments"] == [] and first["gaps"]
    assert second["attachments"][0]["artifact_ref"] == original


async def test_review_receives_cell_and_attachment_content_without_persistent_cas_metadata(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    fixed = snapshot(artifacts, [{"heading": "人员表", "grid": {
        "row_count": 1, "column_count": 2, "cells": [
            {"row": 0, "column": 0, "text": "姓名", "row_span": 1, "col_span": 1},
            {"row": 0, "column": 1, "text": "", "row_span": 1, "col_span": 1}]}}])
    original = artifacts.put_bytes(b"synthetic bidder supplied proof", "text/plain")
    materials = [{"material_id": "proof", "title": "给定证明", "text": "给定证明的人员姓名为张三。",
                  "original_ref": original, "attachment": True, "usage": "proof"}]
    reviews = []

    async def execute(task, request, validate):
        value = payload(request)
        if task == "outline":
            if value.get("operation") == "assemble":
                return validate(assemble_independent_candidates(value))
            return validate(outline(value, kind="form"))
        if task == "authoring_write":
            table = next(entry for entry in value["supports"] if entry.get("grid"))
            return validate({"blocks": [], "cell_fills": [
                {"source_id": table["support_id"], "row": 0, "column": 1, "text": "张三",
                 "target_ids": [value["targets"][0]["target_id"]], "support_ids": ["k1"]}],
                "attachment_ids": ["k1"]})
        if task == "authoring_review":
            if value.get("operation") == "structure":
                return validate({"findings": []})
            reviews.append(value)
            encoded = json.dumps(value, ensure_ascii=False)
            assert "张三" in encoded and "给定证明" in encoded
            assert value["units"][0]["content"]["attachments"][0]["content_id"] == "a1"
            assert not any(key in encoded for key in ('"sha256"', '"relative_path"', '"source_ref"', '"artifact_ref"'))
            return validate({"findings": []})
        return validate(ordinary_response(task, value))

    result = artifacts.read_json(await generate_document(artifacts, fixed, budget(), execute, materials=materials))
    assert reviews
    unit = artifacts.read_json(result["unit_refs"][0])
    assert unit["cell_fills"][0]["source_ref"]["source_index_ref"]["sha256"]
    assert unit["attachments"][0]["artifact_ref"] == original


async def test_source_background_does_not_authorize_a_bidder_proposal_without_knowledge(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    background = "本项目已有三个分中心，业务系统部署于内网。"
    fixed = snapshot(artifacts, [{"text": background, "heading": "项目背景"}])
    outline_inputs = []

    async def execute(task, request, validate):
        value = payload(request)
        assert task not in {"authoring_basis", "authoring_write"}
        if task == "outline":
            outline_inputs.append(json.dumps(value, ensure_ascii=False))
        return validate(ordinary_response(task, value))

    result = artifacts.read_json(await generate_document(artifacts, fixed, budget(), execute))
    assert any(background in value for value in outline_inputs)
    assert artifacts.read_json(result["basis_ref"])["response"]["approach"] == ""
    assert artifacts.read_json(result["unit_refs"][0])["status"] == "placeholder"


async def test_unreadable_frozen_root_keeps_usable_document_and_reports_missing_source(tmp_path):
    artifacts = ArtifactStore(tmp_path)
    fixed = snapshot(artifacts, [{"text": "本项目背景原文"}])
    artifacts.path(fixed["frozen_ref"]).unlink()

    async def execute(task, request, validate):
        return validate(ordinary_response(task, payload(request)))

    result = artifacts.read_json(await generate_document(artifacts, fixed, budget(), execute))
    assert result["docx_ref"] and result["unit_refs"]
    assert any(gap["code"] == "authoring_input_unavailable" for gap in result["gaps"])


@pytest.mark.parametrize("location", ["constraint", "finding"])
async def test_unknown_outline_references_are_rejected_inside_the_executor_validator(tmp_path, location):
    artifacts = ArtifactStore(tmp_path)
    corrections = []

    async def execute(task, request, validate):
        value = payload(request)
        response = ordinary_response(task, value)
        if task == "outline" and value.get("operation") != "assemble":
            invalid = deepcopy(response)
            if location == "constraint":
                invalid["constraint_rule_ids"] = ["not-supplied"]
            else:
                invalid["findings"] = [{"code": "question", "detail": "待核对要求", "rule_ids": ["not-supplied"]}]
            with pytest.raises(ModelContractError) as error:
                validate(invalid)
            corrections.append(error.value.errors)
        return validate(response)

    result = artifacts.read_json(await generate_document(artifacts, snapshot(artifacts), budget(), execute))
    assert corrections and result["docx_ref"]


async def test_source_only_tasks_keep_local_analysis_doubts_and_human_opinions_through_review(tmp_path):
    from test_authoring_contract import snapshot as confirmed_snapshot
    from docreader.tender.evidence.references import EvidenceRepository

    artifacts = ArtifactStore(tmp_path)
    fixed = snapshot(artifacts, [{"text": "本项目覆盖一标段现有机房，服务边界另有说明。", "heading": "项目背景"}])
    frozen = artifacts.read_json(fixed["frozen_ref"])
    index = EvidenceRepository(artifacts).source_index(frozen["documents"][0]["source_index_ref"])
    local = {"issue_id": "local-scope", "code": "scope_unclear", "detail": "一标段机房服务边界仍待核对。",
             "blocking": True, "range_id": index.ranges[0]["range_id"], "item_ids": []}
    global_issue = {"issue_id": "global-source", "code": "source_unclear", "detail": "原件附件数量尚未确认。",
                    "blocking": True}
    resolutions = {"schema_version": 1, "resolutions": [
        {"issue_id": "local-scope", "verdict": "acknowledged", "reason": "人工保留该机房范围疑问。"},
        {"issue_id": "global-source", "verdict": "acknowledged", "reason": "人工知悉原件附件尚待补齐。"}]}
    items = {"schema_version": 2, "packs": []}
    verified = {"schema_version": 2, "items_ref": artifacts.put_json(items), "hard_blocking": False,
                "blocking": True, "item_refs": [], "accepted_decision_ids": [],
                "decisions_ref": artifacts.put_json([]), "relations_ref": artifacts.put_json([]),
                "issues": [local, global_issue], "unverified_relation_ids": [],
                "human_issue_ids": ["local-scope", "global-source"]}
    fixed = confirmed_snapshot(artifacts, roots={"freeze": frozen, "extract": items, "verify": verified},
                               resolutions=resolutions).expected
    seen = {}

    async def execute(task, request, validate):
        value = payload(request)
        seen.setdefault(task, []).append(value)
        if task == "authoring_write":
            assert value["rules"] == [], "This legitimate background task is bound only to original source."
            assert local in value["analysis_issues"]
            assert resolutions["resolutions"][0] in value["issue_resolutions"]
        if task == "authoring_review" and value.get("operation") != "structure":
            for unit in value["units"]:
                assert local in unit["analysis_issues"]
                assert resolutions["resolutions"][0] in unit["issue_resolutions"]
        return validate(ordinary_response(task, value))

    result = artifacts.read_json(await generate_document(artifacts, fixed, budget(), execute, materials=[
        {"material_id": "local-response", "title": "给定项目说明", "text": "待补充实际资料。"}]))
    assert any(resolutions["resolutions"][1] in value["issue_resolutions"]
               for value in seen["authoring_write"])
    assert any(gap["detail"].endswith(local["detail"]) for gap in result["gaps"])
    assert any(gap["detail"].endswith(global_issue["detail"]) for gap in result["gaps"])
    assert seen["authoring_review"]


async def test_unrelated_long_analysis_issues_do_not_repeat_in_other_local_packets(tmp_path):
    from docreader.tender.authoring_pipeline import _outline_payload, _source_records
    from test_outline_organization import context as existing_context

    _, _, context = await existing_context(tmp_path, documents=[
        ("first", "main", [("一标段须提交实施计划。", "obligation")]),
        ("second", "attachment", [("二标段须提交服务计划。", "obligation")])])
    first, second = context.payload["rules"]
    first_item, second_item = first["item"]["item_id"], second["item"]["item_id"]
    local = {"issue_id": "first-doubt", "code": "scope_unclear", "detail": "第一事项疑问。",
             "blocking": True, "item_ids": [first_item]}
    long_unrelated = {"issue_id": "second-doubt", "code": "scope_unclear", "detail": "第二事项长说明。" * 2000,
                      "blocking": True, "item_ids": [second_item]}
    local_opinion = {"issue_id": "first-doubt", "verdict": "acknowledged", "reason": "保留第一事项疑问。"}
    other_opinion = {"issue_id": "second-doubt", "verdict": "acknowledged", "reason": "保留第二事项疑问。"}
    context.payload.update(issues=[local, long_unrelated], relations=[], unverified_relation_ids=[],
                           issue_resolutions={"schema_version": 1, "resolutions": [local_opinion, other_opinion]})
    records = _source_records(context.artifacts, context)
    first_records = [record for record in records if any(r["item"]["item_id"] == first_item
                                                        for r in record.get("rules", []))]
    second_records = [record for record in records if any(r["item"]["item_id"] == second_item
                                                         for r in record.get("rules", []))]
    first_packet = _outline_payload(context, first_records)
    second_packet = _outline_payload(context, second_records)
    assert first_packet["analysis_issues"] == [local]
    assert first_packet["issue_resolutions"] == [local_opinion]
    assert long_unrelated["detail"] not in json.dumps(first_packet, ensure_ascii=False)
    assert second_packet["analysis_issues"] == [long_unrelated]
    assert second_packet["issue_resolutions"] == [other_opinion]


async def test_manual_quote_keeps_original_price_cells_empty_and_never_enters_automatic_revision(tmp_path):
    from io import BytesIO

    from docx import Document

    artifacts = ArtifactStore(tmp_path)
    quote_grid = {"row_count": 4, "column_count": 2, "cells": [
        {"row": 0, "column": 0, "text": "投标报价表", "row_span": 1, "col_span": 2},
        {"row": 1, "column": 0, "text": "招标最高限价", "row_span": 1, "col_span": 1},
        {"row": 1, "column": 1, "text": "200000元", "row_span": 1, "col_span": 1},
        {"row": 2, "column": 0, "text": "投标总价（元）", "row_span": 1, "col_span": 1},
        {"row": 2, "column": 1, "text": "", "row_span": 1, "col_span": 1},
        {"row": 3, "column": 0, "text": "含税单价（元）", "row_span": 1, "col_span": 1},
        {"row": 3, "column": 1, "text": "", "row_span": 1, "col_span": 1}]}
    fixed = snapshot(artifacts, [
        {"heading": "价格模板", "grid": quote_grid},
        {"heading": "技术方案", "text": "技术方案应说明部署、测试和验收方法。"}])
    materials = [{"material_id": "historical-price", "title": "另一项目历史报价",
                  "text": "另一项目曾报价123456元，单价888元；不是本次报价。", "usage": "reference"},
                 {"material_id": "implementation", "title": "项目实施方法",
                  "text": "按需求完成部署、执行功能测试，并以测试记录支持验收。", "usage": "reference"}]
    writing_calls, reviewed_manual_tasks = [], []
    technical_text = "按需求完成部署、执行功能测试，并以测试记录支持验收。"

    async def execute(task, request, validate):
        value = payload(request)
        if task == "outline":
            if value.get("operation") == "assemble":
                return validate(assemble_independent_candidates(value))
            paths = {section["section_id"]: section["heading_path"] for section in value["sections"]}
            quote_sources = [s["range_id"] for s in value["sources"] if paths[s["section_id"]] == "价格模板"]
            technical_sources = [s["range_id"] for s in value["sources"] if paths[s["section_id"]] == "技术方案"]
            assert quote_sources and technical_sources
            return validate({"schema_version": 2, "chapters": [
                {"title": "报价文件", "children": [], "targets": [{
                    "kind": "form", "purpose": "报价原格式留空，最终由人工填写。", "manual_fill": True,
                    "rule_ids": [], "range_ids": quote_sources,
                    "template_part_ids": provided_template_ids(value)}]},
                {"title": "技术方案", "children": [], "targets": [{
                    "kind": "narrative", "purpose": "说明部署、测试和验收方法。",
                    "rule_ids": [], "range_ids": technical_sources}]}], "findings": []})
        if task == "authoring_write":
            writing_calls.append(value)
            assert all(not target.get("manual_fill", False) for target in value["targets"])
            assert "previous_content" not in value, "The manual-price review must not trigger automatic revision."
            return validate({"blocks": [{"kind": "paragraph", "text": technical_text,
                "target_ids": [value["targets"][0]["target_id"]], "support_ids": ["k2"]}]})
        if task == "authoring_review":
            if value.get("operation") == "structure":
                return validate({"findings": []})
            manual = [unit for unit in value["units"] if unit["title_path"] == ["报价文件"]]
            reviewed_manual_tasks.extend(manual)
            # Even a reviewer asking to fill the price must not authorize a model write.
            return validate({"findings": [{"code": "price_pending", "detail": "报价尚未填写。",
                "task_ids": [unit["task_id"] for unit in manual], "suggestion": "填入报价后复核。",
                "can_revise": True}] if manual else []})
        return validate({"approach": "按现有技术要求组织实施；最终报价由人工补充。"})

    result = artifacts.read_json(await generate_document(artifacts, fixed, budget(), execute, materials=materials))
    assert len(writing_calls) == 1 and reviewed_manual_tasks == []
    assert any(entry["text"] == materials[0]["text"] for entry in writing_calls[0]["supports"])
    persisted_outline = artifacts.read_json(result["outline_ref"])
    assert persisted_outline["chapters"][0]["targets"][0]["manual_fill"] is True
    doc = Document(BytesIO(artifacts.read_bytes(result["docx_ref"])))
    assert technical_text in [paragraph.text for paragraph in doc.paragraphs]
    quote = next(table for table in doc.tables if table.cell(0, 0).text == "投标报价表")
    assert quote.cell(1, 0).text == "招标最高限价" and quote.cell(1, 1).text == "200000元"
    assert quote.cell(2, 0).text == "投标总价（元）" and quote.cell(2, 1).text == ""
    assert quote.cell(3, 0).text == "含税单价（元）" and quote.cell(3, 1).text == ""
    assert not any(cell.text in {"0", "123456元", "888元"} for row in quote.rows for cell in row.cells)


async def test_multifragment_original_and_retained_rule_references_survive_outline_task_cas_roundtrips(tmp_path):
    from docreader.tender.authoring import _OutlineContext
    from docreader.tender.authoring_pipeline import _load_task, _organize, plan_writing_tasks
    from docreader.tender.schema import AuthoringBasis, AuthoringOutline, EffectiveItemRef, WritingTask
    from test_analysis_checks import Model, run, scenario
    from test_outline_organization import authoring_input

    data = await scenario(tmp_path, documents=[
        ("main", "main", [("项目一实施范围要求。" * 240, "obligation")]),
        ("annex", "attachment", [("项目二服务安排要求。" * 240, "obligation")])])
    artifacts, plan, items, _ = data
    verified = await run(data, Model())
    original = artifacts.read_json(items["packs"][0]["artifact_ref"])["items"][0]
    retained = [{"text": "仅一标段按分阶段实施。" + "实施说明。" * 120,
                 "scope": "一标段", "evidence": original["evidence"]},
                {"text": "仅二标段按验收安排实施。" + "验收说明。" * 120,
                 "scope": "二标段", "evidence": original["evidence"]}]
    decision = {"decision_id": "retained-scope", "amendment": {
        "target_item_id": original["item_id"], "retained_rules": retained, "replacement_item_ids": []}}
    verified.update(decisions_ref=artifacts.put_json([decision]), accepted_decision_ids=[decision["decision_id"]])
    fixed = authoring_input(artifacts, plan, items, verified)
    fixed.update(confirmation_ref=None, checklist_digest=None)
    context = _OutlineContext(artifacts, fixed)
    local_calls, assembly_calls = [], []

    async def execute(task, request, validate):
        assert task == "outline"
        value = payload(request)
        if value.get("operation") == "assemble":
            assembly_calls.append(value)
            selected = [target["target_id"] for chapter in value["chapters"] for target in chapter["targets"]]
            return validate({"chapters": [{"title": "合并后的实施响应", "children": [], "targets": [
                {"kind": "narrative", "purpose": "落实各标段要求并保留原始依据", "target_ids": selected,
                 "template_part_ids": []}]}]})
        local_calls.append(value)
        return validate(outline(value))

    admission = RequestBudget(model="fixture", context_tokens=60000, prompts={}, output_reserve=4096)
    organized = await _organize(artifacts, context, admission, execute, [])
    assert len(local_calls) > 1 and assembly_calls, "The regression must exercise real budget-driven fragmentation and assembly."
    # Persistence must not inject omitted optional analysis fields as explicit null.
    for _ in range(2):
        organized = artifacts.read_json(artifacts.put_json(organized))
        organized = AuthoringOutline.model_validate(organized).model_dump(exclude_unset=True)
    assert organized["authoring_input"]["confirmation_ref"] is None
    tasks = plan_writing_tasks(organized)
    actual_references = []
    for task in tasks:
        for _ in range(2):
            task = artifacts.read_json(artifacts.put_json(task))
            task = WritingTask.model_validate(task).model_dump(exclude_unset=True)
        actual_references.extend(ref for target in task["targets"] for ref in target["rule_refs"])
        basis = AuthoringBasis(authoring_input=context.input, response={}).model_dump()
        semantic, _, _ = _load_task(artifacts, context, task, basis, [])
        assert len(semantic["rules"]) == len({json.dumps(ref, sort_keys=True) for target in task["targets"] for ref in target["rule_refs"]})
    assert {json.dumps(ref, sort_keys=True) for ref in actual_references} == {
        json.dumps(ref, sort_keys=True) for ref in context.rules.values()}
    originals = [ref for ref in actual_references if "decision_id" not in ref]
    assert originals and all("rule_index" not in ref for ref in originals)
    kept = [ref for ref in actual_references if ref.get("decision_id") == "retained-scope"]
    assert {ref["rule_index"] for ref in kept} == {0, 1}
    for reference in actual_references:
        EffectiveItemRef.model_validate(reference)
