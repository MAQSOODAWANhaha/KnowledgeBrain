"""Final-file profile regressions: no cleanup and no silently missing carriers."""
import hashlib
import json
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO

import grpc
import pytest
from docx import Document as BaseDocument
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from docreader.main import DocReaderServicer
from docreader.parser.parser import Parser
from docreader.proto.docreader_pb2 import ReadConfig, ReadRequest
from docreader.proto.docreader_pb2_grpc import DocReaderStub, add_DocReaderServicer_to_server


def Document():
    # Match the product compiler's supported default presentation.
    doc = BaseDocument()
    defaults = doc.styles.element.find(qn("w:docDefaults"))
    if defaults is not None:
        doc.styles.element.remove(defaults)
    from docx.oxml import parse_xml
    doc.styles.element.insert(0, parse_xml(
        '<w:docDefaults xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:rPrDefault><w:rPr><w:rFonts w:ascii="SimSun" w:hAnsi="SimSun" w:eastAsia="SimSun" w:cs="SimSun"/>'
        '<w:sz w:val="24"/></w:rPr></w:rPrDefault>'
        '<w:pPrDefault><w:pPr><w:spacing w:line="360" w:lineRule="auto"/></w:pPr></w:pPrDefault></w:docDefaults>'
    ))
    section = doc.sections[0]._sectPr
    for child in list(section):
        section.remove(child)
    section.append(parse_xml('<w:pgSz xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:w="11906" w:h="16838" w:orient="portrait"/>'))
    section.append(parse_xml('<w:pgMar xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:top="1417" w:right="1417" w:bottom="1417" w:left="1417"/>'))
    return doc


def parse(name, raw):
    return Parser().parse_file(name, name.rsplit('.', 1)[-1], raw,
                              parser_engine="builtin", engine_overrides={"output_inventory": "v1"})


def manifest(document):
    return json.loads(document.metadata["output_inventory_manifest"])


def docx_fixture():
    doc = Document()
    paragraph = doc.add_paragraph("Unbookmarked body")
    bookmark = OxmlElement("w:bookmarkStart")
    bookmark.set(qn("w:id"), "1")
    bookmark.set(qn("w:name"), "section_one")
    paragraph._p.insert(0, bookmark)
    doc.add_paragraph("")
    doc.sections[0].header.paragraphs[0].text = "Actual header"
    doc.sections[0].footer.paragraphs[0].text = "Actual footer"
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged"
    drawing = OxmlElement("w:drawing")
    doc.add_paragraph().add_run()._r.append(drawing)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    doc.add_paragraph()._p.append(field)
    raw = BytesIO()
    doc.save(raw)
    return raw.getvalue()


def pdf_fixture(blank=False):
    writer = PdfWriter()
    for index in range(4):
        page = writer.add_blank_page(width=300, height=300)
        if not blank:
            font = writer._add_object(DictionaryObject({
                NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'),
                NameObject('/BaseFont'): NameObject('/Helvetica'),
            }))
            page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
            stream = DecodedStreamObject()
            stream.set_data(f"BT /F1 12 Tf 10 280 Td (Repeating header) Tj 0 -30 Td (Body {index}) Tj 0 -200 Td ({index+1}) Tj ET".encode())
            page[NameObject('/Contents')] = writer._add_object(stream)
    raw = BytesIO()
    writer.write(raw)
    return raw.getvalue()


def test_docx_preserves_stories_empty_paragraphs_grid_bookmarks_and_omissions():
    raw = docx_fixture()
    result = parse("final.docx", raw)
    receipt = manifest(result)
    assert receipt["file_sha256"] == hashlib.sha256(raw).hexdigest()
    assert receipt["profile"] == "output_inventory_v1"
    assert len(receipt["units"]) == len(result.structured_source_units)
    assert any(u.text == "Actual header" for u in result.structured_source_units)
    assert any(u.text == "Actual footer" for u in result.structured_source_units)
    assert any(e["bookmarks"] == ["section_one"] for e in receipt["units"])
    assert any(u.grid and u.grid.cells[0].col_span == 2 for u in result.structured_source_units)
    assert any(e["reason"] and "drawing" in e["reason"] for e in receipt["units"])
    assert any(e["reason"] and "field" in e["reason"] for e in receipt["units"])
    assert any(e["kind"] == "paragraphs" and not u.text for e, u in zip(receipt["units"], result.structured_source_units))


@pytest.mark.parametrize("blank", [False, True])
def test_pdf_preserves_all_pages_and_repeated_header_page_numbers(blank):
    result = parse("final.pdf", pdf_fixture(blank))
    receipt = manifest(result)
    pages = [unit for unit in result.structured_source_units if unit.key.endswith(":section:0")]
    assert receipt["page_count"] == len(pages) == 4
    assert len(receipt["image_sha256"]) >= 4
    assert [p.locator.page_ordinal for p in pages] == list(range(4))
    if blank:
        assert all(not page.text for page in pages)
    else:
        assert all("Repeating header" in page.text and str(i + 1) in page.text for i, page in enumerate(pages))


def test_readstream_keeps_all_blank_pdf_profile_receipt():
    server = grpc.server(ThreadPoolExecutor(max_workers=1))
    add_DocReaderServicer_to_server(DocReaderServicer(), server)
    port = server.add_insecure_port("127.0.0.1:0")
    server.start()
    try:
        with grpc.insecure_channel(f"127.0.0.1:{port}") as channel:
            frames = list(DocReaderStub(channel).ReadStream(ReadRequest(
                file_content=pdf_fixture(True), file_name="final.pdf", file_type="pdf",
                config=ReadConfig(parser_engine="builtin", parser_engine_overrides={"output_inventory": "v1"}),
            )))
        assert not frames[0].meta.error
        assert len(frames[0].meta.structured_source_units) == 8
        assert frames[0].meta.image_count == 4
        assert len(frames) == 5
        assert json.loads(frames[0].meta.metadata["output_inventory_manifest"])["page_count"] == 4
    finally:
        server.stop(0).wait()


def test_unknown_inventory_profile_does_not_fall_back():
    with pytest.raises(ValueError, match="unsupported output inventory"):
        Parser().parse_file("final.pdf", "pdf", pdf_fixture(True), engine_overrides={"output_inventory": "v2"})


def test_docx_used_notes_and_picture_bytes_have_immutable_inventory_identity():
    import zipfile
    from lxml import etree
    from PIL import Image
    doc = Document()
    paragraph = doc.add_paragraph("Notes follow")
    for kind in ("footnote", "endnote"):
        reference = OxmlElement(f"w:{kind}Reference")
        reference.set(qn("w:id"), "7")
        paragraph.add_run()._r.append(reference)
    image = BytesIO()
    Image.new("RGB", (10, 10), "blue").save(image, format="PNG")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).paragraphs[0].add_run().add_picture(BytesIO(image.getvalue()))
    raw = BytesIO()
    doc.save(raw)
    with zipfile.ZipFile(BytesIO(raw.getvalue())) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    rels = etree.fromstring(parts["word/_rels/document.xml.rels"])
    content_types = etree.fromstring(parts["[Content_Types].xml"])
    w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    for kind in ("footnote", "endnote"):
        etree.SubElement(rels, "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship", {
            "Id": f"rId{kind}", "Type": f"http://schemas.openxmlformats.org/officeDocument/2006/relationships/{kind}s",
            "Target": f"{kind}s.xml",
        })
        etree.SubElement(content_types, "{http://schemas.openxmlformats.org/package/2006/content-types}Override", {
            "PartName": f"/word/{kind}s.xml",
            "ContentType": f"application/vnd.openxmlformats-officedocument.wordprocessingml.{kind}s+xml",
        })
        parts[f"word/{kind}s.xml"] = (
            f'<w:{kind}s xmlns:w="{w}"><w:{kind} w:id="7"><w:p><w:r><w:t>{kind} content</w:t>'
            f'</w:r></w:p></w:{kind}><w:{kind} w:id="8"><w:p><w:r><w:t>Unused note</w:t>'
            f'</w:r></w:p></w:{kind}></w:{kind}s>'
        ).encode()
    parts["word/_rels/document.xml.rels"] = etree.tostring(rels)
    parts["[Content_Types].xml"] = etree.tostring(content_types)
    output = BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        for name, payload in parts.items():
            archive.writestr(name, payload)
    result = parse("notes.docx", output.getvalue())
    texts = [unit.text for unit in result.structured_source_units]
    assert "footnote content" in texts and "endnote content" in texts
    assert "Unused note" not in texts
    receipt = manifest(result)
    assert hashlib.sha256(image.getvalue()).hexdigest() in receipt["image_sha256"].values()
    pictures = [unit for unit in result.structured_source_units if unit.kind.value == "image_region"]
    assert len(pictures) == 1
    assert pictures[0].locator.compound_parent.parent_kind == "table_cell"


def test_docx_unused_first_story_is_excluded_until_inherited_section_uses_it():
    from docx.enum.section import WD_SECTION_START
    doc = Document()
    doc.add_paragraph("Body")
    doc.sections[0].first_page_header.paragraphs[0].text = "Conditional first header"
    raw = BytesIO()
    doc.save(raw)
    assert all(u.text != "Conditional first header" for u in parse("first.docx", raw.getvalue()).structured_source_units)
    section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    section.different_first_page_header_footer = True
    raw = BytesIO()
    doc.save(raw)
    assert any(u.text == "Conditional first header" for u in parse("first.docx", raw.getvalue()).structured_source_units)


def test_docx_duplicate_bookmarks_are_explicit_not_checked():
    doc = Document()
    for index in range(2):
        p = doc.add_paragraph("Body")
        bookmark = OxmlElement("w:bookmarkStart")
        bookmark.set(qn("w:id"), str(index))
        bookmark.set(qn("w:name"), "duplicate_name")
        p._p.insert(0, bookmark)
    raw = BytesIO()
    doc.save(raw)
    entries = manifest(parse("duplicate.docx", raw.getvalue()))["units"]
    assert len([e for e in entries if "duplicate bookmarks" in (e["reason"] or "")]) == 2


def test_header_footer_and_custom_style_are_not_checked():
    doc = Document()
    doc.add_heading("第一章", level=1)
    branded = doc.add_paragraph("品牌段落")
    branded.style = doc.styles.add_style("BidderBrand", 1)
    doc.sections[0].header.paragraphs[0].text = "页眉内容"
    doc.sections[0].footer.paragraphs[0].text = "页脚内容"
    raw = BytesIO()
    doc.save(raw)
    entries = manifest(parse("styled.docx", raw.getvalue()))["units"]
    assert any("header" in e["part"] for e in entries)
    assert any("footer" in e["part"] for e in entries)
    assert any("custom style" in (e["reason"] or "") for e in entries)
    assert any(e.get("style_name") == "BidderBrand" for e in entries)


@pytest.mark.parametrize("mixed", [False, True])
def test_exact_picture_has_one_visual_carrier_but_mixed_shape_stays_unchecked(mixed):
    from PIL import Image
    doc = Document()
    image = BytesIO()
    Image.new("RGB", (12, 8), "blue").save(image, format="PNG")
    paragraph = doc.add_paragraph("Adjacent text does not prove the picture")
    paragraph.add_run().add_picture(BytesIO(image.getvalue()))
    if mixed:
        # Keep the valid picture but add a second content carrier inside the
        # same graphic. Finding its image must not erase the shape/text story.
        graphic_data = paragraph._p.xpath(".//*[local-name()='graphicData']")[0]
        text_box = OxmlElement("w:txbxContent")
        text = OxmlElement("w:p")
        run = OxmlElement("w:r")
        value = OxmlElement("w:t")
        value.text = "Additional shape text"
        run.append(value)
        text.append(run)
        text_box.append(text)
        graphic_data.append(text_box)
    raw = BytesIO()
    doc.save(raw)
    result = parse("picture.docx", raw.getvalue())
    receipt = manifest(result)
    pictures = [entry for entry in receipt["units"] if entry["kind"] == "image"]
    assert len(pictures) == 1 and pictures[0]["status"] == "not_checked"
    drawings = [entry for entry in receipt["units"] if "}drawing" in (entry["reason"] or "")]
    textboxes = [entry for entry in receipt["units"] if "}txbxContent" in (entry["reason"] or "")]
    if mixed:
        assert drawings and textboxes
    else:
        assert not drawings and not textboxes


def test_reconstructable_toc_has_no_ordinary_field_rejection():
    doc = Document()
    paragraph = doc.add_paragraph()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), 'TOC \\o "1-3"')
    paragraph._p.append(field)
    raw = BytesIO()
    doc.save(raw)
    entries = manifest(parse("toc.docx", raw.getvalue()))["units"]
    assert any(e["field_region"] == "toc" for e in entries)
    assert not any(e["status"] == "not_checked" for e in entries)


def test_table_custom_paragraph_style_is_explicitly_rejected():
    doc = Document()
    doc.add_heading("报价", 1)
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = "报价100"
    cell.paragraphs[0].style = "Intense Quote"
    raw = BytesIO()
    doc.save(raw)
    entries = manifest(parse("table.docx", raw.getvalue()))["units"]
    assert any(e["status"] == "not_checked" and "table paragraph custom style" in e["reason"] for e in entries)


@pytest.mark.parametrize("style_kind", ["character", "table"])
def test_table_nested_custom_styles_are_rejected(style_kind):
    from docx.enum.style import WD_STYLE_TYPE
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    cell.text = "报价100"
    if style_kind == "character":
        style = doc.styles.add_style("BidderCharacter", WD_STYLE_TYPE.CHARACTER)
        cell.paragraphs[0].runs[0].style = style
    else:
        style = doc.styles.add_style("BidderTable", WD_STYLE_TYPE.TABLE)
        table.style = style
    raw = BytesIO()
    doc.save(raw)
    entries = manifest(parse("table.docx", raw.getvalue()))["units"]
    assert any(e["status"] == "not_checked" and "style" in (e["reason"] or "") for e in entries)


def test_toc_does_not_whitelist_adjacent_date_field():
    doc = Document()
    paragraph = doc.add_paragraph()
    for instruction in ['TOC \\o "1-3"', "DATE"]:
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), instruction)
        paragraph._p.append(field)
    raw = BytesIO()
    doc.save(raw)
    entries = manifest(parse("mixed.docx", raw.getvalue()))["units"]
    assert any(e["status"] == "not_checked" and "field" in (e["reason"] or "") for e in entries)


def test_package_root_signature_is_explicitly_rejected():
    doc = Document()
    doc.add_paragraph("已签署报价")
    doc.part.package.rels.get_or_add_ext_rel(
        "http://schemas.openxmlformats.org/package/2006/relationships/digital-signature/origin",
        "https://example.invalid/signature",
    )
    raw = BytesIO()
    doc.save(raw)
    entries = manifest(parse("signed.docx", raw.getvalue()))["units"]
    assert any(e["part"] == "/" and e["status"] == "not_checked"
               and "origin relationship" in (e["reason"] or "") for e in entries)


@pytest.mark.parametrize("change", ["bold", "font_size", "normal", "heading", "numbering", "table_shading"])
def test_reconstruction_rejects_unrepresented_formatting(change):
    from docx.shared import Pt
    doc = Document()
    paragraph = doc.add_paragraph("报价100万元")
    if change == "bold":
        paragraph.runs[0].bold = True
    elif change == "font_size":
        paragraph.runs[0].font.size = Pt(22)
    elif change == "normal":
        doc.styles["Normal"].font.size = Pt(22)
    elif change == "heading":
        paragraph.style = "Heading 1"
        doc.styles["Heading 1"].font.size = Pt(22)
    elif change == "numbering":
        paragraph._p.get_or_add_pPr().append(OxmlElement("w:numPr"))
    else:
        cell = doc.add_table(rows=1, cols=1).cell(0, 0)
        cell.text = "报价"
        cell._tc.get_or_add_tcPr().append(OxmlElement("w:shd"))
    raw = BytesIO()
    doc.save(raw)
    entries = manifest(parse("format.docx", raw.getvalue()))["units"]
    assert any(e["status"] == "not_checked" and any(term in (e["reason"] or "") for term in ("formatting", "style")) for e in entries)


def test_toc_with_ordinary_body_is_not_silently_discarded():
    doc = Document()
    paragraph = doc.add_paragraph("报价100万元")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), 'TOC \\o "1-3"')
    paragraph._p.append(field)
    raw = BytesIO()
    doc.save(raw)
    entries = manifest(parse("mixed-toc.docx", raw.getvalue()))["units"]
    assert any(e["status"] == "not_checked" and "mixed with TOC" in (e["reason"] or "") for e in entries)


def test_standard_style_inheritance_cannot_hide_custom_formatting():
    doc = Document()
    custom = doc.styles.add_style("CustomBase", 1)
    custom.font.bold = True
    doc.styles["Normal"].base_style = custom
    doc.add_paragraph("报价100万元")
    raw = BytesIO()
    doc.save(raw)
    entries = manifest(parse("inherited.docx", raw.getvalue()))["units"]
    assert any(e["status"] == "not_checked" and "inheritance" in (e["reason"] or "") for e in entries)


def test_modified_document_defaults_are_rejected():
    doc = Document()
    doc.add_paragraph("报价100万元")
    defaults = doc.styles.element.find(qn("w:docDefaults"))
    defaults.find(qn("w:rPrDefault")).find(qn("w:rPr")).find(qn("w:sz")).set(qn("w:val"), "60")
    raw = BytesIO()
    doc.save(raw)
    entries = manifest(parse("defaults.docx", raw.getvalue()))["units"]
    assert any(e["status"] == "not_checked" and "default formatting" in (e["reason"] or "") for e in entries)

@pytest.mark.parametrize("attribute,value", [("w:w", "16838"), ("w:orient", "landscape"), ("margin", "283")])
def test_fill_rejects_changed_page_layout(attribute, value):
    doc = Document()
    doc.add_paragraph("保留正文")
    section = doc.sections[0]._sectPr
    if attribute == "margin":
        section.find(qn("w:pgMar")).set(qn("w:left"), value)
    else:
        section.find(qn("w:pgSz")).set(qn(attribute), value)
    stream = BytesIO()
    doc.save(stream)
    entries = manifest(parse("layout.docx", stream.getvalue()))["units"]
    assert any("cannot be reconstructed" in (entry["reason"] or "") for entry in entries)


def test_fill_inventory_carries_unequal_widths_and_repeated_header():
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    for column, width in zip(table._tbl.tblGrid, [2000, 4000]):
        column.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, [2000, 4000]):
            cell._tc.get_or_add_tcPr().find(qn("w:tcW")).set(qn("w:w"), str(width))
    header = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(header)
    stream = BytesIO()
    doc.save(stream)
    entries = manifest(parse("table.docx", stream.getvalue()))["units"]
    entry = next(entry for entry in entries if entry["kind"] == "table")
    assert entry["table_layout"] == {"widths_twips": [2000, 4000], "header_rows": 1}
    assert not any("table layout cannot" in (entry["reason"] or "") for entry in entries)
    table.rows[0].cells[0]._tc.get_or_add_tcPr().find(qn("w:tcW")).set(qn("w:w"), "1000")
    stream = BytesIO()
    doc.save(stream)
    assert any("table layout cannot" in (entry["reason"] or "")
               for entry in manifest(parse("table.docx", stream.getvalue()))["units"])
