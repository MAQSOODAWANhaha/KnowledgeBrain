"""Native carriers and rendered Markdown share one deterministic source map."""
import hashlib
import json
from io import BytesIO

from docx import Document as DocxDocument
from docx.oxml import parse_xml
from openpyxl import Workbook

from docreader.models.document import Document
from docreader.parser.docx_parser import _docx_structured_units, _tc_text
from docreader.parser.excel_parser import ExcelParser
from docreader.parser.output_inventory import parse_output_inventory
from docreader.parser.source_contract import attach_source_contract


def contract(document):
    return json.loads(document.metadata['source_contract'])


def test_ooxml_shared_serializer_preserves_structure_without_run_spaces():
    cell = parse_xml('''<w:tc xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
      <w:p><w:r><w:t>1</w:t></w:r><w:r><w:t>0</w:t><w:br/><w:t>20</w:t><w:tab/><w:t>否</w:t></w:r></w:p>
      <w:p><w:r><w:t>中文</w:t></w:r></w:p>
      <w:tbl><w:tr><w:tc><w:p><w:r><w:t>嵌套1</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>嵌套2</w:t></w:r></w:p></w:tc></w:tr></w:tbl>
      <w:p/>
    </w:tc>''')
    assert _tc_text(cell) == '10\n20\t否\n中文\n嵌套1\t嵌套2\n'


def test_docx_duplicate_text_tables_and_utf8_spans_have_stable_order():
    document = DocxDocument()
    document.add_heading('第一章 条件', 1)
    document.add_paragraph('相同正文😀')
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = '承诺'
    table.cell(1, 0).text = '10'
    table.cell(1, 0).add_paragraph('20')
    table.cell(1, 1).text = '同意\t否'
    document.add_paragraph('相同正文😀')
    raw = BytesIO()
    document.save(raw)
    source = raw.getvalue()
    units = _docx_structured_units(source)
    result = attach_source_contract(Document(structured_source_units=units), source, 'docx')
    receipt = contract(result)
    assert receipt['document_revision'] == hashlib.sha256(source).hexdigest()
    assert receipt['markdown_sha256'] == hashlib.sha256(result.content.encode()).hexdigest()
    assert [unit['unit_id'] for unit in receipt['units']] == [unit.key for unit in units]
    assert len({unit['section_id'] for unit in receipt['units']}) == 1
    assert all(unit['physical_path'] for unit in receipt['units'])
    table_unit = next(unit for unit in units if unit.grid)
    identity = next(unit for unit in receipt['units'] if unit['unit_id'] == table_unit.key)
    assert table_unit.text == ''
    assert table_unit.grid.cells[1].text == '10\n20'
    assert identity['table_id'] == table_unit.key
    span = identity['rendered_spans'][0]
    displayed = result.content.encode()[span['start_byte']:span['end_byte']].decode()
    assert '10<br>20' in displayed
    assert '同意&#9;否' in displayed
    spans = [span for unit in receipt['units'] for span in unit['rendered_spans']]
    assert all(left['end_byte'] <= right['start_byte'] for left, right in zip(spans, spans[1:]))
    reread = parse_output_inventory('final.docx', 'docx', source)
    assert next(unit.grid for unit in reread.structured_source_units if unit.grid) == table_unit.grid


def test_xlsx_grid_maps_to_the_actual_rendered_rows_even_with_empty_text():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(['列1', '列2'])
    sheet.append(['相同', '😀'])
    sheet.append(['相同', '😀'])
    raw = BytesIO()
    workbook.save(raw)
    result = ExcelParser(file_name='source.xlsx', file_type='xlsx').parse(raw.getvalue())
    receipt = contract(result)
    table = next(unit for unit in result.structured_source_units if unit.grid)
    identity = next(unit for unit in receipt['units'] if unit['unit_id'] == table.key)
    assert not table.text
    assert len(identity['rendered_spans']) == 3
    assert len({span['start_byte'] for span in identity['rendered_spans']}) == 3
    for span in identity['rendered_spans']:
        assert result.content.encode()[span['start_byte']:span['end_byte']].decode()


def test_mixed_pdf_inventory_separates_section_and_physical_page_contract(monkeypatch):
    from docreader.scripts.generate_source_contract_fixtures import mixed_pdf
    from docreader.config import CONFIG
    from dataclasses import replace
    monkeypatch.setattr("docreader.parser.pdf_parser.CONFIG", replace(CONFIG, pdf_render_parallelism=1))
    raw = mixed_pdf()
    result = parse_output_inventory('mixed.pdf', 'pdf', raw)
    receipt = contract(result)
    assert [page['classification'] for page in receipt['page_manifest']] == ['text', 'blank', 'scanned']
    assert [page['page_ordinal'] for page in receipt['page_manifest']] == [0, 1, 2]
    native = next(unit for unit in result.structured_source_units if unit.kind.value == 'section')
    assert native.locator.locator_kind == 'document'
    identity = next(unit for unit in receipt['units'] if unit['unit_id'] == native.key)
    assert identity['physical_locator']['locator_kind'] == 'page'
    assert identity['physical_locator']['page_ordinal'] == 0
    assert all(page['image_unit_ids'] for page in receipt['page_manifest'])
    manifest = json.loads(result.metadata['output_inventory_manifest'])
    assert manifest['page_manifest'] == receipt['page_manifest']
    assert all(entry['part'].startswith('pdf:page:') for entry in manifest['units'])
