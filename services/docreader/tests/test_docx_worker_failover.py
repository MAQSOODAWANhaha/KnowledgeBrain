"""A failed DOCX worker may never publish surviving fragments as a full result."""
from concurrent.futures import Future, ThreadPoolExecutor
from concurrent.futures.process import BrokenProcessPool
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_BREAK
from PIL import Image

from docreader.parser import docx_parser
from docreader.parser.docx_parser import Docx, DocxParser, LineData, process_page_multiprocess


def document_bytes(*, image_only=False, with_image=False):
    document = Document()
    if image_only or with_image:
        picture = BytesIO()
        Image.new("RGB", (10, 8), "blue").save(picture, "PNG")
        document.add_paragraph().add_run().add_picture(BytesIO(picture.getvalue()))
    if not image_only:
        document.add_paragraph("Survived paragraph").add_run().add_break(WD_BREAK.PAGE)
        document.add_paragraph("Must not disappear")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "Required table"
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def collect(parser, outcomes, paragraph_groups=None):
    groups = paragraph_groups or [[index] for index in range(len(outcomes))]
    arguments = [(index, groups[index]) for index in range(len(outcomes))]
    futures = {}
    for index, outcome in enumerate(outcomes):
        future = Future()
        if isinstance(outcome, Exception):
            future.set_exception(outcome)
        else:
            future.set_result(outcome)
        futures[future] = index
    parser._collect_process_results(futures, arguments, 0)


def test_broken_worker_and_surviving_fragment_trigger_whole_file_fallback(monkeypatch):
    paths = []
    prepare = Docx._prepare_document_sharing

    def capture_source(self, binary):
        path = prepare(self, binary)
        paths.append(Path(path))
        return path

    def fail_one_group(self, arguments, max_workers):
        assert len(arguments) == 2
        collect(self, [[LineData(text="Survived paragraph", page_num=0)],
                       BrokenProcessPool("worker exited during extraction")])

    monkeypatch.setattr(Docx, "_prepare_document_sharing", capture_source)
    monkeypatch.setattr(Docx, "_execute_multiprocess_tasks", fail_one_group)
    result = DocxParser(file_name="fixture.docx").parse_into_text(document_bytes())
    assert "Survived paragraph" in result.content
    assert "Must not disappear" in result.content
    assert "Required table" in result.content
    assert result.structured_source_units
    assert paths and all(not path.exists() for path in paths)


def test_worker_actual_source_load_failure_is_not_reported_as_empty_success(tmp_path):
    malformed = tmp_path / "malformed.docx"
    malformed.write_bytes(b"not an Office document")
    with pytest.raises(Exception, match="load|source|document"):
        process_page_multiprocess(0, [0], 0, 1, False, 1920, str(malformed), False)


def test_unexpected_empty_result_for_nonempty_input_is_a_failed_group():
    with pytest.raises(Exception, match="incomplete|empty"):
        collect(Docx(), [[LineData(text="survived", page_num=0)], []])


def test_successful_empty_text_and_empty_input_groups_are_not_failures():
    parser = Docx()
    collect(parser, [[LineData(text="", page_num=0)], [],
                     [LineData(text="content", page_num=2)]], [[0], [], [1]])
    assert [(line.page_num, line.text) for line in parser.all_lines] == [(0, ""), (2, "content")]


def test_serial_fallback_preserves_image_only_supported_structure():
    result = DocxParser(file_name="picture.docx")._parse_using_simple_method(document_bytes(image_only=True))
    assert result.is_valid()
    assert result.images
    assert any(unit.kind.value == "image_region" for unit in result.structured_source_units)


def test_failed_full_fallback_cannot_return_partial_success(monkeypatch):
    def fail_workers(self, arguments, max_workers):
        raise BrokenProcessPool("worker exited")

    def fail_structure(content):
        raise ValueError("cannot build complete source structure")

    monkeypatch.setattr(Docx, "_execute_multiprocess_tasks", fail_workers)
    monkeypatch.setattr(docx_parser, "_docx_structured_units", fail_structure)
    with pytest.raises(ValueError, match="complete source structure"):
        DocxParser(file_name="fixture.docx").parse_into_text(document_bytes())


@pytest.fixture
def successful_spawn_groups(monkeypatch):
    """Observe real child-process completion, without permitting fallback success."""
    completed = []
    executor = docx_parser.ProcessPoolExecutor
    collect_results = Docx._collect_process_results

    def spawn_executor(**kwargs):
        assert kwargs["mp_context"].get_start_method() == "spawn"
        return executor(**kwargs)

    def observe_results(self, futures, arguments, started):
        collect_results(self, futures, arguments, started)
        completed.extend(args[0] for args in arguments)

    def unexpected_fallback(self, content):
        pytest.fail("valid DOCX must complete the real spawned-worker path")

    monkeypatch.setattr(docx_parser, "ProcessPoolExecutor", spawn_executor)
    monkeypatch.setattr(Docx, "_collect_process_results", observe_results)
    monkeypatch.setattr(DocxParser, "_parse_using_simple_method", unexpected_fallback)
    return completed


@pytest.mark.parametrize("image_only", [False, True])
def test_real_spawn_workers_complete_from_parent_thread(successful_spawn_groups, image_only):
    raw = document_bytes(image_only=image_only, with_image=True)
    with ThreadPoolExecutor(max_workers=1) as executor:
        if image_only:
            # Textless groups are valid worker output. DocxParser's existing
            # textless fallback is tested separately for image retention.
            lines, _ = executor.submit(Docx(enable_multimodal=True), raw).result(timeout=30)
            assert successful_spawn_groups == [0]
            assert len(lines) == 1
            return
        result = executor.submit(DocxParser(file_name="fixture.docx").parse_into_text, raw).result(timeout=30)
    assert successful_spawn_groups
    assert result.is_valid()
    assert result.images
    assert "Survived paragraph" in result.content
    assert "Must not disappear" in result.content
    assert any(unit.grid and any(cell.text == "Required table" for cell in unit.grid.cells)
               for unit in result.structured_source_units)


def test_readstream_docx_from_rpc_thread_uses_successful_spawn_workers(successful_spawn_groups):
    import grpc

    from docreader.main import DocReaderServicer
    from docreader.proto.docreader_pb2 import ReadConfig, ReadRequest
    from docreader.proto.docreader_pb2_grpc import DocReaderStub, add_DocReaderServicer_to_server

    server = grpc.server(ThreadPoolExecutor(max_workers=1))
    add_DocReaderServicer_to_server(DocReaderServicer(), server)
    port = server.add_insecure_port("127.0.0.1:0")
    server.start()
    try:
        with grpc.insecure_channel(f"127.0.0.1:{port}") as channel:
            frames = list(DocReaderStub(channel).ReadStream(ReadRequest(
                file_content=document_bytes(with_image=True), file_name="fixture.docx", file_type="docx",
                config=ReadConfig(parser_engine="builtin"),
            ), timeout=30))
        assert successful_spawn_groups == [0, 1]
        assert not frames[0].meta.error
        assert "Must not disappear" in frames[0].meta.markdown_content
        assert any(any(cell.text == "Required table" for cell in unit.grid.cells)
                   for unit in frames[0].meta.structured_source_units)
        assert frames[0].meta.image_count > 0
        assert any(frame.HasField("image") for frame in frames[1:])
    finally:
        server.stop(0).wait()
