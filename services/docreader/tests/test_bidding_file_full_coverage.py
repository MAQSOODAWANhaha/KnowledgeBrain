"""Public synthetic document coverage; private PDF acceptance is separate."""
from io import BytesIO

from docx import Document
from docreader.parser.docx_parser import DocxParser

def test_synthetic_document_keeps_text_and_table_content():
    doc = Document()
    doc.add_paragraph("Synthetic heading Alpha")
    table = doc.add_table(rows=2, cols=2)
    for cell, text in zip([c for row in table.rows for c in row.cells], ["Column Alpha", "Column Beta", "Value One", "Value Two"]):
        cell.text = text
    doc.add_paragraph("Synthetic ending Omega")
    stream = BytesIO()
    doc.save(stream)
    parsed = DocxParser(file_name="synthetic.docx", file_type="docx").parse_into_text(stream.getvalue())
    assert all(text in parsed.content for text in ["Synthetic heading Alpha", "Synthetic ending Omega"])
    from docreader.parser.docx_parser import _docx_structured_units
    units = _docx_structured_units(stream.getvalue())
    cell_text = [cell.text for unit in units if unit.grid is not None for cell in unit.grid.cells]
    assert "Value One" in cell_text and "Value Two" in cell_text
