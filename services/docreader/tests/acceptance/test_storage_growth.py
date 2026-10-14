"""Measure physical references and PostgreSQL growth, with no external model calls.

Run pytest with -s to retain the aggregate JSON evidence. Source text, images and
raw model responses are intentionally absent from that report.
"""

import asyncio
from contextlib import asynccontextmanager
import hashlib
import json
import os
from pathlib import Path
import sys
from uuid import uuid4

from docx import Document
from docx.shared import Inches
from PIL import Image, ImageDraw
import psycopg
from psycopg import sql
from psycopg.conninfo import make_conninfo
import pytest
from langgraph.checkpoint.postgres.aio import AsyncPostgresSaver

from docreader.integrations.artifacts import ArtifactStore
from docreader.integrations.store import Store
from docreader.jobs.settings import Settings
from docreader.llm import RequestBudget
from docreader.tender.evidence import register_manifest
from docreader.workflow.nodes import TenderNodes
from docreader.workflow.persistence import initialize_checkpoints
from docreader.workflow.runner import execute_command

sys.path.insert(0, str(Path(__file__).parent))
from fixture_model import FixtureModel

pytestmark = pytest.mark.integration


def encoded(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()


class AuditedFixtureModel(FixtureModel):
    """Record synthetic raw envelopes through exactly the production audit API."""

    def __init__(self, store, lease, artifacts, measurements):
        super().__init__()  # The inherited fixture computes responses, without its own audit.
        self.audit_store, self.audit_lease = store, lease
        self.artifacts, self.measurements = artifacts, measurements
        self.budget = RequestBudget(context_tokens=131072, model="gpt-6-luna")

    def request_budget(self, *, images=False):
        return self.budget

    def execution_spec(self):
        return {"test_double": True, "request_budget": self.budget.spec()}

    async def call(self, task, payload, images=()):
        request_id = uuid4().hex
        body = payload.get("body", [])
        supporting = payload.get("supporting", [])
        body_ids = {item["range_id"] for item in body}
        support_ids = {item["range_id"] for item in supporting}
        measurement = {
            "task": task, "body_ranges": len(body), "body_unique_ranges": len(body_ids),
            "supporting_ranges": len(supporting), "supporting_unique_ranges": len(support_ids),
            "body_support_overlap": len(body_ids & support_ids), "image_count": len(images),
            # These are fixture diagnostics, not reported provider token usage.
            "body_utf8_bytes": len(encoded(body)), "supporting_utf8_bytes": len(encoded(supporting)),
            "request_digest": hashlib.sha256(encoded(self.budget.request(task, payload, images))).hexdigest(),
            "request_budget": self.budget.require(task, payload, images),
        }
        await self.audit_store.begin_request(
            self.audit_lease, request_id, "deterministic-storage-fixture", task,
            {"test_double": True, "external_request": False, **measurement},
        )
        await self.audit_store.finish_request(request_id, "sent")
        result = await super().call(task, payload, images)
        envelope = {"object": "synthetic.chat.completion", "model": "deterministic-storage-fixture",
                    "choices": [{"message": {"role": "assistant", "content": encoded(result).decode()},
                                 "finish_reason": "stop"}], "usage": None}
        response_ref = await asyncio.to_thread(self.artifacts.put_json, envelope)
        await self.audit_store.finish_request(request_id, "succeeded", usage=None, response_ref=response_ref)
        self.measurements.append(measurement)
        return result


def manifest_for(root, count):
    """Each independent source has the same structure and different obligations."""
    image_path = root / "shared-diagram.png"
    image = Image.new("RGB", (160, 60), "white")
    ImageDraw.Draw(image).rectangle((10, 10, 150, 50), outline="black", width=2)
    image.save(image_path)
    documents = []
    for index in range(count):
        label = chr(ord("A") + index)
        path = root / f"source-{label}.docx"
        document = Document()
        document.add_heading(f"Delivery obligations for office {label}", level=1)
        document.add_paragraph(f"The supplier must provide a signed delivery commitment for office {label}.")
        document.add_picture(str(image_path), width=Inches(2))
        document.save(path)
        documents.append({"document_id": f"document-{label}", "role": "main" if index == 0 else "attachment",
                          "document_version": "one", "path": str(path)})
    return {"documents": documents, "expected_document_ids": [item["document_id"] for item in documents]}


@asynccontextmanager
async def runtime(base_dsn, root):
    schema = "tender_growth_" + uuid4().hex
    async with await psycopg.AsyncConnection.connect(base_dsn, autocommit=True) as conn:
        await conn.execute(sql.SQL("CREATE SCHEMA {}").format(sql.Identifier(schema)))
    dsn = make_conninfo(base_dsn, options=f"-c search_path={schema}")
    artifacts = ArtifactStore(root / "artifacts")
    store = Store(dsn)
    try:
        await store.setup()
        await initialize_checkpoints(dsn)
        yield store, artifacts, Settings(dsn, "redis://unused.invalid", artifacts.root)
    finally:
        async with await psycopg.AsyncConnection.connect(base_dsn, autocommit=True) as conn:
            await conn.execute(sql.SQL("DROP SCHEMA {} CASCADE").format(sql.Identifier(schema)))


def references(value):
    if isinstance(value, dict):
        if {"sha256", "relative_path", "size"}.issubset(value):
            yield value
        else:
            for item in value.values():
                yield from references(item)
    elif isinstance(value, list):
        for item in value:
            yield from references(item)


def unique_refs(values):
    return {ref["sha256"]: ref for value in values for ref in references(value)}


def closure(artifacts, initial, excluded=()):
    seen = dict(initial)
    queue = list(initial.values())
    excluded = set(excluded)
    while queue:
        ref = queue.pop()
        if ref["media_type"] != "application/json":
            continue
        for child in references(artifacts.read_json(ref)):
            digest = child["sha256"]
            if digest not in seen and digest not in excluded:
                seen[digest] = child
                queue.append(child)
    return seen


def forbidden_keys(value):
    forbidden = {"text", "content", "grid", "cells", "base64", "messages", "completed_packs"}
    if isinstance(value, dict):
        for key, item in value.items():
            if key in forbidden:
                yield key
            yield from forbidden_keys(item)
    elif isinstance(value, (list, tuple)):
        for item in value:
            yield from forbidden_keys(item)


async def measurements(store, artifacts, row, calls):
    stage_refs = {item["stage"]: item["artifact_ref"] for item in row["artifacts"] if not item["scope"]}
    frozen = artifacts.read_json(stage_refs["freeze"])
    pack_plan = artifacts.read_json(stage_refs["pack-plan"])
    assert pack_plan["frozen_ref"] == stage_refs["freeze"]
    checked = artifacts.read_json(stage_refs["verify"])
    assert "requirements" not in checked, "Verification summaries must reference, not echo, full requirements"
    plan = artifacts.read_json(stage_refs["plan"])
    assert "items" not in plan, "The final plan must index durable per-pack results"
    plan_items = [item for entry in plan["packs"] for item in artifacts.read_json(entry["artifact_ref"])["items"]]
    assert len(plan_items) == plan["counts"]["items"]
    assert all("constraint_snapshot" not in item and "evidence" not in item
               and item["requirement_ref"]["artifact_ref"] and item["constraint_digest"] for item in plan_items)
    originals = unique_refs([document["original_ref"] for document in frozen["documents"]])
    raw_refs = unique_refs([document["raw_ref"] for document in frozen["documents"]])
    raw_documents = [artifacts.read_json(ref) for ref in raw_refs.values()]
    assert all("content" not in raw for raw in raw_documents), "Canonical raw must not duplicate rendered document text"
    pack_refs = unique_refs([item["artifact_ref"] for item in pack_plan["packs"]])
    descriptors = [artifacts.read_json(ref) for ref in pack_refs.values()]
    assert all(not list(forbidden_keys(pack)) for pack in descriptors), "Pack plans must contain selectors, not source copies"
    source_indexes = unique_refs([pack["source_index_ref"] for pack in descriptors])
    pack_plan_files = {**pack_refs, **unique_refs([stage_refs["pack-plan"], pack_plan["budget_spec_ref"]])}
    canonical = closure(artifacts, {**raw_refs, **source_indexes}, excluded=originals)
    images = unique_refs([raw["images"] for raw in raw_documents])
    image_occurrences = sum(sum(unit["kind"] == "image_region" for unit in raw["units"]) for raw in raw_documents)
    # Index files are separate small collections of references, not canonical text.
    frozen_indexes = unique_refs([item["artifact_ref"] for item in row["artifacts"]
                                 if item["stage"] in {"freeze", "freeze-document"}])
    frozen_files = {**canonical, **frozen_indexes}
    response_refs = unique_refs([request["response_ref"] for request in row["usage"]])
    assert all(request["response_ref"] is not None and request["usage"] is None for request in row["usage"])
    assert all(not list(forbidden_keys(request["metadata"])) for request in row["usage"])
    assert len(calls) == len(row["usage"])
    assert all(call["body_ranges"] == call["body_unique_ranges"]
               and call["supporting_ranges"] == call["supporting_unique_ranges"]
               and not call["body_support_overlap"] for call in calls)
    assert sum(call["task"] == "ocr" for call in calls) == 1, "Identical image content reuses its OCR receipt"
    thread_id = f"{row['run_id']}:{row['source_version']}:{row['domain_revision']}"
    checkpoint_tables = {}
    async with store.transaction() as conn:
        for table in ("checkpoints", "checkpoint_blobs", "checkpoint_writes"):
            values = await (await conn.execute(sql.SQL(
                "SELECT count(*) AS rows, COALESCE(sum(pg_column_size(t)), 0) AS bytes FROM {} t WHERE thread_id = %s"
            ).format(sql.Identifier(table)), (thread_id,))).fetchone()
            checkpoint_tables[table] = {key: int(value) for key, value in values.items()}
        audit_bytes = (await (await conn.execute(
            "SELECT COALESCE(sum(pg_column_size(r)), 0) AS bytes FROM tender_model_requests r WHERE run_id = %s",
            (row["run_id"],),
        )).fetchone())["bytes"]
    async with AsyncPostgresSaver.from_conn_string(store.dsn) as saver:
        latest = await saver.aget_tuple({"configurable": {"thread_id": thread_id}})
    state = latest.checkpoint["channel_values"]
    assert not list(forbidden_keys(state)), "Checkpoints cannot contain documents, messages, or growing pack history"
    state_refs = unique_refs([state])
    assert len(state_refs) == 6  # register, native freeze, pack plan, extract, verify, response plan.
    assert {key for key in state if key.endswith("_ref")} == {
        "registered_ref", "frozen_ref", "pack_plan_ref", "requirements_ref", "verification_ref", "response_plan_ref",
    }
    assert not {"version", "registered", "frozen", "requirements", "checked", "plan"} & state.keys()
    assert len(encoded(state)) < 8192
    files = [path for path in artifacts.root.glob("*/*") if path.is_file()]
    return {
        "documents": len(frozen["documents"]), "source_ranges": pack_plan["range_count"],
        "verified_ranges": checked["counts"]["ranges"], "packs": len(pack_plan["packs"]),
        "requests": len(calls), "original_unique_bytes": sum(ref["size"] for ref in originals.values()),
        "actual_provider_usage": None, "response_plan_items": len(plan_items),
        "frozen_unique_bytes": sum(ref["size"] for ref in frozen_files.values()),
        "pack_plan_unique_bytes": sum(ref["size"] for ref in pack_plan_files.values()),
        "audit_response_unique_bytes": sum(ref["size"] for ref in response_refs.values()),
        "audit_database_bytes": int(audit_bytes), "all_artifact_unique_bytes": sum(path.stat().st_size for path in files),
        "artifact_count": len(files), "checkpoint_tables": checkpoint_tables,
        "checkpoint_bytes": sum(value["bytes"] for value in checkpoint_tables.values()),
        "checkpoint_stage_refs": len(state_refs), "checkpoint_state_bytes": len(encoded(state)),
        "image_unique_count": len(images), "image_occurrences": image_occurrences,
        "image_hashes": sorted(images), "freeze_seconds": next(item["metadata"]["elapsed_seconds"]
                    for item in row["artifacts"] if item["stage"] == "freeze"),
        "stage_elapsed_seconds": {item["stage"]: item["metadata"]["elapsed_seconds"]
                                  for item in row["artifacts"]
                                  if not item["scope"] and "elapsed_seconds" in item["metadata"]},
        "requests_diagnostics": calls,
    }


async def test_two_and_four_times_sources_have_linear_storage_and_constant_checkpoint(tmp_path):
    dsn = os.environ.get("TENDER_TEST_DATABASE_URL") or os.environ.get("TENDER_TEST_DSN")
    if not dsn:
        pytest.skip("real PostgreSQL required for storage growth acceptance")
    reports = []
    for count in (4, 8, 16):
        root = tmp_path / f"size-{count}"
        root.mkdir()
        manifest = manifest_for(root, count)
        async with runtime(dsn, root) as (store, artifacts, settings):
            registered = await asyncio.to_thread(register_manifest, manifest, artifacts)
            run = await store.submit(artifacts.put_json(registered), uuid4().hex)
            calls = []
            def factory(s, a, lease):
                return TenderNodes(s, a, lease, AuditedFixtureModel(s, lease, a, calls))
            first = (await store.get_run(run["run_id"]))["commands"][0]
            await execute_command(run["run_id"], first["command_id"], 1, settings, factory)
            paused = await store.get_run(run["run_id"])
            assert paused["status"] == "waiting-confirmation"
            resume = await store.confirm(run["run_id"], 1, paused["checklist_digest"], "approve-fixture", domain_revision=1)
            await execute_command(run["run_id"], resume["command_id"], 1, settings, factory)
            row = await store.get_run(run["run_id"])
            assert row["status"] == "response-plan-ready"
            reports.append(await measurements(store, artifacts, row, calls))
    report_path = tmp_path / "storage-growth.json"
    report_path.write_text(json.dumps(reports, ensure_ascii=False, sort_keys=True, indent=2))
    print("storage-growth=" + json.dumps(reports, ensure_ascii=False, sort_keys=True))
    baseline = reports[0]
    for report, multiplier in zip(reports, (1, 2, 4)):
        assert report["documents"] == baseline["documents"] * multiplier
        assert report["packs"] == baseline["packs"] * multiplier
        assert report["source_ranges"] == report["verified_ranges"] == baseline["source_ranges"] * multiplier
        assert report["image_unique_count"] == 1
        assert report["image_occurrences"] == report["documents"]
        assert report["image_hashes"] == baseline["image_hashes"]
        for table, values in report["checkpoint_tables"].items():
            assert values["rows"] == baseline["checkpoint_tables"][table]["rows"]
        assert report["checkpoint_bytes"] <= baseline["checkpoint_bytes"] * 1.1
        assert report["checkpoint_stage_refs"] == baseline["checkpoint_stage_refs"]
        for field in ("original_unique_bytes", "frozen_unique_bytes", "pack_plan_unique_bytes",
                      "audit_response_unique_bytes", "audit_database_bytes", "all_artifact_unique_bytes"):
            assert report[field] <= baseline[field] * multiplier * 1.3, (field, reports)
