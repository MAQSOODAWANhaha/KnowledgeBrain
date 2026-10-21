"""Analysis v2: real parsers/checkpoints/SQL with explicitly synthetic model output."""

import copy
import hashlib
import json
import os
from uuid import uuid4

from docx import Document
import psycopg
from psycopg import sql
from psycopg.conninfo import make_conninfo
import pytest

from docreader.integrations.artifacts import ArtifactStore
from docreader.integrations.cli import _correction_ref
from docreader.integrations.store import Conflict, InvalidConfirmation, Store
from docreader.jobs.settings import Settings
from docreader.llm import ModelInfrastructureError
from docreader.tender.evidence import register_manifest
from docreader.workflow.nodes import TenderNodes, WORKFLOW_PROTOCOL
from docreader.workflow.persistence import initialize_checkpoints
from docreader.workflow.runner import execute_command
from docreader.workflow.state import STAGES, TenderState


pytestmark = pytest.mark.integration


class AnalysisModel:
    """A contract fixture, never evidence of real model business quality."""

    def __init__(self, store, artifacts, lease, calls, options):
        self.store, self.artifacts, self.lease = store, artifacts, lease
        self.calls, self.options = calls, options

    def execution_spec(self):
        return {"test_fixture": "analysis-v2"}

    async def call(self, task, payload, images=()):
        assert task != "plan", "the analysis workflow must never call legacy Plan"
        self.calls.append((task, copy.deepcopy(payload)))
        views = payload.get("source_views", [])
        if (task == "extract" and any(view.get("document_id") == "annex" for view in views)
                and self.options.pop("fail_annex_once", False)):
            raise ModelInfrastructureError("injected pre-send failure")
        digest = hashlib.sha256(json.dumps({"task": task, "payload": payload}, sort_keys=True).encode()).hexdigest()
        request_id = uuid4().hex
        prior = await self.store.begin_request(self.lease, request_id, "synthetic-analysis", task, {
            "request_digest": digest, "test_double": True, "external_request": False,
        })
        if prior["status"] == "succeeded":
            return self.artifacts.read_json(prior["response_ref"])
        await self.store.finish_request(request_id, "sent")
        if task == "extract":
            items, dispositions = [], []
            snippets = [
                ("负责人须持A证", "obligation"),
                ("持B证加2分", "scoring_rule"),
                ("未附A证复印件则否决投标", "consequence_rule"),
                ("须附A证复印件", "submission_requirement"),
                ("服务响应不超过2小时", "obligation"),
            ]
            for source in views[0]["body"]:
                selected = []
                for quote, kind in snippets:
                    if quote in source["text"]:
                        item_id = f"{payload['work_unit_id']}:i{len(items) + 1}"
                        items.append({"item_id": item_id, "kind": kind, "text": quote,
                                      "evidence": [{"range_id": source["range_id"], "quote": quote}]})
                        if self.options.get("missing_number_quote") and kind == "scoring_rule":
                            items[-1]["numbers"] = [{"value": 2, "unit": "分"}]
                        selected.append(item_id)
                if not selected and self.options.get("omit_non_item_dispositions"):
                    continue
                dispositions.append({"range_id": source["range_id"],
                                     "status": "items" if selected else "non_item",
                                     "reason": "本段给出源文规则。" if selected else "本段为章节名称。",
                                     "item_ids": selected})
                if self.options.get("disposition_item_ids") == "omitted":
                    del dispositions[-1]["item_ids"]
                elif self.options.get("disposition_item_ids") == "redundant" and selected:
                    dispositions[-1]["item_ids"] = [*selected, selected[0]]
            result = {"schema_version": 2, "items": items, "dispositions": dispositions,
                      "relations": [], "issues": []}
        elif payload["operation"] == "locate":
            result = {"schema_version": 2, "candidates": [], "findings": []}
        elif payload["operation"] == "associate":
            result = {"schema_version": 2, "relations": [], "decisions": [], "findings": []}
        else:
            findings = ([{"code": "semantic_ambiguity", "detail": "本例需要人工核对范围。", "blocking": True}]
                        if self.options.get("blocked") else [])
            result = {"schema_version": 2, "accepted_decision_ids": [], "findings": findings}
        await self.store.finish_request(request_id, "succeeded", response_ref=self.artifacts.put_json(result))
        return result


@pytest.fixture
async def analysis_runtime(tmp_path):
    base_dsn = os.environ.get("TENDER_TEST_DATABASE_URL")
    if not base_dsn:
        pytest.skip("TENDER_TEST_DATABASE_URL is required for isolated real PostgreSQL acceptance")
    schema = "analysis_test_" + uuid4().hex
    async with await psycopg.AsyncConnection.connect(base_dsn, autocommit=True) as conn:
        await conn.execute(sql.SQL("CREATE SCHEMA {}").format(sql.Identifier(schema)))
    dsn = make_conninfo(base_dsn, options=f"-c search_path={schema}")
    settings = Settings(dsn, os.environ.get("TENDER_TEST_REDIS_URL", "redis://127.0.0.1:16379/1"), tmp_path / "artifacts")
    store, artifacts = Store(dsn), ArtifactStore(settings.artifact_root)
    await store.setup()
    await initialize_checkpoints(dsn)
    documents = []
    for document_id, text in (
        ("main", "负责人须持A证；持B证加2分；未附A证复印件则否决投标；须附A证复印件。"),
        ("annex", "服务响应不超过2小时。"),
    ):
        path = tmp_path / f"{document_id}.docx"
        document = Document()
        document.add_paragraph(text)
        if document_id == "main":
            document.add_heading("阅读说明", level=1)
            document.add_paragraph("本节为阅读说明。")
        document.save(path)
        documents.append({"document_id": document_id, "document_version": "one",
                          "role": "main" if document_id == "main" else "attachment", "path": str(path)})
    registered = register_manifest({"documents": documents}, artifacts)
    run = await store.submit(artifacts.put_json(registered), uuid4().hex)
    calls = []
    try:
        yield store, artifacts, settings, run, calls
    finally:
        async with await psycopg.AsyncConnection.connect(base_dsn, autocommit=True) as conn:
            await conn.execute(sql.SQL("DROP SCHEMA {} CASCADE").format(sql.Identifier(schema)))


async def execute_analysis(runtime, **options):
    store, artifacts, settings, run, calls = runtime
    command = (await store.get_run(run["run_id"]))["commands"][-1]
    await execute_command(run["run_id"], command["command_id"], command["source_version"], settings,
                          lambda s, a, lease: TenderNodes(s, a, lease, AnalysisModel(s, a, lease, calls, options)))
    return await store.get_run(run["run_id"])


def result(artifacts, row, stage):
    return artifacts.read_json(next(item["artifact_ref"] for item in row["artifacts"]
                                    if item["stage"] == stage and not item["scope"]))


async def test_five_nodes_confirm_to_analysis_without_plan(analysis_runtime):
    store, artifacts, _, run, calls = analysis_runtime
    assert STAGES == ("register", "freeze", "extract", "verify", "confirm")
    assert "items_ref" in TenderState.__annotations__
    assert not {"requirements_ref", "response_plan_ref"} & TenderState.__annotations__.keys()
    waiting = await execute_analysis(analysis_runtime)
    assert waiting["status"] == "waiting-confirmation", result(artifacts, waiting, "verify")
    index = result(artifacts, waiting, "extract")
    assert index["schema_version"] == 2
    items = [item for entry in index["packs"] for item in artifacts.read_json(entry["artifact_ref"])["items"]]
    assert {item["kind"] for item in items} == {"obligation", "scoring_rule", "consequence_rule", "submission_requirement"}
    assert waiting["revision"]["execution_spec"]["workflow_protocol"] == WORKFLOW_PROTOCOL
    before_confirm = len(calls)
    approved = await store.confirm(run["run_id"], 1, waiting["checklist_digest"], "approve", domain_revision=1)
    assert approved == await store.confirm(run["run_id"], 1, waiting["checklist_digest"], "approve", domain_revision=1)
    ready = await execute_analysis(analysis_runtime)
    assert ready["status"] == "analysis-ready"
    assert len(calls) == before_confirm
    assert not any(item["stage"] in {"plan", "plan-pack"} for item in ready["artifacts"])
    assert result(artifacts, ready, "confirm")["checklist_digest"] == waiting["checklist_digest"]


@pytest.mark.parametrize("disposition_item_ids", ["omitted", "redundant"])
async def test_optional_response_recovery_completes_five_nodes_without_paid_repair(analysis_runtime, disposition_item_ids):
    store, artifacts, _, run, calls = analysis_runtime
    waiting = await execute_analysis(analysis_runtime, missing_number_quote=True,
                                     disposition_item_ids=disposition_item_ids)
    assert waiting["status"] == "waiting-confirmation", result(artifacts, waiting, "verify")
    index = result(artifacts, waiting, "extract")
    packs = [artifacts.read_json(entry["artifact_ref"]) for entry in index["packs"]]
    requests = [request for request in waiting["usage"] if request["task"] == "extract"]
    raw = [(request["response_ref"], artifacts.read_json(request["response_ref"])) for request in requests]
    raw_numbers = [number for _, response in raw for item in response["items"] for number in item.get("numbers", [])]
    assert raw_numbers == [{"value": 2, "unit": "分"}]
    assert len(requests) == len(index["packs"]) == 2
    assert all(request["status"] == "succeeded" and request["metadata"]["external_request"] is False
               for request in waiting["usage"])
    for _, response in raw:
        if disposition_item_ids == "omitted":
            assert all("item_ids" not in disposition for disposition in response["dispositions"])
        else:
            assert all(len(d["item_ids"]) > len(set(d["item_ids"]))
                       for d in response["dispositions"] if d["status"] == "items")
    items = [item for pack in packs for item in pack["items"]]
    assert len(items) == 5 and all("numbers" not in item for item in items)
    assert any(item["text"] == "持B证加2分" for item in items)
    for pack in packs:
        assert pack["issues"] == []
        for disposition in pack["dispositions"]:
            expected = {item["item_id"] for item in pack["items"]
                        if any(e["range_id"] == disposition["range_id"] for e in item["evidence"])}
            assert set(disposition["item_ids"]) == expected
            assert len(disposition["item_ids"]) == len(expected)
    verification = result(artifacts, waiting, "verify")
    verify_calls = [payload for task, payload in calls if task == "verify"]
    assert verify_calls and any(payload["operation"] == "review" for payload in verify_calls)
    assert {payload["operation"] for payload in verify_calls} <= {"review", "locate"}
    assert len(verify_calls) == len(verification["operation_refs"])
    assert sum(task == "extract" for task, _ in calls) == len(index["packs"])
    before_confirm = len(calls)
    await store.confirm(run["run_id"], 1, waiting["checklist_digest"], "approve-recovered", domain_revision=1)
    ready = await execute_analysis(analysis_runtime)
    assert ready["status"] == "analysis-ready" and len(calls) == before_confirm
    assert len(ready["usage"]) == len(waiting["usage"])
    assert all(artifacts.read_json(reference) == response for reference, response in raw)


async def test_semantic_issue_needs_human_verdicts_before_ready_without_new_model_calls(analysis_runtime):
    store, artifacts, _, run, calls = analysis_runtime
    waiting = await execute_analysis(analysis_runtime, blocked=True)
    assert waiting["status"] == "waiting-confirmation"
    assert result(artifacts, waiting, "extract")["packs"]
    verification = result(artifacts, waiting, "verify")
    pending = verification["human_issue_ids"]
    assert verification["blocking"] is True and verification["hard_blocking"] is False
    assert pending and {p["code"] for p in verification["issues"] if p["issue_id"] in pending} == {"semantic_ambiguity"}
    digest = waiting["checklist_digest"]

    def verdicts(ids):
        return {"schema_version": 1, "resolutions": [
            {"issue_id": issue_id, "verdict": "acknowledged", "reason": "人工核对原文，保留为已知问题。"} for issue_id in ids]}

    for key, submitted in [("missing", None), ("partial", verdicts(pending[1:])),
                           ("extra", verdicts([*pending, "unrelated-issue"]))]:
        with pytest.raises(InvalidConfirmation):
            await store.confirm(run["run_id"], 1, digest, key, domain_revision=1, issue_resolutions=submitted)
    assert not (await store.get_run(run["run_id"]))["confirmations"]
    before = len(calls)
    await store.confirm(run["run_id"], 1, digest, "resolved", domain_revision=1, issue_resolutions=verdicts(pending))
    ready = await execute_analysis(analysis_runtime)
    assert ready["status"] == "analysis-ready" and len(calls) == before
    confirmed = result(artifacts, ready, "confirm")
    assert confirmed["checklist_digest"] == digest
    assert confirmed["issue_resolutions"] == verdicts(sorted(pending))


async def test_omitted_non_item_disposition_reaches_confirmation_after_complete_independent_review(analysis_runtime):
    store, artifacts, _, run, calls = analysis_runtime
    waiting = await execute_analysis(analysis_runtime, omit_non_item_dispositions=True)
    assert waiting["status"] == "waiting-confirmation"
    index = result(artifacts, waiting, "extract")
    packs = [artifacts.read_json(entry["artifact_ref"]) for entry in index["packs"]]
    assert sum(len(pack["items"]) for pack in packs) == 5
    unresolved = [d for pack in packs for d in pack["dispositions"] if d["status"] == "unresolved"]
    assert len(unresolved) == 1 and unresolved[0]["item_ids"] == []
    verification = result(artifacts, waiting, "verify")
    pending = verification["human_issue_ids"]
    assert verification["blocking"] is True and verification["hard_blocking"] is False
    assert verification["coverage"]["missing"] == {}
    assert unresolved[0]["range_id"] in verification["coverage"]["reviewed"]["range"]
    assert pending and set(pending) == {p["issue_id"] for p in verification["issues"]
                                       if p["code"] == "extraction_disposition_unresolved"}
    with pytest.raises(InvalidConfirmation):
        await store.confirm(run["run_id"], 1, waiting["checklist_digest"], "missing-resolution", domain_revision=1)
    assert not (await store.get_run(run["run_id"]))["confirmations"]

    resolutions = {"schema_version": 1, "resolutions": [
        {"issue_id": identity, "verdict": "dismissed", "reason": "合成验收处理：该段仅为阅读说明，未新增事项。"}
        for identity in pending]}
    before = len(calls)
    await store.confirm(run["run_id"], 1, waiting["checklist_digest"], "fixture-resolved", domain_revision=1,
                        issue_resolutions=resolutions)
    ready = await execute_analysis(analysis_runtime)
    assert ready["status"] == "analysis-ready" and len(calls) == before
    assert len(ready["usage"]) == len(waiting["usage"])
    assert all(request["metadata"]["external_request"] is False for request in ready["usage"])
    assert result(artifacts, ready, "confirm")["issue_resolutions"] == resolutions


async def test_partial_extraction_recovery_reuses_successful_receipts(analysis_runtime):
    store, _, _, run, calls = analysis_runtime
    partial = await execute_analysis(analysis_runtime, fail_annex_once=True)
    assert partial["status"] == "waiting-confirmation"
    assert partial["commands"][-1]["infrastructure_failures"] == 0
    saved = partial["artifacts"]
    assert len([item for item in saved if item["stage"] == "extract-pack"]) == 2
    before = len(calls)
    recovered = await execute_analysis(analysis_runtime)
    assert recovered["status"] == "waiting-confirmation"
    assert recovered["artifacts"] == saved and len(calls) == before
    assert sum(task == "extract" and payload["source_views"][0]["document_id"] == "main"
               for task, payload in calls) == 1
    assert sum(task == "extract" and payload["source_views"][0]["document_id"] == "annex"
               for task, payload in calls) == 2


async def test_confirmation_commit_failure_never_reports_ready_and_replays_safely(analysis_runtime, monkeypatch):
    store, _, _, run, calls = analysis_runtime
    waiting = await execute_analysis(analysis_runtime)
    await store.confirm(run["run_id"], 1, waiting["checklist_digest"], "approve")
    original, attempts = Store.commit_artifact, []

    async def fail_once(self, lease, stage, *args, **kwargs):
        if stage == "confirm" and not attempts:
            attempts.append(stage)
            raise ConnectionError("injected confirmation artifact failure")
        return await original(self, lease, stage, *args, **kwargs)

    monkeypatch.setattr(Store, "commit_artifact", fail_once)
    before = len(calls)
    with pytest.raises(ConnectionError):
        await execute_analysis(analysis_runtime)
    failed = await store.get_run(run["run_id"])
    assert failed["status"] == "pending"
    assert not any(item["stage"] == "confirm" for item in failed["artifacts"])
    assert (await execute_analysis(analysis_runtime))["status"] == "analysis-ready"
    assert len(calls) == before


async def test_correction_revalidates_all_packs_and_invalidates_old_confirmation(analysis_runtime):
    store, artifacts, _, run, _ = analysis_runtime
    waiting = await execute_analysis(analysis_runtime)
    await store.confirm(run["run_id"], 1, waiting["checklist_digest"], "old-confirmation")
    ready = await execute_analysis(analysis_runtime)
    baseline = next(item["artifact_ref"] for item in ready["artifacts"] if item["stage"] == "extract")
    index = artifacts.read_json(baseline)
    inline = {"schema_version": 2, "packs": []}
    for entry in index["packs"]:
        value = artifacts.read_json(entry["artifact_ref"])
        inline["packs"].append({"pack_id": entry["pack_id"], **{key: value[key] for key in
                                ("schema_version", "items", "dispositions", "relations", "issues")}})
    partial = {"schema_version": 2, "packs": inline["packs"][:1]}
    with pytest.raises(ValueError, match="every frozen"):
        _correction_ref(artifacts, ready, 1, 1, baseline["sha256"], partial)
    corrected = _correction_ref(artifacts, ready, 1, 1, baseline["sha256"], inline)
    await store.correct(run["run_id"], 1, 1, baseline["sha256"], corrected, "correction", "Confirm source interpretation")
    revised = await execute_analysis(analysis_runtime)
    assert revised["status"] == "waiting-confirmation"
    assert revised["domain_revision"] == 2
    assert not revised["confirmations"] and revised["history"]["confirmations"]
    assert revised["checklist_digest"] != waiting["checklist_digest"]
    with pytest.raises(InvalidConfirmation):
        await store.confirm(run["run_id"], 1, waiting["checklist_digest"], "stale", domain_revision=1)
    assert result(artifacts, revised, "verify")["items_ref"] == revised["revision"]["corrected_items_ref"]


async def test_cancel_and_real_source_revision_do_not_reuse_old_confirmation(analysis_runtime):
    store, artifacts, _, run, _ = analysis_runtime
    waiting = await execute_analysis(analysis_runtime)
    await store.cancel(run["run_id"])
    with pytest.raises(InvalidConfirmation):
        await store.confirm(run["run_id"], 1, waiting["checklist_digest"], "cancelled")
    manifest = artifacts.read_json(waiting["manifest_ref"])
    manifest["documents"][1]["document_version"] = "two"
    await store.revise(run["run_id"], artifacts.put_json(manifest), "changed-source")
    revised = await execute_analysis(analysis_runtime)
    assert revised["source_version"] == 2 and revised["status"] == "waiting-confirmation"
    assert revised["checklist_digest"] != waiting["checklist_digest"]
    with pytest.raises(InvalidConfirmation):
        await store.confirm(run["run_id"], 1, waiting["checklist_digest"], "old-source")


async def test_legacy_execution_spec_and_results_are_read_only(analysis_runtime):
    store, artifacts, _, run, calls = analysis_runtime
    command = (await store.get_run(run["run_id"]))["commands"][-1]
    lease = await store.acquire(run["run_id"], command["command_id"], 1)
    await store.pin_execution_spec(lease, {"legacy_fixture": "six-stages"})
    ref = artifacts.put_json({"requirements": [], "dispositions": [], "issues": []})
    await store.commit_artifact(lease, "extract", ref)
    await store.commit_artifact(lease, "freeze", artifacts.put_json({"documents": []}))
    await store.commit_artifact(lease, "pack-plan", artifacts.put_json({"packs": []}),
                                metadata={"plan_digest": "legacy-plan"})
    nodes = TenderNodes(store, artifacts, lease, AnalysisModel(store, artifacts, lease, calls, {}))
    with pytest.raises(Conflict, match="read-only"):
        await nodes.prepare()
    await store.release(lease, "waiting-confirmation", checklist_digest="legacy-checklist")
    with pytest.raises(Conflict, match="read-only"):
        await store.confirm(run["run_id"], 1, "legacy-checklist", "must-not-resume")
    with pytest.raises(Conflict, match="read-only"):
        await store.correct(run["run_id"], 1, 1, ref["sha256"], ref, "must-not-correct", "Legacy is historical")
    with pytest.raises(Conflict, match="read-only"):
        await store.review_scope(run["run_id"], 1, 1, "legacy-plan", ref,
                                 "must-not-review", "Legacy is historical")
    with pytest.raises(Conflict, match="read-only"):
        await store.revise(run["run_id"], artifacts.put_json({"documents": ["changed"]}), "must-not-convert")
    row = await store.get_run(run["run_id"])
    assert next(item["artifact_ref"] for item in row["artifacts"] if item["stage"] == "extract") == ref
    assert not calls and row["status"] == "waiting-confirmation"


async def test_legacy_pending_command_cannot_change_status_or_request_audit(analysis_runtime):
    store, _, settings, run, _ = analysis_runtime
    command = (await store.get_run(run["run_id"]))["commands"][-1]
    lease = await store.acquire(run["run_id"], command["command_id"], 1)
    await store.pin_execution_spec(lease, {"legacy_fixture": "six-stages"})
    await store.release(lease, "retry")
    before = await store.get_run(run["run_id"])
    assert not await store.claim_commands()
    with pytest.raises(Conflict, match="read-only"):
        await execute_command(run["run_id"], command["command_id"], 1, settings, TenderNodes)
    after = await store.get_run(run["run_id"])
    assert (after["status"], after["epoch"], after["commands"], after["usage"]) == (
        before["status"], before["epoch"], before["commands"], before["usage"])


async def test_status_projects_current_items_from_verified_refs(analysis_runtime, monkeypatch):
    from argparse import Namespace
    from docreader.integrations.cli import run as cli_run

    _, artifacts, settings, run, _ = analysis_runtime
    waiting = await execute_analysis(analysis_runtime)
    monkeypatch.setenv("TENDER_DATABASE_URL", settings.dsn)
    monkeypatch.setenv("TENDER_REDIS_URL", settings.redis_url)
    monkeypatch.setenv("TENDER_ARTIFACT_ROOT", str(artifacts.root))
    status = await cli_run(Namespace(command="status", run_id=run["run_id"]))
    projected = status["results"]["current_items"]
    assert len(projected) == result(artifacts, waiting, "verify")["counts"]["items"]
    assert all(value["item_ref"]["artifact_ref"] and value["item"]["evidence"] for value in projected)
    observation = status["budget_observation"]
    assert observation["requests"] == len(status["usage"])
    assert {p["phase"] for p in observation["plans"]} == {"extract", "review", "locate", "associate", "joint"}
    assert observation["plans"][0]["planned_requests"] == len(status["results"]["pack-plan"]["packs"])
    assert observation["calibration_status"] == "observations_only"
