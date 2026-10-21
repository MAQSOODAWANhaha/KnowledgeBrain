"""V2 complete-input review, joint reading and recovery on isolated real PostgreSQL.

All model answers are synthetic contract fixtures, with no external requests.
"""
from collections import Counter
from copy import deepcopy
from itertools import combinations
import json
from uuid import uuid4

from docx import Document
import httpx
import psycopg
import pytest

from docreader.integrations.store import InvalidConfirmation, Store
from docreader.llm import ModelConfig, ModelInfrastructureError, OpenAIModel, RequestBudget
from docreader.llm.request import ModelJSONError, ModelWorkUnitExhausted
from docreader.tender.evidence import EvidenceRepository, register_manifest
from docreader.workflow.nodes import TenderNodes
from docreader.workflow.runner import execute_command

from fixture_analysis_model import AnalysisFixtureModel
from fixture_model import fixture_request_digest
from test_phase1 import output, runtime as runtime

pytestmark = pytest.mark.integration


async def _submit(runtime, *, ballast=False, ballast_repeats=70, document_ids=("complete-input",),
                  suffixes="abcdefgh", condition=False):
    store, artifacts, settings, _ = runtime
    documents = []
    for document_id in document_ids:
        path = settings.artifact_root.parent / f"{document_id}.docx"
        document = Document()
        if condition:
            document.add_heading("第一章 适用范围", level=1)
            document.add_paragraph("仅第一标段适用。")
        for suffix in suffixes:
            if ballast:
                document.add_heading(f"Background package {suffix}", level=2 if condition else 1)
                document.add_paragraph("Background explanation of the delivery workflow. " * ballast_repeats)
            document.add_heading(f"Delivery package {suffix}", level=2 if condition else 1)
            document.add_paragraph(f"The supplier must provide a signed commitment for package {suffix} in {document_id}.")
        document.save(path)
        documents.append({"document_id": document_id, "document_version": "one", "role": "main", "path": str(path)})
    registered = register_manifest({"documents": documents}, artifacts)
    return await store.submit(artifacts.put_json(registered), uuid4().hex)


def _executor(runtime, run, *, scope_repeats=12, fail_operation=None, candidates=False,
              unknown_candidate=False, include_conditions=False, accept_decisions=True, relation_repeats=1,
              reason_repeats=10, scattered_failures=False):
    store, artifacts, settings, _ = runtime
    budget = RequestBudget(context_tokens=32768)
    calls, attempts = [], Counter()
    work_units, unit_attempts = {}, Counter()
    failed = False

    class CompleteInputModel(AnalysisFixtureModel):
        def request_budget(self):
            return budget

        def response(self, task, payload, images=()):
            result = super().response(task, payload, images)
            if task == "extract":
                for item in result["items"]:
                    # Deliberate source-backed repetition stresses complete input
                    # admission; it is not a semantic model-quality evaluation.
                    item["scope"] = item["text"] * scope_repeats
                    if include_conditions:
                        condition = "仅第一标段适用。"
                        sources = [source for view in payload["source_views"] for key in ("body", "supporting")
                                   for source in view[key] if condition in source["text"]]
                        if sources:
                            item["conditions"] = [condition]
                            item["evidence"].extend({"range_id": source["range_id"], "quote": condition}
                                                    for source in sources)
                for disposition in result["dispositions"]:
                    supported = [item["item_id"] for item in result["items"]
                                 if any(e["range_id"] == disposition["range_id"] for e in item["evidence"])]
                    if supported:
                        disposition.update(status="items", item_ids=supported)
                    disposition["reason"] += " The fixture source supports this classification." * reason_repeats
            elif payload["operation"] == "locate" and (candidates or unknown_candidate):
                if unknown_candidate:
                    result["candidates"] = [{"candidate_id": payload["work_unit_id"] + ":c1", "issue_ids": [],
                                             "range_ids": ["missing-range"], "item_ids": [],
                                             "reason": "Injected unknown provenance must be rejected."}]
                else:
                    # Adjacent disjoint pairs make multiple independently durable
                    # association/joint operations when the discovery block fits.
                    items = payload["items"]
                    result["candidates"] = [{"candidate_id": f"{payload['work_unit_id']}:c{n + 1}",
                        "issue_ids": [q["issue_id"] for q in payload["issues"]],
                        "range_ids": [], "item_ids": [a["item_id"], b["item_id"]],
                        "reason": "Read both synthetic source statements together."}
                        for n, (a, b) in enumerate(zip(items[::2], items[1::2]))]
            elif payload["operation"] == "associate":
                items = payload["items"]
                if len(items) >= 2:
                    result["relations"] = [{"relation_id": payload["work_unit_id"] + ":l1",
                        "from_item_id": items[0]["item_id"], "to_item_id": items[1]["item_id"],
                        "relation": "The synthetic sources concern separate delivery commitments." * relation_repeats,
                        "evidence": [items[0]["evidence"][0], items[1]["evidence"][0]]}]
                if payload["issues"] or (candidates and items):
                    decision = {"decision_id": payload["work_unit_id"] + ":d1",
                        "issue_ids": [q["issue_id"] for q in payload["issues"]],
                        "reason": "The jointly supplied synthetic sources clarify this reference.",
                        "evidence": [item["evidence"][0] for item in items]}
                    if not decision["issue_ids"] and len(items) >= 2:
                        decision["amendment"] = {
                            "target_item_id": items[0]["item_id"], "effect": "replace",
                            "affected_scope": "Synthetic fixture wording only.",
                            "retained_rules": [], "replacement_item_ids": [items[1]["item_id"]],
                        }
                    if decision["issue_ids"] or decision.get("amendment"):
                        result["decisions"] = [decision]
            elif payload["operation"] == "review" and not accept_decisions:
                result["accepted_decision_ids"] = []
            return result

        async def call(self, task, payload, images=()):
            nonlocal failed
            budget.require(task, payload, images)
            operation = payload.get("operation", task)
            if operation == "review" and payload["work_unit_id"].startswith("joint-"):
                operation = "joint"
            attempts[operation] += 1
            calls.append((self.lease.domain_revision, task, deepcopy(payload)))
            error = None
            if operation == fail_operation and attempts[operation] == 2 and not failed:
                failed = True
                error = ModelInfrastructureError("confirmed fixture rejection after a durable sibling result")
            if scattered_failures:
                key = (operation, payload.get("work_unit_id"))
                if key not in work_units:
                    work_units[key] = 1 + sum(unit[0] == operation for unit in work_units)
                unit_attempts[key] += 1
                ordinal = work_units[key]
                limit = {("extract", 1): 2, ("extract", 2): 1,
                         ("review", 1): 1, ("review", 2): 1}.get((operation, ordinal), 0)
                if unit_attempts[key] <= limit:
                    error = (ModelInfrastructureError("confirmed fixture HTTP 500")
                             if (operation, ordinal) == ("review", 1) else
                             ModelJSONError("json_syntax_error", "Extra data", '{"schema_version":2}]}',
                                            line=1, column=21, offset=20))
            if error is not None:
                request_id = uuid4().hex
                await self.store.begin_request(self.lease, request_id, "synthetic-analysis-fixture", task, {
                    "request_digest": fixture_request_digest(task, payload, images),
                    "test_double": True, "external_request": False,
                })
                await self.store.finish_request(request_id, "sent")
                response_ref = (self.artifacts.put_json({"content": '{"schema_version":2}]}',
                                "finish_reason": "stop", "stream_complete": True})
                                if isinstance(error, ModelJSONError) else None)
                await self.store.finish_request(request_id, "failed", response_ref=response_ref,
                    error={"type": "invalid_json_response" if isinstance(error, ModelJSONError) else
                                    "injected_confirmed_service_rejection"})
                raise error
            return await super().call(task, payload, images)

    async def execute():
        command = (await store.get_run(run["run_id"]))["commands"][-1]
        await execute_command(run["run_id"], command["command_id"], command["source_version"], settings,
                              lambda s, a, lease: TenderNodes(s, a, lease, CompleteInputModel(s, lease, artifacts=a)))
        return await store.get_run(run["run_id"])

    return execute, calls, budget


def _extractions(artifacts, row):
    index = output(artifacts, row, "extract")
    return index, [artifacts.read_json(entry["artifact_ref"]) for entry in index["packs"]]


async def _correct(runtime, row, edited):
    store, artifacts, _, _ = runtime
    index, _ = _extractions(artifacts, row)
    for entry, extraction in zip(index["packs"], edited, strict=True):
        entry["artifact_ref"] = artifacts.put_json(extraction)
    original = next(item["artifact_ref"] for item in row["artifacts"] if item["stage"] == "extract" and not item["scope"])
    await store.correct(row["run_id"], row["source_version"], row["domain_revision"], original["sha256"],
                        artifacts.put_json(index), uuid4().hex, "Review complete analysis against frozen evidence")


def _calls(calls, operation, revision=1):
    return [payload for current, task, payload in calls if task == "verify" and current == revision
            and ((operation == "joint" and payload["work_unit_id"].startswith("joint-")) or
                 (payload["operation"] == operation and not payload["work_unit_id"].startswith("joint-")))]


def _assert_extract_budget(artifacts, row, calls, budget):
    repository = EvidenceRepository(artifacts)
    packs = {entry["pack_id"]: repository.load_pack(entry) for entry in output(artifacts, row, "pack-plan")["packs"]}
    sent = [payload for _, task, payload in calls if task == "extract"]
    assert {payload["work_unit_id"] for payload in sent} == packs.keys()
    for payload in sent:
        assert budget.measure("extract", payload) == packs[payload["work_unit_id"]]["planning_details"]["budget"]


async def test_complete_review_repartitions_with_exact_items_and_source_coverage(runtime):
    store, artifacts, _, _ = runtime
    run = await _submit(runtime)
    execute, calls, budget = _executor(runtime, run)
    row = await execute()
    checked = output(artifacts, row, "verify")
    assert row["status"] == "waiting-confirmation", checked["issues"]
    source_plan = output(artifacts, row, "pack-plan")
    _assert_extract_budget(artifacts, row, calls, budget)
    _, extractions = _extractions(artifacts, row)
    assert len(source_plan["packs"]) == 1
    reviews = _calls(calls, "review")
    assert len(reviews) > len(source_plan["packs"])
    assert Counter(item["item_id"] for p in reviews for item in p["items"]) == Counter(
        item["item_id"] for extraction in extractions for item in extraction["items"])
    assert Counter(d["range_id"] for p in reviews for d in p["dispositions"]) == Counter(
        d["range_id"] for extraction in extractions for d in extraction["dispositions"])
    assert all(budget.fits("verify", p) for p in reviews + _calls(calls, "locate"))
    item_ids = [item["item_id"] for extraction in extractions for item in extraction["items"]]
    discovered_pairs = {frozenset(pair) for payload in _calls(calls, "locate")
                        for pair in combinations([item["item_id"] for item in payload["items"]], 2)}
    assert discovered_pairs == {frozenset(pair) for pair in combinations(item_ids, 2)}
    plan = artifacts.read_json(checked["review_plan_ref"])
    assert len(plan["requests"]) == len(reviews)
    assert all("source_views" not in d and d["input_digest"] and d["budget"]["fits"] for d in plan["requests"])
    await store.confirm(run["run_id"], 1, row["checklist_digest"], "review-complete")
    before = len(calls)
    ready = await execute()
    assert ready["status"] == "analysis-ready"
    assert len(calls) == before and not any(r["task"] == "plan" for r in ready["usage"])


async def test_indivisible_complete_review_blocks_before_affected_model_send(runtime):
    store, artifacts, _, _ = runtime
    run = await _submit(runtime, suffixes="a")
    execute, calls, _ = _executor(runtime, run, scope_repeats=400)
    row = await execute()
    checked = output(artifacts, row, "verify")
    assert row["status"] == "blocked"
    assert all(not payload["items"] for payload in _calls(calls, "review"))
    assert "analysis_budget_exceeded" in {p["code"] for p in checked["issues"]}
    assert "analysis_review_incomplete" in {p["code"] for p in checked["issues"]}
    plan = artifacts.read_json(checked["review_plan_ref"])
    assert plan["requests"] and all(p["budget"]["fits"] is False for p in plan["requests"])
    with pytest.raises(InvalidConfirmation):
        await store.confirm(run["run_id"], 1, row["checklist_digest"], "cannot-force-budget")
    before = len(calls)
    await execute()
    assert len(calls) == before


async def test_joint_review_budget_includes_full_proposed_relation_and_blocks_before_send(runtime):
    _, artifacts, _, _ = runtime
    run = await _submit(runtime, suffixes="ab")
    execute, calls, budget = _executor(runtime, run, scope_repeats=1, candidates=True, relation_repeats=200)
    row = await execute()
    checked = output(artifacts, row, "verify")
    assert row["status"] == "blocked"
    associations = _calls(calls, "associate")
    assert len(associations) == 1 and not _calls(calls, "joint")
    association_plan = artifacts.read_json(checked["association_plan_ref"])
    assert all(entry["budget"]["fits"] is True for entry in association_plan["requests"])
    receipts = [entry for entry in row["artifacts"] if entry["stage"] == "analysis-associate"]
    assert len(receipts) == 1
    proposal = artifacts.read_json(receipts[0]["artifact_ref"])
    joint_payload = {**associations[0], "operation": "review",
                     "work_unit_id": associations[0]["work_unit_id"].replace("associate-", "joint-", 1),
                     "relations": [*associations[0]["relations"], *proposal["relations"]],
                     "decisions": proposal["decisions"]}
    plan = artifacts.read_json(checked["joint_plan_ref"])
    assert len(plan["requests"]) == 1
    descriptor = plan["requests"][0]
    assert descriptor["relation_ids"] == [entry["relation_id"] for entry in proposal["relations"]]
    assert descriptor["budget"] == budget.measure("verify", joint_payload)
    assert descriptor["budget"]["fits"] is False
    without_relations = budget.measure("verify", {**joint_payload, "relations": []})
    assert without_relations["input_tokens"] < descriptor["budget"]["input_tokens"]
    assert "analysis_budget_exceeded" in {problem["code"] for problem in checked["issues"]}
    assert "analysis_review_incomplete" in {problem["code"] for problem in checked["issues"]}


@pytest.mark.parametrize("operation", ["review", "locate", "associate", "joint"])
async def test_partial_analysis_success_reuses_durable_operation_receipts(runtime, operation):
    store, artifacts, _, _ = runtime
    run = await _submit(runtime)
    # Association/joint need complete candidate components that fit together.
    execute, calls, _ = _executor(runtime, run, scope_repeats=2 if operation in {"associate", "joint"} else 12,
                                  fail_operation=operation, candidates=operation in {"associate", "joint"},
                                  reason_repeats=1 if operation in {"associate", "joint"} else 10)
    recovered = await execute()
    assert recovered["status"] == "waiting-confirmation", output(artifacts, recovered, "verify")["issues"]
    assert recovered["commands"][-1]["execution_count"] == 1
    assert recovered["commands"][-1]["infrastructure_failures"] == 0
    assert len([request for request in recovered["usage"] if request["status"] == "failed"]) == 1
    stage = "analysis-" + operation
    saved = [item for item in recovered["artifacts"] if item["stage"] == stage]
    assert len(saved) >= 2
    counts = Counter(payload["work_unit_id"] for payload in _calls(calls, operation))
    assert sorted(counts.values()) == [1] * (len(saved) - 1) + [2]
    before = len(calls)
    unchanged = await execute()
    assert len(calls) == before
    assert unchanged["artifacts"] == recovered["artifacts"]
    assert unchanged["usage"] == recovered["usage"]
    checked = output(artifacts, recovered, "verify")
    if checked.get("human_issue_ids"):
        await store.confirm(run["run_id"], 1, recovered["checklist_digest"], "recovery-confirm",
                            issue_resolutions={"schema_version": 1, "resolutions": [
                                {"issue_id": issue_id, "verdict": "dismissed",
                                 "reason": "Synthetic recovery does not re-open fixture wording."}
                                for issue_id in checked["human_issue_ids"]]})
    else:
        await store.confirm(run["run_id"], 1, recovered["checklist_digest"], "recovery-confirm")
    final_count = len(calls)
    assert (await execute())["status"] == "analysis-ready"
    assert len(calls) == final_count


async def test_same_source_revision_invalidates_analysis_operations_and_confirmation(runtime):
    store, artifacts, _, _ = runtime
    run = await _submit(runtime)
    execute, calls, _ = _executor(runtime, run, scope_repeats=6)
    first = await execute()
    await store.confirm(run["run_id"], 1, first["checklist_digest"], "confirm-one")
    ready = await execute()
    frozen = output(artifacts, ready, "freeze")
    first_plan = output(artifacts, ready, "analysis-review-plan")
    _, extractions = _extractions(artifacts, ready)
    extractions[0]["items"][0]["domains"] = ["delivery"]
    await _correct(runtime, ready, extractions)
    pending = await store.get_run(run["run_id"])
    assert pending["domain_revision"] == 2 and pending["checklist_digest"] is None
    assert not any(item["stage"].startswith("analysis-") for item in pending["artifacts"])
    with pytest.raises(InvalidConfirmation):
        await store.confirm(run["run_id"], 1, first["checklist_digest"], "stale", domain_revision=1)
    second = await execute()
    assert second["status"] == "waiting-confirmation", output(artifacts, second, "verify")["issues"]
    assert second["checklist_digest"] != first["checklist_digest"]
    assert output(artifacts, second, "freeze") == frozen
    assert output(artifacts, second, "analysis-review-plan") != first_plan
    assert all(task != "extract" for revision, task, _ in calls if revision == 2)
    assert any(item.get("domains") == ["delivery"] for p in _calls(calls, "review", 2) for item in p["items"])


async def test_single_item_correction_resends_only_affected_verify_requests(runtime):
    store, artifacts, _, _ = runtime
    run = await _submit(runtime)
    execute, calls, _ = _executor(runtime, run, scope_repeats=6)
    first = await execute()
    assert first["status"] == "waiting-confirmation"
    first_verify = {row["metadata"]["request_digest"]: row
                    for row in first["usage"] if row["task"] == "verify" and row["status"] == "succeeded"}
    _, extractions = _extractions(artifacts, first)
    target = extractions[0]["items"][0]
    target["text"] = target["text"] + " Host-corrected wording."
    await _correct(runtime, first, extractions)
    second = await execute()
    assert second["status"] == "waiting-confirmation", output(artifacts, second, "verify")["issues"]
    second_digests = {row["metadata"]["request_digest"] for row in second["usage"] if row["task"] == "verify"}
    planned = {fixture_request_digest(task, payload)
               for revision, task, payload in calls if revision == 2 and task == "verify"}
    reused = first_verify.keys() & planned
    assert reused
    assert reused.isdisjoint(second_digests)
    assert second_digests
    assert second_digests <= planned


async def test_unknown_candidate_provenance_exhausts_four_attempts_without_outer_retries(runtime):
    store, _, _, _ = runtime
    run = await _submit(runtime, suffixes="ab")
    execute, calls, _ = _executor(runtime, run, scope_repeats=1, unknown_candidate=True)
    with pytest.raises(ModelWorkUnitExhausted):
        await execute()
    row = await store.get_run(run["run_id"])
    assert row["status"] == "failed"
    assert row["commands"][-1]["infrastructure_failures"] == 0
    assert row["commands"][-1]["execution_count"] == 1
    assert len([a for a in row["artifacts"] if a["stage"] == "model-rejection"]) == 4
    assert len(_calls(calls, "locate")) == 4
    assert len([request for request in row["usage"] if request["task"] == "verify"
                and request["metadata"]["request_digest"] in {
                    fixture_request_digest("verify", payload) for payload in _calls(calls, "locate")}]) == 4
    assert _calls(calls, "review") and not _calls(calls, "associate")
    before = len(calls)
    await execute()
    assert len(calls) == before
    assert (await store.get_run(run["run_id"]))["usage"] == row["usage"]


@pytest.mark.parametrize("interruptions", [1, 4], ids=["retry-and-restart", "four-attempt-limit"])
async def test_interrupted_sse_retries_preserve_receipts_and_completed_response_cache(runtime, monkeypatch, interruptions):
    store, artifacts, settings, _ = runtime
    run = await _submit(runtime, suffixes="ab")
    sent, fixture, failed_commit = [], AnalysisFixtureModel(), False
    prompts = {"extract": "Synthetic extraction fixture.", "verify": "Synthetic verification fixture."}
    config = ModelConfig.from_env({
        "KNOWLEDGEBRAIN_CHAT_BASE_URL": "https://synthetic.invalid/v1",
        "KNOWLEDGEBRAIN_CHAT_API_KEY": "synthetic-key",
        "KNOWLEDGEBRAIN_CHAT_MODEL": "synthetic-contract-fixture",
        "KNOWLEDGEBRAIN_CHAT_MAX_CONTEXT_TOKENS": "131072",
    })

    def event(content=None, *, finish=None):
        return ("data: " + json.dumps({"choices": [{"index": 0, "finish_reason": finish,
            "delta": {} if content is None else {"content": content}}]}) + "\n\n").encode()

    class InterruptedStream(httpx.AsyncByteStream):
        async def __aiter__(self):
            yield event('{"schema_version":')
            raise httpx.ReadError("injected SSE interruption after response data")

        async def aclose(self):
            pass

    def handler(request):
        payload = json.loads(json.loads(request.content)["messages"][1]["content"][0]["text"])
        task = "verify" if "operation" in payload else "extract"
        sent.append((task, payload))
        headers = {"content-type": "text/event-stream"}
        if task == "extract" and sum(task == "extract" for task, _ in sent) <= interruptions:
            return httpx.Response(200, stream=InterruptedStream(), headers=headers)
        result = ({"schema_version": 2, "candidates": [], "findings": []}
                  if payload.get("operation") == "locate" else fixture.response(task, payload))
        content = json.dumps(result, ensure_ascii=False)
        return httpx.Response(200, content=event(content) + event(finish="stop") + b"data: [DONE]\n\n",
                              headers=headers)

    commit = Store.commit_artifact

    async def lose_commit_once(self, lease, stage, *args, **kwargs):
        nonlocal failed_commit
        if stage == "extract-pack" and not failed_commit and interruptions == 1:
            failed_commit = True
            raise psycopg.OperationalError("injected loss after complete provider response")
        return await commit(self, lease, stage, *args, **kwargs)

    monkeypatch.setattr(Store, "commit_artifact", lose_commit_once)

    def factory(s, a, lease):
        model = OpenAIModel(config, s, lease, prompts, artifacts=a, transport=httpx.MockTransport(handler))
        # Keep this HTTP recovery test independent of external tokenizer assets.
        model.request_budget = lambda **kwargs: RequestBudget(
            context_tokens=131072, model=config.chat.model, prompts=prompts,
            correction=kwargs.get("correction", False))
        return TenderNodes(s, a, lease, model)

    async def execute():
        command = (await store.get_run(run["run_id"]))["commands"][-1]
        await execute_command(run["run_id"], command["command_id"], 1, settings, factory)
        return await store.get_run(run["run_id"])

    with pytest.raises(psycopg.OperationalError if interruptions == 1 else ModelWorkUnitExhausted):
        await execute()
    first = await store.get_run(run["run_id"])
    unknown = [request for request in first["usage"] if request["status"] == "uncertain"]
    assert len(unknown) == interruptions and all(request["response_ref"] for request in unknown)
    receipts = {request["request_id"]: artifacts.read_json(request["response_ref"]) for request in unknown}
    assert all(not response["done"] and response["usage"] is None for response in receipts.values())
    assert all(response["activity"]["raw_received_chunks"] > 0 and response["activity"]["content_deltas"] == 1
               for response in receipts.values())
    expected_sends = 2 if interruptions == 1 else 4
    assert sum(task == "extract" for task, _ in sent) == expected_sends
    if interruptions == 4:
        assert first["status"] == "failed"
        assert first["commands"][-1]["infrastructure_failures"] == 0
        await execute()
        assert (await store.get_run(run["run_id"]))["usage"] == first["usage"]
        assert len(sent) == 4
        return

    assert first["status"] == "pending" and first["commands"][-1]["infrastructure_failures"] == 1
    waiting = await execute()
    assert waiting["status"] == "waiting-confirmation"
    assert sum(task == "extract" for task, _ in sent) == 2
    for request in first["usage"]:
        assert next(current for current in waiting["usage"] if current["request_id"] == request["request_id"]) == request
    for request in unknown:
        assert artifacts.read_json(request["response_ref"]) == receipts[request["request_id"]]
    assert all(count <= 4 for count in Counter(
        (request["task"], request["metadata"]["logical_request_digest"]) for request in waiting["usage"]).values())
    before_confirm = len(sent)
    await store.confirm(run["run_id"], 1, waiting["checklist_digest"], "confirm-interrupted-sse-recovery")
    assert (await execute())["status"] == "analysis-ready"
    assert len(sent) == before_confirm


async def test_scattered_extract_and_verify_failures_recover_in_one_command(runtime):
    _, artifacts, _, _ = runtime
    run = await _submit(runtime, document_ids=("first-extract", "second-extract"))
    execute, calls, _ = _executor(runtime, run, scope_repeats=6, reason_repeats=2, scattered_failures=True)
    row = await execute()
    checked = output(artifacts, row, "verify")
    assert row["status"] == "waiting-confirmation", checked["issues"]
    assert len(row["commands"]) == 1
    assert row["commands"][0]["execution_count"] == 1
    assert row["commands"][0]["infrastructure_failures"] == 0
    extracts = Counter(payload["work_unit_id"] for _, task, payload in calls if task == "extract")
    reviews = Counter(payload["work_unit_id"] for payload in _calls(calls, "review"))
    assert len(extracts) >= 2 and sorted(extracts.values()) == [1] * (len(extracts) - 2) + [2, 3]
    assert len(reviews) >= 2 and sorted(reviews.values()) == [1] * (len(reviews) - 2) + [2, 2]
    failed = [request for request in row["usage"] if request["status"] == "failed"]
    assert Counter(request["task"] for request in failed) == {"extract": 3, "verify": 2}
    assert all(request["sent_at"] for request in failed)
    assert sum(request["error"]["type"] == "invalid_json_response" for request in failed) == 4
    assert len([entry for entry in row["artifacts"] if entry["stage"] == "model-rejection"]) == 4
    assert not any(request["status"] in {"sent", "uncertain"} for request in row["usage"])
    assert all(request["metadata"]["external_request"] is False for request in row["usage"])


async def test_cross_document_candidate_has_full_sources_and_independent_joint_review(runtime):
    _, artifacts, _, _ = runtime
    run = await _submit(runtime, document_ids=("left-document", "right-document"), suffixes="a")
    execute, calls, _ = _executor(runtime, run, scope_repeats=1, candidates=True, reason_repeats=1)
    row = await execute()
    checked = output(artifacts, row, "verify")
    assert row["status"] == "waiting-confirmation", checked["issues"]
    associated, reviewed = _calls(calls, "associate"), _calls(calls, "joint")
    assert len(associated) == len(reviewed) == 1
    assert {v["document_id"] for v in associated[0]["source_views"]} == {"left-document", "right-document"}
    assert {i["item_id"] for i in associated[0]["items"]} == {i["item_id"] for i in reviewed[0]["items"]}
    assert reviewed[0]["relations"] and checked["coverage"]["joint_candidate_ids"]
    assert checked["counts"]["relations"] == 1


@pytest.mark.parametrize("accept", [False, True])
async def test_explicit_resolution_requires_independent_acceptance(runtime, accept):
    _, artifacts, _, _ = runtime
    run = await _submit(runtime, suffixes="ab")
    execute, calls, _ = _executor(runtime, run, scope_repeats=1, candidates=True,
                                  accept_decisions=accept, reason_repeats=0)
    first = await execute()
    _, extractions = _extractions(artifacts, first)
    left, right = extractions[0]["items"]
    extractions[0]["issues"].append({"code": "reference_target_unresolved", "blocking": True,
        "range_id": left["evidence"][0]["range_id"], "item_ids": [left["item_id"], right["item_id"]],
        "detail": "A synthetic cross-reference requires explicit jointly reviewed resolution."})
    await _correct(runtime, first, extractions)
    row = await execute()
    checked = output(artifacts, row, "verify")
    assert row["status"] == "waiting-confirmation", checked["issues"]
    assert bool(checked["accepted_decision_ids"]) == accept
    assert bool(checked["resolved_issue_ids"]) == accept
    target = next(p["issue_id"] for p in checked["all_issues"] if p["code"] == "reference_target_unresolved")
    assert (target in checked["human_issue_ids"]) is not accept
    assert _calls(calls, "joint", 2)[-1]["decisions"]


async def test_original_condition_evidence_remains_available_to_independent_review(runtime):
    _, artifacts, _, _ = runtime
    run = await _submit(runtime, ballast=True, ballast_repeats=30, condition=True, suffixes="abcd")
    execute, calls, budget = _executor(runtime, run, scope_repeats=1, include_conditions=True)
    row = await execute()
    assert row["status"] == "waiting-confirmation", output(artifacts, row, "verify")["issues"]
    assert len(output(artifacts, row, "pack-plan")["packs"]) > 1
    _assert_extract_budget(artifacts, row, calls, budget)
    _, extractions = _extractions(artifacts, row)
    items = [i for extraction in extractions for i in extraction["items"]]
    assert any(i.get("conditions") for i in items)
    for payload in _calls(calls, "review"):
        supplied = {s["range_id"]: s for v in payload["source_views"] for k in ("body", "supporting") for s in v[k]}
        aliases = {a["range_id"] for v in payload["source_views"] for a in v.get("source_aliases", [])}
        for item in payload["items"]:
            for citation in item["evidence"]:
                assert citation["range_id"] in supplied or citation["range_id"] in aliases
                if citation["range_id"] in supplied:
                    assert citation["quote"] in supplied[citation["range_id"]]["text"]
