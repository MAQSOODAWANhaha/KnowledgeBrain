"""Real parsers + real PostgreSQL + deterministic model, separate from LLM acceptance."""
import asyncio
import os
from pathlib import Path
import sys
from uuid import uuid4

from docx import Document
from openpyxl import Workbook
from openpyxl.drawing.image import Image as SheetImage
from PIL import Image, ImageDraw
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
import pytest

from docreader.tests.database import initialize_schema
import psycopg
from psycopg import sql
from psycopg.conninfo import make_conninfo

from docreader.integrations.artifacts import ArtifactStore
from docreader.integrations.store import InvalidConfirmation, Store
from docreader.jobs.settings import Settings
from docreader.llm import UncertainModelResponse
from docreader.llm.request import ModelWorkUnitExhausted
from docreader.tender.evidence import register_manifest
from docreader.workflow.nodes import TenderNodes
from docreader.workflow.persistence import initialize_checkpoints
from docreader.workflow.runner import execute_command

sys.path.insert(0, str(Path(__file__).parent))
from fixture_model import fixture_request_digest
from fixture_analysis_model import AnalysisFixtureModel

pytestmark = pytest.mark.integration


def mixed_manifest(root):
    document = Document()
    document.add_heading("Delivery obligations", level=1)
    document.add_paragraph("The supplier must provide a signed delivery commitment.")
    document.save(root / "main.docx")
    image = Image.new("RGB", (240, 80), "white")
    ImageDraw.Draw(image).text((10, 20), "Maintenance schedule", fill="black")
    image.save(root / "schedule.png")
    book = Workbook()
    first = book.active
    first.title = "Delivery"
    first.append(["Requirement", "Response"])
    first.append(["The supplier must name a delivery coordinator.", None])
    first.append([None, None])
    first.append(["Descriptive note", None])
    first.add_image(SheetImage(str(root / "schedule.png")), "D2")
    second = book.create_sheet("Service")
    second.append(["Requirement", "Response"])
    second.append(["The supplier must provide a service contact.", None])
    book.save(root / "annex.xlsx")
    writer = PdfWriter()
    page = writer.add_blank_page(width=600, height=800)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                             NameObject("/Subtype"): NameObject("/Type1"),
                             NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 11 Tf 40 740 Td (The supplier must provide a written quality assurance statement.) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    writer.write(root / "quality.pdf")
    return {"documents": [{"document_id": document_id, "role": role, "document_version": "one", "path": str(root / name)}
                          for document_id, role, name in [("main", "main", "main.docx"),
                                                         ("annex", "attachment", "annex.xlsx"),
                                                         ("quality", "attachment", "quality.pdf"),
                                                         ("image", "attachment", "schedule.png")]],
            "expected_document_ids": ["main", "annex", "quality", "image"]}


@pytest.fixture
async def runtime(tmp_path):
    dsn = os.environ.get("TENDER_TEST_DATABASE_URL")
    if not dsn:
        pytest.skip("set TENDER_TEST_DATABASE_URL for real PostgreSQL acceptance")
    schema = "phase1_test_" + uuid4().hex
    async with await psycopg.AsyncConnection.connect(dsn, autocommit=True) as conn:
        await conn.execute(sql.SQL("CREATE SCHEMA {}").format(sql.Identifier(schema)))
    isolated_dsn = make_conninfo(dsn, options=f"-c search_path={schema}")
    settings = Settings(isolated_dsn, os.environ.get("TENDER_TEST_REDIS_URL", "redis://127.0.0.1:16379/1"), tmp_path / "artifacts")
    store, artifacts = Store(isolated_dsn), ArtifactStore(settings.artifact_root)
    await initialize_schema(store.dsn)
    await store.setup()
    await initialize_checkpoints(isolated_dsn)
    registered = await asyncio.to_thread(register_manifest, mixed_manifest(tmp_path), artifacts)
    run = await store.submit(artifacts.put_json(registered), uuid4().hex)
    try:
        yield store, artifacts, settings, run
    finally:
        async with await psycopg.AsyncConnection.connect(dsn, autocommit=True) as conn:
            await conn.execute(sql.SQL("DROP SCHEMA {} CASCADE").format(sql.Identifier(schema)))


async def execute(store, settings, run, **model_options):
    cmd = (await store.get_run(run["run_id"]))["commands"][-1]
    def factory(s, a, lease):
        return TenderNodes(s, a, lease, AnalysisFixtureModel(s, lease, artifacts=a, **model_options))
    await execute_command(run["run_id"], cmd["command_id"], cmd["source_version"], settings, factory)
    return await store.get_run(run["run_id"])


def output(artifacts, run, stage):
    return artifacts.read_json(next(item["artifact_ref"] for item in run["artifacts"]
                                    if item["stage"] == stage and not item["scope"]))


async def test_mixed_formats_confirm_to_analysis_ready(runtime):
    store, artifacts, settings, run = runtime
    row = await execute(store, settings, run)
    checked = output(artifacts, row, "verify")
    assert row["status"] == "waiting-confirmation", checked["issues"]
    frozen = output(artifacts, row, "freeze")
    assert frozen["stats"]["document_count"] == 4
    assert frozen["stats"]["image_count"] == 2  # Excel embedded + standalone; no text PDF render.
    assert output(artifacts, row, "pack-plan")["range_count"] == checked["counts"]["ranges"]
    assert checked["counts"]["items"] >= 5
    assert all(row["usage"][i]["metadata"]["test_double"] for i in range(len(row["usage"])))
    assert all(request["usage"] is None for request in row["usage"])  # Never fabricate paid usage.
    await store.confirm(run["run_id"], 1, row["checklist_digest"], "first-confirm")
    completed = await execute(store, settings, run)
    assert completed["status"] == "analysis-ready"
    extracted = output(artifacts, completed, "extract")
    items = [item for entry in extracted["packs"] for item in artifacts.read_json(entry["artifact_ref"])["items"]]
    assert len(items) == checked["counts"]["items"]
    assert all(item["evidence"] for item in items)
    assert "items" not in checked
    assert not any(item["stage"] == "plan" for item in completed["artifacts"])
    assert not any(request["task"] == "plan" for request in completed["usage"])


async def test_partial_pack_infrastructure_failure_reuses_success_and_frozen_sources(runtime):
    store, artifacts, settings, run = runtime
    with pytest.raises(ModelWorkUnitExhausted):
        await execute(store, settings, run, fail_document="annex")
    failed = await store.get_run(run["run_id"])
    command = failed["commands"][-1]
    assert failed["status"] == "failed"
    assert command["status"] == "completed" and command["receipt"]["status"] == "failed"
    assert command["infrastructure_failures"] == 0
    assert len([request for request in failed["usage"] if request["status"] == "failed"]) == 4
    saved = {(a["stage"], a["scope"]): a for a in failed["artifacts"]}
    assert any(stage == "extract-pack" for stage, _ in saved)
    assert ("freeze", "") in saved
    retried = await store.retry(run["run_id"], 1, 1, command["command_id"], "retry-failed-pack",
                                "Retry the failed pack after infrastructure recovery")
    assert retried["command_id"] != command["command_id"]
    recovered = await execute(store, settings, run)
    assert recovered["status"] == "waiting-confirmation"
    for item in recovered["artifacts"]:
        key = (item["stage"], item["scope"])
        if key in saved:
            assert item["epoch"] == saved[key]["epoch"]
            assert item["artifact_ref"] == saved[key]["artifact_ref"]
    successful_initial = [r for r in failed["usage"] if r["status"] == "succeeded"]
    assert len([r for r in recovered["usage"] if r["task"] == "ocr"]) == 1  # Same PNG at two distinct source positions.
    assert successful_initial


async def test_uncertain_and_refused_siblings_each_stop_after_four_attempts(runtime):
    store, _, settings, run = runtime

    class MixedFailureModel(AnalysisFixtureModel):
        async def call(self, task, payload, images=()):
            if task == "extract" and any(view["document_id"] == "annex" for view in payload.get("source_views", [])):
                request_id = uuid4().hex
                await self.store.begin_request(self.lease, request_id, "deterministic-contract-fixture", task, {
                    "request_digest": fixture_request_digest(task, payload, images),
                    "logical_request_digest": fixture_request_digest(task,
                        {key: value for key, value in payload.items() if key != "contract_feedback"}, images),
                    "test_double": True, "external_request": False,
                })
                await self.store.finish_request(request_id, "sent")
                await self.store.finish_request(request_id, "uncertain", error={"type": "injected_partial_stream_loss"},
                    response_ref=self.artifacts.put_json({"partial_content": "{", "done": False}))
                raise UncertainModelResponse("injected partial stream loss")
            return await super().call(task, payload, images)

    cmd = (await store.get_run(run["run_id"]))["commands"][-1]
    def factory(s, a, lease):
        return TenderNodes(s, a, lease, MixedFailureModel(s, lease, artifacts=a, fail_document="main"))
    with pytest.raises(ModelWorkUnitExhausted):
        await execute_command(run["run_id"], cmd["command_id"], 1, settings, factory)
    row = await store.get_run(run["run_id"])
    assert row["status"] == "failed"
    assert row["commands"][-1]["status"] == "completed"
    assert row["commands"][-1]["receipt"]["status"] == "failed"
    assert row["commands"][-1]["infrastructure_failures"] == 0
    uncertain = [request for request in row["usage"] if request["status"] == "uncertain"]
    assert len(uncertain) == 4 and all(request["response_ref"] for request in uncertain)
    assert len([request for request in row["usage"] if request["status"] == "failed"]) == 4
    assert any(item["stage"] == "extract-pack" for item in row["artifacts"])
    await execute_command(run["run_id"], cmd["command_id"], 1, settings, factory)
    assert (await store.get_run(run["run_id"]))["usage"] == row["usage"]


async def test_protocol_failure_is_failed_and_successful_siblings_survive_source_revision(runtime):
    store, artifacts, settings, run = runtime
    with pytest.raises(ModelWorkUnitExhausted):
        await execute(store, settings, run, malformed_document="annex")
    row = await store.get_run(run["run_id"])
    assert row["status"] == "failed"
    assert row["commands"][-1]["status"] == "completed"
    assert row["commands"][-1]["receipt"]["status"] == "failed"
    assert row["commands"][-1]["infrastructure_failures"] == 0
    await execute(store, settings, run, malformed_document="annex")
    assert (await store.get_run(run["run_id"]))["usage"] == row["usage"]
    assert len([a for a in row["artifacts"] if a["stage"] == "model-rejection"]) == 4
    rejected_digests = {request["metadata"]["request_digest"] for request in row["usage"]
                        if request["task"] == "extract" and artifacts.read_json(request["response_ref"]) == {
                            "not_an_extraction": True}}
    assert len(rejected_digests) == 4
    assert any(item["stage"] == "extract-pack" for item in row["artifacts"])
    assert not any(item["stage"] == "verify" for item in row["artifacts"])
    with pytest.raises(InvalidConfirmation):
        await store.confirm(run["run_id"], 1, "unverified", "cannot-force")
    manifest = artifacts.read_json(row["manifest_ref"])
    manifest["documents"][0]["document_version"] = "two"
    revised = await store.revise(run["run_id"], artifacts.put_json(manifest), "reviewed-source-version")
    assert revised["source_version"] == 2
    recovered = await execute(store, settings, run)
    assert recovered["status"] == "waiting-confirmation"
    with pytest.raises(InvalidConfirmation):
        await store.confirm(run["run_id"], 1, "unverified", "expired-version")


async def test_same_source_correction_revalidates_and_invalidates_previous_confirmation(runtime):
    store, artifacts, settings, run = runtime
    first = await execute(store, settings, run)
    await store.confirm(run["run_id"], 1, first["checklist_digest"], "confirm-r1", domain_revision=1)
    ready = await execute(store, settings, run)
    assert ready["status"] == "analysis-ready"
    original_source = next(a["artifact_ref"] for a in ready["artifacts"] if a["stage"] == "freeze")
    original_extract = next(a["artifact_ref"] for a in ready["artifacts"] if a["stage"] == "extract")
    corrected = artifacts.read_json(original_extract)
    first_pack = corrected["packs"][0]
    pack = artifacts.read_json(first_pack["artifact_ref"])
    pack["items"][0]["domains"] = ["delivery"]  # manual classification, deterministic fixture only
    pack["source_ranges"] = {"forged-source": {"text": "Caller supplied metadata is untrusted."}}
    first_pack["artifact_ref"] = artifacts.put_json(pack)
    corrected_ref = artifacts.put_json(corrected)
    revision = await store.correct(run["run_id"], 1, 1, original_extract["sha256"], corrected_ref,
                                   "correct-r2", "Correct the veto classification for independent re-review")
    replay = await store.correct(run["run_id"], 1, 1, original_extract["sha256"], corrected_ref,
                                 "correct-r2", "Correct the veto classification for independent re-review")
    assert replay["submitted_domain_revision"] == revision["submitted_domain_revision"] == 2
    row = await store.get_run(run["run_id"])
    assert row["source_version"] == 1
    assert row["domain_revision"] == 2
    assert row["manifest_ref"] == first["manifest_ref"]
    assert row["checklist_digest"] is None
    assert not any(a["stage"] in {"verify", "confirm", "plan"} for a in row["artifacts"])
    with pytest.raises(InvalidConfirmation):
        await store.confirm(run["run_id"], 1, first["checklist_digest"], "confirm-r1", domain_revision=1)
    second = await execute(store, settings, run)
    assert second["status"] == "waiting-confirmation"
    assert second["checklist_digest"] != first["checklist_digest"]
    assert next(a["artifact_ref"] for a in second["artifacts"] if a["stage"] == "freeze") == original_source
    reviewed = output(artifacts, second, "verify")
    reviewed_index = artifacts.read_json(reviewed["items_ref"])
    assert artifacts.read_json(reviewed_index["packs"][0]["artifact_ref"])["items"][0]["domains"] == ["delivery"]
    assert "forged-source" not in artifacts.read_json(reviewed_index["packs"][0]["artifact_ref"])["source_ranges"]
    assert all(r["task"] != "extract" for r in second["usage"] if r["domain_revision"] == 2)
    assert any(r["task"] == "verify" for r in second["usage"] if r["domain_revision"] == 2)
    with pytest.raises(InvalidConfirmation):
        await store.confirm(run["run_id"], 1, first["checklist_digest"], "wrong-digest-r2", domain_revision=2)
    await store.confirm(run["run_id"], 1, second["checklist_digest"], "confirm-r2", domain_revision=2)
    final = await execute(store, settings, run)
    assert final["status"] == "analysis-ready"
    assert not any(item["stage"] == "plan" for item in final["artifacts"])


async def test_completed_checkpoint_cannot_bypass_execution_spec_on_recovery(runtime, monkeypatch):
    from docreader.integrations.store import Conflict

    store, artifacts, settings, run = runtime
    async def execute_with_spec(label):
        cmd = (await store.get_run(run["run_id"]))["commands"][-1]
        class ConfiguredFixture(AnalysisFixtureModel):
            def execution_spec(self):
                return {"test_configuration": label}
        def factory(s, a, lease):
            return TenderNodes(s, a, lease, ConfiguredFixture(s, lease, artifacts=a))
        await execute_command(run["run_id"], cmd["command_id"], cmd["source_version"], settings, factory)

    await execute_with_spec("original")
    paused = await store.get_run(run["run_id"])
    await store.confirm(run["run_id"], 1, paused["checklist_digest"], "confirm-before-receipt-fault")
    original_release = Store.release
    failed_once = False
    async def fail_receipt_after_checkpoint(self, lease, status, **kwargs):
        nonlocal failed_once
        if status == "analysis-ready" and not failed_once:
            failed_once = True
            raise psycopg.OperationalError("injected receipt failure after complete checkpoint")
        return await original_release(self, lease, status, **kwargs)
    monkeypatch.setattr(Store, "release", fail_receipt_after_checkpoint)
    with pytest.raises(psycopg.OperationalError):
        await execute_with_spec("original")
    pending = await store.get_run(run["run_id"])
    assert pending["status"] == "pending"
    assert any(item["stage"] == "confirm" for item in pending["artifacts"])
    for attempt in range(1, 4):
        with pytest.raises(Conflict):
            await execute_with_spec("changed")
        current = await store.get_run(run["run_id"])
        assert current["status"] == ("pending" if attempt < 3 else "failed")
        assert current["commands"][-1]["infrastructure_failures"] == attempt + 1
    failed = await store.get_run(run["run_id"])
    assert failed["status"] == "failed"
    assert len(failed["usage"]) == len(pending["usage"])


async def test_partial_recovery_refuses_changed_native_parser_configuration(runtime, monkeypatch):
    from docreader.integrations.store import Conflict
    from docreader.parser import pdf_parser

    store, _, settings, run = runtime
    with pytest.raises(ModelWorkUnitExhausted):
        await execute(store, settings, run, fail_document="annex")
    before = await store.get_run(run["run_id"])
    command = before["commands"][-1]
    assert before["status"] == "failed" and command["infrastructure_failures"] == 0
    retried = await store.retry(run["run_id"], 1, 1, command["command_id"], "retry-with-parser-guard",
                                "Retry the failed pack while preserving the frozen parser configuration")
    assert retried["command_id"] != command["command_id"]
    monkeypatch.setattr(pdf_parser, "SCAN_MIN_CHARS_PER_PAGE", pdf_parser.SCAN_MIN_CHARS_PER_PAGE + 1)
    for attempt in range(1, 5):
        with pytest.raises(Conflict):
            await execute(store, settings, run)
        current = await store.get_run(run["run_id"])
        assert current["status"] == ("pending" if attempt < 4 else "failed")
        assert current["commands"][-1]["infrastructure_failures"] == attempt
    after = await store.get_run(run["run_id"])
    assert after["status"] == "failed"
    assert len(after["usage"]) == len(before["usage"])
    assert after["artifacts"] == before["artifacts"]



async def test_repeated_image_pixels_share_ocr_without_losing_occurrence_coordinates(runtime):
    store, artifacts, settings, _ = runtime
    root = settings.artifact_root.parent
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["The supplier must provide image-backed supporting material."])
    for anchor in ("B2", "E5"):
        sheet.add_image(SheetImage(str(root / "schedule.png")), anchor)
    workbook.save(root / "repeated-images.xlsx")
    registered = register_manifest({"documents": [{"document_id": "repeated", "document_version": "one", "role": "main",
                                                  "path": str(root / "repeated-images.xlsx")}]}, artifacts)
    run = await store.submit(artifacts.put_json(registered), "same-pixels-distinct-occurrences")
    row = await execute(store, settings, run)
    assert row["status"] == "waiting-confirmation"
    frozen = output(artifacts, row, "freeze")
    raw = artifacts.read_json(frozen["documents"][0]["raw_ref"])
    assert frozen["stats"]["image_count"] == len(raw["ocr"]) == 2
    assert len(raw["image_anchors"]) == 2
    assert {item["from"]["column"] for item in raw["image_anchors"].values()} == {2, 5}
    assert len([request for request in row["usage"] if request["task"] == "ocr"]) == 1
    assert len([receipt for receipt in row["artifacts"] if receipt["stage"] == "ocr"]) == 1
    assert all(result["status"] == "complete" for result in raw["ocr"].values())


async def test_evidence_backed_scope_review_creates_revision_replans_and_revalidates(runtime):
    from docreader.integrations.cli import _scope_review_ref
    from docreader.integrations.store import Conflict
    from docreader.tender.evidence import EvidenceRepository

    store, artifacts, settings, _ = runtime
    path = settings.artifact_root.parent / "scope-form.docx"
    document = Document()
    document.add_heading("Response field instructions", level=1)
    condition = "This form applies only to the local response field."
    document.add_paragraph(condition)
    document.add_paragraph("The supplier must provide a signed delivery commitment.")
    document.save(path)
    registered = register_manifest({"documents": [{"document_id": "scope-form", "document_version": "one",
                                                  "role": "main", "path": str(path)}]}, artifacts)
    run = await store.submit(artifacts.put_json(registered), "scope-form")
    flagged = await execute(store, settings, run, scope_ambiguity=True)
    # A model scope doubt waits for a human: confirm a verdict or use the explicit scope review below.
    assert flagged["status"] == "waiting-confirmation"
    first_plan = output(artifacts, flagged, "pack-plan")
    frozen = next(item["artifact_ref"] for item in flagged["artifacts"] if item["stage"] == "freeze")
    assert not any(item["code"] == "condition_scope_unresolved" for item in first_plan["issues"])
    problem = next(item for item in output(artifacts, flagged, "verify")["issues"]
                   if item["code"] == "semantic_scope_unresolved")
    assert problem["issue_id"] in output(artifacts, flagged, "verify")["human_issue_ids"]
    repo = EvidenceRepository(artifacts)
    source = repo.source_index(first_plan["documents"][0]["source_index_ref"])
    source_range = next(entry for entry in source.ranges if condition in entry["text"])
    submitted = {"scope_decisions": [{
        "range_id": problem["range_id"], "decision": "local_only", "target_range_ids": [],
        "rationale": "Human review confirms this condition describes only the named local response field; "
                     "it does not limit the following supplier commitment or the entire section.",
        "evidence": [{"range_id": source_range["range_id"], "quote": condition}],
    }]}
    review_ref = _scope_review_ref(artifacts, flagged, 1, 1, first_plan["plan_digest"], submitted)
    second = await store.review_scope(run["run_id"], 1, 1, first_plan["plan_digest"], review_ref,
                                     "scope-reviewed-r2", "Review the form field boundary")
    replay = await store.review_scope(run["run_id"], 1, 1, first_plan["plan_digest"], review_ref,
                                     "scope-reviewed-r2", "Review the form field boundary")
    assert second["submitted_domain_revision"] == replay["submitted_domain_revision"] == 2
    with pytest.raises(Conflict):
        await store.review_scope(run["run_id"], 1, 1, first_plan["plan_digest"], review_ref,
                                 "expired-scope-baseline", "Old plan cannot revise a newer domain revision")
    reviewed = await execute(store, settings, run, scope_ambiguity=True)  # The same fixture re-reads the human proposal.
    assert reviewed["source_version"] == 1 and reviewed["domain_revision"] == 2
    assert reviewed["status"] == "waiting-confirmation", output(artifacts, reviewed, "verify")["issues"]
    assert next(item["artifact_ref"] for item in reviewed["artifacts"] if item["stage"] == "freeze") == frozen
    updated_plan = output(artifacts, reviewed, "pack-plan")
    assert updated_plan["plan_digest"] != first_plan["plan_digest"]
    assert updated_plan["documents"][0]["raw_ref"] == first_plan["documents"][0]["raw_ref"]
    assert not any(item["code"] == "condition_scope_unresolved" for item in updated_plan["issues"])
    assert any(item["task"] == "extract" and item["domain_revision"] == 2 for item in reviewed["usage"])
    with pytest.raises(InvalidConfirmation):
        await store.confirm(run["run_id"], 1, flagged["checklist_digest"], "old-scope-confirm", domain_revision=1)
    await store.confirm(run["run_id"], 1, reviewed["checklist_digest"], "scope-confirm-r2", domain_revision=2)
    ready = await execute(store, settings, run)
    assert ready["status"] == "analysis-ready"
