"""Document execution persistence with isolated SQL and a small synthetic pipeline.

These cases exercise the real model client and durable execution, not document
content acceptance or the external knowledge provider.
"""

import asyncio
from copy import deepcopy
from io import BytesIO
import hashlib
from importlib.resources import files
import json
from types import SimpleNamespace
from uuid import uuid4
from zipfile import ZipFile

from docx import Document
import httpx
from PIL import Image
import psycopg
import pytest

from docreader.integrations.authoring_store import AuthoringStore
from docreader.integrations.authoring import read_authoring_input
from docreader.integrations.store import Cancelled, Conflict
from docreader.llm.budget import RequestBudget
from docreader.llm.client import OpenAIModel
from docreader.llm.config import ModelConfig, ModelEndpoint
from docreader.llm.request import ModelBudgetError, ModelConfigurationError, ModelContractError, build_request
from docreader.tender import document_execution
from docreader.tender.evidence import register_manifest
from docreader.tender.schema import AuthoringUsage
from docreader.tests.submission_responses import assemble_independent_candidates

from test_analysis import AnalysisModel, analysis_runtime as analysis_runtime, execute_analysis
from test_authoring_execution import config, response
from test_authoring_outline import budget


pytestmark = pytest.mark.integration


@pytest.fixture
def pipeline(monkeypatch):
    state = SimpleNamespace(binding="first", observed=[], snapshots=[], after=None, write_text="same",
                            outline_payload={"operation": "outline"})

    async def generate(artifacts, snapshot, request_budget, execute, *, materials=(), knowledge=None, concurrency=2, on_outline=None):
        state.snapshots.append(deepcopy(snapshot))
        if knowledge is not None:
            result = await knowledge.search([{"target_id": "target", "query": "所需能力"}], artifacts)
            assert result == {"materials": [], "issues": []}
        outputs = []
        for task in ("outline", "authoring_write"):
            payload = deepcopy(state.outline_payload) if task == "outline" else {"operation": task}
            if task == "authoring_write":
                payload["text"] = state.write_text
            prompt = "outline_assemble" if payload.get("operation") == "assemble" else task
            request = build_request(request_budget.model, request_budget.prompts[prompt], payload)

            def validate(value):
                if value.get("value") != "accepted":
                    raise ModelContractError([{"path": ["value"], "code": "bad_value", "detail": "Return accepted."}])
                result = {"binding": state.binding, **value}
                state.observed.append(result)
                return result

            outputs.append(await execute(task, request, validate))
        if state.after is not None:
            await state.after()
        ref = artifacts.put_json({"synthetic_domain_results": outputs})
        return artifacts.put_json({"schema_version": 2, "authoring_input": snapshot,
            "outline_ref": ref, "basis_ref": ref, "unit_refs": [ref], "review_ref": ref,
            "docx_ref": ref, "package_ref": ref, "gaps": [], "usage": AuthoringUsage().model_dump()})

    monkeypatch.setattr(document_execution, "generate_document", generate)
    return state


def identity(runtime):
    row = runtime[3]
    return {key: row[key] for key in ("run_id", "source_version", "domain_revision")}


async def execute(runtime, execution_id, handler, **kwargs):
    return await document_execution.execute_document(runtime[0], runtime[1], config(), budget(),
        execution_id=execution_id, **identity(runtime), transport=httpx.MockTransport(handler), **kwargs)


async def test_unconfirmed_snapshot_runs_multiple_tasks_and_reuses_completed_execution(analysis_runtime, pipeline):
    store, artifacts, settings, run, _ = analysis_runtime
    before = await store.get_run(run["run_id"])
    execution_id, sent = uuid4().hex, []

    def handler(request):
        sent.append(json.loads(request.content))
        return response({"value": "accepted"})

    reference = await execute(analysis_runtime, execution_id, handler)
    value = artifacts.read_json(reference)
    row = await AuthoringStore(settings.dsn).get_execution(execution_id)
    assert row["status"] == "succeeded" and len(row["requests"]) == len(sent) == 2
    assert {request["task"] for request in row["requests"]} == {"outline", "authoring_write"}
    assert value["authoring_input"]["confirmation_ref"] is None
    assert value["usage"]["sent_requests"] == 2 and value["usage"]["retry_requests"] == 0
    assert value["usage"]["prompt_tokens"] == 226 and value["usage"]["completion_tokens"] == 114
    assert value["usage"]["total_tokens"] == 340 and value["usage"]["unknown_usage_requests"] == 0
    assert value["usage"]["estimated_input_tokens"] > 0
    assert len([r for r in row["receipts"] if r["stage"] == "document-model-response"]) == 2
    assert "fixture-key" not in json.dumps(row["execution_spec"])
    assert row["execution_spec"]["prompts"]["outline_assemble"] == hashlib.sha256(
        files("docreader.tender.prompts").joinpath("outline_assemble.md").read_bytes()).hexdigest()
    assert sent[0]["messages"][0]["content"] == files("docreader.tender.prompts").joinpath("outline.md").read_text()
    assert await execute(analysis_runtime, execution_id, handler) == reference
    assert len(sent) == 2 and len(pipeline.snapshots) == 1
    assert await store.get_run(run["run_id"]) == before


async def test_explicit_predecessor_rebinds_cached_json_without_fabricating_paid_calls(analysis_runtime, pipeline):
    sent = []

    def handler(request):
        sent.append(request)
        return response({"value": "accepted"})

    first = uuid4().hex
    await execute(analysis_runtime, first, handler)
    pipeline.binding = "current-host-identities"
    second = uuid4().hex
    reference = await execute(analysis_runtime, second, handler, previous_execution_id=first)
    row = await AuthoringStore(analysis_runtime[2].dsn).get_execution(second)
    assert len(sent) == 2 and row["requests"] == []
    assert analysis_runtime[1].read_json(reference)["usage"] == AuthoringUsage().model_dump()
    assert all(value["binding"] == "current-host-identities" for value in pipeline.observed[-2:])
    inherited = [r for r in row["receipts"] if r["stage"] == "document-model-response"]
    assert len(inherited) == 2
    assert all(r["metadata"]["reused_from_execution_id"] == first for r in inherited)


async def test_explicit_assembly_tokenizer_reaches_http_and_correction_without_enlarging_source_batches(
    analysis_runtime, pipeline,
):
    model = ModelConfig(ModelEndpoint("http://fixture.invalid/v1", "fixture-key", "synthetic-outline", 131072), None)
    request_budget = RequestBudget(context_tokens=131072, model=model.chat.model, prompts={}, output_reserve=32768,
                                   tokenizer="o200k_base", tokenized_outline_assembly=True)
    pipeline.outline_payload = {"operation": "assemble", "text": "服务能力" * 18000}
    sent = []

    def handler(request):
        sent.append(json.loads(request.content))
        return response({"value": "invalid" if len(sent) == 1 else "accepted"})

    async def run(execution_id):
        return await document_execution.execute_document(analysis_runtime[0], analysis_runtime[1], model,
            request_budget, execution_id=execution_id, **identity(analysis_runtime),
            transport=httpx.MockTransport(handler))

    execution_id = uuid4().hex
    await run(execution_id)
    row = await AuthoringStore(analysis_runtime[2].dsn).get_execution(execution_id)
    assert row["status"] == "succeeded" and len(sent) == 3
    assembly_requests = [request for request in row["requests"] if request["task"] == "outline"]
    assert len(assembly_requests) == 2
    assert all(request["metadata"]["method"] == RequestBudget.TOKENIZED_METHOD for request in assembly_requests)
    assert all(request["metadata"]["input_tokens"] < 94208 for request in assembly_requests)
    assert "contract_feedback" in json.loads(sent[1]["messages"][1]["content"][0]["text"])
    frozen_budget = row["execution_spec"]["budget"]
    assert frozen_budget["tokenized"]["tokenizer"] == "o200k_base"
    assert frozen_budget["tokenized"]["operations"] == {"outline": ["assemble"]}

    # Same source content must still use the original conservative admission.
    pipeline.outline_payload["operation"] = "outline"
    with pytest.raises(ModelBudgetError):
        await run(uuid4().hex)
    assert len(sent) == 3


@pytest.mark.parametrize("assembly", [False, True])
async def test_explicit_authoring_tokenizer_survives_execution_budget_reconstruction(
    analysis_runtime, pipeline, assembly,
):
    model = ModelConfig(ModelEndpoint("http://fixture.invalid/v1", "fixture-key", "synthetic-outline", 131072), None)
    prompts = {task: files("docreader.tender.prompts").joinpath(task + ".md").read_text("utf-8")
               for task in document_execution._PROMPTS}
    request_budget = RequestBudget(context_tokens=131072, model=model.chat.model, prompts=prompts,
        output_reserve=32768, tokenizer="o200k_base", tokenized_outline_assembly=assembly, tokenized_authoring=True)
    pipeline.outline_payload = {"operation": "source", "text": "服务能力" * 18000}
    pipeline.write_text = "技术响应" * 18000
    sent = []

    def handler(request):
        sent.append(json.loads(request.content))
        return response({"value": "accepted"})

    execution_id = uuid4().hex
    await document_execution.execute_document(analysis_runtime[0], analysis_runtime[1], model,
        request_budget, execution_id=execution_id, **identity(analysis_runtime),
        transport=httpx.MockTransport(handler))
    row = await AuthoringStore(analysis_runtime[2].dsn).get_execution(execution_id)
    assert row["status"] == "succeeded" and len(sent) == len(row["requests"]) == 2
    assert row["execution_spec"]["budget"] == request_budget.spec()
    for entry in row["requests"]:
        metadata = entry["metadata"]
        assert metadata["method"] == RequestBudget.TOKENIZED_METHOD
        assert metadata["framing_tokens"] == 512 and metadata["output_reserve"] == 32768
        assert metadata["input_tokens"] < metadata["effective_input_budget_tokens"] == 94208


async def test_revision_reuses_only_requests_with_identical_inputs(analysis_runtime, pipeline):
    sent = []

    def handler(request):
        sent.append(request)
        return response({"value": "accepted"})

    first = uuid4().hex
    await execute(analysis_runtime, first, handler)
    pipeline.write_text = "changed writing facts"
    second = uuid4().hex
    reference = await execute(analysis_runtime, second, handler, previous_execution_id=first)
    row = await AuthoringStore(analysis_runtime[2].dsn).get_execution(second)
    assert len(sent) == 3 and len(row["requests"]) == 1
    assert row["requests"][0]["task"] == "authoring_write"
    assert analysis_runtime[1].read_json(reference)["usage"]["sent_requests"] == 1


async def test_rejected_response_is_corrected_then_reused_as_accepted_json(analysis_runtime, pipeline):
    sent = []

    def handler(request):
        payload = json.loads(json.loads(request.content)["messages"][1]["content"][0]["text"])
        sent.append(payload)
        return response({"value": "invalid" if len(sent) == 1 else "accepted"})

    first = uuid4().hex
    reference = await execute(analysis_runtime, first, handler)
    assert len(sent) == 3 and sent[1]["contract_feedback"]["errors"][0]["code"] == "bad_value"
    usage = analysis_runtime[1].read_json(reference)["usage"]
    assert usage["sent_requests"] == 3 and usage["retry_requests"] == 1
    await execute(analysis_runtime, uuid4().hex, handler, previous_execution_id=first)
    assert len(sent) == 3


async def test_publication_failure_resumes_fixed_snapshot_and_knowledge_without_resending(
    analysis_runtime, pipeline, monkeypatch,
):
    sent, queried = [], []

    class Knowledge:
        def spec(self):
            return {"provider": "fixture", "collection": "same"}

        async def search(self, queries, artifacts):
            queried.append(queries)
            return {"materials": [], "issues": []}

    def handler(request):
        sent.append(request)
        return response({"value": "accepted"})

    original = AuthoringStore.release
    failed = []

    async def release(self, lease, status, **kwargs):
        if status == "succeeded" and not failed:
            failed.append(True)
            raise psycopg.OperationalError("synthetic publication outage")
        return await original(self, lease, status, **kwargs)

    execution_id = uuid4().hex
    with monkeypatch.context() as patch:
        patch.setattr(AuthoringStore, "release", release)
        with pytest.raises(psycopg.OperationalError, match="synthetic publication outage"):
            await execute(analysis_runtime, execution_id, handler, knowledge=Knowledge())
    previous = await AuthoringStore(analysis_runtime[2].dsn).get_execution(execution_id)
    assert previous["status"] == "failed" and len(sent) == 2
    reference = await execute(analysis_runtime, execution_id, handler, knowledge=Knowledge())
    assert len(sent) == 2 and len(queried) == 1
    assert pipeline.snapshots[0] == pipeline.snapshots[1] == previous["authoring_input"]
    assert analysis_runtime[1].read_json(reference)["usage"]["sent_requests"] == 2


async def test_model_transport_failures_still_retry_and_count_unknown_usage(analysis_runtime, pipeline, monkeypatch):
    sent = []

    async def no_delay(seconds):
        return None

    monkeypatch.setattr("docreader.llm.recovery.asyncio", SimpleNamespace(sleep=no_delay, to_thread=asyncio.to_thread))

    def handler(request):
        sent.append(request)
        if len(sent) <= 3:
            return httpx.Response(500, json={"error": "synthetic upstream failure"})
        return response({"value": "accepted"})

    reference = await execute(analysis_runtime, uuid4().hex, handler)
    usage = analysis_runtime[1].read_json(reference)["usage"]
    assert usage["sent_requests"] == len(sent) == 5 and usage["retry_requests"] == 3
    assert usage["unknown_usage_requests"] == 3 and usage["total_tokens"] == 340


async def test_persistence_failure_is_not_swallowed_into_a_placeholder(analysis_runtime, pipeline, monkeypatch):
    sent = []
    original = AuthoringStore.commit_artifact

    async def commit(self, lease, stage, *args, **kwargs):
        if stage == "model-request":
            raise psycopg.OperationalError("synthetic request audit outage")
        return await original(self, lease, stage, *args, **kwargs)

    def handler(request):
        sent.append(request)
        return response({"value": "accepted"})

    monkeypatch.setattr(AuthoringStore, "commit_artifact", commit)
    execution_id = uuid4().hex
    with pytest.raises(psycopg.OperationalError, match="synthetic request audit outage"):
        await execute(analysis_runtime, execution_id, handler)
    row = await AuthoringStore(analysis_runtime[2].dsn).get_execution(execution_id)
    assert sent == [] and row["status"] == "failed" and row["result_ref"] is None


async def test_material_snapshot_is_part_of_execution_identity(analysis_runtime, pipeline):
    def handler(request):
        return response({"value": "accepted"})

    execution_id = uuid4().hex
    material = {"material_id": "provided", "title": "企业资料", "text": "已知事实"}
    await execute(analysis_runtime, execution_id, handler, materials=[material])
    with pytest.raises(Conflict):
        await execute(analysis_runtime, execution_id, handler, materials=[{**material, "text": "变更事实"}])


async def test_parent_cancellation_before_publication_preserves_receipts_without_success(analysis_runtime, pipeline):
    async def cancel_parent():
        await analysis_runtime[0].cancel(analysis_runtime[3]["run_id"])

    pipeline.after = cancel_parent
    execution_id = uuid4().hex
    with pytest.raises(Cancelled):
        await execute(analysis_runtime, execution_id, lambda request: response({"value": "accepted"}))
    row = await AuthoringStore(analysis_runtime[2].dsn).get_execution(execution_id)
    assert row["status"] == "failed" and row["result_ref"] is None
    assert len(row["requests"]) == 2


async def test_caller_cancellation_stops_inflight_call_and_keeps_uncertain_receipt(analysis_runtime, pipeline):
    started = asyncio.Event()

    async def handler(request):
        started.set()
        await asyncio.Event().wait()

    execution_id = uuid4().hex
    task = asyncio.create_task(execute(analysis_runtime, execution_id, handler,
                                       lease_seconds=1, heartbeat_seconds=0.05))
    await asyncio.wait_for(started.wait(), timeout=3)
    task.cancel()
    with pytest.raises(asyncio.CancelledError):
        await asyncio.wait_for(task, timeout=3)
    row = await AuthoringStore(analysis_runtime[2].dsn).get_execution(execution_id)
    assert row["status"] == "failed" and row["result_ref"] is None
    assert len(row["requests"]) == 1 and row["requests"][0]["status"] == "uncertain"


@pytest.mark.parametrize("options", [{"concurrency": 0}, {"heartbeat_seconds": 60}, {"lease_seconds": float("nan")}])
async def test_invalid_runtime_parameters_make_no_requests(analysis_runtime, pipeline, options):
    with pytest.raises(ModelConfigurationError):
        await execute(analysis_runtime, uuid4().hex, lambda request: pytest.fail("unexpected send"), **options)


async def test_actual_pipeline_exports_forms_and_evidence_despite_one_exhausted_task(
    analysis_runtime, tmp_path, monkeypatch,
):
    """Real parser -> unconfirmed analysis -> authoring SSE -> physical DOCX/ZIP."""
    store, artifacts, settings, _, _ = analysis_runtime
    source = tmp_path / "synthetic-tender.docx"
    tender = Document()
    tender.add_paragraph("负责人须持A证；持B证加2分；未附A证复印件则否决投标；须附A证复印件。")
    tender.add_paragraph("服务响应不超过2小时。")
    table = tender.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "人员安排表"
    table.cell(1, 0).text = "负责人"
    table.cell(2, 0).text = "证书编号"
    tender.save(source)
    manifest = register_manifest({"documents": [{"document_id": "main", "document_version": "one",
        "role": "main", "path": str(source)}]}, artifacts)
    run = await store.submit(artifacts.put_json(manifest), uuid4().hex)
    runtime = (store, artifacts, settings, run, [])
    analysis_call = AnalysisModel.call

    async def table_aware_fixture(self, task, payload, images=()):
        # The original fixture only handles prose. These synthetic table rows
        # contain field labels, with no additional rule, so their disposition
        # is non_item. Production source/grid data and validation stay intact.
        projected = deepcopy(payload)
        if task == "extract":
            for view in projected.get("source_views", []):
                for entry in view["body"]:
                    entry.setdefault("text", "")
        return await analysis_call(self, task, projected, images)

    monkeypatch.setattr(AnalysisModel, "call", table_aware_fixture)
    waiting = await execute_analysis(runtime)
    assert waiting["status"] == "waiting-confirmation"
    before = await store.get_run(run["run_id"])

    original = BytesIO()
    Image.new("RGB", (40, 30), "white").save(original, format="PNG")
    proof = {"material_id": "synthetic-evidence", "title": "测试证明原件", "text": "仅供合成测试的原件。",
             "original_ref": artifacts.put_bytes(original.getvalue(), "image/png"),
             "attachment": True, "usage": "proof"}
    methods = {"material_id": "synthetic-methods", "title": "测试实施和服务资料", "usage": "reference",
               "text": "按部署、联调和验收三个阶段实施；联调形成测试记录。设置项目负责人，姓名另行补充。"
                       "服务采用受理和分级处置流程。"}
    sent, service_calls, assembly_calls, structure_calls = [], [], [], []
    recovered = False

    async def no_delay(seconds):
        return None

    monkeypatch.setattr("docreader.llm.recovery.asyncio", SimpleNamespace(sleep=no_delay, to_thread=asyncio.to_thread))

    def handler(request):
        payload = json.loads(json.loads(request.content)["messages"][1]["content"][0]["text"])
        sent.append(payload)
        if payload.get("operation") == "assemble":
            assembly_calls.append(payload)
            return response(assemble_independent_candidates(payload))
        if payload.get("operation") == "structure":
            structure_calls.append(payload)
            forms = [target for chapter in payload["chapters"] for target in chapter["targets"]
                     if target["kind"] == "form"]
            assert len(forms) == 1 and forms[0]["template_support_ids"]
            return response({"findings": []})
        if "sources" in payload:
            rules = [entry["rule_id"] for entry in payload["rules"]]
            templates = [entry["range_id"] for entry in payload["sources"] if "人员安排表" in entry["text"]]
            assert templates, payload["sources"]

            parts = [part["part_id"] for part in payload["template_parts"]
                     if set(part["range_ids"]) & set(templates)]
            assert len(parts) == 1

            def chapter(title, kind, purpose, ranges=(), template_parts=()):
                return {"title": title, "children": [], "targets": [{"kind": kind, "purpose": purpose,
                        "rule_ids": rules, "range_ids": list(ranges),
                        "template_part_ids": list(template_parts)}]}

            return response({"schema_version": 2, "chapters": [
                chapter("技术方案", "narrative", "编制项目实施方法"),
                chapter("人员表", "form", "填写负责人并保留未知证号", templates, parts),
                chapter("证明附件", "attachment", "附提供的真实原件"),
                chapter("服务方案", "narrative", "编制项目服务方法")], "findings": []})
        if "supports" in payload and "title_path" not in payload and "units" not in payload:
            return response({"approach": "拟按项目要求组织实施与服务。", "shared_arrangements": []})
        if "title_path" in payload:
            title = payload["title_path"][-1]
            method_support = next(entry["support_id"] for entry in payload["supports"]
                                  if entry["origin"] == "knowledge" and entry["text"] == methods["text"])
            if title == "服务方案":
                service_calls.append(payload)
                if not recovered:
                    return httpx.Response(500, json={"error": "synthetic service writing outage"})
                return response({"blocks": [{"kind": "paragraph", "text": "拟建立服务受理和分级处置流程。",
                    "target_ids": ["t1"], "support_ids": [method_support]}]})
            if title == "人员表":
                template = next(entry for entry in payload["supports"] if entry.get("grid"))
                return response({"blocks": [], "cell_fills": [
                    {"source_id": template["support_id"], "row": 1, "column": 1,
                     "text": "项目负责人（姓名待补）", "target_ids": ["t1"],
                     "support_ids": [method_support]}]})
            if title == "证明附件":
                proof_support = next(entry["support_id"] for entry in payload["supports"]
                                     if entry["origin"] == "knowledge" and entry.get("attachment"))
                return response({"blocks": [], "attachment_ids": [proof_support]})
            return response({"blocks": [
                {"kind": "paragraph", "text": "拟采用部署、联调和验收三个阶段实施本项目。",
                 "target_ids": ["t1"], "support_ids": [method_support]},
                {"kind": "table", "headers": ["阶段", "产物"], "rows": [["联调", "测试记录"]],
                 "target_ids": ["t1"], "support_ids": [method_support]}]})
        if "units" in payload:
            return response({"findings": []})
        pytest.fail("unexpected authoring request shape")

    execution_id = uuid4().hex
    reference = await execute(runtime, execution_id, handler, materials=[proof, methods])
    value = artifacts.read_json(reference)
    assert value["authoring_input"]["confirmation_ref"] is None
    units = [artifacts.read_json(ref) for ref in value["unit_refs"]]
    assert len(units) == 4
    assert sum(unit["status"] == "placeholder" for unit in units) == 1
    assert len(service_calls) == 4
    assert len(assembly_calls) == len(structure_calls) == 1
    outline_value = artifacts.read_json(value["outline_ref"])
    form_targets = [target for chapter in outline_value["chapters"] for target in chapter["targets"]
                    if target["kind"] == "form"]
    assert len(form_targets) == 1 and len(form_targets[0]["template_refs"]) == 1
    document = Document(BytesIO(artifacts.read_bytes(value["docx_ref"])))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    assert any("技术方案" in text for text in paragraphs) and any("服务方案" in text for text in paragraphs)
    assert "拟采用部署、联调和验收三个阶段实施本项目。" in paragraphs
    forms = [table for table in document.tables if table.cell(0, 0).text == "人员安排表"]
    assert len(forms) == 1 and forms[0].cell(1, 1).text == "项目负责人（姓名待补）"
    assert forms[0].cell(2, 1).text == ""
    assert len(document.inline_shapes) == 1
    with ZipFile(BytesIO(artifacts.read_bytes(value["package_ref"]))) as package:
        assert any(name.endswith(".png") and package.read(name) == original.getvalue() for name in package.namelist())
    row = await AuthoringStore(settings.dsn).get_execution(execution_id)
    assert row["status"] == "succeeded"
    assert sum(request["status"] == "failed" for request in row["requests"]) == 4
    assert value["usage"]["retry_requests"] == 3 and value["usage"]["unknown_usage_requests"] == 4
    sent_before = len(sent)
    assert await execute(runtime, execution_id, handler, materials=[proof, methods]) == reference
    assert len(sent) == sent_before
    recovered = True
    second = await execute(runtime, uuid4().hex, handler, materials=[proof, methods], previous_execution_id=execution_id)
    assert len(service_calls) == 5
    # The changed unit, its content review, and the changed final structure summary.
    assert artifacts.read_json(second)["usage"]["sent_requests"] == 3
    assert len(assembly_calls) == 1 and len(structure_calls) == 2
    assert all(artifacts.read_json(ref)["status"] == "generated" for ref in artifacts.read_json(second)["unit_refs"])
    assert await store.get_run(run["run_id"]) == before


@pytest.mark.parametrize("identical", [True, False])
async def test_concurrent_requests_wait_only_for_an_identical_logical_call(analysis_runtime, monkeypatch, identical):
    analysis_store, artifacts, settings, _, _ = analysis_runtime
    store = AuthoringStore(settings.dsn)
    await store.setup()
    snapshot = await read_authoring_input(analysis_store, artifacts, **identity(analysis_runtime))
    execution_id = uuid4().hex
    spec = {"protocol": "synthetic-concurrent-document"}
    await store.submit(execution_id, snapshot, spec)
    lease = await store.acquire(execution_id)
    started, both_started, release = asyncio.Event(), asyncio.Event(), asyncio.Event()
    sends = []
    original_sleep = asyncio.sleep

    async def fast_retry(seconds):
        await original_sleep(0.005)

    monkeypatch.setattr("docreader.llm.recovery.asyncio", SimpleNamespace(sleep=fast_retry, to_thread=asyncio.to_thread))

    async def handler(request):
        sends.append(request)
        started.set()
        if len(sends) == 2:
            both_started.set()
        await release.wait()
        return response({"value": "accepted"})

    async with OpenAIModel(config(), store, lease, prompts={}, artifacts=artifacts,
                           transport=httpx.MockTransport(handler), concurrency=2) as client:
        executor = document_execution._DocumentExecutor(client, store, artifacts, lease, budget(), spec, None)
        first_request = build_request(config().chat.model, "synthetic prompt", {"text": "same"})
        second_request = (deepcopy(first_request) if identical else
                          build_request(config().chat.model, "synthetic prompt", {"text": "different"}))
        tasks = [asyncio.create_task(executor("authoring_write", request,
                    lambda raw, target=target: {**raw, "target_id": target}))
                 for request, target in ((first_request, "first-target"), (second_request, "second-target"))]
        try:
            await asyncio.wait_for(started.wait(), timeout=3)
            if identical:
                # Old behavior spends all four sent-fence retries here while
                # the first paid call is still in progress.
                await original_sleep(0.15)
            else:
                await asyncio.wait_for(both_started.wait(), timeout=3)
            release.set()
            outputs = await asyncio.wait_for(asyncio.gather(*tasks), timeout=3)
        finally:
            release.set()
            for task in tasks:
                task.cancel()
            await asyncio.gather(*tasks, return_exceptions=True)
    assert outputs == [{"value": "accepted", "target_id": "first-target"},
                       {"value": "accepted", "target_id": "second-target"}]
    saved = await store.get_execution(execution_id)
    assert len(sends) == len(saved["requests"]) == (1 if identical else 2)
