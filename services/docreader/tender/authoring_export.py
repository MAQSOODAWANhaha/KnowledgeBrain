"""Assemble editable drafts and original evidence without making missing facts a gate."""

from __future__ import annotations

from collections import Counter, defaultdict
from io import BytesIO
import json
import mimetypes
import re
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from PIL import Image, ImageOps

from docreader.tender.authoring_templates import load_template_part, template_part_key
from docreader.tender.evidence.references import EvidenceRepository
from docreader.tender.word_navigation import add_bookmark


DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PENDING_ANALYSIS_NOTE = "分析整理（待确认）：摘要及适用范围仍需结合原文和有效补遗核对。"


def _source_key(reference):
    return reference["source_index_ref"]["sha256"], reference["unit_id"]


def _requirement_key(reference):
    return json.dumps({key: value for key, value in reference.items() if value is not None},
                      ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def _requirement_value(value, *, parent_label=None):
    if isinstance(value, str):
        return value
    if isinstance(value, list):
        return "\n".join(_requirement_value(item) for item in value)
    if isinstance(value, dict):
        labels = {"value": "数值", "unit": "单位", "quote": "原文", "method": "评分方法",
                  "points": "分值", "max_points": "最高分", "effect": "后果", "trigger": "触发条件",
                  "fields": "填写字段", "documents": "提交材料"}
        values = []
        for index, (field, item) in enumerate(value.items()):
            label = labels.get(field, "待核内容")
            prefix = "" if index == 0 and label == parent_label else f"{label}："
            values.append(prefix + _requirement_value(item))
        return "\n".join(values)
    if isinstance(value, bool):
        return "是" if value else "否"
    return "待核" if value is None else str(value)


def _requirement_text(rule):
    values = [_requirement_value(rule["text"])]
    labels = {"subject": "主体", "scope": "适用范围",
              "conditions": "条件", "exceptions": "例外", "numbers": "量化事实", "scoring": "评分细则",
              "consequence": "后果", "submission": "提交要求"}
    for field, label in labels.items():
        value = rule.get(field)
        if value is None or value == "" or value == [] or value == {}:
            continue
        if ((field in {"subject", "scope"} and not isinstance(value, str))
                or (field in {"conditions", "exceptions"} and (
                    not isinstance(value, list) or any(not isinstance(item, str) for item in value)))
                or (field == "numbers" and not isinstance(value, list))
                or (field in {"scoring", "consequence", "submission"} and not isinstance(value, dict))):
            label += "（待核）"
        values.append(f"{label}：{_requirement_value(value, parent_label=label)}")
    return "\n".join(values)


def _walk(chapters, prefix=()):
    for ordinal, chapter in enumerate(chapters, 1):
        path = (*prefix, ordinal)
        yield chapter, path
        yield from _walk(chapter.get("children", []), path)


def _body_entries(entries):
    notes = [path for chapter, path in entries if chapter.get("presentation_role") == "submission_note"]
    return [(chapter, path) for chapter, path in entries
            if not any(path[:len(parent)] == parent for parent in notes)]


def _display_paths(entries):
    paths, identities, counts = {}, {}, defaultdict(int)
    for chapter, path in _body_entries(entries):
        parent = paths.get(path[:-1], ())
        if chapter.get("presentation_role", "body") == "cover":
            paths[path] = parent
            continue
        counts[parent] += 1
        paths[path] = (*parent, counts[parent])
        identities[chapter["chapter_id"]] = paths[path]
    return identities


def _page_width(document):
    section = document.sections[-1]
    return section.page_width - section.left_margin - section.right_margin


def _table_layout(document, table, widths=None, header_count=0):
    available = _page_width(document)
    widths = widths or [available / len(table.columns)] * len(table.columns)
    factor = min(1, available / sum(widths))
    widths = [int(width * factor) for width in widths]
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = width
    for number, row in enumerate(table.rows):
        if number < header_count:
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            row._tr.get_or_add_trPr().append(header)
        for cell, width in zip(row.cells, widths):
            cell.width = width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    return widths


def _toc(document):
    document.add_paragraph("目录", style="Title")
    paragraph = document.add_paragraph()
    for kind in ("begin", "instruction", "separate", "end"):
        element = OxmlElement("w:instrText" if kind == "instruction" else "w:fldChar")
        if kind == "instruction":
            element.set(qn("xml:space"), "preserve")
            element.text = ' TOC \\o "1-9" \\h \\z \\u '
        else:
            element.set(qn("w:fldCharType"), kind)
            if kind == "begin":
                element.set(qn("w:dirty"), "true")
        paragraph.add_run()._r.append(element)
        if kind == "separate":
            paragraph.add_run("请在文字处理软件中更新目录以显示页码。")
    settings = document.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def _submission_notes(outline, requirements, gaps, artifacts=None):
    entries = list(_walk(outline.get("chapters", [])))
    rules = {_requirement_key(row["rule_ref"]): row for row in (requirements or {}).get("rules", [])}
    assigned = {_requirement_key(ref) for chapter, _ in entries for target in chapter.get("targets", [])
                for ref in target.get("rule_refs", [])}
    constraints = outline.get("constraint_rule_refs", [])
    assigned.update(_requirement_key(ref) for ref in constraints)
    assigned.update(_requirement_key(ref) for issue in outline.get("issues", [])
                    if issue["code"] in {"outline_target_excluded", "outline_rule_excluded"}
                    for ref in issue.get("rule_refs", []))
    lines = ["# 提交要求与待办", "", "以下内容来自本次固定分析和编制结果，供人工整理提交；不表示已完成履约或企业应答。", ""]
    if requirements is not None and not requirements.get("authoring_input", {}).get("confirmation_ref"):
        lines.extend([PENDING_ANALYSIS_NOTE, ""])

    def append_rules(title, references):
        lines.extend(["## " + title, ""])
        references = {_requirement_key(ref): ref for ref in references}
        for number, reference in enumerate(references.values(), 1):
            rule = rules.get(_requirement_key(reference), {}).get("rule")
            if not isinstance(rule, dict) or not rule.get("text"):
                lines.extend([f"{number}. 具体要求待核对，正式引用保留在大纲和要求底稿中。", ""])
                continue
            lines.extend([f"{number}. " + _requirement_text(rule).replace("\n", "\n   "), ""])
            evidence = rule.get("evidence") or []
            if not isinstance(evidence, list):
                evidence = [evidence]
            for item in evidence:
                quote = _requirement_value(item.get("quote", item) if isinstance(item, dict) else item)
                lines.extend(["> " + quote.replace("\n", "\n> "), ""])

    append_rules("已知提交与编制约束", constraints)
    body_paths = {path for _, path in _body_entries(entries)}
    for chapter, path in entries:
        if path not in body_paths:
            lines.extend(["## 提交待办：" + chapter["title"], ""])
            for target in chapter.get("targets", []):
                lines.extend(["- " + target["purpose"], ""])
                if artifacts is not None:
                    from docreader.tender.authoring_excerpts import load_source_excerpt, whole_source_excerpts
                    excerpts = list(target.get("source_excerpts", []))
                    precise = {_source_key(ref["source_ref"]) for ref in excerpts}
                    try:
                        for reference in target.get("source_refs", []):
                            if _source_key(reference) not in precise:
                                excerpts.extend(whole_source_excerpts(artifacts, reference))
                        for excerpt in excerpts:
                            text = load_source_excerpt(artifacts, excerpt)["text"]
                            lines.extend(["> " + text.replace("\n", "\n> "), ""])
                    except (OSError, ValueError, KeyError):
                        lines.extend(["原文暂不可读，来源引用保留在本次大纲中，待人工核对。", ""])
    paths = _display_paths(entries)
    for chapter, _ in entries:
        references = [ref for target in chapter.get("targets", []) for ref in target.get("rule_refs", [])]
        if references:
            location = " ".join(part for part in (
                ".".join(map(str, paths.get(chapter["chapter_id"], ()))), chapter["title"]) if part)
            label = "封面填写要求：" if chapter.get("presentation_role", "body") == "cover" else "章节要求："
            append_rules(label + location, references)
    for issue in outline.get("issues", []):
        if issue["code"] in {"outline_target_excluded", "outline_rule_excluded"}:
            excluded_target = issue["code"] == "outline_target_excluded"
            append_rules("已排除候选及关联要求" if excluded_target else "已明确排除的要求",
                         issue.get("rule_refs", []))
            if excluded_target:
                lines.extend(["此记录仅排除该候选；关联要求在其他响应位置的安排仍保留。", ""])
            if issue.get("detail"):
                lines.extend(["排除理由：" + issue["detail"], ""])
    append_rules("尚未安排的要求", [row["rule_ref"] for key, row in rules.items() if key not in assigned])
    lines.extend(["## 人工填写位置", ""])
    for chapter, _ in entries:
        for target in chapter.get("targets", []):
            if target.get("manual_fill"):
                lines.extend([f"- {chapter['title']}：{target['purpose']}", ""])
    lines.extend(["## 当前核对意见与待办", "", "各项保持原核对含义；不统一视为资料缺失，也不阻止草稿导出。", ""])
    for issue in gaps:
        lines.extend(["- " + issue["detail"].replace("\n", "\n  "), ""])
    return "\n".join(lines)


def _image(document, content):
    """Fit real pixels into the page without cropping or changing their aspect ratio."""
    with Image.open(BytesIO(content)) as original:
        oriented = ImageOps.exif_transpose(original)
        try:
            output = BytesIO()
            oriented.save(output, format="PNG")
            section = document.sections[-1]
            width = section.page_width - section.left_margin - section.right_margin
            height = section.page_height - section.top_margin - section.bottom_margin - Mm(20)
            ratio = min(width / oriented.width, height / oriented.height)
            document.add_picture(BytesIO(output.getvalue()),
                                 width=int(oriented.width * ratio), height=int(oriented.height * ratio))
        finally:
            oriented.close()


def _pdf(document, path):
    import pypdfium2 as pdfium
    from docreader.parser.pdf_parser import _PDFIUM_LOCK, _render_page_to_jpeg

    # PDFium is process-global; use the same lock as ingest/source-view rendering.
    with _PDFIUM_LOCK, pdfium.PdfDocument(str(path)) as original:
        if not len(original):
            raise ValueError("attachment PDF has no pages")
        for ordinal in range(len(original)):
            page = original[ordinal]
            try:
                encoded = _render_page_to_jpeg(page, 150 / 72, 90, 2400)
            finally:
                page.close()
            _image(document, encoded)


def _block(document, block, *, width_ratios=None):
    if block.get("text"):
        document.add_paragraph(block["text"])
    rows = ([block["headers"]] if block.get("headers") else []) + block.get("rows", [])
    if not rows or not any(rows):
        return
    table = document.add_table(rows=len(rows), cols=max(map(len, rows)))
    table.style = "Table Grid"
    widths = ([_page_width(document) * ratio / sum(width_ratios) for ratio in width_ratios]
              if width_ratios else None)
    _table_layout(document, table, widths, 1 if block.get("headers") else 0)
    for row_index, values in enumerate(rows):
        for column_index, text in enumerate(values):
            table.cell(row_index, column_index).text = text


class _Assembler:
    def __init__(self, artifacts, writing_units, review_issues, entries=()):
        self.artifacts = artifacts
        body_targets = {target['target_id'] for chapter, _ in _body_entries(entries)
                        for target in chapter.get('targets', [])}
        self.note_targets = {target['target_id'] for chapter, _ in entries
                             for target in chapter.get('targets', [])} - body_targets
        self.units = defaultdict(list)
        for unit in writing_units:
            if unit.get('target_ids') and set(unit['target_ids']).issubset(self.note_targets):
                continue
            self.units[unit["chapter_id"]].append(unit)
        self.gaps = list(review_issues or [])
        self.attachments = {}
        self.template_parts = {}
        self.repository = EvidenceRepository(artifacts)
        self.display_paths = _display_paths(entries)
        self.locations = {target["target_id"]: " ".join(part for part in (
            ".".join(map(str, self.display_paths.get(chapter["chapter_id"], ()))), chapter["title"]) if part)
            for chapter, _ in _body_entries(entries) for target in chapter.get("targets", [])}

    def gap(self, code, detail, chapter_id, target_id=None):
        value = {"code": code, "detail": detail, "chapter_id": chapter_id}
        if target_id is not None:
            value["target_id"] = target_id
        self.gaps.append(value)

    def load_template(self, part):
        key = template_part_key(part)
        if key not in self.template_parts:
            self.template_parts[key] = load_template_part(self.artifacts, part, repository=self.repository)
        return self.template_parts[key]

    def template(self, document, part, fills, chapter_id, target_id):
        try:
            loaded = self.load_template(part)
            unit = loaded["unit"]
            grid = unit.get("grid")
            if unit.get("text"):
                # Display parser heading syntax without changing frozen evidence.
                document.add_paragraph(re.sub(r"(?m)^( {0,3})#{1,6}[ \t]+", r"\1", unit["text"]))
            if grid:
                count = part.get("header_row_count")
                if count is None:
                    rows = {cell["row"] for cell in unit.get("header_cells", [])}
                    count = 0
                    while count in rows:
                        count += 1
                self.table(document, grid, fills, chapter_id, target_id, header_count=count)
            elif loaded.get("image_ref"):
                _image(document, self.artifacts.read_bytes(loaded["image_ref"]))
                self.gap("export_template_not_editable", "该原格式以真实图片保留，填写位置仍待补。",
                         chapter_id, target_id)
            elif loaded.get("spreadsheet_cells") or unit.get("locator", {}).get("cells"):
                cells = list(loaded.get("spreadsheet_cells", {}).values()) or unit["locator"]["cells"]
                rows = defaultdict(dict)
                for cell in cells:
                    rows[cell["row"]][cell["column"]] = cell.get("text", "")
                columns = sorted({cell["column"] for cell in cells})
                _block(document, {"kind": "table", "rows": [
                    [rows[row].get(column, "") for column in columns] for row in sorted(rows)]})
                self.gap("export_template_structure_incomplete", "仅恢复已有原生单元格，需核对原表合并及填写位置。",
                         chapter_id, target_id)
            elif not unit.get("text"):
                document.add_paragraph("（原格式待补）")
                self.gap("export_template_unrendered", "原格式缺少可渲染的文字、网格或图片。", chapter_id, target_id)
            if fills and not grid:
                self.gap("export_cell_fill_unplaced", "该来源没有可填写的原生网格，填充值未写入原格式。",
                         chapter_id, target_id)
            self.gap("export_template_layout_unverified", "已保留可读取的原文和网格；字体、分页和原版式仍需核对。",
                     chapter_id, target_id)
        except Exception as exc:
            document.add_paragraph("（原格式待补）")
            self.gap("export_source_unavailable", f"原格式暂不可用：{type(exc).__name__}。", chapter_id, target_id)

    def table(self, document, grid, fills, chapter_id, target_id, *, header_count=0):
        table = document.add_table(rows=grid["row_count"], cols=grid["column_count"])
        table.style = "Table Grid"
        widths = _table_layout(document, table,
            [Mm(width) for width in grid["widths_mm"]] if grid.get("widths_mm") else None, header_count)
        choices = defaultdict(set)
        for fill in fills:
            choices[(fill["row"], fill["column"])].add(fill["text"])
        anchors = {(cell["row"], cell["column"]) for cell in grid["cells"]}
        for position in choices.keys() - anchors:
            self.gap("export_cell_fill_unplaced", f"填充值不对应原网格锚点：{position}。", chapter_id, target_id)
        for cell in grid["cells"]:
            row, column = cell["row"], cell["column"]
            target = table.cell(row, column)
            row_span, col_span = cell.get("row_span", 1), cell.get("col_span", 1)
            if row_span > 1 or col_span > 1:
                target = target.merge(table.cell(row + row_span - 1, column + col_span - 1))
            target.width = sum(widths[column:column + col_span])
            text = cell.get("text", "")
            values = choices.get((row, column), set())
            if text.strip():
                if values - {text}:
                    self.gap("export_fixed_cell_preserved", f"原固定文字已保留，未覆盖单元格({row}, {column})。",
                             chapter_id, target_id)
            elif len(values) > 1:
                self.gap("export_conflicting_cell_fills", f"单元格({row}, {column})有冲突填值，保留原空白。",
                         chapter_id, target_id)
            elif values:
                text = next(iter(values))
            target.text = text

    def response_table(self, document, target, units, chapter_id):
        from docreader.tender.authoring_excerpts import load_source_excerpt

        rows = target.get("response_table", {}).get("rows", [])
        if not rows:
            return False
        context_texts = []
        for excerpt in target["response_table"].get("context_source_excerpts", []):
            try:
                context_texts.append(load_source_excerpt(self.artifacts, excerpt, repository=self.repository)["text"])
            except (OSError, ValueError, KeyError):
                context_texts.append("（原表共同条件待核对，来源引用已保留）")
                self.gap("export_response_source_unavailable", "该表部分共同条件不可读，保留来源和待核位置。",
                         chapter_id, target["target_id"])
        if context_texts:
            document.add_paragraph(re.sub(r"(?m)^( {0,3})#{1,6}[ \t]+", r"\1", "\n".join(context_texts)))
        choices = defaultdict(set)
        row_ids = {row["row_id"] for row in rows}
        for unit in units:
            for fill in unit.get("row_fills", []):
                if fill["row_id"] not in row_ids:
                    continue
                if target.get("manual_fill", False):
                    self.gap("export_manual_content_omitted", "人工填写响应表未采用自动填值。",
                             chapter_id, target["target_id"])
                elif fill.get("text", "").strip() and not fill.get("material_ids"):
                    self.gap("export_row_fill_unsupported", "该行应答没有知识依据，保留要求和空白响应。",
                             chapter_id, target["target_id"])
                else:
                    choices[fill["row_id"]].add(fill["text"])
        table = document.add_table(rows=len(rows) + 1, cols=2)
        table.style = "Table Grid"
        _table_layout(document, table, [_page_width(document) * 0.65, _page_width(document) * 0.35], 1)
        table.cell(0, 0).text, table.cell(0, 1).text = "招标要求", "投标响应"
        for ordinal, row in enumerate(rows, 1):
            texts = []
            for excerpt in row["source_excerpts"]:
                try:
                    texts.append(load_source_excerpt(self.artifacts, excerpt, repository=self.repository)["text"])
                except (OSError, ValueError, KeyError):
                    texts.append("（原文要求待核对，来源引用已保留）")
                    self.gap("export_response_source_unavailable", "该行部分原文不可读，保留来源和待核位置。",
                             chapter_id, target["target_id"])
            table.cell(ordinal, 0).text = re.sub(r"(?m)^( {0,3})#{1,6}[ \t]+", r"\1", "\n".join(texts))
            values = choices.get(row["row_id"], set())
            if len(values) > 1:
                self.gap("export_conflicting_row_fills", "该要求行存在冲突应答，保留完整要求及空白响应。",
                         chapter_id, target["target_id"])
            elif values:
                table.cell(ordinal, 1).text = next(iter(values))
        return True

    def attachment(self, document, attachment, chapter_id):
        reference = attachment["artifact_ref"]
        try:
            path = self.artifacts.path(reference)
        except Exception as exc:
            self.gap("export_attachment_unavailable", f"{attachment['title']}原件不可用：{type(exc).__name__}。", chapter_id)
            document.add_paragraph(f"{attachment['title']}：待补原件。")
            return
        key = reference["sha256"]
        if key not in self.attachments:
            extension = mimetypes.guess_extension(reference["media_type"]) or ".bin"
            self.attachments[key] = {**attachment,
                "package_path": f"attachments/{len(self.attachments) + 1:03d}-{key[:12]}{extension}",
                "embedded": False, "chapter_ids": [], "occurrences": []}
        manifest = self.attachments[key]
        if chapter_id not in manifest["chapter_ids"]:
            manifest["chapter_ids"].append(chapter_id)
        occurrence = {"material_id": attachment["material_id"], "title": attachment["title"], "chapter_id": chapter_id}
        if occurrence not in manifest["occurrences"]:
            manifest["occurrences"].append(occurrence)
        document.add_paragraph(attachment["title"])
        try:
            media_type = reference["media_type"]
            if media_type.startswith("image/"):
                _image(document, self.artifacts.read_bytes(reference))
            elif media_type == "application/pdf":
                _pdf(document, path)
            else:
                document.add_paragraph("（真实原件随交付包附送）")
                self.gap("export_attachment_not_embedded", f"{attachment['title']}原件已附送，尚未嵌入正文。", chapter_id)
                return
            manifest["embedded"] = True
        except Exception as exc:
            document.add_paragraph("（真实原件随交付包附送）")
            self.gap("export_attachment_not_embedded", f"{attachment['title']}未完整嵌入：{type(exc).__name__}；原件仍随包附送。",
                     chapter_id)

    def writing_unit(self, document, unit, chapter_id, blocked_forms=(), manual_forms=(), response_tables=()):
        blocks = []
        for block in unit.get("blocks", []):
            target_ids = set(block.get("target_ids") or unit.get("target_ids", []))
            if target_ids and target_ids.issubset(self.note_targets):
                continue
            if target_ids and target_ids.issubset(blocked_forms):
                manual = target_ids.issubset(manual_forms)
                self.gap("export_manual_content_omitted" if manual else "export_form_content_unplaced",
                         "人工填写位置未采用生成正文。" if manual else
                         "该原格式没有可填写网格，未另行追加生成声明；补充说明须安排独立正文位置。", chapter_id)
            else:
                blocks.append(block)
        for block in blocks:
            _block(document, block)
            if block["kind"] == "paragraph" and (block.get("headers") or block.get("rows")):
                self.gap("block_kind_mismatch", "段落块同时携带表格字段，已完整保留段落及表格内容。", chapter_id)
        for attachment in unit.get("attachments", []):
            self.attachment(document, attachment, chapter_id)
        for gap in unit.get("gaps", []):
            self.gap("authoring_content_gap", gap, chapter_id)
        has_blocks = any(block.get("text", "").strip() or block.get("headers") or block.get("rows") for block in blocks)
        table_only = bool(unit.get("target_ids")) and set(unit["target_ids"]).issubset(response_tables)
        if not table_only and (unit["status"] == "placeholder"
                               or not (has_blocks or unit.get("cell_fills") or unit.get("attachments"))):
            document.add_paragraph("（本部分内容待补）")
            self.gap("export_content_pending", "写作任务尚未形成完整内容，保留当前位置。", chapter_id)

    def chapter(self, document, chapter, path, *, heading=True):
        chapter_id = chapter["chapter_id"]
        chapter_start = len(document.element.body) - 1
        if heading:
            path = self.display_paths.get(chapter_id, path)
            paragraph = document.add_heading(f"{'.'.join(map(str, path))} {chapter['title']}", level=min(len(path), 9))
            add_bookmark(paragraph._p, chapter_id)
        units = self.units.get(chapter_id, [])
        targets = chapter.get("targets", [])
        reference_only = {target['target_id'] for target in targets
            if target['kind'] == 'narrative' and not target.get('manual_fill', False)
            and not any(target.get(field) for field in ('rule_refs', 'source_refs', 'source_excerpts',
                                                       'template_refs', 'response_table'))
            and target.get('related_target_ids')
            and all(identity in self.locations and identity != target['target_id']
                    for identity in target['related_target_ids'])}
        forms, grid_forms = {}, set()
        manual_forms = {target["target_id"] for target in targets
                        if target["kind"] == "form" and target.get("manual_fill", False)}
        for target in targets:
            if target["kind"] != "form":
                continue
            selected = forms[target["target_id"]] = {}
            for part in target.get("template_refs", []):
                try:
                    selected[template_part_key(part)] = part
                except (KeyError, TypeError, ValueError):
                    self.gap("export_source_unavailable", "原格式选择无效，保留待补位置。",
                             chapter_id, target["target_id"])
            for part in selected.values():
                try:
                    if self.load_template(part)["unit"].get("grid"):
                        grid_forms.add(target["target_id"])
                except Exception:
                    # Rendering below records the unavailable part without substituting its whole unit.
                    continue
        part_uses = Counter(key for selected in forms.values() for key in selected)
        blocked_forms = (forms.keys() - grid_forms) | manual_forms
        row_owners = {row["row_id"]: target["target_id"] for target in targets
                      for row in (target.get("response_table") or {}).get("rows", [])}
        for unit in units:
            allowed = {_source_key(part["source_ref"]) for target_id in unit.get("target_ids", [])
                       for part in forms.get(target_id, {}).values()}
            for fill in unit.get("cell_fills", []):
                if _source_key(fill["source_ref"]) not in allowed:
                    self.gap("export_cell_fill_unplaced", "填值来源不属于本写作任务响应位置的原格式。", chapter_id)
            for fill in unit.get("row_fills", []):
                if row_owners.get(fill["row_id"]) not in unit.get("target_ids", []):
                    self.gap("export_row_fill_unplaced", "行应答无法绑定当前任务的实际响应行，要求保持不变。", chapter_id)
        rendered, unit_anchors = set(), {}
        response_tables = {target["target_id"] for target in targets
                           if (target.get("response_table") or {}).get("rows")}
        for target in targets:
            start_index = len(document.element.body) - 1
            target_id = target["target_id"]
            related = []
            for identity in dict.fromkeys(target.get("related_target_ids", [])):
                if identity in self.locations and identity != target_id:
                    related.append(self.locations[identity])
                else:
                    self.gap("export_related_target_unavailable", "相关响应位置暂不能定位，原关系保留待核。",
                             chapter_id, target_id)
            if related:
                document.add_paragraph("参见 " + "、".join(dict.fromkeys(related)))
            linked = [(ordinal, unit) for ordinal, unit in enumerate(units) if target_id in unit.get("target_ids", [])]
            if target_id in response_tables:
                self.response_table(document, target, [unit for _, unit in linked], chapter_id)
            if target["kind"] == "form":
                if target.get("subject") or any(part_uses[key] > 1 for key in forms[target_id]):
                    label = (target["subject"] + "：" if target.get("subject") else "") + target["purpose"]
                    document.add_paragraph(label, style="Caption")
                for part in forms[target_id].values():
                    key = _source_key(part["source_ref"])
                    fills = [fill for _, unit in linked for fill in unit.get("cell_fills", [])
                             if _source_key(fill["source_ref"]) == key]
                    if target_id in manual_forms and fills:
                        fills = []
                        self.gap("export_manual_content_omitted", "人工填写表格未采用自动填值。", chapter_id, target_id)
                    self.template(document, part, fills, chapter_id, target_id)
                if not forms[target_id] and target_id not in response_tables:
                    document.add_paragraph("（原格式待补）")
                    self.gap("export_template_missing", "该填写位置尚无可用原格式。", chapter_id, target_id)
            for ordinal, unit in linked:
                if ordinal not in rendered:
                    has_content = (any(block.get('text', '').strip() or block.get('headers') or block.get('rows')
                                       for block in unit.get('blocks', []))
                                   or unit.get('cell_fills') or unit.get('row_fills') or unit.get('attachments'))
                    if set(unit.get('target_ids', [])).issubset(reference_only) and not has_content:
                        for gap in unit.get('gaps', []):
                            self.gap('authoring_content_gap', gap, chapter_id)
                        rendered.add(ordinal)
                        continue
                    unit_start = len(document.element.body) - 1
                    self.writing_unit(document, unit, chapter_id, blocked_forms, manual_forms, response_tables)
                    if len(document.element.body) - 1 > unit_start:
                        unit_anchors[ordinal] = document.element.body[unit_start]
                    rendered.add(ordinal)
            if not linked and target_id not in reference_only and target_id not in response_tables:
                document.add_paragraph(f"（待补：{target['purpose']}）")
                self.gap("export_content_pending", "响应位置尚无写作内容。", chapter_id, target_id)
            if len(document.element.body) - 1 > start_index:
                add_bookmark(document.element.body[start_index], target_id, "target")
            else:
                for ordinal, _ in linked:
                    if ordinal in unit_anchors:
                        add_bookmark(unit_anchors[ordinal], target_id, "target")
                        break
        for ordinal, unit in enumerate(units):
            if ordinal not in rendered:
                self.writing_unit(document, unit, chapter_id, blocked_forms, manual_forms, response_tables)
                self.gap("export_unassigned_content", "写作内容未绑定当前章节的响应位置，已保留在章节末尾。", chapter_id)
        if not heading and len(document.element.body) - 1 > chapter_start:
            add_bookmark(document.element.body[chapter_start], chapter_id)

    def render(self, entries, title):
        entries = _body_entries(entries)
        document = Document()
        section = document.sections[0]
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.left_margin = section.right_margin = Mm(25)
        section.top_margin = section.bottom_margin = Mm(22)
        normal = document.styles["Normal"]
        normal.font.name, normal.font.size = "宋体", Pt(11)
        normal.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
        volumes = [path for chapter, path in entries if chapter.get("is_volume")
                   and chapter.get("presentation_role", "body") != "cover"]
        scope = (entries[0][1] if entries and entries[0][0].get("is_volume")
                 and entries[0][0].get("presentation_role", "body") != "cover"
                 and all(path[:len(entries[0][1])] == entries[0][1] for _, path in entries) else ())
        covers = defaultdict(list)
        for chapter, path in entries:
            if chapter.get("presentation_role", "body") == "cover":
                owner = next((volume for volume in reversed(volumes)
                              if len(volume) < len(path) and path[:len(volume)] == volume), ())
                covers[owner].append((chapter, path))

        def render_covers(owner):
            for chapter, path in covers[owner]:
                self.chapter(document, chapter, path, heading=False)
                document.add_page_break()

        if covers[scope]:
            render_covers(scope)
        else:
            document.add_heading(title, 0)
        _toc(document)
        document.add_page_break()
        for chapter, path in entries:
            if chapter.get("presentation_role", "body") == "cover":
                continue
            if chapter.get("is_volume") and path != scope:
                render_covers(path)
            self.chapter(document, chapter, path)
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        # Word rejects a simple field that has no result run.
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "1"
        run.append(text)
        field.append(run)
        section.footer.paragraphs[0]._p.append(field)
        output = BytesIO()
        document.save(output)
        return self.artifacts.put_bytes(output.getvalue(), DOCX_MEDIA_TYPE)


def export_authoring_document(artifacts, outline: dict, writing_units: list[dict],
                              review_issues: list[dict] | None = None, *, requirements: dict | None = None) -> dict:
    """Return immutable DOCX/ZIP references; source and business gaps remain deliverable."""
    entries = list(_walk(outline.get("chapters", [])))
    assembler = _Assembler(artifacts, writing_units, review_issues, entries)
    if requirements is not None:
        rules = {_requirement_key(entry["rule_ref"]): entry.get("rule") for entry in requirements.get("rules", [])}
        locations = [(ref, chapter["chapter_id"], target["target_id"]) for chapter, _ in entries
                     for target in chapter.get("targets", []) for ref in target.get("rule_refs", [])]
        located = {_requirement_key(ref) for ref, _, _ in locations}
        locations.extend((entry["rule_ref"], None, None) for entry in requirements.get("rules", [])
                         if _requirement_key(entry["rule_ref"]) not in located)
        for reference, chapter_id, target_id in locations:
            rule = rules.get(_requirement_key(reference))
            if not isinstance(rule, dict) or not rule.get("text"):
                assembler.gap("export_requirement_unavailable", "已保留要求位置及引用，具体内容尚待核对。",
                              chapter_id, target_id)
                continue
            evidence = rule.get("evidence")
            if not evidence or not isinstance(evidence, list) or any(
                    not isinstance(item, dict) or not isinstance(item.get("quote"), str) for item in evidence):
                assembler.gap("export_requirement_evidence_unverified", "已有原文依据结构需核对，原值保留在底稿中。",
                              chapter_id, target_id)
        assigned = {_requirement_key(ref) for chapter, _ in entries
                    for target in chapter.get("targets", []) for ref in target.get("rule_refs", [])}
        assigned.update(_requirement_key(ref) for ref in outline.get("constraint_rule_refs", []))
        assigned.update(_requirement_key(ref) for issue in outline.get("issues", [])
                        if issue["code"] in {"outline_target_excluded", "outline_rule_excluded"}
                        for ref in issue.get("rule_refs", []))
        for entry in requirements.get("rules", []):
            if _requirement_key(entry["rule_ref"]) not in assigned:
                assembler.gap("export_requirement_unassigned", "底稿中有尚未安排正文或交付位置的要求，完整内容保留在底稿中。", None)
    chapter_ids = {chapter["chapter_id"] for chapter, _ in entries}
    for chapter_id in assembler.units.keys() - chapter_ids:
        assembler.gap("export_unassigned_content", "内容的章节已不在当前大纲中，保留在补充内容中。", chapter_id)
        entries.append(({"chapter_id": chapter_id, "title": "补充内容", "targets": []}, (len(entries) + 1,)))
    docx_ref = assembler.render(entries, "投标文件")
    volumes = []
    for chapter, path in _body_entries(entries):
        if chapter.get("is_volume"):
            subtree = [(value, value_path) for value, value_path in entries if value_path[:len(path)] == path]
            reference = assembler.render(subtree, chapter["title"])
            volumes.append({"chapter_id": chapter["chapter_id"], "title": chapter["title"],
                "artifact_ref": reference, "package_path": f"volumes/{len(volumes) + 1:03d}.docx"})
    # Repeated uses in a main file and a volume share a single diagnostic entry.
    gaps = list({json.dumps(gap, ensure_ascii=False, sort_keys=True): gap for gap in assembler.gaps}.values())
    attachments = list(assembler.attachments.values())
    manifest = {"schema_version": 2, "document": {"artifact_ref": docx_ref, "package_path": "投标文件.docx"},
                "volumes": volumes, "attachments": attachments}
    requirement_files = []
    if requirements is not None:
        manifest["requirements"] = {"artifact_ref": artifacts.put_json(requirements), "package_path": "招标要求底稿.json"}
        requirement_files.append(manifest["requirements"])
    notes = artifacts.put_bytes(_submission_notes(outline, requirements, gaps, artifacts).encode(), "text/markdown")
    manifest["submission_notes"] = {"artifact_ref": notes, "package_path": "提交要求与待办.md"}
    output = BytesIO()
    with ZipFile(output, "w", compression=ZIP_DEFLATED) as archive:
        for entry in [manifest["document"], *volumes, *attachments, *requirement_files, manifest["submission_notes"]]:
            archive.write(artifacts.path(entry["artifact_ref"]), entry["package_path"])
        archive.writestr("manifest.json", json.dumps(manifest, ensure_ascii=False, indent=2))
        archive.writestr("待补事项.json", json.dumps(gaps, ensure_ascii=False, indent=2))
    return {"docx_ref": docx_ref, "package_ref": artifacts.put_bytes(output.getvalue(), "application/zip"),
            "attachments": [{key: attachment[key] for key in ("material_id", "title", "artifact_ref")}
                            for attachment in attachments], "gaps": gaps}
