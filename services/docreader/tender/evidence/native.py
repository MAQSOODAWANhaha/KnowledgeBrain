"""New-path adapters; the existing parser and gRPC paths are unchanged."""
from __future__ import annotations

import base64
import io
import json
from collections import Counter
from pathlib import Path

from docreader.models.document import AttachmentLocator, DocumentLocator, ImageLocator, StructuredSourceUnit, StructuredSourceUnitKind
from docreader.parser.parser import Parser, detect_effective_file_type
from docreader.parser.source_contract import attach_source_contract


def canonical_image_keys(value) -> None:
    """Adapt the native image string ID once; original-file refs stay ArtifactRef."""
    if isinstance(value, dict):
        if isinstance(value.get("original_ref"), str):
            if "image_key" in value:
                raise ValueError("image identity supplies two field names")
            value["image_key"] = value.pop("original_ref")
        for child in value.values():
            canonical_image_keys(child)
    elif isinstance(value, list):
        for child in value:
            canonical_image_keys(child)


def _supplement_pdf_images(document, content: bytes) -> None:
    """Retain actual image occurrences filtered out by ingest heuristics."""
    import pypdfium2 as pdfium
    import pypdfium2.raw as raw
    from docreader.parser.pdf_parser import _PDFIUM_LOCK, _pdf_image_unit

    known = Counter()
    for unit in document.structured_source_units:
        loc = unit.locator
        if isinstance(loc, ImageLocator) and loc.left is not None:
            known[(loc.page_ordinal, tuple(round(v, 3) for v in (loc.left, loc.top, loc.right, loc.bottom)))] += 1
    additions = []
    problems = []
    with _PDFIUM_LOCK:
        pdf = pdfium.PdfDocument(content)
        try:
            for page_index, classification in enumerate(document.metadata.get("page_classifications", [])):
                if classification != "text":
                    continue  # A scanned-page image already covers this page.
                page = pdf[page_index]
                try:
                    for occurrence, obj in enumerate(page.get_objects()):
                        if obj.type != raw.FPDF_PAGEOBJ_IMAGE:
                            continue
                        try:
                            bounds = tuple(obj.get_bounds())
                            key = (page_index, tuple(round(v, 3) for v in bounds))
                            if known[key]:
                                known[key] -= 1
                                continue
                            bitmap = obj.get_bitmap()
                            try:
                                pixels = bitmap.to_pil()
                                output = io.BytesIO()
                                pixels.save(output, format="PNG")
                            finally:
                                bitmap.close()
                            ref = f"images/native-page-{page_index}-object-{occurrence}.png"
                            encoded = base64.b64encode(output.getvalue()).decode("ascii")
                            document.images[ref] = encoded
                            unit = _pdf_image_unit(f"page:{page_index}:native-image:{occurrence}", 0,
                                                   page_index, ref, encoded, bounds)
                            unit.physical_path = f"pdf:page:{page_index}/image-object:{occurrence}"
                            additions.append(unit)
                        except Exception as exc:
                            problems.append({"code": "pdf_image_unreadable", "page_ordinal": page_index,
                                             "detail": type(exc).__name__})
                finally:
                    page.close()
        finally:
            pdf.close()
    document.structured_source_units.extend(additions)
    def page_number(unit):
        return getattr(unit.physical_locator or unit.locator, "page_ordinal", 0) or 0
    document.structured_source_units.sort(key=page_number)
    for ordinal, unit in enumerate(document.structured_source_units):
        unit.ordinal = ordinal
    document.metadata["evidence_issues"] = problems
    from pypdf import PdfReader
    # Embedded files are explicit unresolved sources until separately registered.
    # Keeping their carrier prevents apparently complete coverage from losing them.
    for attachment_index, name in enumerate(PdfReader(io.BytesIO(content)).attachments):
        document.structured_source_units.append(StructuredSourceUnit(
            key=f"attachment:{attachment_index}", ordinal=len(document.structured_source_units),
            kind=StructuredSourceUnitKind.ATTACHMENT_REGION, text=name,
            locator=AttachmentLocator(part_name=name, relationship_type="pdf_embedded_file"),
            physical_path=f"pdf:embedded-file:{name}",
        ))
    if document.metadata.get("repeating_lines"):
        document.metadata["evidence_issues"].append({
            "code": "pdf_repeating_edge_lines_removed",
            "detail": "Native parser removed repeated page-edge text; original remains available for review.",
            "lines": document.metadata["repeating_lines"],
        })


def _xlsx_anchors(content: bytes) -> dict:
    """Read original drawing anchors, independent of row/text parser rendering."""
    from openpyxl import load_workbook
    from docreader.parser.xlsx_repair import repair_xlsx_bytes
    workbook = load_workbook(io.BytesIO(repair_xlsx_bytes(content) or content), read_only=False)
    result = {}
    try:
        for sheet_index, sheet in enumerate(workbook.worksheets):
            for image_index, image in enumerate(sheet._images):
                anchor = image.anchor
                value = {"sheet_ordinal": sheet_index, "sheet_name": sheet.title,
                         "anchor_kind": type(anchor).__name__}
                for name, attribute in (("from", "_from"), ("to", "to")):
                    marker = getattr(anchor, attribute, None)
                    if marker is not None:
                        value[name] = {"row": marker.row + 1, "column": marker.col + 1,
                                       "row_offset_emu": marker.rowOff, "column_offset_emu": marker.colOff}
                for name in ("pos", "ext"):
                    item = getattr(anchor, name, None)
                    if item is not None:
                        value[name] = dict(item)
                result[f"sheet:{sheet_index}:image:{image_index}"] = value
    finally:
        workbook.close()
    return result


def parse_native(file_name: str, content: bytes):
    extension = detect_effective_file_type(Path(file_name).suffix, content)
    if extension in {"docx", "xlsx", "xlsm"}:
        import zipfile
        if not zipfile.is_zipfile(io.BytesIO(content)):
            raise ValueError("invalid native Office container")
    if extension == "pdf":
        from docreader.parser.pdf_parser import PDFParser
        # _route is the exact existing native parser, without its silent fallback
        # or force-scanned switch. No output_inventory/full-page text rendering.
        document = PDFParser(file_name=file_name, file_type="pdf", pdf_force_scanned=False)._route(content)
        _supplement_pdf_images(document, content)
        document = attach_source_contract(document, content, "pdf")
    else:
        document = Parser().parse_file(file_name, extension, content)
    if extension in {"xlsx", "xlsm"}:
        document.metadata["image_anchors"] = _xlsx_anchors(content)
    if not document.structured_source_units and document.content.strip():
        # Preserve text from non-native engines but never claim a complete native map.
        document.structured_source_units = [StructuredSourceUnit(
            key="opaque:0", ordinal=0, kind=StructuredSourceUnitKind.SECTION,
            text=document.content, locator=DocumentLocator(section_ordinal=0),
            source_issues=["native_structure_unavailable"],
        )]
        document = attach_source_contract(document, content, extension, render=False)
    if "source_contract" not in document.metadata:
        document.metadata["source_contract"] = json.dumps({"units": [], "page_manifest": []})
    return document


def compact_native_snapshot(raw: dict) -> None:
    """Remove reconstructible display copies; original native grid values remain."""
    tables, grid_values = {}, {}
    for unit in raw["units"]:
        locator = unit["locator"]
        if unit.get("grid"):
            unit["text"] = ""  # Grid text is rendered from the one canonical grid.
            if locator.get("locator_kind") == "spreadsheet":
                region = locator["region"]
                for cell in unit["grid"]["cells"]:
                    coordinate = (locator["sheet_ordinal"], region["start_row"] + cell["row"], region["start_column"] + cell["column"])
                    tables[coordinate] = {"unit_id": unit["unit_id"], "row": cell["row"], "column": cell["column"]}
                    grid_values[coordinate] = cell["text"]
        if locator.get("locator_kind") == "spreadsheet" and unit["kind"] == "table_row":
            unit["text"] = ""  # Per-cell formula/type/display metadata is retained below.
    spreadsheet_cells = {}
    for unit in raw["units"]:
        locator = unit["locator"]
        if locator.get("locator_kind") == "spreadsheet":
            cell_refs = []
            for cell in locator.pop("cells", []):
                coordinate = (locator["sheet_ordinal"], cell["row"], cell["column"])
                key = f"sheet:{coordinate[0]}:cell:{coordinate[1]}:{coordinate[2]}"
                cell_refs.append(key)
                if coordinate in tables:
                    inherited = [field for field in ("text", "raw_value", "display_text", "cached_value")
                                 if field in cell and cell[field] == grid_values[coordinate]]
                    for field in inherited:
                        cell.pop(field)
                    cell["grid_value_fields"] = inherited
                    cell["grid_cell_ref"] = tables[coordinate]
                spreadsheet_cells.setdefault(key, cell)
            locator["cell_refs"] = cell_refs
        # The grid definition already owns every cell of table locators.
        if unit.get("grid") and locator.get("locator_kind") != "spreadsheet":
            locator.pop("cells", None)
        physical = unit.get("physical_locator")
        if isinstance(physical, dict):
            physical.pop("cells", None)
    for identity in raw["source_contract"].get("units", []):
        identity.pop("header_cells", None)  # Canonical unit retains native header coordinates.
        identity.pop("rendered_spans", None)  # Refers to omitted rendered content.
        if isinstance(identity.get("physical_locator"), dict):
            identity["physical_locator"].pop("cells", None)
    if spreadsheet_cells:
        raw["spreadsheet_cells"] = spreadsheet_cells
