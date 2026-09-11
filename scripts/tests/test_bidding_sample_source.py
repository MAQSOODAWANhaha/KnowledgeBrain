"""Exercise the sample freeze through the shared service's real Office parser."""
import importlib.util
import json
from pathlib import Path
import sys

from docx import Document
from openpyxl import Workbook
import pytest

spec = importlib.util.spec_from_file_location(
    'sample_source', Path(__file__).parents[1] / 'bidding_sample_source.py')
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)


@pytest.mark.parametrize('extension', ['docx', 'xlsx'])
def test_freeze_preserves_v3_sparse_grid_and_source_identity(tmp_path, monkeypatch, extension):
    source = tmp_path / f'source.{extension}'
    if extension == 'docx':
        document = Document()
        document.add_paragraph('表前说明')
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).merge(table.cell(0, 1)).text = '合并表头'
        table.cell(1, 0).text = '名称：'
        document.add_paragraph('表后说明')
        document.save(source)
    else:
        document = Workbook()
        sheet = document.active
        sheet.merge_cells('A1:B1')
        sheet['A1'] = '合并表头'
        sheet['A2'] = '名称：'
        sheet['B2'] = '待填'
        document.save(source)
    limits = tmp_path / 'limits.json'
    limits.write_text(json.dumps({'extraction': {}}))
    output = tmp_path / 'frozen'
    monkeypatch.setattr(sys, 'argv', ['bidding_sample_source', '--source', str(source),
                        '--output', str(output), '--limits', str(limits)])
    module.main()
    frozen_path = output / 'frozen-input.json'
    original = frozen_path.read_bytes()
    frozen = json.loads(original)
    parsed = json.loads((output / 'parsed-source.json').read_text())
    parsed_grids = [unit['grid'] for unit in parsed['structured_source_units'] if unit['grid']]
    assert len(parsed_grids) == len(frozen['structured_forms']) == 1
    sources = {s['source_unit_revision_id']: s for s in frozen['source_units']}
    form = frozen['structured_forms'][0]
    definition = form['definition']
    grid = parsed_grids[0]
    assert definition['schema_version'] == 3
    assert definition['form_definition_revision_id'] == form['form_definition_revision_id']
    assert definition['source_unit_revision_id'] == form['source_unit_revision_id']
    assert definition['cells'] == grid['cells']
    assert definition['row_count'] == definition['column_count'] == 2
    assert len(definition['cells']) == 3
    assert definition['cells'][0]['col_span'] == 2
    assert definition.get('widths_mm') == grid.get('widths_mm')
    table_source = sources[form['source_unit_revision_id']]
    assert table_source['text'] == ''
    assert not table_source['locator'].get('cells')
    assert table_source['locator']['locator_kind'] == ('document' if extension == 'docx' else 'spreadsheet')
    with pytest.raises(SystemExit, match='frozen input exists'):
        module.main()
    assert frozen_path.read_bytes() == original
