"""Office may make OOXML defaults explicit without changing table semantics."""
from pathlib import Path
import sys
import unittest

from docx import Document
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from onlyoffice_rich_probe import tables


class TableSemanticsTests(unittest.TestCase):
    def test_explicit_merge_continuation_preserves_structure(self):
        document = Document()
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Merged label"
        table.cell(0, 0).merge(table.cell(1, 0))
        original = tables(document)
        continuation = table._tbl.tr_lst[1].tc_lst[0].tcPr.find(qn("w:vMerge"))
        continuation.set(qn("w:val"), "continue")
        self.assertEqual(original, tables(document))
        continuation.set(qn("w:val"), "restart")
        self.assertNotEqual(original, tables(document))
        continuation.getparent().remove(continuation)
        self.assertNotEqual(original, tables(document))

    def test_unrelated_cell_text_and_width_changes_are_detected(self):
        document = Document()
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Original"
        original = tables(document)
        table.cell(0, 1).text = "Unexpected"
        self.assertNotEqual(original, tables(document))
        table.cell(0, 1).text = ""
        self.assertEqual(original, tables(document))
        table._tbl.tr_lst[0].tc_lst[1].tcPr.find(qn("w:tcW")).set(qn("w:w"), "1")
        self.assertNotEqual(original, tables(document))


if __name__ == "__main__":
    unittest.main()
