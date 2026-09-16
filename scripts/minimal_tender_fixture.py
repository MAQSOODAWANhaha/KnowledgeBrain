"""Generate a deterministic synthetic procurement source, never a bid answer.

Run with services/docreader/.venv/bin/python. Only the DOCX is model source;
the independent acceptance matrix lives under artifacts/minimal-bid-fixture.
"""

from datetime import datetime, timezone
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "testdata/bid/minimal/source/minimal-security-tender.docx"
EVIDENCE = ROOT / "artifacts/minimal-bid-fixture"


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Noto Sans CJK SC", Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15
    for name in ["Title", "Heading 1", "Heading 2"]:
        doc.styles[name].font.name = "Noto Sans CJK SC"
        doc.styles[name].element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    section.header.paragraphs[0].text = "星河研究中心日志管理平台采购 · 采购文件"
    footer = section.footer.paragraphs[0]
    footer.alignment = 1
    footer.add_run("第 ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.add_run(" 页")

    def text(value):
        return doc.add_paragraph(value)

    def heading(value, level=1):
        return doc.add_heading(value, level)

    def page(value):
        doc.add_page_break()
        heading(value)

    def table(headers, rows):
        result = doc.add_table(rows=1, cols=len(headers))
        result.style = "Table Grid"
        for cell, value in zip(result.rows[0].cells, headers):
            cell.text = value
        prop = result.rows[0]._tr.get_or_add_trPr()
        prop.append(OxmlElement("w:tblHeader"))
        for row in rows:
            for cell, value in zip(result.add_row().cells, row):
                cell.text = value
        return result

    heading("星河研究中心日志管理平台采购文件", 0)
    text("项目编号：XH-CG-2026-014\n采购人：星河研究中心\n发布日期：2026年9月1日")
    heading("第一章 项目及投标前附表")
    text("1.1 本项目采购一套日志管理软件及部署、培训和维护服务，使用采购人现有服务器，不采购新增硬件。软件部署在采购人内网，采购数量及服务期限见本表。")
    table(["事项", "本项目规定"], [
        ["采购标的", "日志管理软件1套，部署培训1项，维护服务12个月"],
        ["投标截止时间", "2026年10月9日09时00分（北京时间）"],
        ["交付地点与期限", "采购人指定内网环境；合同生效后30个自然日内完成部署并申请验收"],
        ["递交方式", "采用电子递交，不要求纸质正副本；提交一个完整PDF和内容一致的可编辑DOCX"],
        ["部署及运维", "本地部署；不启用云托管；维护期每季度开展一次现场巡检"],
        ["联合体", "不接受联合体，不允许转包"],
        ["签署", "投标声明由法定代表人或授权代表签字并加盖单位公章；授权代表签署时附授权书"],
        ["报价", "人民币含税总价；费用按附表丙列明，不得另列未说明的必选费用"],
    ])
    text("1.2 本前附表对本项目的选择优先于通用条款。本文第二章规定的是投标文件的组成；本采购文件各章标题不构成投标文件目录。")

    page("第二章 投标文件组成及评审")
    text("2.1 投标文件必须按以下顺序编排为一本，设置封面及目录，封面写明项目名称、项目编号、投标人名称和递交日期。各章内文件顺序也按本条执行。")
    table(["顺序", "投标文件章节", "必须包含的内容"], [
        ["一", "投标声明与资格文件", "附表甲投标声明；主体资格证明；签署人为授权代表时的授权书"],
        ["二", "交付方案及逐项响应", "部署及验收安排；附表乙（乙1技术响应、乙2服务响应）；技术证明材料索引"],
        ["三", "费用明细与服务承诺", "附表丙费用明细；维护服务承诺"],
    ])
    text("2.2 附表甲、乙、丙是本项目指定格式，须保留标题、固定声明、栏目、单位、说明及签署位置。乙1和乙2是附表乙的两个连续子表，不可缺一；其间说明仍属于附表乙。每一响应须写明具体内容及偏差，不得仅以‘满足’替代。证明材料可附在第二章末，并在索引中标注对应条款和材料名称。")
    text("2.3 主体资格证明为有效营业执照或事业单位法人证书复印件。投标人独立承担法律责任；不要求未获选的联合体提交协议。法定代表人本人签署时无需授权书。")
    text("2.4 资格、递交方式、指定组成和签署要求采用符合性评审；第三章技术及服务条款须逐项响应。总分100分：技术响应50分（第3.2.1至3.2.4条每条12.5分，满足且证明完整得分，否则该条不得分）；交付安排20分（里程碑和验收方法各10分）；价格30分（评审基准价/有效投标含税总价×30，保留两位小数）。评审基准价为通过符合性评审的最低有效含税总价。")

    page("第三章 技术要求及服务条件")
    text("3.1 软件应支持采购人内网日志集中管理，实际软件型号、版本及配置由投标人如实填报。以下阈值是采购要求，不是投标人已达到的事实。")
    text("3.2.1 持续日志接收能力不低于2000条/秒，且在线可查询保存期不少于180天；须同时满足两项。在乙1分别填写接收能力与保存期，并提供同一投标版本的测试报告或产品手册页作为证明，标明材料名称及位置。")
    text("3.2.2 日志接入至少支持Syslog或HTTPS API中的一种，两者任选其一即可；采用HTTPS API时应支持TLS 1.2及以上。乙1写明所选接入方式及适用的传输保护，不得将两种接入方式误列为同时必选。")
    text("3.2.3 应至少设置系统管理员、审计员、只读用户三类相互独立的角色；管理操作日志应保留不少于180天。乙1须分别说明角色隔离与操作日志保存。")
    text("3.2.4 应支持按来源、时间范围及事件级别组合查询，并将结果导出为CSV；提交界面说明或产品手册作为证明。")
    text("3.3.1 交付方案列明安装、联调、培训、试运行及验收里程碑，整体期限按前附表；培训不少于2次，每次不少于2小时。乙2分别填写培训次数与单次时长，并说明验收方法。")
    text("3.3.2 维护期内故障响应不超过2小时，现场巡检安排按前附表；乙2须填服务期限、响应时间和巡检频率，并在第三章提供维护服务承诺。")
    text("3.3.3 仅选择云托管时，须提交云平台数据驻留说明和租赁费明细。本项目选择见前附表；未选择云托管时，不须提交这两项文件，也不计云资源费用。")
    text("3.4 采用纸质递交的项目，副本数量和装订要求由相应项目的前附表规定。本项目采用电子递交，纸质副本和装订条款不适用；电子文件签署仍按本项目第1.1条前附表。")

    page("第四章 指定格式：附表甲")
    heading("附表甲 投标声明", 2)
    text("致：星河研究中心")
    text("项目名称：星河研究中心日志管理平台采购\n项目编号：XH-CG-2026-014")
    text("我方已阅读采购文件，现就本项目提交投标文件。我方所提供资格证明和投标资料真实、准确；具体技术响应、偏差及报价以本投标文件相应章节为准。本声明不得替代逐项技术和服务响应。")
    text("投标人名称：________________\n统一社会信用代码：________________\n联系人：________________  联系电话：________________")
    text("法定代表人或授权代表（签字）：________________\n投标人（盖章）：________________\n日期：________年____月____日")
    text("填写说明：固定声明和日期的‘年、月、日’标签须保留，横线处由投标人填写；不在此表填写含税总价。授权代表签署时，在本表之后附授权书。")
    heading("资格文件及授权书编排说明", 2)
    text("按第二章第2.1条在投标声明之后编排主体资格证明；如需授权书，须写明委托人、被授权人、项目名称、授权事项和期限，并由法定代表人签字及单位盖章。授权书样式由投标人编制，采购人不指定另一张表格。")

    page("第四章 指定格式（续）：附表乙")
    heading("附表乙 逐项响应表", 2)
    text("项目编号：XH-CG-2026-014\n投标人名称：________________")
    heading("乙1 技术响应", 2)
    table(["对应条款及项目", "单位", "投标响应及偏差", "证明材料名称及位置"], [
        ["3.2.1 持续接收能力", "条/秒", "", ""],
        ["3.2.1 在线可查询保存期", "天", "", ""],
        ["3.2.2 接入方式及适用的传输保护", "—", "", ""],
        ["3.2.3 角色隔离", "—", "", ""],
        ["3.2.3 管理操作日志保存", "天", "", ""],
        ["3.2.4 组合查询与CSV导出", "—", "", ""],
    ])
    text("说明一：乙1的条款名称和单位是固定内容；响应、偏差与证明位置由投标人填写。一个条款含多个指标时须逐行说明，不得合并遗漏。")
    text("说明二：下一页乙2为本附表的续表。乙1与乙2共同构成完整逐项响应；证明材料索引应覆盖需提交证明的条款，说明无需证明的条款不受此限制。")
    text("技术联系人：________________  日期：________年____月____日")

    page("第四章 指定格式（续）：乙2与材料索引")
    heading("乙2 服务响应（附表乙续表）", 2)
    table(["对应条款及项目", "单位", "投标响应及偏差", "验收方法或承诺位置"], [
        ["3.3.1 总体交付期限", "自然日", "", ""],
        ["3.3.1 培训次数", "次", "", ""],
        ["3.3.1 单次培训时长", "小时/次", "", ""],
        ["3.3.2 维护服务期限", "个月", "", ""],
        ["3.3.2 故障响应时间", "小时", "", ""],
        ["3.3.2 现场巡检频率", "次/季度", "", ""],
    ])
    text("说明三：巡检频率与维护期限须符合前附表；此处应写服务安排，不得用技术性能指标替代。签署确认覆盖乙1及乙2两表。")
    text("法定代表人或授权代表（签字）：________________\n投标人（盖章）：________________\n日期：________年____月____日")
    heading("技术证明材料索引", 2)
    text("在下表之后编排证明材料；索引可按实际材料增行，但不得删除固定栏目。")
    table(["对应条款", "材料名称", "投标文件中的位置"], [["", "", ""], ["", "", ""]])
    text("编排说明：材料内容须与投标软件版本一致；未知页码可在定稿后填写，不得预填无依据的页码。")

    page("第四章 指定格式（续）：附表丙")
    heading("附表丙 费用明细", 2)
    text("项目名称：星河研究中心日志管理平台采购\n币种：人民币；金额单位：元")
    costs = table(["费用项目", "数量及单位", "含税单价（元）", "含税合价（元）"], [
        ["日志管理软件", "1套", "", ""],
        ["部署及培训", "1项", "", ""],
        ["维护服务", "12个月", "", ""],
        ["含税总价", "", "", ""],
    ])
    costs.rows[-1].cells[0].merge(costs.rows[-1].cells[2])
    text("说明：含税合价按对应数量计算，含税总价为各项合价之和；应列明适用税率：________。本项目不计云资源租赁费用。空白金额由投标人填写，采购文件未给出投标价格。")
    heading("维护服务承诺", 2)
    text("投标人应在本处列明维护联系人、故障受理方式、响应安排及季度巡检安排，并说明与乙2响应的对应关系。服务期限与故障响应时间须与乙2一致。")
    text("承诺内容：________________________________________________\n________________________________________________________")
    text("法定代表人或授权代表（签字）：________________\n投标人（盖章）：________________\n日期：________年____月____日")
    props = doc.core_properties
    props.title = "星河研究中心日志管理平台采购文件"
    props.author = "星河研究中心"
    props.created = props.modified = datetime(2026, 9, 1, tzinfo=timezone.utc)
    raw = BytesIO()
    doc.save(raw)
    normalized = BytesIO()
    with ZipFile(raw) as original, ZipFile(normalized, "w", ZIP_DEFLATED) as output:
        for name in sorted(original.namelist()):
            info = ZipInfo(name, date_time=(2026, 9, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            output.writestr(info, original.read(name))
    return normalized.getvalue()


def main():
    data = build()
    SOURCE.parent.mkdir(parents=True, exist_ok=True)
    SOURCE.write_bytes(data)
    EVIDENCE.mkdir(parents=True, exist_ok=True)
    manifest = {
        "kind": "synthetic_procurement_source_fixture",
        "source": str(SOURCE.relative_to(ROOT)),
        "sha256": sha256(data).hexdigest(),
        "byte_length": len(data),
        "generator": str(Path(__file__).resolve().relative_to(ROOT)),
        "explicit_page_breaks": 6,
        "intended_pages": 7,
        "rendered_page_count": None,
        "source_pdf": None,
        "parsed_by_unified_docreader": False,
        "model_calls": 0,
        "acceptance": "not_run",
    }
    (EVIDENCE / "source-manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(f"{manifest['source']} {manifest['sha256']}")


if __name__ == "__main__":
    main()
