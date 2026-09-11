"""User-visible table/image edits and DOCX assertions for the native web probe."""
import hashlib
import io
import json
import re
import uuid
import zipfile

from docx import Document
from docx.oxml.ns import qn


def tables(document, remove=""):
    def prop(element, name, attribute="val"):
        found = element.find(qn("w:" + name)) if element is not None else None
        if found is None:
            return None
        # OOXML defines an omitted vMerge value as "continue". Document Server
        # writes the explicit value when saving; both encode the same merge.
        default = "continue" if name == "vMerge" else "present"
        return found.get(qn("w:" + attribute), default)
    return [{"grid": [col.get(qn("w:w")) for col in table._tbl.tblGrid],
        "rows": [{"repeat_header": prop(row.trPr, "tblHeader"), "cells": [
            {"span": prop(cell.tcPr, "gridSpan"), "merge": prop(cell.tcPr, "vMerge"),
             "width": prop(cell.tcPr, "tcW", "w"),
             "text": "".join(cell.xpath(".//w:t/text()")).replace(remove, "") if remove
                else "".join(cell.xpath(".//w:t/text()"))} for cell in row.tc_lst]}
            for row in table._tbl.tr_lst]} for table in document.tables]


def media(data):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return {hashlib.sha256(archive.read(name)).hexdigest() for name in archive.namelist()
            if name.startswith("word/media/") and not name.endswith("/")}


class RichEdits:
    def __init__(self, root, original):
        self.root = root
        self.original = Document(io.BytesIO(original))
        self.original_media = media(original)
        if not self.original.tables:
            raise ValueError("rich edit sample must contain a table")
        self.needle = self.original.tables[0].cell(0, 0).text
        if not self.needle.strip():
            raise ValueError("rich edit sample needs a nonempty first table cell")
        self.marker = "KB-CELL-" + uuid.uuid4().hex
        from PIL import Image
        self.image = root / (uuid.uuid4().hex + ".png")
        pixels = bytes.fromhex(uuid.uuid4().hex) * (96 * 64 * 3 // 16)
        fixture = Image.frombytes("RGB", (96, 64), pixels)
        fixture.save(self.image, dpi=(96, 96))
        self.image_sha = hashlib.sha256(self.image.read_bytes()).hexdigest()
        self.edited = False
        self.extent = None

    def edit(self, page, web):
        frame = page.frame_locator('iframe[name="frameEditor"]')
        try:
            frame.locator("#id_viewer_overlay").click(position={"x": 350, "y": 200})
            page.keyboard.press("Control+Home")
            page.keyboard.press("Control+f")
            page.keyboard.type(self.needle, delay=30)
            page.keyboard.press("Enter")
            page.keyboard.press("Escape")
            page.wait_for_timeout(300)  # Allow the SDK search panel to return keyboard focus.
            page.keyboard.press("ArrowRight")
            page.wait_for_timeout(100)
            page.keyboard.type(self.marker, delay=30)
            web.edited(page)
            frame.locator("#id_viewer_overlay").click(position={"x": 350, "y": 200})
            page.keyboard.press("Control+Home")
            frame.get_by_text("插入", exact=True).click()
            frame.get_by_role("button", name="图片", exact=True).click()
            controls = frame.locator("button, [role=menuitem], [role=button], [data-tab], input").evaluate_all(
                "nodes => nodes.map(n=>({tag:n.tagName,id:n.id,title:n.getAttribute('title'),"
                "text:n.innerText,placeholder:n.getAttribute('placeholder')}))")
            (self.root / "insert-controls.json").write_text(json.dumps(controls, ensure_ascii=False, indent=2))
            page.screenshot(path=str(self.root / "insert-toolbar.png"))
            with page.expect_file_chooser() as chooser:
                frame.get_by_text("来自文件的图片", exact=True).filter(visible=True).click()
            chooser.value.set_files(self.image)
            frame.locator("#id-right-menu-image").click()
            frame.locator("#image-button-original-size").wait_for(state="visible")
            frame.locator("#image-advanced-link").click()
            width = frame.locator("#image-advanced-spin-width input")
            old_width = width.input_value()
            new_width = re.sub(r"[0-9]+(?:\.[0-9]+)?", lambda match: str(float(match[0]) * 1.5), old_width, count=1)
            assert new_width != old_width
            width.fill(new_width)
            width.press("Tab")
            self.dimensions = {"before":old_width,"after":width.input_value()}
            assert self.dimensions["before"] != self.dimensions["after"]
            frame.locator('button[result="ok"]:visible').click()
            frame.locator("#id-right-menu-image").click()
            web.edited(page)
            page.screenshot(path=str(self.root / "image-inserted.png"))
            self.edited = True
        except Exception:
            page.screenshot(path=str(self.root / "rich-failed.png"))
            raise

    def verify(self, data):
        if not self.edited:
            return
        document = Document(io.BytesIO(data))
        (self.root / "rich-inspected.json").write_text(json.dumps(tables(document), ensure_ascii=False, indent=2))
        assert self.marker in document.tables[0].cell(0, 0).text, "table edit did not persist in the target cell"
        assert tables(document, self.marker) == tables(self.original), "table structure or unrelated cell text changed"
        assert self.original_media | {self.image_sha} <= media(data), "original or inserted image bytes missing"
        # Check a drawing references the inserted bytes; an orphan ZIP part is
        # insufficient evidence that the image is in the saved document.
        embedded = {hashlib.sha256(document.part.related_parts[blip.get(qn("r:embed"))].blob).hexdigest()
            for blip in document.element.xpath(".//a:blip") if blip.get(qn("r:embed"))}
        assert self.image_sha in embedded, "inserted picture has no document drawing relationship"
        pictures = [shape for shape in document.element.xpath(".//wp:inline | .//wp:anchor")
            if any(hashlib.sha256(document.part.related_parts[blip.get(qn("r:embed"))].blob).hexdigest() == self.image_sha
                for blip in shape.xpath(".//a:blip") if blip.get(qn("r:embed")))]
        assert len(pictures) == 1, "expected one test picture drawing"
        size = pictures[0].find(qn("wp:extent"))
        extent = (int(size.get("cx")), int(size.get("cy")))
        from docx.shared import Inches
        assert extent[0] > Inches(96 / 96), "saved picture retained its original width"
        if self.extent is not None:
            assert extent == self.extent, "picture dimensions changed after a later save/reopen"
        self.extent = extent
        (self.root / "rich-result.json").write_text(json.dumps({"table_marker":self.marker,
            "tables_preserved":len(document.tables), "image_sha256":self.image_sha,
            "original_media_preserved":len(self.original_media), "picture_extent_emu":extent,
            "picture_width_ui":self.dimensions}, indent=2))
