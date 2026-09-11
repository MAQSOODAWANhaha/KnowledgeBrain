#!/usr/bin/env python3
"""Run the required Rust replay contract against the existing Python service.

Default inputs are generated locally for CI. Explicit --docx/--pdf exercise real
tenders locally; the Rust contract mocks vision and never invokes an LLM.
"""
import argparse
from concurrent.futures import ThreadPoolExecutor
import hashlib
import json
import os
from pathlib import Path
import secrets
import subprocess
import sys
import tempfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "services"))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--docx", type=Path)
    parser.add_argument("--pdf", type=Path)
    args = parser.parse_args()
    if bool(args.docx) != bool(args.pdf):
        parser.error("--docx and --pdf must be supplied together")

    import grpc
    from docreader.auth import AuthInterceptor
    from docreader.config import CONFIG
    from docreader.main import DocReaderServicer
    from docreader.proto.docreader_pb2_grpc import add_DocReaderServicer_to_server

    with tempfile.TemporaryDirectory(prefix="kb-docreader-replay-") as temporary:
        directory = Path(temporary)
        if args.docx:
            docx_path, pdf_path = args.docx.resolve(strict=True), args.pdf.resolve(strict=True)
        else:
            from docx import Document
            from pypdf import PdfWriter
            from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

            docx_path, pdf_path = directory / "input.docx", directory / "input.pdf"
            document = Document()
            document.add_heading("Replay acceptance input A", 1)
            document.add_paragraph("The same uploaded bytes must reuse the successful conversion.")
            table = document.add_table(rows=2, cols=2)
            for cell, text in zip((c for row in table.rows for c in row.cells), ["Item", "Response", "A", ""]):
                cell.text = text
            document.save(docx_path)
            document = PdfWriter()
            page = document.add_blank_page(width=612, height=792)
            font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                     NameObject("/Subtype"): NameObject("/Type1"),
                                     NameObject("/BaseFont"): NameObject("/Helvetica")})
            page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({
                NameObject("/F1"): document._add_object(font)})})
            content = DecodedStreamObject()
            content.set_data(b"BT /F1 12 Tf 72 720 Td (Different PDF input B must trigger a new conversion.) Tj ET")
            page[NameObject("/Contents")] = document._add_object(content)
            document.write(pdf_path)

        token = secrets.token_hex()
        old_token = os.environ.get("GRPC_AUTH_TOKEN")
        os.environ["GRPC_AUTH_TOKEN"] = token
        executor = ThreadPoolExecutor(max_workers=CONFIG.grpc_max_workers)
        server = grpc.server(executor, interceptors=[AuthInterceptor()], options=[
            ("grpc.max_send_message_length", CONFIG.grpc_max_file_size_mb),
            ("grpc.max_receive_message_length", CONFIG.grpc_max_file_size_mb),
        ])
        add_DocReaderServicer_to_server(DocReaderServicer(), server)
        port = server.add_insecure_port("127.0.0.1:0")
        channel = None
        try:
            server.start()
            channel = grpc.insecure_channel(f"127.0.0.1:{port}")
            grpc.channel_ready_future(channel).result(timeout=20)
            env = dict(os.environ, DOCREADER_ADDR=f"127.0.0.1:{port}", GRPC_AUTH_TOKEN=token,
                       KNOWLEDGEBRAIN_REQUIRE_DOCREADER_TESTS="1",
                       KB_DOCREADER_TEST_DOCX=str(docx_path), KB_DOCREADER_TEST_PDF=str(pdf_path))
            command = ["cargo", "test", "--locked", "-p", "bidding", "--features", "docreader-contract-tests",
                       "--test", "tender_document_process_real_parse_counts", "--", "--nocapture"]
            result = subprocess.run(command, cwd=ROOT, env=env, capture_output=True, text=True)
            sys.stdout.write(result.stdout)
            sys.stderr.write(result.stderr)
            if result.returncode:
                raise SystemExit(result.returncode)
            if "1 passed; 0 failed; 0 ignored" not in result.stdout:
                raise SystemExit("required DocReader replay contract did not execute exactly one passing test")
            print(json.dumps({"status": "passed", "inputs": "real" if args.docx else "generated",
                              "docx_sha256": hashlib.sha256(docx_path.read_bytes()).hexdigest(),
                              "pdf_sha256": hashlib.sha256(pdf_path.read_bytes()).hexdigest(),
                              "service": "existing_docreader_servicer", "external_model_calls": 0}))
        finally:
            if channel is not None:
                channel.close()
            server.stop(0).wait()
            executor.shutdown(wait=True)
            if old_token is None:
                os.environ.pop("GRPC_AUTH_TOKEN", None)
            else:
                os.environ["GRPC_AUTH_TOKEN"] = old_token


if __name__ == "__main__":
    main()
